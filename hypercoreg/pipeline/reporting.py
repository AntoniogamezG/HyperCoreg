"""Manifest and summary reporting helpers.

This module is the migration target for reporting/runtime metadata helpers
previously hosted in ``_legacy_coreg.py``.
"""

from __future__ import annotations

import json
import hashlib
import math
import logging
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger("COREG_PROCESSING")

RUN_MANIFEST_SCHEMA_VERSION = 1
METADATA_SCHEMA_VERSION = 2

DATASET_XLSX_COLUMNS: List[str] = [
    "folder_name",
    "prisma_date",
    "prisma_cloud_pct",
    "prisma_sea_pct",
    "observation_angle",
    "rel_azimuth_angle",
    "sun_azimuth_angle",
    "solar_zenith_angle",
    "bbox_top_left_x",
    "bbox_top_left_y",
    "bbox_bottom_right_x",
    "bbox_bottom_right_y",
    "s2_product_id",
    "s2_date",
    "s2_cloud_cover_pct",
    "tie_points_count",
    "accuracy_pct",
    "residual_mean_m",
    "residual_median_m",
    "residual_rmse_m",
    "residual_p90_m",
    "polynomial_warp_used",
    "polynomial_n_gcps",
    "status",
    "notes",
    "rmse_global",
    "rmse_local",
    "rmse_improvement_pct",
    "spatial_spread_score",
    "hull_bbox_ratio",
    "is_clustered",
    "quality_tier",
    "quality_score",
    "ssim_before",
    "ssim_after",
    "ssim_delta",
    "multiband_tiepoint_counts.B02",
    "multiband_tiepoint_counts.B03",
    "multiband_tiepoint_counts.B04",
    "multiband_tiepoint_counts.B08",
    "multiband_tiepoint_counts.B11",
    "multiband_tiepoint_counts.B12",
    "multiband_tiepoint_counts",
]
DATASET_MULTIBAND_COLUMNS: Tuple[str, ...] = ("B02", "B03", "B04", "B08", "B11", "B12")
BATCH_SUMMARY_CONTEXT_COLUMNS: Tuple[str, ...] = (
    "filename",
    "hyp_type",
    "output_path",
    "dataset_xlsx_path",
    "run_manifest_path",
    "displacement_vectors_path",
)
BATCH_SUMMARY_XLSX_COLUMNS: List[str] = list(DATASET_XLSX_COLUMNS) + list(BATCH_SUMMARY_CONTEXT_COLUMNS)


def _fmt_issue(scope: str, message: str) -> str:
    return f"[{scope}] {message}"


def _atomic_write_text(path: str, text: str, encoding: str = "utf-8") -> None:
    target = os.path.abspath(path)
    parent = os.path.dirname(target)
    if parent:
        os.makedirs(parent, exist_ok=True)

    tmp_path = None
    try:
        with NamedTemporaryFile("w", encoding=encoding, dir=parent or None, delete=False) as tmp:
            tmp.write(text)
            tmp.flush()
            os.fsync(tmp.fileno())
            tmp_path = tmp.name
        os.replace(tmp_path, target)
        tmp_path = None
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


def _sanitize_manifest_value(key: str, value: Any) -> Any:
    key_l = str(key).lower()
    if any(token in key_l for token in ("password", "secret", "token", "client_secret")):
        return "***REDACTED***"
    if callable(value):
        name = getattr(value, "__name__", "callable")
        return f"<callable:{name}>"
    if isinstance(value, Path):
        return str(value)
    if hasattr(value, "tolist") and not isinstance(value, (str, bytes)):
        try:
            return _sanitize_manifest_value(key, value.tolist())
        except Exception:
            pass
    if isinstance(value, bool) or value is None or isinstance(value, str):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (list, tuple)):
        return [_sanitize_manifest_value(key, item) for item in value]
    if isinstance(value, dict):
        out = {}
        for k, v in value.items():
            out[str(k)] = _sanitize_manifest_value(str(k), v)
        return out
    try:
        numeric = float(value)
    except Exception:
        return str(value)
    return numeric if math.isfinite(numeric) else None


def _sanitize_config_for_manifest(config: Dict[str, Any]) -> Dict[str, Any]:
    safe: Dict[str, Any] = {}
    for key in sorted(config.keys()):
        if str(key).startswith("_"):
            continue
        safe[str(key)] = _sanitize_manifest_value(str(key), config.get(key))
    return safe


def _write_shift_report(report_path, title, coreg_info, extra_lines=None):
    try:
        lines = [f"{'=' * 60}", f"{title}", f"{'=' * 60}"]
        lines.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("")
        if coreg_info:
            for key, val in coreg_info.items():
                lines.append(f"{key}: {val}")
        if extra_lines:
            for line in extra_lines:
                lines.append(f"{line}")
        _atomic_write_text(report_path, "\n".join(lines), encoding="utf-8")
        logger.debug("Shift report written: %s", report_path)
    except Exception as exc:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write shift report: {exc}"))


def _write_per_scene_metrics_json(metrics_dict, json_path):
    try:
        payload = json.dumps(
            _sanitize_manifest_value("metrics", metrics_dict),
            indent=2,
            allow_nan=False,
            default=str,
        )
        _atomic_write_text(json_path, payload, encoding="utf-8")
        logger.debug("Metrics JSON written: %s", json_path)
    except Exception as exc:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write metrics JSON: {exc}"))


def _write_scene_run_manifest(manifest_dict: Dict[str, Any], json_path: str) -> None:
    try:
        payload = json.dumps(
            _sanitize_manifest_value("manifest", manifest_dict),
            indent=2,
            allow_nan=False,
            default=str,
        )
        _atomic_write_text(json_path, payload, encoding="utf-8")
        logger.debug("Run manifest JSON written: %s", json_path)
    except Exception as exc:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write run manifest JSON: {exc}"))


def _safe_dataset_file_stem(raw_value: Any, fallback: str = "scene") -> str:
    text = str(raw_value).strip() if raw_value is not None else ""
    if not text:
        text = fallback
    text = text.encode("ascii", errors="ignore").decode("ascii")
    text = re.sub(r'[<>:"/\\|?*]+', "_", text)
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
    text = text.strip(" ._")
    return text or fallback


def _build_scene_output_name(source_path: str, hyp_type: Optional[str], acquisition_time: Any = None) -> str:
    sensor = _infer_sensor_from_identifiers(hyp_type, None, Path(str(source_path)).name)
    sensor_tag = str(sensor or hyp_type or "SCENE").upper().strip() or "SCENE"
    source_text = os.path.abspath(str(source_path)) if source_path else ""
    token_src = source_text or str(source_path) or sensor_tag
    token = hashlib.sha1(token_src.encode("utf-8", errors="ignore")).hexdigest()[:8].upper()
    stem = _safe_dataset_file_stem(Path(str(source_path)).stem if source_path else "scene", fallback="scene")
    if isinstance(acquisition_time, datetime):
        date_tag = acquisition_time.strftime("%y%m%d")
        return f"{sensor_tag}_{date_tag}_{stem}_{token}"
    return f"{sensor_tag}_{stem}_{token}"


def _row_identity_key(row: Dict[str, Any]) -> str:
    for field in ("source_path", "run_manifest_path", "dataset_xlsx_path"):
        value = row.get(field)
        if value:
            return str(os.path.normcase(os.path.abspath(str(value))))
    for field in ("scene_name", "filename"):
        value = row.get(field)
        if value:
            return f"{field}:{str(value)}"
    return ""


def _resolve_dataset_scene_name(metrics_dict: Dict[str, Any]) -> str:
    scene_name = metrics_dict.get("scene_name") or metrics_dict.get("folder_name")
    if scene_name:
        return str(scene_name)
    filename = metrics_dict.get("filename")
    if filename:
        return Path(str(filename)).stem
    return "scene"


def _build_dataset_xlsx_filename(scene_name: Optional[str], filename: Optional[str] = None) -> str:
    label = scene_name
    if not label and filename:
        label = Path(str(filename)).stem
    safe_label = _safe_dataset_file_stem(label, fallback="scene")
    return f"{safe_label}_DATASET.xlsx"


def _infer_sensor_from_identifiers(
    hyp_type: Optional[str],
    scene_name: Optional[str],
    filename: Optional[str],
) -> Optional[str]:
    if hyp_type is not None:
        sensor = str(hyp_type).upper().strip()
        if sensor in {"PRISMA", "ENMAP"}:
            return sensor

    for candidate in (scene_name, filename):
        if candidate is None:
            continue
        text = str(candidate).upper()
        if "PRISMA" in text:
            return "PRISMA"
        if "ENMAP" in text or "SPECTRAL_IMAGE" in text:
            return "ENMAP"

    if filename:
        suffix = Path(str(filename)).suffix.lower()
        if suffix == ".he5":
            return "PRISMA"

    return None


def _build_dataset_row(metrics_dict: Dict[str, Any]) -> Dict[str, Any]:
    row = {column: None for column in DATASET_XLSX_COLUMNS}
    row["folder_name"] = _resolve_dataset_scene_name(metrics_dict)

    for column in DATASET_XLSX_COLUMNS:
        if column == "folder_name" or column.startswith("multiband_tiepoint_counts."):
            continue
        if column in metrics_dict:
            row[column] = metrics_dict.get(column)

    if row.get("status") is None and metrics_dict.get("status") is not None:
        row["status"] = metrics_dict.get("status")

    if row.get("notes") is None:
        for key in ("notes", "reason", "error", "message"):
            val = metrics_dict.get(key)
            if val is not None:
                row["notes"] = str(val)
                break

    if row.get("polynomial_n_gcps") is None:
        poly_decision = metrics_dict.get("polynomial_order_decision")
        if isinstance(poly_decision, dict) and poly_decision.get("n_gcps") is not None:
            row["polynomial_n_gcps"] = poly_decision.get("n_gcps")

    tp_residuals = metrics_dict.get("tp_residuals")
    if isinstance(tp_residuals, dict):
        for key in ("residual_mean_m", "residual_median_m", "residual_rmse_m", "residual_p90_m"):
            if row.get(key) is None and tp_residuals.get(key) is not None:
                row[key] = tp_residuals.get(key)

    multiband_counts = metrics_dict.get("multiband_tiepoint_counts")
    total_count = None
    if isinstance(multiband_counts, dict):
        running_total = 0
        running_total_valid = False
        for band in DATASET_MULTIBAND_COLUMNS:
            val = multiband_counts.get(band)
            if val is None:
                val = multiband_counts.get(band.lower())
            if val is not None:
                row[f"multiband_tiepoint_counts.{band}"] = val
                try:
                    running_total += int(val)
                    running_total_valid = True
                except Exception:
                    pass
        if multiband_counts.get("total") is not None:
            total_count = multiband_counts.get("total")
        elif multiband_counts.get("TOTAL") is not None:
            total_count = multiband_counts.get("TOTAL")
        elif running_total_valid:
            total_count = running_total
    elif multiband_counts is not None:
        total_count = multiband_counts

    if total_count is None and metrics_dict.get("multiband_tiepoint_counts_total") is not None:
        total_count = metrics_dict.get("multiband_tiepoint_counts_total")

    if total_count is not None:
        row["multiband_tiepoint_counts"] = total_count

    return row


def _write_single_scene_dataset_xlsx(metrics_dict: Dict[str, Any], xlsx_path: str) -> bool:
    try:
        from openpyxl import Workbook
    except ImportError:
        logger.debug("openpyxl not available, skipping per-scene dataset workbook.")
        return False

    try:
        parent = os.path.dirname(xlsx_path)
        if parent:
            os.makedirs(parent, exist_ok=True)

        row = _build_dataset_row(metrics_dict)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        for col_idx, column in enumerate(DATASET_XLSX_COLUMNS, 1):
            ws.cell(row=1, column=col_idx, value=column)
            ws.cell(row=2, column=col_idx, value=row.get(column))

        wb.save(xlsx_path)
        logger.info("Per-scene dataset workbook written: %s", xlsx_path)
        return True
    except Exception as exc:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write per-scene dataset workbook: {exc}"))
        return False


def _build_batch_summary_row(metrics_dict: Dict[str, Any]) -> Dict[str, Any]:
    row = _build_dataset_row(metrics_dict)
    for column in BATCH_SUMMARY_CONTEXT_COLUMNS:
        row[column] = metrics_dict.get(column)
    return row


def _collect_batch_summary_rows(results: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    success_rows: List[Dict[str, Any]] = []
    for metrics in results.get("metrics", []):
        if isinstance(metrics, dict):
            success_rows.append(_build_batch_summary_row(metrics))

    skipped_rows: List[Dict[str, Any]] = []
    for metrics in results.get("skipped", []):
        if isinstance(metrics, dict):
            skipped_rows.append(_build_batch_summary_row(metrics))

    failed_rows: List[Dict[str, Any]] = []
    seen_failed_keys: set[str] = set()
    for metrics in results.get("failed_details", []):
        if not isinstance(metrics, dict):
            continue
        failed_rows.append(_build_batch_summary_row(metrics))
        identity_key = _row_identity_key(metrics)
        if identity_key:
            seen_failed_keys.add(identity_key)

    errors = results.get("errors", {})
    for failed_name in results.get("failed", []):
        if failed_name is None:
            continue
        failed_key = str(os.path.normcase(os.path.abspath(str(failed_name))))
        if failed_key in seen_failed_keys:
            continue
        error_message = None
        if isinstance(errors, dict):
            error_message = errors.get(failed_name)
            if error_message is None:
                error_message = errors.get(str(failed_name))
        fallback = _build_failed_scene_metrics(
            hs_file=str(failed_name),
            hyp_type=None,
            error_message=str(error_message if error_message is not None else "Unknown"),
        )
        failed_rows.append(_build_batch_summary_row(fallback))

    return {"success": success_rows, "skipped": skipped_rows, "failed": failed_rows}


def _write_batch_summary_txt(
    summary_path: str,
    results: Dict[str, Any],
    grouped_rows: Dict[str, List[Dict[str, Any]]],
) -> None:
    with open(summary_path, "w", encoding="utf-8") as f:
        f.write("Batch Processing Summary\n")
        f.write("=" * 60 + "\n")
        f.write(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write(f"Total: {results.get('total', 0)}\n")
        f.write(f"  Succeeded: {len(results.get('succeeded', []))}\n")
        f.write(f"  Skipped: {len(results.get('skipped', []))}\n")
        f.write(f"  Failed: {len(results.get('failed', []))}\n\n")

        f.write("Scene details:\n")
        f.write("-" * 60 + "\n")
        section_defs: Tuple[Tuple[str, str], ...] = (("Succeeded", "success"), ("Skipped", "skipped"), ("Failed", "failed"))
        for section_label, section_key in section_defs:
            rows = grouped_rows.get(section_key, [])
            if not rows:
                continue
            f.write(f"{section_label} ({len(rows)}):\n")
            for row in rows:
                scene_name = row.get("folder_name") or row.get("filename") or "n/a"
                status = row.get("status") or "UNKNOWN"
                sensor = row.get("hyp_type") or "n/a"
                note = row.get("notes")
                note_text = " ".join(str(note).split()) if note is not None else ""
                detail = f"  {scene_name} | {status} | {sensor}"
                if note_text:
                    detail += f" | {note_text}"
                f.write(detail + "\n")
            f.write("\n")


def _write_batch_summary_xlsx(
    xlsx_path: str,
    results: Dict[str, Any],
    grouped_rows: Dict[str, List[Dict[str, Any]]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    overview_rows: Tuple[Tuple[str, Any], ...] = (
        ("Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total", results.get("total", 0)),
        ("Succeeded", len(results.get("succeeded", []))),
        ("Skipped", len(results.get("skipped", []))),
        ("Failed", len(results.get("failed", []))),
    )
    for row_idx, (label, value) in enumerate(overview_rows, 1):
        label_cell = ws_overview.cell(row=row_idx, column=1, value=label)
        label_cell.font = Font(bold=True)
        ws_overview.cell(row=row_idx, column=2, value=value)

    header_fill = PatternFill(start_color="CCCCCC", fill_type="solid")
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center")

    def _write_status_sheet(title: str, rows: List[Dict[str, Any]]) -> None:
        ws = wb.create_sheet(title=title)
        ws.freeze_panes = "A2"
        for col_idx, header in enumerate(BATCH_SUMMARY_XLSX_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        for row_idx, row_values in enumerate(rows, 2):
            for col_idx, header in enumerate(BATCH_SUMMARY_XLSX_COLUMNS, 1):
                ws.cell(row=row_idx, column=col_idx, value=row_values.get(header))

    _write_status_sheet("Success", grouped_rows.get("success", []))
    _write_status_sheet("Skipped", grouped_rows.get("skipped", []))
    _write_status_sheet("Failed", grouped_rows.get("failed", []))

    wb.save(xlsx_path)


def _build_failed_scene_metrics(hs_file: str, hyp_type: Optional[str], error_message: str) -> Dict[str, Any]:
    filename = os.path.basename(hs_file)
    inferred_sensor = _infer_sensor_from_identifiers(hyp_type, None, filename)
    scene_name = _build_scene_output_name(hs_file, inferred_sensor)
    return {
        "scene_name": scene_name,
        "filename": filename,
        "source_path": hs_file,
        "hyp_type": inferred_sensor,
        "status": "FAIL",
        "error": error_message,
        "notes": error_message,
    }


def _build_skip_result(
    hs_file: str,
    hyp_type: str,
    output_dir: str,
    scene_idx: int,
    scene_total: int,
    hs_time: Any,
    bbox: Any,
    config: Dict[str, Any],
    status: str,
    reason: str,
    normalization_mode: str,
    normalization_params: Dict[str, Any],
    build_overviews: bool,
    remove_detector_overlap_bands: bool,
    strict_metadata: bool,
    metadata_extension_level: str,
    metadata_stats_mode: str,
    metadata_stats_sample_windows: int,
    metadata_stats_seed: int,
    metadata_histogram_buckets: int,
    metadata_label_precision: int,
    validation_max_windows: int,
    extra_summary: Optional[Dict[str, Any]] = None,
    extra_metrics: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    scene_name = _build_scene_output_name(hs_file, hyp_type, hs_time)
    skip_manifest_path = os.path.join(output_dir, f"{scene_name}_run_manifest.json")
    skip_dataset_xlsx_path = os.path.join(
        output_dir,
        _build_dataset_xlsx_filename(scene_name=scene_name, filename=os.path.basename(hs_file)),
    )
    skip_metrics = {
        "status": status,
        "scene_name": scene_name,
        "filename": os.path.basename(hs_file),
        "source_path": hs_file,
        "hyp_type": hyp_type,
        "reason": reason,
        "metadata_schema_version": METADATA_SCHEMA_VERSION,
        "metadata_status": "skipped",
        "metadata_warnings": [],
        "normalization_mode": str(normalization_mode),
        "normalization_params": dict(normalization_params),
        "build_overviews": bool(build_overviews),
        "remove_overlapping_bands": bool(remove_detector_overlap_bands),
        "strict_metadata": bool(strict_metadata),
        "metadata_extension_level": metadata_extension_level,
        "metadata_stats_mode": str(metadata_stats_mode),
        "metadata_stats_sample_windows": int(metadata_stats_sample_windows),
        "metadata_stats_seed": int(metadata_stats_seed),
        "metadata_histogram_buckets": metadata_histogram_buckets,
        "metadata_label_precision": metadata_label_precision,
        "validation_max_windows": int(validation_max_windows),
        "ancillary_status": "not_processed",
        "ancillary_warp_method": None,
        "ancillary_tiepoints_used": 0,
        "ancillary_outputs": {
            "pan": None,
            "quality_vnir": None,
            "quality_swir": None,
            "enmap_ql": {},
            "enmap_sidecars": {},
        },
        "ancillary_warnings": [],
        "run_manifest_path": skip_manifest_path,
        "displacement_vectors_path": None,
        "dataset_xlsx_path": skip_dataset_xlsx_path,
    }
    if extra_metrics:
        skip_metrics.update(dict(extra_metrics))

    skip_manifest = {
        "manifest_schema_version": RUN_MANIFEST_SCHEMA_VERSION,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "scene": {
            "scene_name": scene_name,
            "filename": os.path.basename(hs_file),
            "sensor_type": hyp_type,
            "scene_idx": int(scene_idx),
            "scene_total": int(scene_total),
        },
        "paths": {
            "output_dir": output_dir,
            "run_manifest": skip_manifest_path,
            "displacement_vectors_path": None,
            "dataset_xlsx": skip_dataset_xlsx_path,
        },
        "input": {
            "hs_file": hs_file,
            "acquisition_time": hs_time.isoformat() if isinstance(hs_time, datetime) else None,
            "bbox": list(bbox) if bbox is not None else None,
        },
        "config": _sanitize_config_for_manifest(config),
        "summary": {"reason": reason},
    }
    if extra_summary:
        skip_manifest["summary"].update(dict(extra_summary))

    _write_scene_run_manifest(skip_manifest, skip_manifest_path)
    _write_single_scene_dataset_xlsx(skip_metrics, skip_dataset_xlsx_path)
    return skip_metrics


__all__ = [
    "_build_batch_summary_row",
    "_build_dataset_row",
    "_build_dataset_xlsx_filename",
    "_build_failed_scene_metrics",
    "_build_skip_result",
    "_collect_batch_summary_rows",
    "_infer_sensor_from_identifiers",
    "_resolve_dataset_scene_name",
    "_safe_dataset_file_stem",
    "_sanitize_config_for_manifest",
    "_sanitize_manifest_value",
    "_write_batch_summary_txt",
    "_write_batch_summary_xlsx",
    "_write_per_scene_metrics_json",
    "_write_scene_run_manifest",
    "_write_shift_report",
    "_write_single_scene_dataset_xlsx",
]
