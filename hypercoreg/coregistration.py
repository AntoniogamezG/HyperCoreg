"""
Core coregistration workflow for HyperCoreg.

This module contains the main coregistration functions that orchestrate
the entire pipeline from input to output for PRISMA and EnMAP hyperspectral
imagery alignment to Sentinel-2 reference data.
"""

import os
import sys
import uuid
import shutil
import logging
import subprocess
import io
import contextlib
import re
import json
import inspect
import zipfile
import textwrap
import tempfile
import threading
import warnings
import getpass
from pathlib import Path
from datetime import datetime, timezone, timedelta
from time import perf_counter, sleep
from typing import Dict, Any, List, Optional, Tuple, Callable, Sequence
from xml.etree import ElementTree as ET
from xml.dom import minidom

import numpy as np
import requests
from requests.auth import AuthBase
import rasterio
from rasterio.errors import RasterioIOError
from rasterio import warp
from rasterio.enums import Resampling
from rasterio.transform import Affine, from_bounds as transform_from_bounds
from rasterio.windows import Window, from_bounds
from rasterio.plot import plotting_extent
from pyproj import Transformer, CRS
from shapely.geometry import Polygon, shape
import shapely.wkt as shapely_wkt
from geoarray import GeoArray
from arosics import COREG_LOCAL, COREG

from hypercoreg.logging_config import log_section_header, log_subsection_header, setup_logging
from hypercoreg.utils import (
    detect_hyp_type, SentinelNotFoundError, resolve_gdalwarp_exe,
    create_narrowband_average, diagnose_raster
)
from hypercoreg.config import (
    S2_BANDS,
    MULTIBAND_S2_WAVELENGTHS,
    DEFAULT_CONFIG,
    PRISMA_FIXED_BAND_PAIRS,
)
from hypercoreg.spectral import build_prisma_band_table, build_enmap_band_table
from hypercoreg.normalization import (
    NormalizationParams,
    normalize_mode,
    sanitize_normalization_params,
    normalize_raster_to_path,
    stream_copy_raster_to_path,
)
from hypercoreg.readers.prisma import (
    estimate_prisma_geotransform, read_prisma_cube_and_meta,
    read_prisma_pan_and_geo, read_prisma_quality_mask,
    extract_prisma_extended_metadata, check_cloud_threshold
)
from hypercoreg.readers.enmap import (
    find_enmap_metadata_for_spectral_image, read_enmap_metadata,
    inject_metadata_into_raster, read_enmap_metadata_from_raster,
    derive_enmap_bbox_from_raster,
)

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import cm
    from matplotlib.colors import ListedColormap
    from matplotlib.lines import Line2D
    HAS_MATPLOTLIB = True
except ImportError:
    HAS_MATPLOTLIB = False

try:
    import geopandas as gpd
    HAS_GEOPANDAS = True
except ImportError:
    gpd = None
    HAS_GEOPANDAS = False

logger = logging.getLogger("COREG_PROCESSING")

# Processing constants
PROCESSING_DTYPE = np.float32
PROCESSING_NODATA = -9999.0
CPUS_FOR_AROSICS = 1
LOCAL_GRID_RES_M = 150
S2_BAND08_CENTER_WL_NM = 842.0
MIN_TIE_POINTS_FOR_POLYNOMIAL = 10
MIN_RELIABILITY_THRESHOLD = 75.0
POLYNOMIAL_FALLBACK_TO_AROSICS = True
PRISMA_OUTPUT_RESOLUTION = 30.0
DEFAULT_QUICKLOOK_MAX_DIM = 1800
CATALOGUE_URL = "https://catalogue.dataspace.copernicus.eu/odata/v1/Products"
TOKEN_URL = "https://identity.dataspace.copernicus.eu/auth/realms/CDSE/protocol/openid-connect/token"
HTTP_TIMEOUT_S = 30
HTTP_TIMEOUT_CONNECT_READ = (30, 180)
HTTP_AUTH_RETRY_ATTEMPTS = 2
HTTP_QUERY_RETRY_ATTEMPTS = 3
HTTP_DOWNLOAD_RETRY_ATTEMPTS = 3
HTTP_RETRY_BACKOFF_BASE_S = 1.0
HTTP_RETRY_BACKOFF_FACTOR = 2.0
HTTP_RETRYABLE_STATUS_CODES = {408, 429, 500, 502, 503, 504}
PROGRESS_HEARTBEAT_INTERVAL_S = 20.0
RUN_MANIFEST_SCHEMA_VERSION = 1

DATASET_XLSX_COLUMNS: List[str] = [
    "folder_name",
    "prisma_date",
    "prisma_cloud_pct",
    "prisma_sea_pct",
    "observation_angle",
    "rel_azimuth_angle",
    "sun_azimuth_angle",
    "solar_zenith_angle",
    "bbox_top_left_x",
    "bbox_top_left_y",
    "bbox_bottom_right_x",
    "bbox_bottom_right_y",
    "s2_product_id",
    "s2_date",
    "s2_cloud_cover_pct",
    "tie_points_count",
    "accuracy_pct",
    "residual_mean_m",
    "residual_median_m",
    "residual_rmse_m",
    "residual_p90_m",
    "polynomial_warp_used",
    "polynomial_n_gcps",
    "status",
    "notes",
    "rmse_global",
    "rmse_local",
    "rmse_improvement_pct",
    "spatial_spread_score",
    "hull_bbox_ratio",
    "is_clustered",
    "quality_tier",
    "quality_score",
    "ssim_before",
    "ssim_after",
    "ssim_delta",
    "multiband_tiepoint_counts.B02",
    "multiband_tiepoint_counts.B03",
    "multiband_tiepoint_counts.B04",
    "multiband_tiepoint_counts.B08",
    "multiband_tiepoint_counts.B11",
    "multiband_tiepoint_counts.B12",
    "multiband_tiepoint_counts",
]
DATASET_MULTIBAND_COLUMNS: Tuple[str, ...] = ("B02", "B03", "B04", "B08", "B11", "B12")

# Scene Classification Layer (SCL) classes to exclude
SCL_EXCLUDE_CLASSES = {
    0: "No Data", 1: "Saturated/Defective", 3: "Cloud Shadows",
    8: "Cloud Medium Probability", 9: "Cloud High Probability", 10: "Thin Cirrus"
}

METADATA_SCHEMA_VERSION = 2
ENVI_DTYPE_MAP = {
    "uint8": 1,
    "int16": 2,
    "int32": 3,
    "float32": 4,
    "float64": 5,
    "complex64": 6,
    "complex128": 9,
    "uint16": 12,
    "uint32": 13,
    "int64": 14,
    "uint64": 15,
}

PromptUserpassFn = Callable[[], Tuple[str, str, Optional[str]]]
_CDSE_CREDENTIAL_CACHE_LOCK = threading.Lock()
_CDSE_CREDENTIAL_CACHE: Dict[str, Optional[str]] = {
    "username": None,
    "password": None,
    "totp": None,
}


def _fmt_issue(scope: str, message: str) -> str:
    """Standardized issue message format for logs and raised exceptions."""
    return f"[{scope}] {message}"


def _retry_delay_seconds(retry_index: int) -> float:
    """Compute exponential backoff delay for retry index (1-based)."""
    idx = max(1, int(retry_index))
    return float(HTTP_RETRY_BACKOFF_BASE_S * (HTTP_RETRY_BACKOFF_FACTOR ** (idx - 1)))


def _is_retryable_http_status(status_code: int) -> bool:
    """Return True when an HTTP status code should be retried."""
    return int(status_code) in HTTP_RETRYABLE_STATUS_CODES


def _sanitize_manifest_value(key: str, value: Any) -> Any:
    """Redact sensitive values and coerce non-serializable config values."""
    key_l = str(key).lower()
    if any(token in key_l for token in ("password", "secret", "token", "client_secret")):
        return "***REDACTED***"
    if callable(value):
        name = getattr(value, "__name__", "callable")
        return f"<callable:{name}>"
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, (list, tuple)):
        return [_sanitize_manifest_value(key, item) for item in value]
    if isinstance(value, dict):
        out = {}
        for k, v in value.items():
            out[str(k)] = _sanitize_manifest_value(str(k), v)
        return out
    return str(value)


def _sanitize_config_for_manifest(config: Dict[str, Any]) -> Dict[str, Any]:
    """Build a JSON-safe and redacted copy of effective run configuration."""
    safe: Dict[str, Any] = {}
    for key in sorted(config.keys()):
        if str(key).startswith("_"):
            continue
        safe[str(key)] = _sanitize_manifest_value(str(key), config.get(key))
    return safe


def _cache_cdse_public_credentials(
    username: Optional[str],
    password: Optional[str],
    totp: Optional[str] = None,
) -> None:
    """Cache CDSE username/password in-memory for the current process."""
    user = str(username or "").strip()
    pwd = str(password or "")
    otp = str(totp).strip() if totp is not None else ""
    with _CDSE_CREDENTIAL_CACHE_LOCK:
        _CDSE_CREDENTIAL_CACHE["username"] = user or None
        _CDSE_CREDENTIAL_CACHE["password"] = pwd if pwd else None
        _CDSE_CREDENTIAL_CACHE["totp"] = otp or None


def _get_cached_cdse_public_credentials() -> Optional[Tuple[str, str, Optional[str]]]:
    """Return cached CDSE username/password, if present."""
    with _CDSE_CREDENTIAL_CACHE_LOCK:
        user = str(_CDSE_CREDENTIAL_CACHE.get("username") or "").strip()
        pwd = _CDSE_CREDENTIAL_CACHE.get("password") or ""
        otp = str(_CDSE_CREDENTIAL_CACHE.get("totp") or "").strip() or None
    if user and pwd:
        return user, str(pwd), otp
    return None


def _clear_cached_cdse_public_credentials() -> None:
    """Clear in-memory cached CDSE username/password."""
    with _CDSE_CREDENTIAL_CACHE_LOCK:
        _CDSE_CREDENTIAL_CACHE["username"] = None
        _CDSE_CREDENTIAL_CACHE["password"] = None
        _CDSE_CREDENTIAL_CACHE["totp"] = None


class _BearerAuth(AuthBase):
    """Attach Bearer token auth to every outgoing request."""

    def __init__(self, token: str):
        self.token = token

    def __call__(self, request):
        request.headers["Authorization"] = f"Bearer {self.token}"
        return request


def _create_cdse_bearer_session(access_token: str) -> requests.Session:
    """Create session from an issued bearer token."""
    token = (access_token or "").strip()
    if not token:
        raise ValueError(_fmt_issue("AUTH", "Empty CDSE bearer token provided."))
    s = requests.Session()
    s.auth = _BearerAuth(token)
    s.headers.update({"Accept": "application/json"})
    return s


def _request_cdse_access_token(payload: Dict[str, str], flow_name: str) -> str:
    """Request CDSE access token with retry logic."""
    attempts = max(1, int(HTTP_AUTH_RETRY_ATTEMPTS))
    last_error: Optional[Exception] = None
    for attempt in range(1, attempts + 1):
        try:
            r = requests.post(TOKEN_URL, data=payload, timeout=HTTP_TIMEOUT_S)
            if r.status_code >= 400:
                if _is_retryable_http_status(r.status_code) and attempt < attempts:
                    delay_s = _retry_delay_seconds(attempt)
                    logger.warning(
                        _fmt_issue(
                            "AUTH",
                            f"{flow_name} token retryable HTTP {r.status_code} "
                            f"(attempt {attempt}/{attempts}); retrying in {delay_s:.1f}s.",
                        )
                    )
                    sleep(delay_s)
                    continue
                msg = r.text[:500]
                raise RuntimeError(
                    _fmt_issue("AUTH", f"{flow_name} token generation failed ({r.status_code}): {msg}")
                )

            token = r.json().get("access_token")
            if not token:
                raise RuntimeError(_fmt_issue("AUTH", f"{flow_name} token response missing access_token."))
            return token
        except requests.RequestException as e:
            last_error = e
            if attempt < attempts:
                delay_s = _retry_delay_seconds(attempt)
                logger.warning(
                    _fmt_issue(
                        "AUTH",
                        f"{flow_name} token request failed (attempt {attempt}/{attempts}): {e}; "
                        f"retrying in {delay_s:.1f}s.",
                    )
                )
                sleep(delay_s)
                continue
            raise RuntimeError(_fmt_issue("AUTH", f"{flow_name} token request failed: {e}")) from e
        except Exception as e:
            last_error = e
            if attempt < attempts:
                delay_s = _retry_delay_seconds(attempt)
                logger.warning(
                    _fmt_issue(
                        "AUTH",
                        f"{flow_name} token generation failed (attempt {attempt}/{attempts}): {e}; "
                        f"retrying in {delay_s:.1f}s.",
                    )
                )
                sleep(delay_s)
                continue
            raise
    if last_error is not None:
        raise RuntimeError(_fmt_issue("AUTH", f"{flow_name} token generation failed: {last_error}"))
    raise RuntimeError(_fmt_issue("AUTH", f"{flow_name} token generation failed unexpectedly."))


def _generate_cdse_public_access_token(username: str, password: str, totp: Optional[str] = None) -> str:
    """Generate CDSE access token using cdse-public password grant."""
    payload = {
        "client_id": "cdse-public",
        "grant_type": "password",
        "username": username,
        "password": password,
    }
    if totp:
        payload["totp"] = str(totp).strip()
    return _request_cdse_access_token(payload, flow_name="cdse-public")

def _generate_cdse_client_access_token(client_id: str, client_secret: str) -> str:
    """Generate CDSE access token using client credentials."""
    payload = {
        "client_id": client_id,
        "client_secret": client_secret,
        "grant_type": "client_credentials",
    }
    return _request_cdse_access_token(payload, flow_name="client-credentials")


def _create_cdse_public_session(username: str, password: str, totp: Optional[str] = None) -> requests.Session:
    """Create bearer session from CDSE username/password token generation."""
    token = _generate_cdse_public_access_token(username=username, password=password, totp=totp)
    return _create_cdse_bearer_session(token)


def _create_cdse_session_from_environment() -> Optional[requests.Session]:
    """Create CDSE session from non-interactive environment credentials if available."""
    access_token = (os.environ.get("CDSE_ACCESS_TOKEN") or "").strip()
    if access_token:
        logger.info("Using CDSE bearer token from environment (CDSE_ACCESS_TOKEN).")
        return _create_cdse_bearer_session(access_token)

    client_id = (os.environ.get("CDSE_CLIENT_ID") or "").strip()
    client_secret = (os.environ.get("CDSE_CLIENT_SECRET") or "").strip()
    if client_id and client_secret:
        logger.info("Using CDSE client credentials from environment.")
        token = _generate_cdse_client_access_token(client_id=client_id, client_secret=client_secret)
        return _create_cdse_bearer_session(token)
    if client_id or client_secret:
        logger.warning(
            _fmt_issue(
                "AUTH",
                "Incomplete CDSE client credentials in environment; set both "
                "CDSE_CLIENT_ID and CDSE_CLIENT_SECRET.",
            )
        )

    username = (os.environ.get("CDSE_USERNAME") or "").strip()
    password = os.environ.get("CDSE_PASSWORD") or ""
    if username and password:
        logger.info("Using CDSE username/password from environment.")
        totp = (os.environ.get("CDSE_TOTP") or "").strip() or None
        sess = _create_cdse_public_session(username=username, password=password, totp=totp)
        _cache_cdse_public_credentials(username=username, password=password, totp=totp)
        return sess
    if username or password:
        logger.warning(
            _fmt_issue(
                "AUTH",
                "Incomplete CDSE username/password in environment; set both "
                "CDSE_USERNAME and CDSE_PASSWORD.",
            )
        )

    return None


def _prompt_cdse_userpass_cli(max_prompt_attempts: int = 2) -> Tuple[str, str, Optional[str]]:
    """Prompt for CDSE username/password in terminal (CLI mode)."""
    attempts = max(1, int(max_prompt_attempts))
    for attempt in range(1, attempts + 1):
        try:
            user = input("CDSE username: ").strip()
            pwd = getpass.getpass("CDSE password: ")
            totp = input("CDSE TOTP (optional): ").strip() or None
        except (EOFError, KeyboardInterrupt) as e:
            raise RuntimeError(_fmt_issue("AUTH", "CDSE login cancelled by user.")) from e

        if user and pwd:
            return user, pwd, totp

        logger.warning(
            _fmt_issue(
                "AUTH",
                f"Username and password are required (attempt {attempt}/{attempts}).",
            )
        )

    raise RuntimeError(_fmt_issue("AUTH", "CDSE username/password were not provided."))


def _request_cdse_userpass(
    allow_gui_prompt: bool,
    prompt_userpass_fn: Optional[PromptUserpassFn] = None,
) -> Tuple[str, str, Optional[str]]:
    """Request CDSE username/password from GUI callback or terminal."""
    if allow_gui_prompt:
        if prompt_userpass_fn is not None:
            return prompt_userpass_fn()
        if threading.current_thread() is not threading.main_thread():
            raise RuntimeError(
                _fmt_issue(
                    "AUTH",
                    "GUI credential prompt requested from a worker thread. "
                    "Provide a main-thread prompt callback.",
                )
            )
        return _prompt_cdse_userpass_gui()

    return _prompt_cdse_userpass_cli()


def _create_public_session_with_retry(
    allow_gui_prompt: bool = False,
    max_prompt_attempts: int = 2,
    prompt_userpass_fn: Optional[PromptUserpassFn] = None,
) -> requests.Session:
    """
    Create cdse-public session from per-run user prompt.
    """
    cached_userpass = _get_cached_cdse_public_credentials()
    if cached_userpass is not None:
        user, pwd, totp = cached_userpass
        try:
            logger.info("Reusing cached CDSE username/password for session refresh.")
            return _create_cdse_public_session(user, pwd, totp=totp)
        except Exception as e:
            logger.warning(
                _fmt_issue(
                    "AUTH",
                    f"Cached CDSE credentials failed ({e}); requesting credentials again.",
                )
            )
            _clear_cached_cdse_public_credentials()

    last_error: Optional[Exception] = None
    attempts = max(1, int(max_prompt_attempts))
    for attempt in range(1, attempts + 1):
        try:
            user, pwd, totp = _request_cdse_userpass(
                allow_gui_prompt=allow_gui_prompt,
                prompt_userpass_fn=prompt_userpass_fn,
            )
            sess = _create_cdse_public_session(user, pwd, totp=totp)
            _cache_cdse_public_credentials(username=user, password=pwd, totp=totp)
            return sess
        except RuntimeError as e:
            # User cancelled.
            if "cancelled" in str(e).lower():
                raise
            last_error = e
        except Exception as e:
            last_error = e

        if attempt < attempts:
            logger.warning(
                _fmt_issue(
                    "AUTH",
                    f"CDSE login failed (attempt {attempt}/{attempts}): {last_error}",
                )
            )

    raise RuntimeError(_fmt_issue("AUTH", f"CDSE username/password authentication failed: {last_error}"))


def _prompt_cdse_userpass_gui() -> Tuple[str, str, Optional[str]]:
    """Prompt for CDSE username/password (and optional TOTP) via GUI."""
    import tkinter as tk
    from tkinter import messagebox

    parent = tk._default_root
    owns_root = False
    if parent is None:
        parent = tk.Tk()
        parent.withdraw()
        owns_root = True

    dialog = tk.Toplevel(parent)
    dialog.title("CDSE Access Token Login")
    dialog.geometry("480x260")
    dialog.resizable(False, False)
    dialog.transient(parent)
    dialog.grab_set()

    dialog.update_idletasks()
    x = (dialog.winfo_screenwidth() // 2) - 240
    y = (dialog.winfo_screenheight() // 2) - 130
    dialog.geometry(f"480x260+{x}+{y}")

    result = {"ok": False, "username": "", "password": "", "totp": ""}

    tk.Label(dialog, text="Copernicus Data Space Login", font=("Arial", 12, "bold")).pack(pady=(12, 8))
    tk.Label(dialog, text="Used to generate a temporary cdse-public access token.", fg="gray").pack(pady=(0, 10))

    frame = tk.Frame(dialog)
    frame.pack(fill="x", padx=20)

    tk.Label(frame, text="Username:", width=14, anchor="w").grid(row=0, column=0, pady=5, sticky="w")
    user_entry = tk.Entry(frame, width=38)
    user_entry.grid(row=0, column=1, pady=5, sticky="w")

    tk.Label(frame, text="Password:", width=14, anchor="w").grid(row=1, column=0, pady=5, sticky="w")
    pass_entry = tk.Entry(frame, width=38, show="*")
    pass_entry.grid(row=1, column=1, pady=5, sticky="w")

    tk.Label(frame, text="TOTP (optional):", width=14, anchor="w").grid(row=2, column=0, pady=5, sticky="w")
    totp_entry = tk.Entry(frame, width=20)
    totp_entry.grid(row=2, column=1, pady=5, sticky="w")

    btns = tk.Frame(dialog)
    btns.pack(pady=16)

    def on_ok():
        # Read directly from widgets to avoid stale StringVar edge cases.
        u = user_entry.get().strip()
        p = pass_entry.get()
        if not u or not p:
            messagebox.showerror("Missing credentials", "Username and password are required.", parent=dialog)
            return
        result["ok"] = True
        result["username"] = u
        result["password"] = p
        result["totp"] = totp_entry.get().strip()
        dialog.destroy()

    def on_cancel():
        dialog.destroy()

    tk.Button(btns, text="OK", width=12, command=on_ok).pack(side="left", padx=8)
    tk.Button(btns, text="Cancel", width=12, command=on_cancel).pack(side="left", padx=8)
    dialog.bind("<Return>", lambda _evt: on_ok())
    pass_entry.bind("<Return>", lambda _evt: on_ok())
    user_entry.focus_set()

    parent.wait_window(dialog)
    if owns_root:
        try:
            parent.destroy()
        except Exception:
            pass

    if not result["ok"]:
        raise RuntimeError(_fmt_issue("AUTH", "CDSE login cancelled by user."))

    return result["username"], result["password"], (result["totp"] or None)


def _get_attr(attributes, name, default=100.0):
    """Get attribute value from CDSE product attributes."""
    for a in attributes:
        if a.get("Name") == name:
            try:
                return float(a.get("Value"))
            except Exception:
                return default
    return default


def _bbox_to_wkt(bbox):
    """Convert bounding box to WKT polygon."""
    minx, miny, maxx, maxy = bbox
    return f"POLYGON(({minx} {miny},{minx} {maxy},{maxx} {maxy},{maxx} {miny},{minx} {miny}))"


def _create_cdse_session_with_retry(
    allow_gui_prompt: bool = False,
    prompt_userpass_fn: Optional[PromptUserpassFn] = None,
):
    """
    Create CDSE session preferring non-interactive env credentials, then prompt fallback.
    """
    env_session = _create_cdse_session_from_environment()
    if env_session is not None:
        return env_session

    return _create_public_session_with_retry(
        allow_gui_prompt=allow_gui_prompt,
        max_prompt_attempts=2,
        prompt_userpass_fn=prompt_userpass_fn,
    )


def _query_s2_with_retry(
    session,
    center_time,
    bbox,
    days_window=30,
    max_cloud=20,
    allow_gui_prompt=False,
    prompt_userpass_fn: Optional[PromptUserpassFn] = None,
):
    """Query Sentinel-2 and retry with a fresh prompt-based session if auth fails."""
    try:
        return _query_s2(session, center_time, bbox, days_window, max_cloud), session
    except requests.HTTPError as e:
        status = e.response.status_code if getattr(e, "response", None) is not None else None
        if status in (400, 401, 403):
            logger.warning(
                _fmt_issue(
                    "CDSE_QUERY",
                    f"Sentinel-2 query failed with HTTP {status}. "
                    "Prompting for CDSE username/password and retrying.",
                )
            )
            public_session = _create_public_session_with_retry(
                allow_gui_prompt=allow_gui_prompt,
                prompt_userpass_fn=prompt_userpass_fn,
            )
            return _query_s2(public_session, center_time, bbox, days_window, max_cloud), public_session
        raise


def _query_s2(session, center_time, bbox, days_window=30, max_cloud=20):
    """Query Sentinel-2 L2A products from CDSE OData API."""
    logger.info("Querying Sentinel-2...")
    logger.info(f"  Search window: +/-{days_window} days from {center_time.strftime('%Y-%m-%d')}")
    logger.info(f"  Max cloud cover: {max_cloud}%, Bounding box: {bbox}")

    start = (center_time - timedelta(days=days_window)).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    end = (center_time + timedelta(days=days_window)).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    wkt = _bbox_to_wkt(bbox)

    filter_expr = (
        "Collection/Name eq 'SENTINEL-2' and "
        "Attributes/OData.CSC.StringAttribute/any(a:a/Name eq 'productType' and a/OData.CSC.StringAttribute/Value eq 'S2MSI2A') and "
        f"Attributes/OData.CSC.DoubleAttribute/any(a:a/Name eq 'cloudCover' and a/OData.CSC.DoubleAttribute/Value le {max_cloud}) and "
        f"ContentDate/Start gt {start} and ContentDate/Start lt {end} and "
        f"OData.CSC.Intersects(area=geography'SRID=4326;{wkt}')"
    )

    params = {"$filter": filter_expr, "$expand": "Attributes"}
    attempts = max(1, int(HTTP_QUERY_RETRY_ATTEMPTS))
    resp = None
    last_error: Optional[Exception] = None
    for attempt in range(1, attempts + 1):
        try:
            resp = session.get(CATALOGUE_URL, params=params, timeout=HTTP_TIMEOUT_S)
            if resp.status_code == 200:
                break
            if _is_retryable_http_status(resp.status_code) and attempt < attempts:
                delay_s = _retry_delay_seconds(attempt)
                logger.warning(
                    _fmt_issue(
                        "CDSE_QUERY",
                        f"Retryable HTTP {resp.status_code} on attempt {attempt}/{attempts}; "
                        f"retrying in {delay_s:.1f}s.",
                    )
                )
                sleep(delay_s)
                continue
            logger.error(_fmt_issue("CDSE_QUERY", f"Query failed: {resp.status_code} - {resp.text[:500]}"))
            resp.raise_for_status()
        except requests.RequestException as e:
            if isinstance(e, requests.HTTPError):
                status = e.response.status_code if getattr(e, "response", None) is not None else None
                if status in (400, 401, 403):
                    # Preserve auth HTTP errors so _query_s2_with_retry can trigger re-auth flow.
                    raise
            last_error = e
            if attempt < attempts:
                delay_s = _retry_delay_seconds(attempt)
                logger.warning(
                    _fmt_issue(
                        "CDSE_QUERY",
                        f"Request failed on attempt {attempt}/{attempts}: {e}; retrying in {delay_s:.1f}s.",
                    )
                )
                sleep(delay_s)
                continue
            raise RuntimeError(_fmt_issue("CDSE_QUERY", f"Request failed after retries: {e}")) from e
    if resp is None:
        raise RuntimeError(_fmt_issue("CDSE_QUERY", f"No response returned: {last_error}"))

    items = resp.json().get("value", [])
    logger.info(f"  Found {len(items)} Sentinel-2 products")
    return items


def _rank_s2_candidates(items, center_time, bbox, min_overlap=0.5):
    """Rank Sentinel-2 candidates by quality score and filter by overlap."""
    logger.info(f"Ranking S2 candidates (min overlap: {min_overlap:.1%})...")

    minx, miny, maxx, maxy = bbox
    bbox_poly = Polygon([(minx, miny), (minx, maxy), (maxx, maxy), (maxx, miny)])
    scored_candidates = []

    for it in items:
        geom = it["GeoFootprint"]
        geom_poly = shapely_wkt.loads(geom) if isinstance(geom, str) else shape(geom)
        inter = bbox_poly.intersection(geom_poly)
        ov = inter.area / bbox_poly.area if bbox_poly.area > 0 else 0.0

        dt = datetime.fromisoformat(it["ContentDate"]["Start"].replace("Z", "+00:00")).astimezone(timezone.utc)
        cloud = _get_attr(it["Attributes"], "cloudCover", 100.0)

        if ov < min_overlap:
            logger.debug(f"  {it['Name']}: overlap={ov:.1%} < {min_overlap:.1%} - REJECTED")
            continue

        logger.info(f"  {it['Name']}: overlap={ov:.1%}, cloud={cloud:.1f}%")
        tdiff_h = abs((dt - center_time).total_seconds()) / 3600.0
        score = cloud + 0.5 * tdiff_h - 20.0 * ov
        scored_candidates.append({'item': it, 'score': score, 'cloud': cloud, 'overlap': ov, 'time': dt})

    if not scored_candidates:
        logger.error(_fmt_issue("CDSE_QUERY", "No suitable Sentinel-2 products found."))
        raise SentinelNotFoundError("No valid Sentinel-2 match found for this scene")

    scored_candidates.sort(key=lambda x: x['score'])
    logger.info(f"Found {len(scored_candidates)} valid candidates. Top: {scored_candidates[0]['item']['Name']}")
    return [x['item'] for x in scored_candidates]


def _download_s2_product(
    session,
    product,
    out_dir,
    allow_gui_prompt=False,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
    scene_idx: int = 1,
    scene_total: int = 1,
    prompt_userpass_fn: Optional[PromptUserpassFn] = None,
):
    """Download Sentinel-2 product from CDSE."""
    pid = product["Id"]
    download_urls = [
        f"https://download.dataspace.copernicus.eu/odata/v1/Products({pid})/$value",
        f"https://catalogue.dataspace.copernicus.eu/odata/v1/Products({pid})/$value",
        f"https://zipper.dataspace.copernicus.eu/odata/v1/Products({pid})/$value",
    ]
    zip_path = os.path.join(out_dir, f"{pid}.zip")

    def _is_valid_zip_quick(path: str) -> bool:
        """Fast ZIP validity check without full CRC scan."""
        if not os.path.exists(path):
            return False
        if not zipfile.is_zipfile(path):
            return False
        try:
            with zipfile.ZipFile(path) as z:
                return len(z.infolist()) > 0
        except Exception:
            return False

    if os.path.exists(zip_path):
        if _is_valid_zip_quick(zip_path):
            logger.info("S2 already downloaded (valid zip file).")
            return zip_path, session
        try:
            os.remove(zip_path)
        except Exception:
            pass

    logger.info("Downloading S2...")
    _emit_progress(
        progress_callback,
        "Downloading Sentinel-2 reference",
        scene_idx=scene_idx,
        scene_total=scene_total,
    )

    def _attempt_download(sess: requests.Session):
        attempt_errors = []
        audience_error = False
        download_attempts = max(1, int(HTTP_DOWNLOAD_RETRY_ATTEMPTS))
        for url in download_urls:
            for attempt in range(1, download_attempts + 1):
                try:
                    logger.info(f"  Trying endpoint: {url} (attempt {attempt}/{download_attempts})")
                    _emit_progress(
                        progress_callback,
                        "Downloading Sentinel-2 reference",
                        scene_idx=scene_idx,
                        scene_total=scene_total,
                        download_endpoint=url,
                        download_percent=0.0,
                        download_mb=0.0,
                    )
                    with sess.get(url, stream=True, timeout=HTTP_TIMEOUT_CONNECT_READ) as r:
                        if r.status_code >= 400:
                            msg = r.text[:500]
                            if "DAT-ZIP-609" in msg or "Token audience not allowed" in msg:
                                audience_error = True
                            if _is_retryable_http_status(r.status_code) and attempt < download_attempts:
                                delay_s = _retry_delay_seconds(attempt)
                                logger.warning(
                                    _fmt_issue(
                                        "S2_DOWNLOAD",
                                        f"Retryable HTTP {r.status_code} for {url} "
                                        f"(attempt {attempt}/{download_attempts}); "
                                        f"retrying in {delay_s:.1f}s.",
                                    )
                                )
                                sleep(delay_s)
                                continue
                            attempt_errors.append(f"{url} -> {r.status_code}: {msg}")
                            break

                        total_size = int(r.headers.get("Content-Length", "0") or "0")
                        if total_size > 0:
                            logger.info(f"  Download size: {total_size / (1024 * 1024):.1f} MB")
                        else:
                            logger.info("  Download size: unknown")

                        bytes_written = 0
                        next_log = 200 * 1024 * 1024
                        last_emit_t = perf_counter()
                        last_emit_bytes = 0
                        with open(zip_path, "wb") as f:
                            for chunk in r.iter_content(1024 * 1024):
                                if chunk:
                                    f.write(chunk)
                                    bytes_written += len(chunk)
                                    now_t = perf_counter()
                                    if (now_t - last_emit_t) >= 1.0 or (bytes_written - last_emit_bytes) >= (25 * 1024 * 1024):
                                        download_mb = bytes_written / (1024 * 1024)
                                        event_kwargs = {
                                            "download_endpoint": url,
                                            "download_mb": round(download_mb, 2),
                                        }
                                        if total_size > 0:
                                            event_kwargs["download_total_mb"] = round(total_size / (1024 * 1024), 2)
                                            event_kwargs["download_percent"] = round((100.0 * bytes_written) / total_size, 2)
                                        _emit_progress(
                                            progress_callback,
                                            "Downloading Sentinel-2 reference",
                                            scene_idx=scene_idx,
                                            scene_total=scene_total,
                                            **event_kwargs,
                                        )
                                        last_emit_t = now_t
                                        last_emit_bytes = bytes_written
                                    if bytes_written >= next_log:
                                        logger.info(f"  Downloaded {bytes_written / (1024 * 1024):.1f} MB...")
                                        next_log += 200 * 1024 * 1024
                        logger.info(f"  Download completed: {bytes_written / (1024 * 1024):.1f} MB")
                        final_kwargs = {
                            "download_endpoint": url,
                            "download_mb": round(bytes_written / (1024 * 1024), 2),
                        }
                        if total_size > 0:
                            final_kwargs["download_total_mb"] = round(total_size / (1024 * 1024), 2)
                            final_kwargs["download_percent"] = 100.0
                        _emit_progress(
                            progress_callback,
                            "Downloading Sentinel-2 reference",
                            scene_idx=scene_idx,
                            scene_total=scene_total,
                            **final_kwargs,
                        )

                    # Validate zip before returning
                    if _is_valid_zip_quick(zip_path):
                        return True, attempt_errors, audience_error
                    attempt_errors.append(f"{url} -> downloaded file is not a valid zip")
                    try:
                        os.remove(zip_path)
                    except Exception:
                        pass
                    break
                except requests.RequestException as e:
                    if attempt < download_attempts:
                        delay_s = _retry_delay_seconds(attempt)
                        logger.warning(
                            _fmt_issue(
                                "S2_DOWNLOAD",
                                f"Request error for {url} (attempt {attempt}/{download_attempts}): {e}; "
                                f"retrying in {delay_s:.1f}s.",
                            )
                        )
                        sleep(delay_s)
                        continue
                    attempt_errors.append(f"{url} -> request error: {e}")
                    break
                except Exception as e:
                    attempt_errors.append(f"{url} -> request error: {e}")
                    break
        return False, attempt_errors, audience_error

    ok, errors, saw_audience_error = _attempt_download(session)
    if ok:
        return zip_path, session

    # Fallback: generate cdse-public token from username/password if available.
    if saw_audience_error:
        fallback_session = None
        try:
            logger.warning(_fmt_issue("S2_DOWNLOAD", "Token audience rejected. Trying username/password token flow."))
            fallback_session = _create_public_session_with_retry(
                allow_gui_prompt=allow_gui_prompt,
                prompt_userpass_fn=prompt_userpass_fn,
            )
        except Exception as e:
            logger.warning(_fmt_issue("S2_DOWNLOAD", f"Username/password fallback unavailable: {e}"))

        if fallback_session is not None:
            ok2, errors2, audience2 = _attempt_download(fallback_session)
            if ok2:
                return zip_path, fallback_session
            errors.extend(errors2)
            saw_audience_error = saw_audience_error or audience2

    if saw_audience_error:
        raise RuntimeError(
            _fmt_issue(
                "S2_DOWNLOAD",
                "CDSE token audience not allowed for product download (DAT-ZIP-609). "
                "Please re-run and log in again with CDSE username/password. "
                f"Endpoint attempts: {' | '.join(errors[:3])}",
            )
        )

    if errors:
        raise RuntimeError(_fmt_issue("S2_DOWNLOAD", f"S2 download failed on all endpoints: {' | '.join(errors[:3])}"))
    return zip_path, session


def _build_s2_stack(
    zip_path,
    bbox,
    out_s2_path,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
    scene_idx: int = 1,
    scene_total: int = 1,
):
    """Build 7-band Sentinel-2 reference stack from L2A ZIP product."""
    from osgeo import gdal
    gdal.UseExceptions()

    log_section_header("BUILDING SENTINEL-2 STACK")
    _emit_progress(
        progress_callback,
        "Building Sentinel-2 stack",
        scene_idx=scene_idx,
        scene_total=scene_total,
        substage="Inspecting granule metadata",
    )

    # Find reference band and granule
    reference_band = "B04"
    granule_prefix = None

    with zipfile.ZipFile(zip_path) as z:
        all_files = z.namelist()
        b04_hits = [n for n in all_files if f"_{reference_band}_10m.jp2" in n and "/GRANULE/" in n]
        if not b04_hits:
            reference_band = "B08"
            b04_hits = [n for n in all_files if f"_{reference_band}_10m.jp2" in n and "/GRANULE/" in n]
        if not b04_hits:
            raise ValueError(f"No reference band found in S2 product.")

        reference_jp2_path = b04_hits[0]
        if "/GRANULE/" in reference_jp2_path:
            parts = reference_jp2_path.split("/GRANULE/")
            granule_id = parts[1].split("/")[0]
            granule_prefix = parts[0] + "/GRANULE/" + granule_id + "/"
            logger.info(f"  Selected granule: {granule_id}")
            _emit_progress(
                progress_callback,
                "Building Sentinel-2 stack",
                scene_idx=scene_idx,
                scene_total=scene_total,
                substage=f"Selected granule {granule_id}",
            )

    bands_to_extract = {
        "B02": "_B02_10m.jp2", "B03": "_B03_10m.jp2", "B04": "_B04_10m.jp2",
        "B08": "_B08_10m.jp2", "B11": "_B11_20m.jp2", "B12": "_B12_20m.jp2",
        "SCL": "_SCL_20m.jp2"
    }

    bands_jp2_paths = {}
    bands_extracted_paths = {}

    with zipfile.ZipFile(zip_path) as z:
        total_extract = len(bands_to_extract)
        for idx, (band_name, suffix) in enumerate(bands_to_extract.items(), 1):
            _emit_progress(
                progress_callback,
                "Building Sentinel-2 stack",
                scene_idx=scene_idx,
                scene_total=scene_total,
                substage=f"Extracting {band_name} ({idx}/{total_extract})",
                step=idx,
                step_total=total_extract,
            )
            hits = [n for n in z.namelist() if n.startswith(granule_prefix) and suffix in n]
            if band_name == "SCL" and not hits:
                hits = [n for n in z.namelist() if n.startswith(granule_prefix) and "_SCL_60m.jp2" in n]
            if not hits:
                if band_name == "SCL":
                    continue
                raise ValueError(f"Band {band_name} not found in granule.")

            jp2_path_in_zip = hits[0]
            bands_jp2_paths[band_name] = jp2_path_in_zip
            z.extract(jp2_path_in_zip, os.path.dirname(out_s2_path))
            bands_extracted_paths[band_name] = os.path.join(os.path.dirname(out_s2_path), jp2_path_in_zip)

    # Convert JP2 to GeoTIFF
    bands_paths = {}
    for idx, (band_name, jp2_path) in enumerate(bands_extracted_paths.items(), 1):
        _emit_progress(
            progress_callback,
            "Building Sentinel-2 stack",
            scene_idx=scene_idx,
            scene_total=scene_total,
            substage=f"Translating {band_name} JP2 to GeoTIFF ({idx}/{len(bands_extracted_paths)})",
            step=idx,
            step_total=len(bands_extracted_paths),
        )
        if jp2_path.lower().endswith('.jp2'):
            tif_path = jp2_path.replace('.jp2', '_temp.tif').replace('.JP2', '_temp.tif')
            try:
                ds = gdal.Open(jp2_path)
                if ds:
                    gdal.Translate(tif_path, ds, format='GTiff', creationOptions=['COMPRESS=LZW', 'TILED=YES'])
                    ds = None
                    bands_paths[band_name] = tif_path
                else:
                    bands_paths[band_name] = jp2_path
            except Exception:
                bands_paths[band_name] = jp2_path
        else:
            bands_paths[band_name] = jp2_path

    # Define reference grid
    _emit_progress(
        progress_callback,
        "Building Sentinel-2 stack",
        scene_idx=scene_idx,
        scene_total=scene_total,
        substage="Defining target grid",
    )
    ref_band_for_grid = reference_band if reference_band in bands_paths else "B04"
    with rasterio.open(bands_paths[ref_band_for_grid]) as ref:
        s2_crs = ref.crs
        tr = Transformer.from_crs("EPSG:4326", s2_crs, always_xy=True)
        minx, miny = tr.transform(bbox[0], bbox[1])
        maxx, maxy = tr.transform(bbox[2], bbox[3])
        buffer = 200
        window = from_bounds(minx - buffer, miny - buffer, maxx + buffer, maxy + buffer, ref.transform)
        window = window.round_offsets().round_shape()
        ref_transform = ref.window_transform(window)
        ref_height, ref_width = int(window.height), int(window.width)

    # Resample all bands
    stack = []
    resample_bands = ["B02", "B03", "B04", "B08", "B11", "B12"]
    for idx, band_name in enumerate(resample_bands, 1):
        _emit_progress(
            progress_callback,
            "Building Sentinel-2 stack",
            scene_idx=scene_idx,
            scene_total=scene_total,
            substage=f"Resampling {band_name} ({idx}/{len(resample_bands)})",
            step=idx,
            step_total=len(resample_bands),
        )
        if band_name not in bands_paths:
            stack.append(np.zeros((ref_height, ref_width), dtype=np.uint16))
            continue
        with rasterio.open(bands_paths[band_name]) as src:
            data = np.zeros((ref_height, ref_width), dtype=np.uint16)
            rasterio.warp.reproject(
                source=rasterio.band(src, 1), destination=data,
                src_transform=src.transform, src_crs=src.crs,
                dst_transform=ref_transform, dst_crs=s2_crs,
                resampling=Resampling.bilinear, dst_nodata=0
            )
            stack.append(data)

    # Handle SCL
    _emit_progress(
        progress_callback,
        "Building Sentinel-2 stack",
        scene_idx=scene_idx,
        scene_total=scene_total,
        substage="Resampling SCL mask",
    )
    if "SCL" in bands_paths:
        with rasterio.open(bands_paths["SCL"]) as src:
            scl_data = np.zeros((ref_height, ref_width), dtype=np.uint8)
            rasterio.warp.reproject(
                source=rasterio.band(src, 1), destination=scl_data,
                src_transform=src.transform, src_crs=src.crs,
                dst_transform=ref_transform, dst_crs=s2_crs,
                resampling=Resampling.nearest, dst_nodata=0
            )
            stack.append(scl_data)
    else:
        stack.append(np.zeros((ref_height, ref_width), dtype=np.uint8))

    stack = np.stack(stack)

    # Apply cloud mask
    if "SCL" in bands_paths:
        scl_band = stack[6]
        cloud_shadow_mask = np.isin(scl_band, list(SCL_EXCLUDE_CLASSES.keys()))
        for i in range(6):
            stack[i][cloud_shadow_mask] = 0
        logger.info(f"  Applied cloud/shadow mask: {100 * np.sum(cloud_shadow_mask) / cloud_shadow_mask.size:.1f}% masked")

    # Write stack
    _emit_progress(
        progress_callback,
        "Building Sentinel-2 stack",
        scene_idx=scene_idx,
        scene_total=scene_total,
        substage="Writing stacked reference raster",
    )
    with rasterio.open(out_s2_path, "w", driver="GTiff", width=ref_width, height=ref_height,
                       count=7, dtype=stack.dtype, crs=s2_crs, transform=ref_transform,
                       compress="lzw", tiled=True, BIGTIFF="YES") as dst:
        dst.write(stack)

    # Cleanup
    for p in bands_paths.values():
        try:
            os.remove(p)
        except Exception:
            pass

    logger.info("S2 stack created successfully")
    _emit_progress(
        progress_callback,
        "Building Sentinel-2 stack",
        scene_idx=scene_idx,
        scene_total=scene_total,
        substage="Stack ready",
    )
    return out_s2_path, s2_crs


def _crs_equivalent(lhs: Any, rhs: Any) -> bool:
    """Return True if two CRS definitions represent the same projection."""
    if lhs is None or rhs is None:
        return False
    try:
        return CRS.from_user_input(lhs) == CRS.from_user_input(rhs)
    except Exception:
        return str(lhs) == str(rhs)


def _reproject_reference_stack_to_target_crs(
    source_path: str,
    target_crs: Any,
    output_path: str,
) -> Dict[str, Any]:
    """
    Reproject a Sentinel-2 reference stack into a target CRS.

    The last band (typically SCL) is reprojected with nearest-neighbor;
    all other bands use bilinear interpolation.
    """
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "path": source_path,
        "reprojected": False,
        "source_crs": None,
        "target_crs": None,
    }
    try:
        if not source_path or not os.path.exists(source_path):
            out["error"] = f"Reference stack missing: {source_path}"
            return out

        target_crs_obj = CRS.from_user_input(target_crs)
        out["target_crs"] = target_crs_obj.to_string()

        with rasterio.open(source_path) as src:
            if src.crs is None:
                out["error"] = "Reference stack has no CRS."
                return out
            source_crs_obj = CRS.from_user_input(src.crs)
            out["source_crs"] = source_crs_obj.to_string()

            if source_crs_obj == target_crs_obj:
                out["ok"] = True
                out["path"] = source_path
                out["reprojected"] = False
                return out

            dst_transform, dst_width, dst_height = rasterio.warp.calculate_default_transform(
                src.crs,
                target_crs_obj,
                src.width,
                src.height,
                *src.bounds,
                resolution=src.res,
            )

            dst_nodata = src.nodata if src.nodata is not None else 0
            profile = src.profile.copy()
            profile.update(
                {
                    "crs": target_crs_obj,
                    "transform": dst_transform,
                    "width": int(dst_width),
                    "height": int(dst_height),
                    "nodata": dst_nodata,
                    "compress": "LZW",
                    "tiled": True,
                    "BIGTIFF": "YES",
                }
            )

            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            with rasterio.open(output_path, "w", **profile) as dst:
                src_tags = src.tags()
                if src_tags:
                    dst.update_tags(**src_tags)

                for bidx in range(1, int(src.count) + 1):
                    band_dtype = np.dtype(src.dtypes[bidx - 1])
                    dest = np.full(
                        (int(dst_height), int(dst_width)),
                        fill_value=dst_nodata,
                        dtype=band_dtype,
                    )

                    band_desc = src.descriptions[bidx - 1]
                    is_scl_band = (
                        (band_desc is not None and "scl" in str(band_desc).lower())
                        or (bidx == int(src.count) and int(src.count) >= 7)
                    )
                    band_resampling = Resampling.nearest if is_scl_band else Resampling.bilinear

                    rasterio.warp.reproject(
                        source=rasterio.band(src, bidx),
                        destination=dest,
                        src_transform=src.transform,
                        src_crs=src.crs,
                        dst_transform=dst_transform,
                        dst_crs=target_crs_obj,
                        src_nodata=src.nodata,
                        dst_nodata=dst_nodata,
                        resampling=band_resampling,
                    )
                    dst.write(dest, bidx)

                    if band_desc:
                        dst.set_band_description(bidx, band_desc)
                    band_tags = src.tags(bidx)
                    if band_tags:
                        dst.update_tags(bidx, **band_tags)

        out["ok"] = True
        out["path"] = output_path
        out["reprojected"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _create_scene_folder_structure(base_output_dir, sensor_tag, date_tag, unique_hex):
    """Create organized subfolder structure for scene outputs."""
    scene_root = os.path.join(base_output_dir, f"{sensor_tag}_{date_tag}_{unique_hex}")
    folders = {
        'scene_root': scene_root,
        'inputs': os.path.join(scene_root, "00_inputs"),
        'reference': os.path.join(scene_root, "01_reference"),
        'temp': os.path.join(scene_root, "02_temp"),
        'coreg': os.path.join(scene_root, "03_coreg"),
        'reports': os.path.join(scene_root, "04_reports"),
        'quicklooks': os.path.join(scene_root, "05_quicklooks"),
    }
    for folder_path in folders.values():
        os.makedirs(folder_path, exist_ok=True)
    return folders


def _parse_arosics_output(stdout_text):
    """Extract metrics from AROSICS console output."""
    result = {
        'reliability': None, 'ssim_before': None, 'ssim_after': None,
        'ssim_delta': None, 'x_shift_m': None, 'y_shift_m': None,
        'shift_magnitude': None, 'parsed_success': False
    }
    rel_match = re.search(r"Estimated reliability[^0-9\n\r]*?([0-9]+\.?[0-9]*)%?", stdout_text)
    if rel_match:
        result['reliability'] = float(rel_match.group(1))
        result['parsed_success'] = True
    ssim_match = re.search(r"SSIM.*?([0-9]+\.[0-9]+)\s*=>\s*([0-9]+\.[0-9]+)", stdout_text)
    if ssim_match:
        result['ssim_before'] = float(ssim_match.group(1))
        result['ssim_after'] = float(ssim_match.group(2))
        result['ssim_delta'] = result['ssim_after'] - result['ssim_before']
    shift_match = re.search(r"Calculated map shifts \(X,Y\):\s*(-?[0-9]+\.?[0-9]*)/(-?[0-9]+\.?[0-9]*)", stdout_text)
    if shift_match:
        x_shift, y_shift = float(shift_match.group(1)), float(shift_match.group(2))
        result['x_shift_m'], result['y_shift_m'] = x_shift, y_shift
        result['shift_magnitude'] = np.sqrt(x_shift**2 + y_shift**2)
        result['parsed_success'] = True
    return result


def _validate_coreg_hybrid(CRG, stdout_text, sensor_type="PRISMA", max_displacement=350.0):
    """Validate coregistration using AROSICS reliability metric."""
    result = {
        'is_valid': False, 'confidence': 0.0, 'shift_m': 0.0, 'message': '',
        'reliability_parsed': None, 'ssim_before': None, 'ssim_after': None, 'ssim_delta': None
    }
    try:
        coreg_info = getattr(CRG, "coreg_info", {})
        if not coreg_info.get("success", False):
            result['message'] = "AROSICS reported success=False"
            return result

        parsed = _parse_arosics_output(stdout_text)
        reliability = parsed.get('reliability')
        shift_magnitude = parsed.get('shift_magnitude')

        if shift_magnitude is None:
            x_shift = coreg_info.get("x_shift_m")
            y_shift = coreg_info.get("y_shift_m")
            if x_shift is not None and y_shift is not None:
                shift_magnitude = np.sqrt(float(x_shift)**2 + float(y_shift)**2)

        shift_magnitude = shift_magnitude if shift_magnitude is not None else 0.0
        result['shift_m'] = shift_magnitude
        result['ssim_before'] = parsed.get('ssim_before')
        result['ssim_after'] = parsed.get('ssim_after')
        result['ssim_delta'] = parsed.get('ssim_delta')

        if shift_magnitude > max_displacement:
            result['message'] = f"Shift {shift_magnitude:.1f}m exceeds max {max_displacement:.0f}m"
            return result

        if reliability is not None:
            result['reliability_parsed'] = reliability
            result['confidence'] = reliability / 100.0
            if reliability >= 40.0:
                result['is_valid'] = True
                result['message'] = f"Reliability {reliability:.1f}% >= 40% (shift: {shift_magnitude:.1f}m)"
            else:
                result['message'] = f"Reliability {reliability:.1f}% < 40% (shift: {shift_magnitude:.1f}m)"
        else:
            result['message'] = f"No reliability available (shift: {shift_magnitude:.1f}m)"
        return result
    except Exception as e:
        result['message'] = f"Validation error: {e}"
        return result


def _compute_tiepoint_residuals(tie_points_df, pixel_size_m=30.0):
    """Compute residual metrics from tie points."""
    result = {
        'n_tiepoints_used': 0, 'residual_mean_m': None, 'residual_median_m': None,
        'residual_rmse_m': None, 'residual_p90_m': None, 'residual_source': 'unavailable'
    }
    if tie_points_df is None:
        return result
    try:
        df = tie_points_df.copy()
    except Exception:
        return result
    if hasattr(df, "empty") and df.empty:
        return result

    for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
        if col in df.columns:
            try:
                outlier_mask = df[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
                df = df[~outlier_mask]
            except Exception:
                pass
    if hasattr(df, "empty") and df.empty:
        return result

    result['n_tiepoints_used'] = len(df)

    if 'ABS_SHIFT_M' in df.columns:
        try:
            residuals = df['ABS_SHIFT_M'].dropna().astype(float).values
            if len(residuals) > 0:
                result['residual_source'] = 'ABS_SHIFT_M'
                result['residual_mean_m'] = float(np.mean(residuals))
                result['residual_median_m'] = float(np.median(residuals))
                result['residual_rmse_m'] = float(np.sqrt(np.mean(residuals**2)))
                result['residual_p90_m'] = float(np.percentile(residuals, 90))
                return result
        except Exception:
            pass

    if 'X_SHIFT_M' in df.columns and 'Y_SHIFT_M' in df.columns:
        try:
            x_shift = df['X_SHIFT_M'].fillna(0).astype(float).values
            y_shift = df['Y_SHIFT_M'].fillna(0).astype(float).values
            residuals = np.sqrt(x_shift**2 + y_shift**2)
            valid_mask = np.isfinite(residuals)
            residuals = residuals[valid_mask]
            if len(residuals) > 0:
                result['residual_source'] = 'XY_SHIFT_M'
                result['residual_mean_m'] = float(np.mean(residuals))
                result['residual_median_m'] = float(np.median(residuals))
                result['residual_rmse_m'] = float(np.sqrt(np.mean(residuals**2)))
                result['residual_p90_m'] = float(np.percentile(residuals, 90))
        except Exception:
            pass
    return result


def _normalize_sensor_name(sensor_type: Optional[str]) -> str:
    """Normalize sensor labels used in profile dictionaries."""
    name = str(sensor_type or "").strip().upper()
    return name if name else "DEFAULT"


def _coerce_window_size(value: Any, default: Tuple[int, int]) -> Tuple[int, int]:
    """Parse matcher window size configuration with safe fallback."""
    if isinstance(value, str):
        cleaned = value.replace("x", ",")
        parts = [p.strip() for p in cleaned.split(",") if p.strip()]
        if len(parts) == 2:
            try:
                return (max(32, int(parts[0])), max(32, int(parts[1])))
            except Exception:
                return default
        return default
    if isinstance(value, (tuple, list)) and len(value) == 2:
        try:
            return (max(32, int(value[0])), max(32, int(value[1])))
        except Exception:
            return default
    return default


def _supports_constructor_kwarg(target: Any, kwarg_name: str) -> bool:
    """Return True if class/function constructor appears to support a keyword argument."""
    try:
        sig = inspect.signature(target.__init__)
    except Exception:
        try:
            sig = inspect.signature(target)
        except Exception:
            return False
    if kwarg_name in sig.parameters:
        return True
    return any(p.kind == inspect.Parameter.VAR_KEYWORD for p in sig.parameters.values())


def _coerce_global_coreg_attempt_ladder(raw: Any, fallback: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Normalize global COREG attempt definitions into ws/max_shift entries."""
    out: List[Dict[str, Any]] = []
    if not isinstance(raw, list):
        raw = fallback
    for item in raw:
        if not isinstance(item, dict):
            continue
        ws = item.get("window_size", item.get("ws"))
        ws_tuple = _coerce_window_size(ws, default=(256, 256))
        try:
            max_shift = float(item.get("max_shift", 150))
        except Exception:
            max_shift = 150.0
        out.append({"ws": ws_tuple, "max_shift": max(5.0, max_shift)})
    if not out:
        out = [{"ws": (256, 256), "max_shift": 150.0}]
    return out


def _resolve_sensor_matcher_profile(config: Optional[Dict[str, Any]], sensor_type: str) -> Dict[str, Any]:
    """Resolve sensor-aware matcher profile for global/local AROSICS usage."""
    cfg = config or {}
    sensor_key = _normalize_sensor_name(sensor_type)

    default_profiles = DEFAULT_CONFIG.get("global_coreg_profiles_by_sensor", {})
    profile_dict = cfg.get("global_coreg_profiles_by_sensor", default_profiles)
    if not isinstance(profile_dict, dict):
        profile_dict = default_profiles
    profile_source = "global_coreg_profiles_by_sensor"

    ladder_override = cfg.get("global_coreg_attempt_ladder", None)
    if isinstance(ladder_override, list) and ladder_override:
        global_ladder = _coerce_global_coreg_attempt_ladder(ladder_override, [])
        profile_source = "global_coreg_attempt_ladder override"
    else:
        sensor_profiles = profile_dict.get(sensor_key) or profile_dict.get("DEFAULT") or []
        if not sensor_profiles and isinstance(default_profiles, dict):
            sensor_profiles = default_profiles.get(sensor_key) or default_profiles.get("DEFAULT") or []
        global_ladder = _coerce_global_coreg_attempt_ladder(sensor_profiles, [])
        profile_source = f"{profile_source}:{sensor_key if profile_dict.get(sensor_key) else 'DEFAULT'}"

    local_shift_defaults = DEFAULT_CONFIG.get("local_max_shift_by_sensor", {"DEFAULT": 50})
    local_shift_map = cfg.get("local_max_shift_by_sensor", local_shift_defaults)
    if not isinstance(local_shift_map, dict):
        local_shift_map = local_shift_defaults
    try:
        local_max_shift = float(
            local_shift_map.get(sensor_key, local_shift_map.get("DEFAULT", local_shift_defaults.get("DEFAULT", 50)))
        )
    except Exception:
        local_max_shift = float(local_shift_defaults.get("DEFAULT", 50))
    local_max_shift = max(5.0, local_max_shift)

    try:
        local_grid_res = int(cfg.get("local_coreg_grid_res", DEFAULT_CONFIG.get("local_coreg_grid_res", LOCAL_GRID_RES_M)))
    except Exception:
        local_grid_res = int(DEFAULT_CONFIG.get("local_coreg_grid_res", LOCAL_GRID_RES_M))
    local_grid_res = max(30, local_grid_res)

    local_window = _coerce_window_size(
        cfg.get("local_coreg_window_size", DEFAULT_CONFIG.get("local_coreg_window_size", (256, 256))),
        default=(256, 256),
    )
    try:
        tiep_filter = int(
            cfg.get(
                "local_coreg_tieP_filter_level",
                DEFAULT_CONFIG.get("local_coreg_tieP_filter_level", 1),
            )
        )
    except Exception:
        tiep_filter = int(DEFAULT_CONFIG.get("local_coreg_tieP_filter_level", 1))
    tiep_filter = max(0, tiep_filter)

    max_iter_raw = cfg.get("local_coreg_max_iter", DEFAULT_CONFIG.get("local_coreg_max_iter"))
    try:
        local_max_iter = None if max_iter_raw in (None, "", False) else max(1, int(max_iter_raw))
    except Exception:
        local_max_iter = None

    return {
        "sensor": sensor_key,
        "global_attempt_ladder": global_ladder,
        "global_profile_source": profile_source,
        "local_max_shift": local_max_shift,
        "local_grid_res": local_grid_res,
        "local_window_size": local_window,
        "local_tieP_filter_level": tiep_filter,
        "local_max_iter": local_max_iter,
    }


def _resolve_fixed_band_pair(
    hs_wavelengths: np.ndarray,
    band_label: str,
    sensor_type: str,
    prefer_fixed_band_pairs: bool = True,
    fixed_band_pairs_by_sensor: Optional[Dict[str, Dict[str, int]]] = None,
) -> Dict[str, Any]:
    """Resolve curated fixed band-pair indices for the requested sensor/band."""
    sensor_key = _normalize_sensor_name(sensor_type)
    fixed_map = fixed_band_pairs_by_sensor or {}
    sensor_map: Dict[str, int] = {}
    if isinstance(fixed_map, dict):
        sensor_map = fixed_map.get(sensor_key) or fixed_map.get("DEFAULT") or {}
    if not sensor_map and sensor_key == "PRISMA":
        sensor_map = dict(PRISMA_FIXED_BAND_PAIRS)

    out: Dict[str, Any] = {
        "resolved": False,
        "mode": "window",
        "indices": [],
        "reason": None,
        "band_label": str(band_label),
        "sensor": sensor_key,
    }
    if not prefer_fixed_band_pairs:
        out["reason"] = "fixed mapping disabled by config"
        return out
    if not sensor_map:
        out["reason"] = f"no fixed map for sensor {sensor_key}"
        return out

    idx = sensor_map.get(str(band_label))
    if idx is None:
        out["reason"] = f"no fixed mapping for band {band_label}"
        return out

    try:
        idx_1based = int(idx)
    except Exception:
        out["reason"] = f"invalid fixed index {idx!r}"
        return out
    if idx_1based < 1 or idx_1based > int(len(hs_wavelengths)):
        out["reason"] = f"fixed index {idx_1based} out of range for {len(hs_wavelengths)} bands"
        return out

    out["resolved"] = True
    out["mode"] = "fixed"
    out["indices"] = [idx_1based]
    return out


def _resolve_band_indices_for_matching(
    hs_wavelengths: np.ndarray,
    target_wavelength_nm: float,
    band_label: str,
    sensor_type: str,
    prefer_fixed_band_pairs: bool = True,
    fixed_band_pairs_by_sensor: Optional[Dict[str, Dict[str, int]]] = None,
    wavelength_window_nm: float = 20.0,
) -> Dict[str, Any]:
    """Resolve source HS band indices via fixed map first, then wavelength fallback."""
    hs_wl = np.asarray(hs_wavelengths).flatten().astype(float, copy=False)
    out = _resolve_fixed_band_pair(
        hs_wavelengths=hs_wl,
        band_label=band_label,
        sensor_type=sensor_type,
        prefer_fixed_band_pairs=prefer_fixed_band_pairs,
        fixed_band_pairs_by_sensor=fixed_band_pairs_by_sensor,
    )
    if out.get("resolved", False):
        return out

    win = float(max(0.1, wavelength_window_nm))
    half = win / 2.0
    band_mask = (hs_wl >= float(target_wavelength_nm) - half) & (hs_wl <= float(target_wavelength_nm) + half)
    indices = (np.where(band_mask)[0] + 1).astype(int).tolist()
    if indices:
        out["resolved"] = True
        out["mode"] = "window"
        out["indices"] = indices
        return out

    finite = np.isfinite(hs_wl)
    if not np.any(finite):
        out["reason"] = "no finite wavelengths in source band table"
        return out
    nearest = int(np.nanargmin(np.abs(hs_wl - float(target_wavelength_nm))))
    out["resolved"] = True
    out["mode"] = "nearest"
    out["indices"] = [nearest + 1]
    out["reason"] = out.get("reason") or "window match unavailable; using nearest wavelength"
    return out


def _build_consensus_group_ids(df, rounding_px: float = 1.0):
    """Build deterministic consensus group IDs using POINT_ID or image-space fallback."""
    import pandas as pd

    round_px = float(max(0.1, rounding_px))
    if "POINT_ID" in df.columns:
        ids = df["POINT_ID"].astype(str)
        missing = ids.isna() | (ids.str.strip() == "") | (ids.str.lower() == "nan")
        if not missing.any():
            return ids
    else:
        ids = pd.Series([""] * len(df), index=df.index, dtype=str)
        missing = pd.Series([True] * len(df), index=df.index)

    if "X_IM" in df.columns and "Y_IM" in df.columns:
        x = pd.to_numeric(df["X_IM"], errors="coerce")
        y = pd.to_numeric(df["Y_IM"], errors="coerce")
        xq = np.round(x / round_px).astype("Int64")
        yq = np.round(y / round_px).astype("Int64")
        fallback = xq.astype(str) + "_" + yq.astype(str)
        ids = ids.where(~missing, fallback)
        missing = ids.isna() | (ids.str.strip() == "") | (ids.str.lower() == "nan")

    if missing.any():
        seq = pd.Series(np.arange(len(df)), index=df.index).astype(str)
        ids = ids.where(~missing, "row_" + seq)
    return ids.astype(str)


def _compute_quality_score(df):
    """Compute composite tie-point quality score with graceful metric degradation."""
    import pandas as pd

    work = df.copy()

    def _normalize_series(series, higher_is_better=True):
        values = pd.to_numeric(series, errors="coerce").astype(float)
        valid = np.isfinite(values)
        out = np.zeros(len(values), dtype=float)
        if not np.any(valid):
            return out
        valid_vals = values[valid]
        lo = float(np.nanmin(valid_vals))
        hi = float(np.nanmax(valid_vals))
        if not np.isfinite(lo) or not np.isfinite(hi) or abs(hi - lo) < 1e-9:
            out[valid] = 0.5
            return out if higher_is_better else (1.0 - out)
        norm = (values - lo) / (hi - lo)
        norm = np.clip(norm, 0.0, 1.0)
        out[valid] = norm[valid]
        return out if higher_is_better else (1.0 - out)

    metric_specs = [
        ("RELIABILITY", 0.40, True),
        ("SSIM_IMPRO", 0.20, True),
        ("SSIM_AFTER", 0.15, True),
        ("LAST_ERR", 0.15, False),
        ("ABS_SHIFT", 0.05, False),
        ("ABS_SHIFT_M", 0.05, False),
    ]
    score = np.zeros(len(work), dtype=float)
    weight_sum = 0.0
    for col, weight, higher_is_better in metric_specs:
        if col in work.columns:
            comp = _normalize_series(work[col], higher_is_better=higher_is_better)
            score += weight * comp
            weight_sum += weight
    score = (score / weight_sum) if weight_sum > 0 else np.zeros(len(work), dtype=float)

    if "BAND_SUPPORT" in work.columns:
        support = pd.to_numeric(work["BAND_SUPPORT"], errors="coerce").fillna(1.0).astype(float)
        score += np.clip((support - 1.0) * 0.03, 0.0, 0.30)

    penalty = np.zeros(len(work), dtype=float)
    for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
        if col in work.columns:
            flagged = work[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
            penalty[flagged] -= 10.0
    work["QUALITY_SCORE"] = score + penalty
    return work


def _count_occupied_cells(
    df,
    grid_rows: int,
    grid_cols: int,
    x_col: str = "X_IM",
    y_col: str = "Y_IM",
) -> int:
    """Count occupied image-space stratification cells for the given tie points."""
    if df is None or len(df) == 0 or x_col not in df.columns or y_col not in df.columns:
        return 0
    x = np.asarray(df[x_col], dtype=float)
    y = np.asarray(df[y_col], dtype=float)
    valid = np.isfinite(x) & np.isfinite(y)
    if not np.any(valid):
        return 0
    x = x[valid]
    y = y[valid]
    gx = max(1, int(grid_cols))
    gy = max(1, int(grid_rows))
    x_min, x_max = float(np.min(x)), float(np.max(x))
    y_min, y_max = float(np.min(y)), float(np.max(y))
    x_span = max(1e-6, x_max - x_min)
    y_span = max(1e-6, y_max - y_min)
    cx = np.floor((x - x_min) / x_span * gx).astype(int)
    cy = np.floor((y - y_min) / y_span * gy).astype(int)
    cx = np.clip(cx, 0, gx - 1)
    cy = np.clip(cy, 0, gy - 1)
    return int(len(np.unique(cy * gx + cx)))


def _apply_spatial_stratification(
    df,
    grid_rows: int,
    grid_cols: int,
    max_points_per_cell: int,
    min_points_required: int,
    min_distance: float = 60.0,
) -> Dict[str, Any]:
    """Apply image-space balancing with quota relaxation and distance-constrained fallback."""
    import pandas as pd

    selected_default = df.copy() if hasattr(df, "copy") else df
    result: Dict[str, Any] = {
        "selected_df": selected_default,
        "used_stratification": False,
        "quota_used": int(max(1, max_points_per_cell)),
        "fallback_notes": [],
        "occupied_cells_selected": 0,
        "occupied_cells_total": 0,
    }
    if df is None or len(df) == 0:
        return result
    if "X_IM" not in df.columns or "Y_IM" not in df.columns:
        result["fallback_notes"].append("stratification disabled: missing X_IM/Y_IM")
        return result

    rank_col = "QUALITY_SCORE" if "QUALITY_SCORE" in df.columns else None
    if rank_col is None and "RELIABILITY" in df.columns:
        rank_col = "RELIABILITY"
    if rank_col is not None:
        work = df.copy().sort_values(rank_col, ascending=False, kind="mergesort")
    else:
        work = df.copy()
    x = np.asarray(work["X_IM"], dtype=float)
    y = np.asarray(work["Y_IM"], dtype=float)
    valid = np.isfinite(x) & np.isfinite(y)
    if not np.any(valid):
        result["fallback_notes"].append("stratification disabled: no finite coordinates")
        return result

    gx = max(1, int(grid_cols))
    gy = max(1, int(grid_rows))
    x_valid = x[valid]
    y_valid = y[valid]
    x_min, x_max = float(np.min(x_valid)), float(np.max(x_valid))
    y_min, y_max = float(np.min(y_valid)), float(np.max(y_valid))
    x_span = max(1e-6, x_max - x_min)
    y_span = max(1e-6, y_max - y_min)

    cell_x = np.floor((x - x_min) / x_span * gx).astype(int)
    cell_y = np.floor((y - y_min) / y_span * gy).astype(int)
    cell_x = np.clip(cell_x, 0, gx - 1)
    cell_y = np.clip(cell_y, 0, gy - 1)
    work["_STRAT_CELL"] = (cell_y * gx + cell_x).astype(int)
    result["occupied_cells_total"] = int(work["_STRAT_CELL"].nunique())

    def _select_with_quota(quota: int):
        return work.groupby("_STRAT_CELL", sort=False).head(int(max(1, quota)))

    quota = int(max(1, max_points_per_cell))
    selected = _select_with_quota(quota)
    result["used_stratification"] = True

    while len(selected) < int(max(1, min_points_required)) and quota < len(work):
        prev_quota = quota
        quota = min(len(work), max(quota + 1, quota * 2))
        if quota == prev_quota:
            break
        selected = _select_with_quota(quota)
        result["fallback_notes"].append(
            f"stratification quota relaxed from {prev_quota} to {quota} due to low point count"
        )
        if len(selected) >= int(max(1, min_points_required)):
            break

    min_dist = float(max(0.0, min_distance))
    required = int(max(1, min_points_required))
    fallback_exhausted = False

    if len(selected) < required:
        needed = required - len(selected)
        remainder = work.loc[~work.index.isin(selected.index)].copy()

        if needed > 0 and len(remainder) > 0:
            rem_rank_col = "QUALITY_SCORE" if "QUALITY_SCORE" in remainder.columns else None
            if rem_rank_col is None and "RELIABILITY" in remainder.columns:
                rem_rank_col = "RELIABILITY"
            if rem_rank_col is not None:
                remainder = remainder.sort_values(rem_rank_col, ascending=False, kind="mergesort")

            use_map_space = False
            if (
                "X_MAP" in selected.columns
                and "Y_MAP" in selected.columns
                and "X_MAP" in remainder.columns
                and "Y_MAP" in remainder.columns
            ):
                sel_x_map = pd.to_numeric(selected["X_MAP"], errors="coerce").to_numpy(dtype=float)
                sel_y_map = pd.to_numeric(selected["Y_MAP"], errors="coerce").to_numpy(dtype=float)
                use_map_space = bool(np.any(np.isfinite(sel_x_map) & np.isfinite(sel_y_map)))

            x_col = "X_MAP" if use_map_space else "X_IM"
            y_col = "Y_MAP" if use_map_space else "Y_IM"

            if min_dist <= 0.0:
                added_df = remainder.head(needed)
                added_count = int(len(added_df))
                if added_count > 0:
                    selected = pd.concat([selected, added_df], axis=0)
                needed -= added_count
                result["fallback_notes"].append(
                    f"global ranked fallback (compat mode) added {added_count} points"
                )
            else:
                added_indices: List[Any] = []
                sel_x = pd.to_numeric(selected[x_col], errors="coerce").to_numpy(dtype=float)
                sel_y = pd.to_numeric(selected[y_col], errors="coerce").to_numpy(dtype=float)
                valid_sel = np.isfinite(sel_x) & np.isfinite(sel_y)

                active_x = sel_x[valid_sel].tolist()
                active_y = sel_y[valid_sel].tolist()
                min_dist_sq = min_dist * min_dist

                for idx, row in remainder.iterrows():
                    if needed <= 0:
                        break

                    try:
                        cand_x = float(row[x_col])
                        cand_y = float(row[y_col])
                    except (TypeError, ValueError):
                        continue

                    if not (np.isfinite(cand_x) and np.isfinite(cand_y)):
                        continue

                    if not active_x:
                        added_indices.append(idx)
                        active_x.append(cand_x)
                        active_y.append(cand_y)
                        needed -= 1
                        continue

                    dx = np.asarray(active_x, dtype=float) - cand_x
                    dy = np.asarray(active_y, dtype=float) - cand_y
                    dists_sq = dx * dx + dy * dy
                    if float(np.min(dists_sq)) > min_dist_sq:
                        added_indices.append(idx)
                        active_x.append(cand_x)
                        active_y.append(cand_y)
                        needed -= 1

                if added_indices:
                    added_df = remainder.loc[added_indices]
                    selected = pd.concat([selected, added_df], axis=0)
                    result["fallback_notes"].append(
                        f"distance-constrained fallback added {len(added_indices)} points "
                        f"(min_dist={min_dist:g})"
                    )

            if needed > 0:
                fallback_exhausted = True
                result["fallback_notes"].append(
                    f"distance-constrained fallback exhausted with {needed} points deficit"
                )
        elif needed > 0:
            fallback_exhausted = True
            result["fallback_notes"].append(
                f"distance-constrained fallback exhausted with {needed} points deficit"
            )

    min_gcps_order2 = _minimum_gcps_for_polynomial_order(2)
    if fallback_exhausted and len(selected) < int(min_gcps_order2):
        logger.warning(
            "Fallback exhausted: selected tie points (%d) < order-2 minimum (%d). "
            "Recommend downgrading to order-1 affine warp.",
            int(len(selected)),
            int(min_gcps_order2),
        )

    if "QUALITY_SCORE" in selected.columns:
        selected = selected.sort_values("QUALITY_SCORE", ascending=False, kind="mergesort")
    elif "RELIABILITY" in selected.columns:
        selected = selected.sort_values("RELIABILITY", ascending=False, kind="mergesort")
    if "_STRAT_CELL" in selected.columns:
        selected = selected.drop(columns=["_STRAT_CELL"])
    result["quota_used"] = quota
    result["selected_df"] = selected
    result["occupied_cells_selected"] = _count_occupied_cells(selected, gy, gx)
    return result


def _collect_multiband_tiepoints(
    s2_path,
    hs_path,
    hs_wl,
    temp_folder,
    sensor_tag,
    date_tag,
    cand_idx,
    s2_band_subset: Optional[Sequence[str]] = None,
    config: Optional[Dict[str, Any]] = None,
):
    """Collect tie points from multiple S2 bands with fixed-pair and profile-aware controls."""
    result = {
        'success': False,
        'all_tiepoints': [],
        'tiepoint_counts': {},
        'band_match_modes': {},
        'error_message': None,
        'matcher_profile': {},
    }
    hs_wl = np.asarray(hs_wl).flatten()
    cfg = config or {}
    profile = _resolve_sensor_matcher_profile(cfg, sensor_tag)
    result["matcher_profile"] = profile

    prefer_fixed_pairs = bool(cfg.get("prefer_fixed_band_pairs", DEFAULT_CONFIG.get("prefer_fixed_band_pairs", True)))
    fixed_pairs = cfg.get("fixed_band_pairs_by_sensor", DEFAULT_CONFIG.get("fixed_band_pairs_by_sensor", {}))
    band_window_nm = float(
        max(
            0.1,
            float(cfg.get("bandpair_wavelength_window_nm", DEFAULT_CONFIG.get("bandpair_wavelength_window_nm", 20.0))),
        )
    )
    logger.info(
        "Multiband matcher profile sensor=%s global_source=%s local_shift=%.1f local_grid=%s local_ws=%s filter=%s",
        profile.get("sensor"),
        profile.get("global_profile_source"),
        float(profile.get("local_max_shift", 0.0)),
        int(profile.get("local_grid_res", LOCAL_GRID_RES_M)),
        profile.get("local_window_size"),
        int(profile.get("local_tieP_filter_level", 1)),
    )

    try:
        with rasterio.open(hs_path) as src:
            hs_profile = src.profile.copy()
    except Exception as e:
        result['error_message'] = f"Failed to open HS GeoTIFF: {e}"
        return result

    supports_local_max_iter = _supports_constructor_kwarg(COREG_LOCAL, "max_iter")

    selected_band_names: List[str]
    if s2_band_subset:
        requested = [str(x).upper().strip() for x in s2_band_subset if str(x).strip()]
        selected_band_names = [b for b in requested if b in MULTIBAND_S2_WAVELENGTHS]
        unknown = [b for b in requested if b not in MULTIBAND_S2_WAVELENGTHS]
        if unknown:
            logger.warning(
                _fmt_issue(
                    "TIEPOINTS",
                    f"Ignoring unknown Sentinel-2 band labels in subset: {unknown}",
                )
            )
        if not selected_band_names:
            result["error_message"] = "No valid Sentinel-2 band labels in s2_band_subset."
            return result
    else:
        selected_band_names = list(MULTIBAND_S2_WAVELENGTHS.keys())

    for band_name in selected_band_names:
        band_info = MULTIBAND_S2_WAVELENGTHS[band_name]
        target_wl = float(band_info['wavelength'])
        s2_stack_idx = int(band_info['stack_idx'])
        temp_nb_path = os.path.join(temp_folder, f"{sensor_tag}_{date_tag}_nb_{band_name}_c{cand_idx}.tif")
        temp_nb_coreg_path = os.path.join(temp_folder, f"{sensor_tag}_{date_tag}_nb_coreg_{band_name}_c{cand_idx}.tif")
        try:
            resolution = _resolve_band_indices_for_matching(
                hs_wavelengths=hs_wl,
                target_wavelength_nm=target_wl,
                band_label=band_name,
                sensor_type=sensor_tag,
                prefer_fixed_band_pairs=prefer_fixed_pairs,
                fixed_band_pairs_by_sensor=fixed_pairs,
                wavelength_window_nm=band_window_nm,
            )
            result["band_match_modes"][band_name] = resolution.get("mode", "window")
            band_indices_1based = [int(i) for i in resolution.get("indices", []) if int(i) >= 1]
            if not band_indices_1based:
                result['tiepoint_counts'][band_name] = 0
                continue

            with rasterio.open(hs_path) as src:
                selected_bands = src.read(band_indices_1based)
            nb_avg = np.nanmean(selected_bands.astype(PROCESSING_DTYPE), axis=0)
            del selected_bands

            profile_hs = hs_profile.copy()
            profile_hs.update(count=1, dtype=PROCESSING_DTYPE)
            with rasterio.open(temp_nb_path, 'w', **profile_hs) as dst:
                dst.write(nb_avg.astype(PROCESSING_DTYPE), 1)
            del nb_avg

            hs_nodata = hs_profile.get('nodata') if hs_profile.get('nodata') is not None else PROCESSING_NODATA
            
            local_kwargs = {
                "grid_res": int(profile.get("local_grid_res", LOCAL_GRID_RES_M)),
                "window_size": _coerce_window_size(profile.get("local_window_size"), (256, 256)),
                "path_out": temp_nb_coreg_path,
                "fmt_out": "GTiff",
                "out_crea_options": ["COMPRESS=LZW", "BIGTIFF=YES"],
                "r_b4match": s2_stack_idx,
                "s_b4match": 1,
                "max_shift": float(profile.get("local_max_shift", 50.0)),
                "resamp_alg_deshift": 'cubic',
                "tieP_filter_level": int(profile.get("local_tieP_filter_level", 1)),
                "outFillVal": hs_nodata,
                "CPUs": CPUS_FOR_AROSICS,
                "nodata": (0, hs_nodata),
                "progress": False,
                "v": False,
                "q": True,
                "ignore_errors": True,
            }
            local_max_iter = profile.get("local_max_iter")
            if local_max_iter is not None:
                if supports_local_max_iter:
                    local_kwargs["max_iter"] = int(local_max_iter)
                else:
                    logger.info("COREG_LOCAL max_iter override ignored (installed AROSICS signature has no max_iter)")

            CRL = COREG_LOCAL(
                s2_path,
                temp_nb_path,
                **local_kwargs,
            )
            try:
                CRL.correct_shifts()
            except Exception:
                pass

            tie_points_df = getattr(CRL, "CoRegPoints_table", None)
            if tie_points_df is not None and len(tie_points_df) > 0:
                tie_points_df = tie_points_df.copy()
                tie_points_df['BAND_LABEL'] = band_name
                tie_points_df['MATCH_MODE'] = resolution.get("mode", "window")
                tie_points_df['MATCH_BAND_INDICES'] = ",".join(str(v) for v in band_indices_1based)
                if 'RELIABILITY' in tie_points_df.columns:
                    tie_points_df['RELIABILITY'] = tie_points_df['RELIABILITY'].replace([-9999, -9998], np.nan)
                    tie_points_df = tie_points_df[tie_points_df['RELIABILITY'].notna()]
                    if len(tie_points_df) > 0 and float(tie_points_df['RELIABILITY'].max()) <= 1.0:
                        tie_points_df['RELIABILITY'] = tie_points_df['RELIABILITY'] * 100.0
                tp_count = len(tie_points_df)
                if tp_count > 0:
                    result['all_tiepoints'].append(tie_points_df)
                    result['tiepoint_counts'][band_name] = tp_count
                else:
                    result['tiepoint_counts'][band_name] = 0
            else:
                result['tiepoint_counts'][band_name] = 0
        except Exception as e:
            logger.warning(f"  {band_name} error: {e}")
            result['tiepoint_counts'][band_name] = 0
        finally:
            for temp_path in [temp_nb_path, temp_nb_coreg_path]:
                if os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except Exception:
                        pass

    total = sum(result['tiepoint_counts'].values())
    result['success'] = total > 0
    if not result['success']:
        result['error_message'] = "No tie points collected"
    return result


def _merge_tiepoints(
    tiepoint_dfs,
    min_reliability: float = 75.0,
    trim_by_residual: bool = True,
    residual_mad_factor: float = 3.0,
    min_band_support: int = 2,
    allow_single_band_fallback: bool = True,
    consensus_group_rounding_px: float = 1.0,
    grid_rows: int = 4,
    grid_cols: int = 4,
    max_points_per_cell: int = 3,
    min_points_required: int = MIN_TIE_POINTS_FOR_POLYNOMIAL,
    spatial_fallback_min_distance: float = 60.0,
) -> Dict[str, Any]:
    """Merge tiepoints using consensus, composite scoring, and spatial stratification."""
    import pandas as pd

    result = {
        'success': False,
        'merged_df': None,
        'visualization_df': None,
        'n_total_before_merge': 0,
        'n_after_outlier_filter': 0,
        'n_after_reliability_filter': 0,
        'n_after_residual_trim': 0,
        'reliability_threshold_used': None,
        'residual_threshold_m': None,
        'error_message': None,
        'stage_counts': {},
        'fallback_notes': [],
        'consensus_fallback_used': False,
        'occupied_cells_selected': 0,
        'occupied_cells_total': 0,
    }
    if not tiepoint_dfs:
        result['error_message'] = "No tie point DataFrames provided"
        return result

    try:
        combined_df = pd.concat(tiepoint_dfs, ignore_index=True)
        result['n_total_before_merge'] = len(combined_df)
        result['stage_counts']['01_concat_sanitize'] = int(len(combined_df))
        if len(combined_df) == 0:
            result['error_message'] = "No tiepoints after concatenation"
            return result

        for col in ("X_IM", "Y_IM", "X_SHIFT_M", "Y_SHIFT_M", "LAST_ERR", "ABS_SHIFT", "ABS_SHIFT_M"):
            if col in combined_df.columns:
                combined_df[col] = pd.to_numeric(combined_df[col], errors='coerce')
        if 'RELIABILITY' in combined_df.columns:
            combined_df['RELIABILITY'] = pd.to_numeric(combined_df['RELIABILITY'], errors='coerce')
            combined_df = combined_df[combined_df['RELIABILITY'].notna() & (combined_df['RELIABILITY'] > -9990)]
            if len(combined_df) > 0 and float(combined_df['RELIABILITY'].max()) <= 1.0:
                combined_df['RELIABILITY'] = combined_df['RELIABILITY'] * 100.0
        result['stage_counts']['01_concat_sanitize'] = int(len(combined_df))
        if len(combined_df) == 0:
            result['error_message'] = "All tie points removed during sanitation"
            return result

        for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
            if col in combined_df.columns:
                outlier_mask = combined_df[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
                combined_df = combined_df[~outlier_mask]
        result['n_after_outlier_filter'] = int(len(combined_df))
        result['stage_counts']['02_remove_outliers'] = int(len(combined_df))
        if len(combined_df) == 0:
            result['error_message'] = "All tie points removed by outlier filter"
            return result

        combined_df = combined_df.copy()
        combined_df["CONSENSUS_GROUP_ID"] = _build_consensus_group_ids(
            combined_df,
            rounding_px=float(max(0.1, consensus_group_rounding_px)),
        )
        if "BAND_LABEL" in combined_df.columns:
            band_support = combined_df.groupby("CONSENSUS_GROUP_ID")["BAND_LABEL"].nunique()
        else:
            band_support = combined_df.groupby("CONSENSUS_GROUP_ID").size()
        combined_df["BAND_SUPPORT"] = combined_df["CONSENSUS_GROUP_ID"].map(band_support).astype(int)

        min_support = max(1, int(min_band_support))
        consensus_df = combined_df[combined_df["BAND_SUPPORT"] >= min_support]
        if len(consensus_df) == 0:
            if allow_single_band_fallback:
                consensus_df = combined_df
                result["consensus_fallback_used"] = True
                result["fallback_notes"].append(
                    f"consensus fallback enabled: no points met min_band_support={min_support}"
                )
            else:
                result['error_message'] = (
                    f"No tie points meet min_band_support={min_support} and fallback is disabled"
                )
                return result
        result['stage_counts']['03_consensus_filter'] = int(len(consensus_df))

        scored_df = _compute_quality_score(consensus_df)
        result['stage_counts']['04_quality_score'] = int(len(scored_df))
        result['visualization_df'] = scored_df.copy()

        strat = _apply_spatial_stratification(
            scored_df,
            grid_rows=max(1, int(grid_rows)),
            grid_cols=max(1, int(grid_cols)),
            max_points_per_cell=max(1, int(max_points_per_cell)),
            min_points_required=max(1, int(min_points_required)),
            min_distance=float(spatial_fallback_min_distance),
        )
        selected_df = strat.get("selected_df", scored_df)
        result["fallback_notes"].extend(strat.get("fallback_notes", []))
        result["occupied_cells_selected"] = int(strat.get("occupied_cells_selected", 0))
        result["occupied_cells_total"] = int(strat.get("occupied_cells_total", 0))
        result['stage_counts']['05_spatial_stratification'] = int(len(selected_df))

        selected_df = (
            selected_df.sort_values("QUALITY_SCORE", ascending=False)
            .drop_duplicates(subset="CONSENSUS_GROUP_ID", keep="first")
        )
        result['stage_counts']['06_select_best_per_group_cell'] = int(len(selected_df))

        selected_df = selected_df.sort_values("QUALITY_SCORE", ascending=False)
        if 'RELIABILITY' in selected_df.columns:
            thresholds = [float(min_reliability), 65.0, 55.0, 45.0, 30.0, 20.0]
            for threshold in thresholds:
                filtered = selected_df[selected_df['RELIABILITY'] >= threshold]
                if len(filtered) >= max(1, int(min_points_required)):
                    selected_df = filtered
                    result['reliability_threshold_used'] = threshold
                    break
                if threshold == thresholds[-1] and len(filtered) > 0:
                    selected_df = filtered
                    result['reliability_threshold_used'] = threshold
                    result['fallback_notes'].append(
                        f"reliability threshold relaxed to {threshold:.1f} due to low retained points"
                    )
        result['n_after_reliability_filter'] = int(len(selected_df))
        result['stage_counts']['07_final_reliability_filter'] = int(len(selected_df))

        if trim_by_residual and len(selected_df) > max(3, int(min_points_required)):
            if 'X_SHIFT_M' in selected_df.columns and 'Y_SHIFT_M' in selected_df.columns:
                residuals = np.sqrt(selected_df['X_SHIFT_M'].values**2 + selected_df['Y_SHIFT_M'].values**2)
                selected_df = selected_df.copy()
                selected_df['_RESIDUAL_M'] = residuals
                median_res = np.median(residuals)
                mad = np.median(np.abs(residuals - median_res))
                if mad > 0:
                    threshold_m = median_res + float(residual_mad_factor) * 1.4826 * mad
                    result['residual_threshold_m'] = float(threshold_m)
                    trimmed_df = selected_df[selected_df['_RESIDUAL_M'] <= threshold_m]
                    if len(trimmed_df) < int(max(1, min_points_required)):
                        trimmed_df = selected_df.nsmallest(int(max(1, min_points_required)), '_RESIDUAL_M')
                    selected_df = trimmed_df
                if '_RESIDUAL_M' in selected_df.columns:
                    selected_df = selected_df.drop(columns=['_RESIDUAL_M'])

        result['n_after_residual_trim'] = int(len(selected_df))
        result['stage_counts']['08_optional_residual_trim'] = int(len(selected_df))
        result['merged_df'] = selected_df

        logger.info(
            "Tie-point merge counts: %s",
            ", ".join(
                f"{stage}={count}"
                for stage, count in result.get("stage_counts", {}).items()
            ),
        )
        for note in result.get("fallback_notes", []):
            logger.info("Tie-point merge fallback: %s", note)

        if len(selected_df) >= int(max(1, min_points_required)):
            result['success'] = True
        else:
            result['error_message'] = (
                f"Insufficient tiepoints after merge: {len(selected_df)} < {int(max(1, min_points_required))}"
            )
    except Exception as e:
        result['error_message'] = f"Merge failed: {e}"
    return result


def _build_gcps_from_tiepoints(merged_df):
    """Build GDAL GCPs from tie points."""
    result = {'success': False, 'gcps': [], 'n_gcps': 0, 'error_message': None}
    if merged_df is None or len(merged_df) == 0:
        result['error_message'] = "No tie points provided"
        return result
    try:
        from osgeo import gdal
        required = ['X_IM', 'Y_IM', 'X_MAP', 'Y_MAP', 'X_SHIFT_M', 'Y_SHIFT_M']
        missing = [c for c in required if c not in merged_df.columns]
        if missing:
            result['error_message'] = f"Missing columns: {missing}"
            return result

        gcps = []
        for idx, row in merged_df.iterrows():
            pixel = float(row['X_IM'])
            line = float(row['Y_IM'])
            x_map_corrected = float(row['X_MAP']) + float(row['X_SHIFT_M'])
            y_map_corrected = float(row['Y_MAP']) + float(row['Y_SHIFT_M'])
            gcp = gdal.GCP(x_map_corrected, y_map_corrected, 0.0, pixel, line)
            gcps.append(gcp)

        result['gcps'] = gcps
        result['n_gcps'] = len(gcps)
        result['success'] = True
        logger.info(f"Built {len(gcps)} GCPs from tie points")
    except Exception as e:
        result['error_message'] = f"GCP construction failed: {e}"
    return result


def _minimum_gcps_for_polynomial_order(order: int) -> int:
    """Minimum algebraic GCP count for a polynomial order."""
    order_i = max(1, int(order))
    return int((order_i + 1) * (order_i + 2) / 2)


def _assess_gcp_geometry_for_order2(merged_df) -> Dict[str, Any]:
    """Assess whether tiepoint geometry is suitable for robust order-2 fitting."""
    out = {"ok": False, "reason": "missing geometry columns", "condition_number": None}
    if merged_df is None or len(merged_df) == 0:
        out["reason"] = "no merged tiepoints"
        return out
    if "X_IM" not in merged_df.columns or "Y_IM" not in merged_df.columns:
        return out
    x = np.asarray(merged_df["X_IM"], dtype=float)
    y = np.asarray(merged_df["Y_IM"], dtype=float)
    valid = np.isfinite(x) & np.isfinite(y)
    if int(np.count_nonzero(valid)) < 6:
        out["reason"] = "insufficient finite image-space coordinates"
        return out
    x = x[valid]
    y = y[valid]

    x_span = float(np.percentile(x, 95) - np.percentile(x, 5))
    y_span = float(np.percentile(y, 95) - np.percentile(y, 5))
    if x_span < 20.0 or y_span < 20.0:
        out["reason"] = f"point spread too clustered (x_span={x_span:.2f}, y_span={y_span:.2f})"
        return out

    try:
        design = np.column_stack([np.ones_like(x), x, y, x * y, x * x, y * y])
        cond = float(np.linalg.cond(design))
        out["condition_number"] = cond
        if np.isfinite(cond) and cond < 1.0e8:
            out["ok"] = True
            out["reason"] = "geometry conditioning acceptable"
        else:
            out["reason"] = f"poor geometry conditioning (cond={cond:.2e})"
    except Exception as exc:
        out["reason"] = f"geometry conditioning failed: {exc}"
    return out


def _decide_polynomial_order(
    merged_df,
    preferred_order: int = 2,
    auto_downgrade: bool = True,
    min_gcps_order2: int = 12,
    min_cells_order2: int = 6,
    grid_rows: int = 4,
    grid_cols: int = 4,
) -> Dict[str, Any]:
    """Choose polynomial order with optional safety downgrade from 2->1."""
    order_target = 1 if int(preferred_order) <= 1 else 2
    n_gcps = int(len(merged_df)) if merged_df is not None else 0
    occupied_cells = _count_occupied_cells(merged_df, grid_rows=max(1, int(grid_rows)), grid_cols=max(1, int(grid_cols)))
    geom = _assess_gcp_geometry_for_order2(merged_df)

    decision = {
        "preferred_order": order_target,
        "order_used": order_target,
        "downgraded": False,
        "reason": "preferred order honored",
        "n_gcps": n_gcps,
        "occupied_cells": occupied_cells,
        "geometry_ok": bool(geom.get("ok", False)),
        "geometry_reason": geom.get("reason"),
        "geometry_condition_number": geom.get("condition_number"),
    }
    if order_target <= 1:
        decision["reason"] = "preferred_polynomial_order=1"
        return decision
    if not auto_downgrade:
        decision["reason"] = "auto-downgrade disabled"
        return decision

    fail_reasons = []
    if n_gcps < int(max(6, min_gcps_order2)):
        fail_reasons.append(f"gcps {n_gcps} < min_gcps_order2 {int(max(6, min_gcps_order2))}")
    if occupied_cells < int(max(1, min_cells_order2)):
        fail_reasons.append(f"occupied_cells {occupied_cells} < min_cells_order2 {int(max(1, min_cells_order2))}")
    if not bool(geom.get("ok", False)):
        fail_reasons.append(str(geom.get("reason", "geometry check failed")))

    if fail_reasons:
        decision["order_used"] = 1
        decision["downgraded"] = True
        decision["reason"] = "; ".join(fail_reasons)
    return decision


def _apply_polynomial_warp(
    input_raster,
    output_raster,
    gcps,
    target_crs,
    polynomial_order: int = 2,
    output_resolution: float = 30.0,
    nodata: float = -9999.0,
    resampling: str = "cubic",
    s2_bounds=None,
):
    """Apply polynomial warp using GDAL with configurable order."""
    order = 1 if int(polynomial_order) <= 1 else 2
    result = {
        'success': False,
        'output_path': None,
        'error_message': None,
        'order_used': order,
    }

    min_needed = _minimum_gcps_for_polynomial_order(order)
    if not gcps or len(gcps) < min_needed:
        result['error_message'] = (
            f"Insufficient GCPs for order {order}: {len(gcps) if gcps else 0} < {min_needed}"
        )
        return result

    temp_vrt = None
    try:
        from osgeo import gdal
        gdal.UseExceptions()

        temp_vrt = output_raster.replace('.tif', '_gcps.vrt')
        ds_in = gdal.Open(input_raster, gdal.GA_ReadOnly)
        if ds_in is None:
            result['error_message'] = f"Failed to open input: {input_raster}"
            return result

        driver_vrt = gdal.GetDriverByName('VRT')
        ds_vrt = driver_vrt.CreateCopy(temp_vrt, ds_in)
        crs_wkt = target_crs.to_wkt() if hasattr(target_crs, 'to_wkt') else str(target_crs)
        ds_vrt.SetGCPs(gcps, crs_wkt)
        ds_vrt.FlushCache()
        ds_vrt = None
        ds_in = None

        gdalwarp_exe = resolve_gdalwarp_exe()
        resampling_alg = str(resampling or "cubic").strip().lower()
        if resampling_alg not in {"near", "bilinear", "cubic", "cubicspline", "lanczos"}:
            resampling_alg = "cubic"

        cmd = [
            gdalwarp_exe, "-order", str(order), "-t_srs", crs_wkt,
            "-tr", str(output_resolution), str(output_resolution),
            "-r", str(resampling_alg), "-of", "GTiff",
            "-co", "COMPRESS=LZW", "-co", "BIGTIFF=YES", "-co", "TILED=YES",
            "-srcnodata", str(nodata), "-dstnodata", str(nodata),
        ]
        if s2_bounds is not None:
            minx, miny, maxx, maxy = s2_bounds
            cmd.extend(["-te", str(minx), str(miny), str(maxx), str(maxy)])
        cmd.extend([temp_vrt, output_raster])

        subprocess.run(cmd, check=True, capture_output=True, text=True)

        if os.path.exists(output_raster):
            result['success'] = True
            result['output_path'] = output_raster
        else:
            result['error_message'] = "Output file not created"
    except subprocess.CalledProcessError as e:
        result['error_message'] = f"gdalwarp failed: {e.stderr}"
    except Exception as e:
        result['error_message'] = f"Polynomial warp failed: {e}"
    finally:
        if temp_vrt and os.path.exists(temp_vrt):
            try:
                os.remove(temp_vrt)
            except Exception:
                pass
    return result


def _estimate_transform_from_corner_coords(corner_info: Dict[str, Any], target_crs) -> Affine:
    """Estimate affine transform from UL/UR/LL/LR lon/lat corners."""
    required_keys = (
        "ul_lon", "ul_lat", "ur_lon", "ur_lat",
        "ll_lon", "ll_lat", "lr_lon", "lr_lat",
        "rows", "cols",
    )
    missing = [k for k in required_keys if corner_info.get(k) is None]
    if missing:
        raise RuntimeError(
            _fmt_issue("ANCILLARY", f"Missing corner georeference keys: {missing}")
        )

    rows = int(corner_info["rows"])
    cols = int(corner_info["cols"])
    if rows <= 1 or cols <= 1:
        raise RuntimeError(
            _fmt_issue("ANCILLARY", f"Invalid raster shape for corner transform ({rows}, {cols}).")
        )

    corner_lonlat = [
        (float(corner_info["ul_lon"]), float(corner_info["ul_lat"])),
        (float(corner_info["ur_lon"]), float(corner_info["ur_lat"])),
        (float(corner_info["ll_lon"]), float(corner_info["ll_lat"])),
        (float(corner_info["lr_lon"]), float(corner_info["lr_lat"])),
    ]

    tr = Transformer.from_crs("EPSG:4326", target_crs, always_xy=True)
    x_map = [tr.transform(x, y)[0] for x, y in corner_lonlat]
    y_map = [tr.transform(x, y)[1] for x, y in corner_lonlat]

    # Solve affine from image corners in pixel coordinates.
    A = np.array(
        [
            [0, 0, 1],
            [cols - 1, 0, 1],
            [0, rows - 1, 1],
            [cols - 1, rows - 1, 1],
        ],
        dtype=float,
    )
    sol_x = np.linalg.lstsq(A, np.array(x_map, dtype=float), rcond=None)[0]
    sol_y = np.linalg.lstsq(A, np.array(y_map, dtype=float), rcond=None)[0]
    a, b, c = sol_x
    d, e, f = sol_y

    # Shift from pixel-center fit to pixel-corner geotransform convention.
    c -= (a * 0.5 + b * 0.5)
    f -= (d * 0.5 + e * 0.5)
    return Affine(a, b, c, d, e, f)


def _prepare_bands_first(data: np.ndarray, rows: int, cols: int, label: str) -> np.ndarray:
    """Convert 2D/3D raster array to (bands, rows, cols)."""
    arr = np.asarray(data)
    if arr.ndim == 2:
        if arr.shape != (rows, cols):
            raise RuntimeError(
                _fmt_issue(
                    "ANCILLARY",
                    f"{label} shape mismatch: expected {(rows, cols)}, got {tuple(arr.shape)}.",
                )
            )
        return arr[np.newaxis, :, :]

    if arr.ndim != 3:
        raise RuntimeError(
            _fmt_issue("ANCILLARY", f"{label} array must be 2D/3D, got ndim={arr.ndim}.")
        )

    # Find row/col axes matching the geolocation grid.
    for row_ax in range(3):
        for col_ax in range(3):
            if row_ax == col_ax:
                continue
            if (arr.shape[row_ax], arr.shape[col_ax]) == (rows, cols):
                band_ax = [ax for ax in range(3) if ax not in (row_ax, col_ax)][0]
                return np.moveaxis(arr, (band_ax, row_ax, col_ax), (0, 1, 2))

    raise RuntimeError(
        _fmt_issue(
            "ANCILLARY",
            f"Could not map {label} dimensions {tuple(arr.shape)} to geolocation shape {(rows, cols)}.",
        )
    )


def _write_georeferenced_raster(
    output_path: str,
    data_bands_first: np.ndarray,
    crs,
    transform: Affine,
    dtype: str,
    nodata: Optional[float] = None,
) -> None:
    """Write bands-first data to GeoTIFF with georeference."""
    data = np.asarray(data_bands_first)
    if data.ndim != 3:
        raise RuntimeError(_fmt_issue("ANCILLARY", "Expected bands-first 3D array."))
    bands, rows, cols = data.shape
    profile = {
        "driver": "GTiff",
        "height": int(rows),
        "width": int(cols),
        "count": int(bands),
        "dtype": dtype,
        "crs": crs,
        "transform": transform,
        "compress": "lzw",
        "tiled": True,
        "BIGTIFF": "YES",
    }
    if nodata is not None:
        profile["nodata"] = nodata
    with rasterio.open(output_path, "w", **profile) as dst:
        dst.write(data)


def _infer_raster_native_resolution(raster_path: str, fallback: float = 30.0) -> Tuple[float, float]:
    """Estimate raster native resolution from affine transform."""
    with rasterio.open(raster_path) as src:
        tr = src.transform
        x_res = float(np.hypot(tr.a, tr.b))
        y_res = float(np.hypot(tr.d, tr.e))
    if not np.isfinite(x_res) or x_res <= 0:
        x_res = float(fallback)
    if not np.isfinite(y_res) or y_res <= 0:
        y_res = float(fallback)
    return x_res, y_res


def _resolve_pan_window_size_for_raster(
    raster_path: str,
    requested_window_size: Tuple[int, int] = (512, 512),
) -> Tuple[int, int]:
    """
    Clamp PAN matcher window to raster dimensions.

    This enforces the PAN minimum strategy (512x512 where possible) while
    preventing AROSICS from receiving a window larger than the image bounds.
    """
    req_w, req_h = _coerce_window_size(requested_window_size, default=(512, 512))
    req_w = max(512, int(req_w))
    req_h = max(512, int(req_h))
    with rasterio.open(raster_path) as src:
        width = int(src.width)
        height = int(src.height)
    return (
        int(max(1, min(req_w, width))),
        int(max(1, min(req_h, height))),
    )


def _create_synthetic_s2_pan(
    s2_stack_path: str,
    output_path: str,
    s2_band_indices: Sequence[int] = (1, 2, 3, 4),
) -> Dict[str, Any]:
    """
    Create synthetic Sentinel-2 PAN proxy as mean(B02,B03,B04,B08) in windowed IO.

    Key behavior:
    - nodata-safe per-pixel mean (ignores nodata/NaN values)
    - all-nodata pixels become output nodata
    - preserves compact source dtype (e.g., uint16) on disk
    """
    out: Dict[str, Any] = {
        "ok": False,
        "path": output_path,
        "error": None,
        "dtype": None,
        "nodata": None,
        "window_count": 0,
    }
    try:
        if not s2_stack_path or not os.path.exists(s2_stack_path):
            out["error"] = f"S2 stack not found for synthetic PAN: {s2_stack_path}"
            return out

        band_ids = [int(b) for b in list(s2_band_indices)]
        if not band_ids:
            out["error"] = "Synthetic PAN requires at least one Sentinel-2 band index."
            return out

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with rasterio.open(s2_stack_path) as src:
            invalid = [b for b in band_ids if b < 1 or b > int(src.count)]
            if invalid:
                out["error"] = (
                    f"Synthetic PAN band index out of range for stack count={int(src.count)}: {invalid}"
                )
                return out

            out_dtype_name = str(src.dtypes[band_ids[0] - 1])
            out_dtype = np.dtype(out_dtype_name)

            src_nodata = src.nodata
            if src_nodata is not None and np.isfinite(float(src_nodata)):
                nodata_value = float(src_nodata)
            else:
                nodata_value = 0.0 if np.issubdtype(out_dtype, np.integer) else float(PROCESSING_NODATA)

            if np.issubdtype(out_dtype, np.integer):
                info = np.iinfo(out_dtype)
                nodata_value = float(np.clip(np.rint(nodata_value), info.min, info.max))

            profile = src.profile.copy()
            profile.update(
                count=1,
                dtype=out_dtype_name,
                nodata=nodata_value,
                compress="LZW",
                BIGTIFF="YES",
            )

            with rasterio.open(output_path, "w", **profile) as dst:
                src_tags = src.tags()
                if src_tags:
                    dst.update_tags(**src_tags)
                dst.set_band_description(1, "S2_SYNTHETIC_PAN_B02_B03_B04_B08")

                window_count = 0
                for _, window in src.block_windows(1):
                    window_count += 1
                    block = src.read(band_ids, window=window).astype(np.float64, copy=False)
                    valid = np.isfinite(block)
                    if src_nodata is not None and np.isfinite(float(src_nodata)):
                        valid &= (block != float(src_nodata))

                    all_invalid = ~np.any(valid, axis=0)
                    block = np.where(valid, block, np.nan)
                    with warnings.catch_warnings():
                        warnings.simplefilter("ignore", category=RuntimeWarning)
                        with np.errstate(invalid="ignore", divide="ignore"):
                            mean_block = np.nanmean(block, axis=0)

                    if np.issubdtype(out_dtype, np.integer):
                        info = np.iinfo(out_dtype)
                        out_block_f = np.full(mean_block.shape, float(nodata_value), dtype=np.float64)
                        valid_mean = (~all_invalid) & np.isfinite(mean_block)
                        if np.any(valid_mean):
                            out_block_f[valid_mean] = np.rint(mean_block[valid_mean])
                        out_block_f = np.clip(out_block_f, info.min, info.max)
                        out_block = out_block_f.astype(out_dtype, copy=False)
                    else:
                        out_block = np.where(
                            (~all_invalid) & np.isfinite(mean_block),
                            mean_block,
                            float(nodata_value),
                        ).astype(out_dtype, copy=False)

                    dst.write(out_block, 1, window=window)

            out["window_count"] = int(window_count)
            out["dtype"] = out_dtype_name
            out["nodata"] = float(nodata_value)

        out["ok"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _collect_pan_tiepoints_with_synthetic_reference(
    synthetic_s2_pan_path: str,
    pan_source_path: str,
    pan_global_path: str,
    pan_local_path: str,
    ws: Tuple[int, int],
    grid_res: int,
    max_shift: float,
    tiep_filter_level: int,
    local_max_iter: Optional[int],
    reference_nodata: float,
    source_nodata: float,
) -> Dict[str, Any]:
    """Run PAN-specific global+local AROSICS against synthetic S2 PAN reference."""
    import pandas as pd

    out: Dict[str, Any] = {
        "ok": False,
        "warnings": [],
        "global_path": pan_source_path,
        "local_path": None,
        "tiepoints_df": None,
        "n_tiepoints_raw": 0,
    }

    # Global pass (best-effort) to reduce large translation before local matching.
    try:
        global_kwargs = {
            "r_b4match": 1,
            "s_b4match": 1,
            "path_out": pan_global_path,
            "fmt_out": "GTiff",
            "out_crea_options": ["COMPRESS=LZW", "BIGTIFF=YES", "TILED=YES"],
            "max_shift": float(max(5.0, max_shift)),
            "ws": tuple(ws),
            "resamp_alg_deshift": "cubic",
            "nodata": (float(reference_nodata), float(source_nodata)),
            "progress": False,
            "ignore_errors": True,
            "v": False,
            "q": True,
        }
        with contextlib.redirect_stdout(io.StringIO()):
            CRG = COREG(
                synthetic_s2_pan_path,
                pan_source_path,
                **global_kwargs,
            )
            CRG.correct_shifts()
        if os.path.exists(pan_global_path):
            out["global_path"] = pan_global_path
    except Exception as e:
        out["warnings"].append(f"PAN global synthetic-reference alignment failed: {e}")

    if not os.path.exists(out["global_path"]):
        copy_res = stream_copy_raster_to_path(
            source_path=pan_source_path,
            output_path=pan_global_path,
            nodata_fallback=float(source_nodata),
            collect_band_stats=False,
        )
        if not copy_res.get("ok", False):
            out["warnings"].append(
                f"PAN fallback copy after global failure also failed: {copy_res.get('errors', ['unknown'])}"
            )
            return out
        out["global_path"] = pan_global_path

    supports_local_max_iter = _supports_constructor_kwarg(COREG_LOCAL, "max_iter")
    local_kwargs = {
        "grid_res": int(max(30, int(grid_res))),
        "window_size": tuple(ws),
        "path_out": pan_local_path,
        "fmt_out": "GTiff",
        "out_crea_options": ["COMPRESS=LZW", "BIGTIFF=YES", "TILED=YES"],
        "r_b4match": 1,
        "s_b4match": 1,
        "max_shift": float(max(5.0, max_shift)),
        "resamp_alg_deshift": "cubic",
        "tieP_filter_level": int(max(0, int(tiep_filter_level))),
        "outFillVal": float(source_nodata),
        "CPUs": CPUS_FOR_AROSICS,
        "nodata": (float(reference_nodata), float(source_nodata)),
        "progress": False,
        "v": False,
        "q": True,
        "ignore_errors": True,
    }
    if local_max_iter is not None and supports_local_max_iter:
        local_kwargs["max_iter"] = int(max(1, int(local_max_iter)))

    try:
        with contextlib.redirect_stdout(io.StringIO()):
            CRL = COREG_LOCAL(
                synthetic_s2_pan_path,
                out["global_path"],
                **local_kwargs,
            )
            CRL.correct_shifts()
        tie_points_df = getattr(CRL, "CoRegPoints_table", None)
        if tie_points_df is not None and len(tie_points_df) > 0:
            tie_points_df = tie_points_df.copy()
            tie_points_df["BAND_LABEL"] = "S2_SYNTH_PAN"
            if "RELIABILITY" in tie_points_df.columns:
                tie_points_df["RELIABILITY"] = pd.to_numeric(tie_points_df["RELIABILITY"], errors="coerce")
            out["tiepoints_df"] = tie_points_df
            out["n_tiepoints_raw"] = int(len(tie_points_df))
            out["local_path"] = pan_local_path if os.path.exists(pan_local_path) else None
            out["ok"] = True
            return out
        out["warnings"].append("PAN local synthetic-reference matching returned zero tie points.")
        return out
    except Exception as e:
        out["warnings"].append(f"PAN local synthetic-reference matching failed: {e}")
        return out


def _normalize_pan_gcp_mode(mode: Any) -> str:
    """Normalize PAN GCP mode."""
    token = str(mode or "").strip().lower()
    if token in {"map_inverse", "scaled_image"}:
        return token
    if token:
        logger.warning(
            _fmt_issue(
                "ANCILLARY",
                f"Invalid pan_gcp_mode '{mode}'; using default 'map_inverse'.",
            )
        )
    return "map_inverse"


def _normalize_pan_dxdy_source(source: Any) -> str:
    """Normalize PAN dx/dy source selection."""
    token = str(source or "").strip().lower()
    if token in {"auto", "xy_shift_m", "zero"}:
        return token
    if token:
        logger.warning(
            _fmt_issue(
                "ANCILLARY",
                f"Invalid pan_map_dxdy_source '{source}'; using default 'auto'.",
            )
        )
    return "auto"


def _build_gdalwarp_tps_command(
    gdalwarp_exe: str,
    temp_vrt: str,
    output_raster: str,
    crs_wkt: str,
    x_res: float,
    y_res: float,
    resampling: str = "near",
    nodata: Optional[float] = None,
    target_aligned_pixels: bool = False,
    target_extent: Optional[Tuple[float, float, float, float]] = None,
) -> List[str]:
    """Build gdalwarp command for TPS-based warps."""
    cmd = [
        gdalwarp_exe,
        "-tps",
        "-t_srs",
        crs_wkt,
        "-tr",
        str(float(x_res)),
        str(float(y_res)),
        "-r",
        str(resampling),
        "-of",
        "GTiff",
        "-co",
        "COMPRESS=LZW",
        "-co",
        "BIGTIFF=YES",
        "-co",
        "TILED=YES",
    ]
    if bool(target_aligned_pixels):
        cmd.append("-tap")
    if target_extent is not None and len(target_extent) == 4:
        minx, miny, maxx, maxy = [float(v) for v in target_extent]
        cmd.extend(["-te", str(minx), str(miny), str(maxx), str(maxy)])
    if nodata is not None:
        cmd.extend(["-srcnodata", str(nodata), "-dstnodata", str(nodata)])
    cmd.extend([temp_vrt, output_raster])
    return cmd


def build_pan_gcps_from_tiepoints(
    tiepoints_df,
    hs_pixel_size_m: float,
    pan_pixel_size_m: float,
    map_dx_dy_source: str = "auto",
    nodata: float = PROCESSING_NODATA,
    crs_wkt_or_epsg: Optional[str] = None,
    pan_shape: Optional[Tuple[int, int]] = None,
):
    """
    Build PAN GCPs from tiepoint image coordinates scaled from HS pixel-space.

    Pixel/line on PAN are derived from tiepoint image-space coordinates:
        col_pan = X_IM * (hs_pixel_size_m / pan_pixel_size_m)
        row_pan = Y_IM * (hs_pixel_size_m / pan_pixel_size_m)

    Map destination coordinates use tiepoint map shifts:
        X = X_MAP + dx
        Y = Y_MAP + dy
    """
    if tiepoints_df is None or len(tiepoints_df) == 0:
        logger.warning(_fmt_issue("ANCILLARY", "No tiepoints available for PAN scaled-image GCP mode."))
        return []

    try:
        from osgeo import gdal
    except Exception as e:
        logger.warning(_fmt_issue("ANCILLARY", f"GDAL Python bindings unavailable for PAN GCP build: {e}"))
        return []

    mode = _normalize_pan_dxdy_source(map_dx_dy_source)
    try:
        hs_px = float(hs_pixel_size_m)
        pan_px = float(pan_pixel_size_m)
    except Exception:
        hs_px = np.nan
        pan_px = np.nan
    if not (np.isfinite(hs_px) and hs_px > 0 and np.isfinite(pan_px) and pan_px > 0):
        logger.warning(
            _fmt_issue(
                "ANCILLARY",
                f"Invalid HS/PAN pixel sizes for scaled-image PAN GCPs (hs={hs_px}, pan={pan_px}).",
            )
        )
        return []
    scale = float(hs_px / pan_px)

    required = ["X_IM", "Y_IM", "X_MAP", "Y_MAP"]
    missing = [c for c in required if c not in tiepoints_df.columns]
    if missing:
        logger.warning(_fmt_issue("ANCILLARY", f"Missing tiepoint columns for PAN scaled-image GCPs: {missing}"))
        return []

    df = tiepoints_df.copy()
    for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
        if col in df.columns:
            try:
                bad = df[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
            except Exception:
                bad = np.asarray(df[col], dtype=bool)
            df = df[~bad]

    n_total = int(len(df))
    if n_total < 1:
        logger.warning(_fmt_issue("ANCILLARY", "No inlier tiepoints available after outlier filtering."))
        return []

    x_im = np.asarray(df["X_IM"], dtype=float)
    y_im = np.asarray(df["Y_IM"], dtype=float)
    x_map = np.asarray(df["X_MAP"], dtype=float)
    y_map = np.asarray(df["Y_MAP"], dtype=float)

    if mode == "zero":
        dx = np.zeros_like(x_map, dtype=float)
        dy = np.zeros_like(y_map, dtype=float)
    else:
        if "X_SHIFT_M" in df.columns and "Y_SHIFT_M" in df.columns:
            dx = np.asarray(df["X_SHIFT_M"], dtype=float)
            dy = np.asarray(df["Y_SHIFT_M"], dtype=float)
        elif mode == "xy_shift_m":
            logger.warning(
                _fmt_issue(
                    "ANCILLARY",
                    "pan_map_dxdy_source=xy_shift_m requested but X_SHIFT_M/Y_SHIFT_M are missing.",
                )
            )
            return []
        else:
            logger.warning(
                _fmt_issue(
                    "ANCILLARY",
                    "pan_map_dxdy_source=auto: X_SHIFT_M/Y_SHIFT_M missing; falling back to zero shifts.",
                )
            )
            dx = np.zeros_like(x_map, dtype=float)
            dy = np.zeros_like(y_map, dtype=float)

    finite = (
        np.isfinite(x_im)
        & np.isfinite(y_im)
        & np.isfinite(x_map)
        & np.isfinite(y_map)
        & np.isfinite(dx)
        & np.isfinite(dy)
    )

    col_pan = x_im * scale
    row_pan = y_im * scale

    if pan_shape is not None and len(pan_shape) == 2:
        h = int(pan_shape[0])
        w = int(pan_shape[1])
        finite &= (col_pan >= -1.0) & (row_pan >= -1.0) & (col_pan <= (w + 1.0)) & (row_pan <= (h + 1.0))

    if np.isfinite(float(nodata)):
        finite &= (x_map != float(nodata)) & (y_map != float(nodata))

    idxs = np.where(finite)[0]
    gcps = [
        gdal.GCP(
            float(x_map[i] + dx[i]),
            float(y_map[i] + dy[i]),
            0.0,
            float(col_pan[i]),
            float(row_pan[i]),
        )
        for i in idxs
    ]

    logger.info(
        _fmt_issue(
            "ANCILLARY",
            (
                "PAN scaled-image GCP build: "
                f"total={n_total}, used={len(gcps)}, skipped={n_total - len(gcps)}, "
                f"scale={scale:.6f}, dxdy_source={mode}, crs={crs_wkt_or_epsg or 'n/a'}"
            ),
        )
    )
    return gcps


def _apply_tps_warp_from_gcps(
    input_raster: str,
    output_raster: str,
    gcps,
    target_crs,
    x_res: float,
    y_res: float,
    resampling: str = "near",
    nodata: Optional[float] = None,
    target_aligned_pixels: bool = False,
    target_extent: Optional[Tuple[float, float, float, float]] = None,
) -> Dict[str, Any]:
    """Apply thin-plate-spline warp from GCPs to raster using GDAL."""
    result = {"success": False, "output_path": None, "error_message": None}
    if not gcps:
        result["error_message"] = "No GCPs provided for TPS warp."
        return result

    temp_vrt = None
    try:
        from osgeo import gdal

        gdal.UseExceptions()
        temp_vrt = output_raster.replace(".tif", "_gcps.vrt")
        ds_in = gdal.Open(input_raster, gdal.GA_ReadOnly)
        if ds_in is None:
            result["error_message"] = f"Failed to open input raster: {input_raster}"
            return result

        driver_vrt = gdal.GetDriverByName("VRT")
        ds_vrt = driver_vrt.CreateCopy(temp_vrt, ds_in)
        crs_wkt = target_crs.to_wkt() if hasattr(target_crs, "to_wkt") else str(target_crs)
        ds_vrt.SetGCPs(gcps, crs_wkt)
        ds_vrt.FlushCache()
        ds_vrt = None
        ds_in = None

        gdalwarp_exe = resolve_gdalwarp_exe()
        cmd = _build_gdalwarp_tps_command(
            gdalwarp_exe=gdalwarp_exe,
            temp_vrt=temp_vrt,
            output_raster=output_raster,
            crs_wkt=crs_wkt,
            x_res=x_res,
            y_res=y_res,
            resampling=resampling,
            nodata=nodata,
            target_aligned_pixels=target_aligned_pixels,
            target_extent=target_extent,
        )
        subprocess.run(cmd, check=True, capture_output=True, text=True)

        if os.path.exists(output_raster):
            result["success"] = True
            result["output_path"] = output_raster
        else:
            result["error_message"] = "TPS warp completed but output was not created."
    except subprocess.CalledProcessError as e:
        stderr = (e.stderr or "").strip()
        stdout = (e.stdout or "").strip()
        msg = stderr if stderr else stdout
        result["error_message"] = f"gdalwarp TPS failed: {msg}"
    except Exception as e:
        result["error_message"] = f"TPS warp failed: {e}"
    finally:
        if temp_vrt and os.path.exists(temp_vrt):
            try:
                os.remove(temp_vrt)
            except Exception:
                pass

    return result


def _build_tps_gcps_for_source_raster(
    tie_points_df,
    source_raster_path: str,
    min_gcps: int = 10,
) -> Dict[str, Any]:
    """Build raster-specific TPS GCPs from map-space tiepoint shifts."""
    out = {"success": False, "gcps": [], "n_gcps": 0, "error_message": None}
    if tie_points_df is None or len(tie_points_df) == 0:
        out["error_message"] = "No tie points available."
        return out

    required = ["X_MAP", "Y_MAP", "X_SHIFT_M", "Y_SHIFT_M"]
    missing = [c for c in required if c not in tie_points_df.columns]
    if missing:
        out["error_message"] = f"Missing tiepoint columns for TPS: {missing}"
        return out

    try:
        from osgeo import gdal

        with rasterio.open(source_raster_path) as src:
            inv = ~src.transform
            width = int(src.width)
            height = int(src.height)

        gcps = []
        for _, row in tie_points_df.iterrows():
            try:
                x_map = float(row["X_MAP"])
                y_map = float(row["Y_MAP"])
                dx = float(row["X_SHIFT_M"])
                dy = float(row["Y_SHIFT_M"])
                if not (np.isfinite(x_map) and np.isfinite(y_map) and np.isfinite(dx) and np.isfinite(dy)):
                    continue
                pixel, line = inv * (x_map, y_map)
                if not (np.isfinite(pixel) and np.isfinite(line)):
                    continue
                if pixel < -1 or line < -1 or pixel > (width + 1) or line > (height + 1):
                    continue
                gcps.append(gdal.GCP(x_map + dx, y_map + dy, 0.0, float(pixel), float(line)))
            except Exception:
                continue

        out["gcps"] = gcps
        out["n_gcps"] = len(gcps)
        if len(gcps) >= max(3, int(min_gcps)):
            out["success"] = True
            return out

        out["error_message"] = f"Insufficient TPS GCPs after reprojection: {len(gcps)}"
        return out
    except Exception as e:
        out["error_message"] = f"Failed to build TPS GCPs for source raster: {e}"
        return out


def _validate_ancillary_raster(path: str, target_crs) -> Dict[str, Any]:
    """Lightweight validation of ancillary output raster georeference and content."""
    out = {"ok": False, "error": None}
    try:
        if not os.path.exists(path):
            out["error"] = f"Output missing: {path}"
            return out
        with rasterio.open(path) as src:
            if src.count < 1 or src.width < 1 or src.height < 1:
                out["error"] = "Invalid raster dimensions/count."
                return out
            if src.crs is None:
                out["error"] = "Output CRS is missing."
                return out
            if target_crs is not None and src.crs != target_crs:
                out["error"] = f"CRS mismatch (expected {target_crs}, got {src.crs})."
                return out
            sample = src.read(1, masked=True)
            valid_count = int(np.count_nonzero(~sample.mask)) if np.ma.isMaskedArray(sample) else sample.size
            if valid_count <= 0:
                out["error"] = "No valid pixels found in output."
                return out
        out["ok"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _estimate_translation_phasecorr(
    reference_raster_path: str,
    candidate_raster_path: str,
    max_dim: int = 1024,
    reference_band: int = 1,
    candidate_band: int = 1,
    reference_nodata: float = PROCESSING_NODATA,
    candidate_nodata: float = PROCESSING_NODATA,
) -> Dict[str, Any]:
    """Estimate global translation between two rasters via phase correlation."""
    out: Dict[str, Any] = {
        "ok": False,
        "shift_x_px": None,
        "shift_y_px": None,
        "shift_magnitude_px": None,
        "valid_pixels": 0,
        "error": None,
    }
    try:
        with rasterio.open(candidate_raster_path) as cand_ds:
            cand_band = int(max(1, min(int(candidate_band), int(cand_ds.count))))
            cw = int(cand_ds.width)
            ch = int(cand_ds.height)
            longest = max(cw, ch)
            tgt = int(max(64, int(max_dim)))
            if longest > tgt:
                scale = float(tgt) / float(longest)
                out_w = max(64, int(round(cw * scale)))
                out_h = max(64, int(round(ch * scale)))
            else:
                out_w = cw
                out_h = ch
            cand = cand_ds.read(
                cand_band,
                out_shape=(out_h, out_w),
                resampling=Resampling.bilinear,
            ).astype(np.float32, copy=False)
            nodata_c = cand_ds.nodata
            if nodata_c is None or not np.isfinite(float(nodata_c)):
                nodata_c = float(candidate_nodata)
            else:
                nodata_c = float(nodata_c)

        with rasterio.open(reference_raster_path) as ref_ds:
            ref_band = int(max(1, min(int(reference_band), int(ref_ds.count))))
            ref = ref_ds.read(
                ref_band,
                out_shape=(cand.shape[0], cand.shape[1]),
                resampling=Resampling.bilinear,
            ).astype(np.float32, copy=False)
            nodata_r = ref_ds.nodata
            if nodata_r is None or not np.isfinite(float(nodata_r)):
                nodata_r = float(reference_nodata)
            else:
                nodata_r = float(nodata_r)

        valid = np.isfinite(ref) & np.isfinite(cand)
        if np.isfinite(nodata_r):
            valid &= (ref != nodata_r)
        if np.isfinite(nodata_c):
            valid &= (cand != nodata_c)

        valid_count = int(np.count_nonzero(valid))
        out["valid_pixels"] = valid_count
        if valid_count < 256:
            out["error"] = "Insufficient overlapping valid pixels for phase-correlation residual check."
            return out

        ref_work = ref.copy()
        cand_work = cand.copy()
        ref_fill = float(np.median(ref_work[valid]))
        cand_fill = float(np.median(cand_work[valid]))
        ref_work[~valid] = ref_fill
        cand_work[~valid] = cand_fill
        ref_work -= float(np.mean(ref_work))
        cand_work -= float(np.mean(cand_work))

        h, w = ref_work.shape
        win_y = np.hanning(h).astype(np.float32)
        win_x = np.hanning(w).astype(np.float32)
        window = np.outer(win_y, win_x)
        ref_work *= window
        cand_work *= window

        f_ref = np.fft.fft2(ref_work)
        f_cand = np.fft.fft2(cand_work)
        cross = f_ref * np.conj(f_cand)
        denom = np.abs(cross)
        denom[denom == 0] = 1.0
        corr = np.fft.ifft2(cross / denom)
        corr_abs = np.abs(corr)
        peak_idx = np.unravel_index(np.argmax(corr_abs), corr_abs.shape)
        py = int(peak_idx[0])
        px = int(peak_idx[1])

        shift_y = py if py <= (h // 2) else py - h
        shift_x = px if px <= (w // 2) else px - w
        shift_mag = float(np.hypot(float(shift_x), float(shift_y)))

        out["ok"] = True
        out["shift_x_px"] = float(shift_x)
        out["shift_y_px"] = float(shift_y)
        out["shift_magnitude_px"] = shift_mag
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _run_postwarp_phasecorr_qa(
    enabled: bool,
    reference_raster_path: Optional[str],
    candidate_raster_path: Optional[str],
    reference_band: int,
    candidate_band: int,
    warn_threshold_px: float,
    reject_threshold_px: float,
    reject_bad: bool,
    max_dim: int,
    estimator: Optional[Callable[..., Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """Run optional post-warp phase-correlation QA and return warning/reject flags."""
    result: Dict[str, Any] = {
        "enabled": bool(enabled),
        "ok": False,
        "warning": False,
        "reject": False,
        "shift_magnitude_px": None,
        "error": None,
    }
    if not bool(enabled):
        return result
    if not reference_raster_path or not candidate_raster_path:
        result["error"] = "missing reference or candidate raster path"
        return result
    if not os.path.exists(reference_raster_path) or not os.path.exists(candidate_raster_path):
        result["error"] = "reference or candidate raster does not exist"
        return result

    fn = estimator or _estimate_translation_phasecorr
    qa = fn(
        reference_raster_path=reference_raster_path,
        candidate_raster_path=candidate_raster_path,
        max_dim=int(max(128, int(max_dim))),
        reference_band=int(max(1, int(reference_band))),
        candidate_band=int(max(1, int(candidate_band))),
        reference_nodata=0.0,
        candidate_nodata=PROCESSING_NODATA,
    )
    if isinstance(qa, dict):
        result.update(qa)
    shift = float(result.get("shift_magnitude_px", 0.0) or 0.0)
    result["warning"] = bool(result.get("ok", False) and shift > float(max(0.0, warn_threshold_px)))
    reject_limit = max(float(max(0.0, warn_threshold_px)), float(max(0.0, reject_threshold_px)))
    result["reject"] = bool(result.get("ok", False) and bool(reject_bad) and shift > reject_limit)
    return result


def _sanitize_raster_nonfinite_inplace(path: str, nodata: float = PROCESSING_NODATA) -> Dict[str, Any]:
    """Replace non-finite values in a raster with nodata in-place."""
    result = {
        "ok": False,
        "error": None,
        "nonfinite_replaced": 0,
        "bands": 0,
    }
    try:
        if not os.path.exists(path):
            result["error"] = f"Raster not found: {path}"
            return result
        with rasterio.open(path, "r+") as dst:
            result["bands"] = int(dst.count)
            if dst.count < 1:
                result["error"] = "Raster has no bands."
                return result

            src_nodata = dst.nodata
            write_nodata = float(src_nodata) if src_nodata is not None else float(nodata)
            if src_nodata is None or not np.isfinite(float(src_nodata)):
                dst.nodata = float(write_nodata)

            replaced_total = 0
            for bidx in range(1, dst.count + 1):
                for _, window in dst.block_windows(bidx):
                    band = dst.read(bidx, window=window)
                    bad = ~np.isfinite(band)
                    if not np.any(bad):
                        continue
                    band = band.astype(np.float32, copy=False)
                    band[bad] = write_nodata
                    dst.write(band, bidx, window=window)
                    replaced_total += int(np.count_nonzero(bad))

        result["ok"] = True
        result["nonfinite_replaced"] = replaced_total
        return result
    except Exception as e:
        result["error"] = str(e)
        return result


def _resolve_enmap_band_selection(
    raster_band_count: int,
    wl,
    fwhm,
    band_names,
    band_detectors,
    sort_idx: Optional[np.ndarray],
) -> Dict[str, Any]:
    """
    Resolve a safe EnMAP band mapping against an actual raster band count.

    Returns a mapping that can be used to copy/reorder bands without loading
    the full cube in memory. If metadata and raster band counts disagree,
    wavelength metadata is trimmed as needed to keep the pipeline consistent.
    """
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "warnings": [],
        "source_bands_1based": [],
        "wl": None,
        "fwhm": None,
        "band_names": None,
        "band_detectors": None,
        "sort_idx": None,
    }

    try:
        src_count = int(raster_band_count)
    except Exception:
        out["error"] = f"Invalid EnMAP raster band count: {raster_band_count}"
        return out
    if src_count < 1:
        out["error"] = f"EnMAP raster has no bands (count={src_count})."
        return out

    wl_arr = np.asarray(wl, dtype=float).reshape(-1)
    if wl_arr.size < 1:
        out["error"] = "Missing EnMAP wavelength metadata."
        return out

    fwhm_arr = np.asarray(fwhm, dtype=float).reshape(-1) if fwhm is not None else None
    names_list = [str(x) for x in band_names] if band_names is not None else None
    detectors_list = [str(x).upper() if x is not None else "" for x in band_detectors] if band_detectors is not None else None

    def _apply_keep_mask(mask: np.ndarray) -> None:
        nonlocal wl_arr, fwhm_arr, names_list, detectors_list
        keep = np.asarray(mask, dtype=bool).reshape(-1)
        if keep.size != wl_arr.size:
            return
        wl_arr = wl_arr[keep]
        if fwhm_arr is not None and fwhm_arr.size == keep.size:
            fwhm_arr = fwhm_arr[keep]
        if names_list is not None and len(names_list) == keep.size:
            names_list = [names_list[i] for i in range(len(names_list)) if keep[i]]
        if detectors_list is not None and len(detectors_list) == keep.size:
            detectors_list = [detectors_list[i] for i in range(len(detectors_list)) if keep[i]]

    if sort_idx is not None:
        idx_raw = np.asarray(sort_idx).reshape(-1)
        if idx_raw.size < 1:
            out["error"] = "EnMAP band ordering is empty."
            return out

        if idx_raw.size != wl_arr.size:
            n_keep = min(int(idx_raw.size), int(wl_arr.size))
            out["warnings"].append(
                f"EnMAP band-order length mismatch ({int(idx_raw.size)} vs metadata {int(wl_arr.size)}); "
                f"truncating to {n_keep}."
            )
            idx_raw = idx_raw[:n_keep]
            _apply_keep_mask(np.arange(int(wl_arr.size)) < n_keep)

        if np.issubdtype(idx_raw.dtype, np.floating):
            finite_int = np.isfinite(idx_raw) & (np.floor(idx_raw) == idx_raw)
            if not np.all(finite_int):
                dropped = int(np.count_nonzero(~finite_int))
                out["warnings"].append(
                    f"EnMAP band-order contains {dropped} non-integer/non-finite entries; dropping them."
                )
                _apply_keep_mask(finite_int)
                idx_raw = idx_raw[finite_int]

        idx = np.asarray(idx_raw, dtype=np.int64).reshape(-1)
        valid = (idx >= 0) & (idx < src_count)
        if not np.all(valid):
            dropped = int(np.count_nonzero(~valid))
            out["warnings"].append(
                f"EnMAP band-order references {dropped} bands outside [0, {src_count - 1}]; dropping them."
            )
            _apply_keep_mask(valid)
            idx = idx[valid]

        if idx.size < 1 or wl_arr.size < 1:
            out["error"] = "No usable EnMAP bands remain after harmonizing metadata with raster bands."
            return out

        out["sort_idx"] = idx
        out["source_bands_1based"] = [int(i) + 1 for i in idx.tolist()]
    else:
        meta_count = int(wl_arr.size)
        if src_count < meta_count:
            out["warnings"].append(
                f"EnMAP raster has fewer bands than metadata ({src_count} vs {meta_count}); "
                "truncating spectral metadata."
            )
            _apply_keep_mask(np.arange(meta_count) < src_count)
            meta_count = int(wl_arr.size)
        if meta_count < 1:
            out["error"] = "No usable EnMAP spectral metadata after harmonization."
            return out
        if src_count > meta_count:
            out["warnings"].append(
                f"EnMAP raster has extra bands not described by metadata ({src_count} vs {meta_count}); "
                "ignoring trailing bands."
            )
        keep = min(src_count, meta_count)
        out["source_bands_1based"] = list(range(1, int(keep) + 1))

    out["wl"] = wl_arr
    out["fwhm"] = fwhm_arr
    out["band_names"] = names_list
    out["band_detectors"] = detectors_list
    out["ok"] = True
    return out


def _stream_copy_raster_with_band_order(
    source_path: str,
    output_path: str,
    source_bands_1based: Sequence[int],
    out_dtype: Any = PROCESSING_DTYPE,
    compress: Optional[str] = "NONE",
    tile_size: int = 512,
) -> Dict[str, Any]:
    """Copy and optionally reorder raster bands in a streaming, windowed pass."""
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "source_bands": None,
        "output_bands": 0,
        "window_strategy": None,
        "window_count": 0,
    }
    try:
        src_bands = [int(b) for b in source_bands_1based]
    except Exception:
        out["error"] = f"Invalid band-order specification: {source_bands_1based}"
        return out
    if not src_bands:
        out["error"] = "Band-order specification is empty."
        return out

    try:
        with rasterio.open(source_path) as src:
            out["source_bands"] = int(src.count)
            invalid = [b for b in src_bands if b < 1 or b > int(src.count)]
            if invalid:
                out["error"] = (
                    f"Band-order contains out-of-range indices for source count {int(src.count)}: {invalid[:5]}"
                )
                return out

            out_dtype_name = np.dtype(out_dtype).name
            nodata_value = src.nodata
            if nodata_value is None or not np.isfinite(float(nodata_value)):
                nodata_value = float(PROCESSING_NODATA)
            else:
                nodata_value = float(nodata_value)

            profile = src.profile.copy()
            profile.update(
                driver="GTiff",
                count=len(src_bands),
                dtype=out_dtype_name,
                nodata=nodata_value,
                tiled=True,
                interleave="pixel",
                BIGTIFF="YES",
            )
            # ENVI/VRT source profiles may expose non-GTIFF block sizes (e.g., 4x1)
            # that are invalid for tiled GTiff creation.
            profile.pop("blockxsize", None)
            profile.pop("blockysize", None)
            if compress is not None:
                comp = str(compress).strip().upper()
                profile["compress"] = comp
                if comp == "NONE":
                    # Compression predictors are invalid without compression.
                    profile.pop("predictor", None)

            read_indexes = tuple(src_bands)

            windows_iter = None
            window_strategy = "fixed_tile"
            try:
                bw_iter = src.block_windows(1)
                first_item = next(bw_iter, None)
                if first_item is not None:
                    window_strategy = "block_windows"
                    first_window = first_item[1]

                    def _iter_block_windows():
                        yield first_window
                        for _, win in bw_iter:
                            yield win

                    windows_iter = _iter_block_windows()
            except Exception:
                windows_iter = None
            if windows_iter is None:
                step = int(max(64, tile_size))

                def _iter_fixed_windows():
                    for row_off in range(0, int(src.height), step):
                        win_h = min(step, int(src.height) - row_off)
                        for col_off in range(0, int(src.width), step):
                            win_w = min(step, int(src.width) - col_off)
                            yield Window(
                                col_off=int(col_off),
                                row_off=int(row_off),
                                width=int(win_w),
                                height=int(win_h),
                            )

                windows_iter = _iter_fixed_windows()
            out["window_strategy"] = window_strategy

            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            with rasterio.open(output_path, "w", **profile) as dst:
                src_tags = src.tags()
                if src_tags:
                    dst.update_tags(**src_tags)

                for out_bidx, src_bidx in enumerate(src_bands, start=1):
                    desc = src.descriptions[src_bidx - 1]
                    if desc:
                        dst.set_band_description(out_bidx, desc)
                    band_tags = src.tags(src_bidx)
                    if band_tags:
                        dst.update_tags(out_bidx, **band_tags)

                window_count = 0
                for window in windows_iter:
                    window_count += 1
                    window_txt = (
                        f"row_off={int(window.row_off)}, col_off={int(window.col_off)}, "
                        f"height={int(window.height)}, width={int(window.width)}"
                    )
                    try:
                        block = src.read(indexes=read_indexes, window=window, out_dtype=out_dtype_name)
                    except (RasterioIOError, OSError) as io_err:
                        raise RuntimeError(
                            "Window read failed "
                            f"(source={source_path}, output={output_path}, source_driver={src.driver}, "
                            f"output_driver={profile.get('driver')}, window={window_txt}): {io_err}"
                        ) from io_err

                    if np.issubdtype(block.dtype, np.floating):
                        bad = ~np.isfinite(block)
                        if np.any(bad):
                            block = block.copy()
                            block[bad] = nodata_value

                    try:
                        dst.write(block, window=window)
                    except (RasterioIOError, OSError) as io_err:
                        raise RuntimeError(
                            "Window write failed "
                            f"(source={source_path}, output={output_path}, source_driver={src.driver}, "
                            f"output_driver={profile.get('driver')}, window={window_txt}): {io_err}"
                        ) from io_err
                out["window_count"] = int(window_count)

        out["output_bands"] = int(len(src_bands))
        out["ok"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _prepare_enmap_processing_source(
    source_path: str,
    output_path: str,
    metadata: Optional[Dict[str, Any]] = None,
    inject_metadata: bool = True,
) -> Dict[str, Any]:
    """Prepare EnMAP source raster for processing, including BSQ-to-GTIFF conversion."""
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "path": source_path,
        "converted": False,
        "metadata_injection_attempted": False,
        "metadata_injected": False,
        "metadata_injection_error": None,
        "metadata_readback": None,
    }
    try:
        if not source_path or not os.path.exists(source_path):
            out["error"] = f"EnMAP source raster not found: {source_path}"
            return out

        src_suffix = Path(source_path).suffix.lower()
        prepared_path = str(source_path)
        if src_suffix == ".bsq":
            with rasterio.open(source_path) as src:
                src_band_count = int(src.count)
            src_bands_1based = list(range(1, src_band_count + 1))
            copy_res = _stream_copy_raster_with_band_order(
                source_path=source_path,
                output_path=output_path,
                source_bands_1based=src_bands_1based,
                out_dtype=PROCESSING_DTYPE,
                compress="LZW",
                tile_size=512,
            )
            if not copy_res.get("ok", False):
                out["error"] = (
                    "Failed BSQ-to-GeoTIFF preparation copy: "
                    f"{copy_res.get('error', 'unknown error')}"
                )
                return out
            prepared_path = str(output_path)
            out["converted"] = True

        meta = metadata if isinstance(metadata, dict) else {}
        with rasterio.open(prepared_path, "r+") as dst:
            meta_crs = meta.get("crs")
            if dst.crs is None and meta_crs is not None and str(meta_crs).strip():
                try:
                    dst.crs = CRS.from_user_input(str(meta_crs).strip())
                except Exception as crs_exc:
                    logger.warning(
                        _fmt_issue(
                            "METADATA",
                            f"Could not apply metadata CRS '{meta_crs}' to EnMAP prep raster: {crs_exc}",
                        )
                    )

            transform_vals = np.asarray(tuple(dst.transform), dtype=float)
            transform_identity_like = (
                not np.all(np.isfinite(transform_vals))
                or _transforms_equivalent(dst.transform, Affine.identity())
                or _transforms_equivalent(dst.transform, Affine(1.0, 0.0, 0.0, 0.0, -1.0, 0.0))
            )
            bbox_val = meta.get("bbox")
            if transform_identity_like and isinstance(bbox_val, (list, tuple)) and len(bbox_val) == 4:
                try:
                    west, south, east, north = [float(v) for v in bbox_val]
                    if (
                        np.all(np.isfinite([west, south, east, north]))
                        and east > west
                        and north > south
                        and int(dst.width) > 0
                        and int(dst.height) > 0
                    ):
                        dst.transform = transform_from_bounds(
                            west,
                            south,
                            east,
                            north,
                            int(dst.width),
                            int(dst.height),
                        )
                except Exception as tx_exc:
                    logger.warning(
                        _fmt_issue(
                            "METADATA",
                            f"Could not apply metadata transform to EnMAP prep raster: {tx_exc}",
                        )
                    )

        if inject_metadata and meta:
            out["metadata_injection_attempted"] = True
            inject_res = inject_metadata_into_raster(prepared_path, meta)
            if inject_res.get("ok", False):
                out["metadata_injected"] = True
                out["metadata_readback"] = read_enmap_metadata_from_raster(prepared_path)
            else:
                out["metadata_injection_error"] = inject_res.get("error")

        out["path"] = prepared_path
        out["ok"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _build_vrt_with_band_order(
    source_path: str,
    output_path: str,
    source_bands_1based: Sequence[int],
) -> Dict[str, Any]:
    """Build a VRT that remaps band order without materializing a new GeoTIFF."""
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "source_bands": None,
        "output_bands": 0,
        "driver": "VRT",
    }
    try:
        src_bands = [int(b) for b in source_bands_1based]
    except Exception:
        out["error"] = f"Invalid band-order specification: {source_bands_1based}"
        return out
    if not src_bands:
        out["error"] = "Band-order specification is empty."
        return out

    try:
        with rasterio.open(source_path) as src:
            out["source_bands"] = int(src.count)
            invalid = [b for b in src_bands if b < 1 or b > int(src.count)]
            if invalid:
                out["error"] = (
                    f"Band-order contains out-of-range indices for source count {int(src.count)}: {invalid[:5]}"
                )
                return out
    except Exception as e:
        out["error"] = f"Failed to open source raster: {e}"
        return out

    try:
        from osgeo import gdal

        gdal.UseExceptions()
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        ds = gdal.Open(str(source_path), gdal.GA_ReadOnly)
        if ds is None:
            out["error"] = f"Failed to open source raster with GDAL: {source_path}"
            return out
        try:
            translate_opts = gdal.TranslateOptions(format="VRT", bandList=src_bands)
            vrt_ds = gdal.Translate(str(output_path), ds, options=translate_opts)
            if vrt_ds is None:
                out["error"] = "GDAL Translate returned no dataset."
                return out
            vrt_ds.FlushCache()
            vrt_ds = None
        finally:
            ds = None

        with rasterio.open(output_path) as vrt_src:
            out_count = int(vrt_src.count)
            out["output_bands"] = out_count
            if out_count != len(src_bands):
                out["error"] = (
                    f"VRT output band count mismatch ({out_count} vs expected {len(src_bands)})."
                )
                return out

        out["ok"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _transforms_equivalent(lhs: Any, rhs: Any, tol: float = 1e-9) -> bool:
    """Return True when two affine transforms are numerically equivalent."""
    try:
        lvals = tuple(float(v) for v in lhs)
        rvals = tuple(float(v) for v in rhs)
        if len(lvals) != len(rvals):
            return False
        return bool(np.allclose(np.asarray(lvals), np.asarray(rvals), atol=float(tol), rtol=0.0))
    except Exception:
        return str(lhs) == str(rhs)


def _build_detector_branch_plan(
    wl,
    fwhm,
    band_names,
    band_detectors,
    sensor_type: str,
) -> Dict[str, Any]:
    """
    Build VNIR/SWIR branch index maps and reordered metadata.

    Overlap bands are intentionally preserved: final order is always VNIR indices
    followed by SWIR indices, without wavelength deduplication.
    """
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "warnings": [],
        "vnir_idx_0based": [],
        "swir_idx_0based": [],
        "final_order_idx_0based": [],
        "final_wl": None,
        "final_fwhm": None,
        "final_band_names": None,
        "final_band_detectors": None,
        "branches": {},
    }
    try:
        wl_arr = np.asarray(wl, dtype=float).reshape(-1)
    except Exception:
        out["error"] = "Invalid wavelength metadata for detector split planning."
        return out

    n_bands = int(wl_arr.size)
    if n_bands < 1:
        out["error"] = "No spectral bands available for detector split planning."
        return out

    if not np.all(np.isfinite(wl_arr)):
        out["error"] = "Non-finite wavelengths encountered in detector split planning."
        return out

    if fwhm is not None:
        fwhm_arr = np.asarray(fwhm, dtype=float).reshape(-1)
        if int(fwhm_arr.size) != n_bands:
            out["warnings"].append(
                f"FWHM count mismatch in detector split planning ({int(fwhm_arr.size)} vs {n_bands}); dropping FWHM."
            )
            fwhm_arr = None
    else:
        fwhm_arr = None

    if band_names is not None:
        names = [str(x) for x in list(band_names)]
        if len(names) != n_bands:
            if len(names) > n_bands:
                out["warnings"].append(
                    f"Band-name count mismatch in detector split planning ({len(names)} vs {n_bands}); truncating."
                )
                names = names[:n_bands]
            else:
                out["warnings"].append(
                    f"Band-name count mismatch in detector split planning ({len(names)} vs {n_bands}); padding."
                )
                sensor_tag = str(sensor_type).upper()
                names = names + [f"{sensor_tag}_{i + 1:03d}" for i in range(len(names), n_bands)]
    else:
        sensor_tag = str(sensor_type).upper()
        names = [f"{sensor_tag}_{i + 1:03d}" for i in range(n_bands)]

    detectors: List[str] = []
    det_in = list(band_detectors) if band_detectors is not None else []
    for i in range(n_bands):
        raw = str(det_in[i]).upper().strip() if i < len(det_in) and det_in[i] is not None else ""
        if raw not in {"VNIR", "SWIR"}:
            raw = _infer_detector_fallback(float(wl_arr[i]), names[i], str(sensor_type))
        detectors.append(raw)

    vnir_idx = [i for i, d in enumerate(detectors) if str(d).upper() == "VNIR"]
    swir_idx = [i for i, d in enumerate(detectors) if str(d).upper() == "SWIR"]
    if not vnir_idx or not swir_idx:
        out["error"] = (
            "Detector split planning requires both VNIR and SWIR groups; "
            f"found VNIR={len(vnir_idx)}, SWIR={len(swir_idx)}."
        )
        return out

    final_order_idx = [int(i) for i in (vnir_idx + swir_idx)]
    final_wl = wl_arr[final_order_idx]
    final_fwhm = fwhm_arr[final_order_idx] if fwhm_arr is not None else None
    final_names = [names[i] for i in final_order_idx]
    final_detectors = [detectors[i] for i in final_order_idx]

    out["vnir_idx_0based"] = [int(i) for i in vnir_idx]
    out["swir_idx_0based"] = [int(i) for i in swir_idx]
    out["final_order_idx_0based"] = list(final_order_idx)
    out["final_wl"] = final_wl
    out["final_fwhm"] = final_fwhm
    out["final_band_names"] = final_names
    out["final_band_detectors"] = final_detectors
    out["branches"] = {
        "VNIR": {
            "indices_0based": [int(i) for i in vnir_idx],
            "indices_1based": [int(i) + 1 for i in vnir_idx],
            "wl": wl_arr[np.asarray(vnir_idx, dtype=int)],
            "fwhm": (fwhm_arr[np.asarray(vnir_idx, dtype=int)] if fwhm_arr is not None else None),
            "band_names": [names[i] for i in vnir_idx],
            "band_detectors": [detectors[i] for i in vnir_idx],
            "s2_subset": ("B02", "B03", "B04", "B08"),
            "global_s2_band_label": "B08",
            "global_s2_stack_idx": int(MULTIBAND_S2_WAVELENGTHS["B08"]["stack_idx"]),
            "global_target_wl_nm": float(MULTIBAND_S2_WAVELENGTHS["B08"]["wavelength"]),
        },
        "SWIR": {
            "indices_0based": [int(i) for i in swir_idx],
            "indices_1based": [int(i) + 1 for i in swir_idx],
            "wl": wl_arr[np.asarray(swir_idx, dtype=int)],
            "fwhm": (fwhm_arr[np.asarray(swir_idx, dtype=int)] if fwhm_arr is not None else None),
            "band_names": [names[i] for i in swir_idx],
            "band_detectors": [detectors[i] for i in swir_idx],
            "s2_subset": ("B11", "B12"),
            "global_s2_band_label": "B11",
            "global_s2_stack_idx": int(MULTIBAND_S2_WAVELENGTHS["B11"]["stack_idx"]),
            "global_target_wl_nm": float(MULTIBAND_S2_WAVELENGTHS["B11"]["wavelength"]),
        },
    }
    out["ok"] = True
    return out


def _write_branch_raster_windowed(
    source_path: str,
    output_path: str,
    source_bands_1based: Sequence[int],
    out_dtype: Any = PROCESSING_DTYPE,
) -> Dict[str, Any]:
    """
    Write detector branch raster using strict windowed IO.

    This is a thin wrapper around streaming copy to guarantee split operations
    do not load full cubes in memory.
    """
    return _stream_copy_raster_with_band_order(
        source_path=source_path,
        output_path=output_path,
        source_bands_1based=source_bands_1based,
        out_dtype=out_dtype,
        compress="LZW",
        tile_size=512,
    )


def _recombine_detector_branches_windowed(
    vnir_raster_path: str,
    swir_raster_path: str,
    output_raster_path: str,
    nodata: float = PROCESSING_NODATA,
    out_dtype: Any = PROCESSING_DTYPE,
) -> Dict[str, Any]:
    """
    Recombine VNIR/SWIR warped branches with strict windowed IO.

    Critical spectral-integrity rule:
    - If a pixel is invalid (nodata/non-finite) in either branch, it is forced
      to nodata across *all* output bands (union nodata mask).
    """
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "output_path": output_raster_path,
        "vnir_bands": 0,
        "swir_bands": 0,
        "output_bands": 0,
        "window_count": 0,
        "union_nodata_pixels": 0,
    }
    try:
        if not os.path.exists(vnir_raster_path):
            out["error"] = f"VNIR branch raster missing: {vnir_raster_path}"
            return out
        if not os.path.exists(swir_raster_path):
            out["error"] = f"SWIR branch raster missing: {swir_raster_path}"
            return out

        with rasterio.open(vnir_raster_path) as vsrc, rasterio.open(swir_raster_path) as ssrc:
            if int(vsrc.count) < 1 or int(ssrc.count) < 1:
                out["error"] = "VNIR/SWIR branch rasters must each contain at least one band."
                return out
            if int(vsrc.width) != int(ssrc.width) or int(vsrc.height) != int(ssrc.height):
                out["error"] = (
                    f"Branch grid size mismatch: VNIR={vsrc.width}x{vsrc.height}, "
                    f"SWIR={ssrc.width}x{ssrc.height}."
                )
                return out
            if not _crs_equivalent(vsrc.crs, ssrc.crs):
                out["error"] = f"Branch CRS mismatch: VNIR={vsrc.crs}, SWIR={ssrc.crs}."
                return out
            if not _transforms_equivalent(vsrc.transform, ssrc.transform):
                out["error"] = "Branch affine transforms differ; cannot safely recombine."
                return out

            out_dtype_name = np.dtype(out_dtype).name
            nodata_value = float(nodata)
            profile = vsrc.profile.copy()
            profile.update(
                count=int(vsrc.count + ssrc.count),
                dtype=out_dtype_name,
                nodata=nodata_value,
                compress="LZW",
                tiled=True,
                BIGTIFF="YES",
            )

            os.makedirs(os.path.dirname(output_raster_path) or ".", exist_ok=True)
            with rasterio.open(output_raster_path, "w", **profile) as dst:
                src_tags = vsrc.tags()
                if src_tags:
                    dst.update_tags(**src_tags)

                v_nodata = vsrc.nodata
                s_nodata = ssrc.nodata
                if v_nodata is None or not np.isfinite(float(v_nodata)):
                    v_nodata = nodata_value
                else:
                    v_nodata = float(v_nodata)
                if s_nodata is None or not np.isfinite(float(s_nodata)):
                    s_nodata = nodata_value
                else:
                    s_nodata = float(s_nodata)

                union_count = 0
                window_count = 0
                for _, window in vsrc.block_windows(1):
                    window_count += 1
                    v_block = vsrc.read(window=window, out_dtype=out_dtype_name)
                    s_block = ssrc.read(window=window, out_dtype=out_dtype_name)

                    # Branch-valid requires all branch bands finite and non-nodata.
                    v_valid = np.all(np.isfinite(v_block), axis=0)
                    s_valid = np.all(np.isfinite(s_block), axis=0)
                    if np.isfinite(v_nodata):
                        v_valid &= np.all(v_block != v_nodata, axis=0)
                    if np.isfinite(s_nodata):
                        s_valid &= np.all(s_block != s_nodata, axis=0)

                    invalid_union = (~v_valid) | (~s_valid)
                    if np.any(invalid_union):
                        union_count += int(np.count_nonzero(invalid_union))

                    out_block = np.concatenate([v_block, s_block], axis=0)
                    out_block[:, invalid_union] = nodata_value
                    dst.write(out_block, window=window)

                out["window_count"] = int(window_count)
                out["union_nodata_pixels"] = int(union_count)
                out["vnir_bands"] = int(vsrc.count)
                out["swir_bands"] = int(ssrc.count)
                out["output_bands"] = int(vsrc.count + ssrc.count)

        out["ok"] = True
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _process_detector_branch_candidate(
    *,
    candidate_idx: int,
    scene_name: str,
    sensor_tag: str,
    date_tag: str,
    detector_plan: Dict[str, Any],
    source_raster_path: str,
    s2_raster_path: str,
    s2_crs,
    hyp_type: str,
    folder_struct: Dict[str, str],
    matcher_profile: Dict[str, Any],
    prefer_fixed_band_pairs: bool,
    fixed_band_pairs_by_sensor: Dict[str, Any],
    bandpair_wavelength_window_nm: float,
    min_tie_points: int,
    min_accuracy: float,
    max_displacement: float,
    residual_mad_factor: float,
    min_band_support: int,
    allow_single_band_fallback: bool,
    consensus_group_rounding_px: float,
    spatial_grid_rows: int,
    spatial_grid_cols: int,
    max_points_per_cell: int,
    preferred_polynomial_order: int,
    auto_downgrade_polynomial_order: bool,
    min_gcps_order2: int,
    min_cells_order2: int,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]],
    scene_idx: int,
    scene_total: int,
    progress_heartbeat_interval_s: float,
    postwarp_phasecorr_check: bool,
    postwarp_phasecorr_warn_threshold_px: float,
    postwarp_phasecorr_reject_threshold_px: float,
    postwarp_phasecorr_reject_bad: bool,
    postwarp_phasecorr_max_dim: int,
    s2_ref_stack_idx: int,
    s2_ref_wavelength_nm: float,
    config: Dict[str, Any],
) -> Dict[str, Any]:
    """Process one candidate by splitting VNIR/SWIR branches and recombining with nodata union."""
    import pandas as pd

    out: Dict[str, Any] = {
        "validation": {
            "is_valid": False,
            "confidence": 0.0,
            "shift_m": 0.0,
            "message": "branch processing not executed",
            "branch_validations": {},
            "ssim_before": None,
            "ssim_after": None,
            "ssim_delta": None,
        },
        "final_quality_pass": False,
        "local_valid_tp_count": 0,
        "tp_residuals": {"n_tiepoints_used": 0},
        "polynomial_warp_used": False,
        "tps_warp_used": False,
        "polynomial_order_decision": {},
        "polynomial_order_used": int(preferred_polynomial_order),
        "local_tiepoints_df": None,
        "local_tiepoints_visualization_df": None,
        "output_path": None,
        "output_content_valid": False,
        "output_validation": {},
        "postwarp_phasecorr_qa": {},
        "merged_tiepoint_stages": {},
        "multiband_tiepoint_counts": {},
        "branch_results": {},
        "recombine_result": {"ok": False, "error": "not run"},
        "temp_global_path": {},
        "temp_local_path": {},
        "pre_coreg_source_bands_1based": [],
        "output_wl": None,
        "output_fwhm": None,
        "output_band_names": [],
        "output_band_detectors": [],
    }

    plan_branches = detector_plan.get("branches", {})
    for branch_name in ("VNIR", "SWIR"):
        if branch_name not in plan_branches:
            raise RuntimeError(_fmt_issue("HS_PREP", f"Detector branch plan missing '{branch_name}' branch."))
    final_order_idx = [int(i) for i in detector_plan.get("final_order_idx_0based", [])]
    if not final_order_idx:
        raise RuntimeError(_fmt_issue("HS_PREP", "Detector branch plan returned empty final_order_idx_0based."))
    out["pre_coreg_source_bands_1based"] = [int(i) + 1 for i in final_order_idx]
    out["output_wl"] = np.asarray(detector_plan.get("final_wl"), dtype=float).reshape(-1)
    if detector_plan.get("final_fwhm") is not None:
        out["output_fwhm"] = np.asarray(detector_plan.get("final_fwhm"), dtype=float).reshape(-1)
    out["output_band_names"] = list(detector_plan.get("final_band_names") or [])
    out["output_band_detectors"] = list(detector_plan.get("final_band_detectors") or [])

    with rasterio.open(s2_raster_path) as s2_src:
        s2_bounds = s2_src.bounds
        s2_extent = (
            float(s2_bounds.left),
            float(s2_bounds.bottom),
            float(s2_bounds.right),
            float(s2_bounds.top),
        )

    _emit_progress(
        progress_callback,
        "Global/local coregistration by detector branch",
        scene_idx=scene_idx,
        scene_total=scene_total,
    )
    log_section_header("DETECTOR-BRANCH COREGISTRATION (VNIR + SWIR)")

    branch_results: Dict[str, Dict[str, Any]] = {}
    branch_multiband_counts: Dict[str, int] = {}
    branch_stage_counts: Dict[str, Dict[str, Any]] = {}
    branch_order_decisions: Dict[str, Dict[str, Any]] = {}
    branch_paths_for_recombine: Dict[str, Optional[str]] = {"VNIR": None, "SWIR": None}
    branch_quality_flags: List[bool] = []

    for branch_name in ("VNIR", "SWIR"):
        branch_cfg = dict(plan_branches.get(branch_name, {}))
        branch_indices = [int(v) for v in branch_cfg.get("indices_1based", [])]
        if not branch_indices:
            raise RuntimeError(_fmt_issue("HS_PREP", f"Candidate {candidate_idx + 1}: no {branch_name} source bands."))

        branch_input = os.path.join(
            folder_struct["temp"], f"{scene_name}_{branch_name}_SRC_c{candidate_idx}.tif"
        )
        split_res = _write_branch_raster_windowed(
            source_path=source_raster_path,
            output_path=branch_input,
            source_bands_1based=branch_indices,
            out_dtype=PROCESSING_DTYPE,
        )
        if not split_res.get("ok", False):
            raise RuntimeError(
                _fmt_issue(
                    "HS_PREP",
                    f"Candidate {candidate_idx + 1}: failed {branch_name} split: "
                    f"{split_res.get('error', 'unknown error')}",
                )
            )

        branch_wl = np.asarray(branch_cfg.get("wl", []), dtype=float).reshape(-1)
        if branch_wl.size < 1:
            raise RuntimeError(
                _fmt_issue("HS_PREP", f"Candidate {candidate_idx + 1}: branch {branch_name} has no wavelengths.")
            )
        branch_target_wl = float(branch_cfg.get("global_target_wl_nm", s2_ref_wavelength_nm))
        branch_match_bidx = int(np.argmin(np.abs(branch_wl - branch_target_wl))) + 1
        branch_ref_bidx = int(branch_cfg.get("global_s2_stack_idx", s2_ref_stack_idx))
        branch_s2_subset = tuple(branch_cfg.get("s2_subset", ()))

        branch_global_out = os.path.join(
            folder_struct["temp"], f"{scene_name}_{branch_name}_GLOBAL_c{candidate_idx}.tif"
        )
        branch_local_out = os.path.join(
            folder_struct["temp"], f"{scene_name}_{branch_name}_LOCAL_c{candidate_idx}.tif"
        )
        branch_poly_out = os.path.join(
            folder_struct["temp"], f"{scene_name}_{branch_name}_POLY_c{candidate_idx}.tif"
        )
        branch_tps_out = os.path.join(
            folder_struct["temp"], f"{scene_name}_{branch_name}_TPS_c{candidate_idx}.tif"
        )

        out["temp_global_path"][branch_name] = branch_global_out
        out["temp_local_path"][branch_name] = branch_local_out

        branch_geo = GeoArray(branch_input)
        branch_nodata = branch_geo.nodata if branch_geo.nodata is not None else PROCESSING_NODATA

        global_ladder = matcher_profile.get("global_attempt_ladder", [{"ws": (256, 256), "max_shift": 150.0}])
        branch_validation = {
            "is_valid": False,
            "confidence": 0.0,
            "shift_m": 0.0,
            "message": f"{branch_name}: global validation unavailable",
            "branch": branch_name,
        }
        branch_global_success = False

        for g_idx, g_cfg in enumerate(global_ladder, 1):
            try:
                with _ProgressHeartbeat(
                    progress_callback,
                    "Global coregistration",
                    scene_idx=scene_idx,
                    scene_total=scene_total,
                    substage=(
                        f"Candidate {candidate_idx + 1} [{branch_name}]: attempt {g_idx}/{len(global_ladder)} "
                        f"(ws={g_cfg['ws'][0]}x{g_cfg['ws'][1]})"
                    ),
                    interval_s=progress_heartbeat_interval_s,
                ) as hb_global:
                    captured_output = io.StringIO()
                    with contextlib.redirect_stdout(captured_output):
                        CRG = COREG(
                            s2_raster_path,
                            branch_geo,
                            r_b4match=branch_ref_bidx,
                            s_b4match=branch_match_bidx,
                            path_out=branch_global_out,
                            fmt_out="GTiff",
                            out_crea_options=["COMPRESS=LZW", "BIGTIFF=YES", "TILED=YES"],
                            max_shift=g_cfg["max_shift"],
                            ws=g_cfg["ws"],
                            resamp_alg_deshift="cubic",
                            nodata=(0, branch_nodata),
                            progress=True,
                            ignore_errors=True,
                            v=False,
                            q=False,
                        )
                        CRG.correct_shifts()

                    hb_global.update(
                        substage=f"Candidate {candidate_idx + 1} [{branch_name}]: validating global shift"
                    )
                    validation_candidate = _validate_coreg_hybrid(
                        CRG,
                        captured_output.getvalue(),
                        sensor_type=hyp_type,
                        max_displacement=max_displacement,
                    )
                if float(validation_candidate.get("confidence", 0.0)) >= float(branch_validation.get("confidence", 0.0)):
                    branch_validation = dict(validation_candidate)
                    branch_validation["branch"] = branch_name
                if validation_candidate.get("is_valid", False):
                    branch_global_success = True
                    break
            except Exception as g_exc:
                logger.warning("Candidate %d [%s] global error: %s", candidate_idx + 1, branch_name, g_exc)

        if not branch_global_success:
            forward_res = stream_copy_raster_to_path(
                source_path=branch_input,
                output_path=branch_global_out,
                nodata_fallback=PROCESSING_NODATA,
                expected_band_count=int(branch_wl.size),
                source_bands_1based=None,
                tile_size=256,
                collect_band_stats=False,
            )
            if not forward_res.get("ok", False):
                err = "; ".join(list(forward_res.get("errors", []) or ["unknown copy error"]))
                raise RuntimeError(
                    _fmt_issue(
                        "HS_PREP",
                        f"Candidate {candidate_idx + 1} [{branch_name}] failed global bypass copy: {err}",
                    )
                )

        merged_tp_result: Dict[str, Any] = {"merged_df": None, "stage_counts": {}}
        polynomial_order_decision: Dict[str, Any] = {
            "preferred_order": int(preferred_polynomial_order),
            "order_used": int(preferred_polynomial_order),
            "downgraded": False,
            "reason": "polynomial path not used",
            "n_gcps": 0,
        }
        multiband_result: Dict[str, Any] = {}
        poly_success = False
        tps_success = False
        CRL = None

        with _ProgressHeartbeat(
            progress_callback,
            "Local coregistration (multi-band)",
            scene_idx=scene_idx,
            scene_total=scene_total,
            substage=f"Candidate {candidate_idx + 1} [{branch_name}]: collecting tie points",
            interval_s=progress_heartbeat_interval_s,
        ) as hb_local:
            multiband_result = _collect_multiband_tiepoints(
                s2_path=s2_raster_path,
                hs_path=branch_global_out,
                hs_wl=branch_wl,
                temp_folder=folder_struct["temp"],
                sensor_tag=f"{sensor_tag}_{branch_name}",
                date_tag=date_tag,
                cand_idx=candidate_idx,
                s2_band_subset=branch_s2_subset,
                config={
                    "prefer_fixed_band_pairs": prefer_fixed_band_pairs,
                    "fixed_band_pairs_by_sensor": fixed_band_pairs_by_sensor,
                    "bandpair_wavelength_window_nm": bandpair_wavelength_window_nm,
                    "local_coreg_grid_res": matcher_profile.get("local_grid_res"),
                    "local_coreg_window_size": matcher_profile.get("local_window_size"),
                    "local_coreg_tieP_filter_level": matcher_profile.get("local_tieP_filter_level"),
                    "local_coreg_max_iter": matcher_profile.get("local_max_iter"),
                    "local_max_shift_by_sensor": {
                        matcher_profile.get("sensor", sensor_tag): matcher_profile.get("local_max_shift", 50.0),
                        "DEFAULT": matcher_profile.get("local_max_shift", 50.0),
                    },
                    "global_coreg_profiles_by_sensor": config.get("global_coreg_profiles_by_sensor"),
                    "global_coreg_attempt_ladder": config.get("global_coreg_attempt_ladder"),
                },
            )

            if multiband_result.get("success") and multiband_result.get("all_tiepoints"):
                merged_tp_result = _merge_tiepoints(
                    multiband_result["all_tiepoints"],
                    min_reliability=MIN_RELIABILITY_THRESHOLD,
                    trim_by_residual=True,
                    residual_mad_factor=residual_mad_factor,
                    min_band_support=min_band_support,
                    allow_single_band_fallback=allow_single_band_fallback,
                    consensus_group_rounding_px=consensus_group_rounding_px,
                    grid_rows=spatial_grid_rows,
                    grid_cols=spatial_grid_cols,
                    max_points_per_cell=max_points_per_cell,
                    min_points_required=max(min_tie_points, _minimum_gcps_for_polynomial_order(2)),
                )
                if merged_tp_result.get("success") and merged_tp_result.get("merged_df") is not None:
                    polynomial_order_decision = _decide_polynomial_order(
                        merged_df=merged_tp_result["merged_df"],
                        preferred_order=preferred_polynomial_order,
                        auto_downgrade=auto_downgrade_polynomial_order,
                        min_gcps_order2=min_gcps_order2,
                        min_cells_order2=min_cells_order2,
                        grid_rows=spatial_grid_rows,
                        grid_cols=spatial_grid_cols,
                    )
                    gcp_result = _build_gcps_from_tiepoints(merged_tp_result["merged_df"])
                    min_gcps_for_order = _minimum_gcps_for_polynomial_order(
                        int(polynomial_order_decision.get("order_used", 2))
                    )
                    if gcp_result.get("success") and int(gcp_result.get("n_gcps", 0)) >= int(min_gcps_for_order):
                        hb_local.update(
                            substage=f"Candidate {candidate_idx + 1} [{branch_name}]: applying polynomial warp"
                        )
                        poly_result = _apply_polynomial_warp(
                            input_raster=branch_global_out,
                            output_raster=branch_poly_out,
                            gcps=gcp_result["gcps"],
                            target_crs=s2_crs,
                            polynomial_order=int(polynomial_order_decision.get("order_used", 2)),
                            output_resolution=PRISMA_OUTPUT_RESOLUTION,
                            nodata=PROCESSING_NODATA,
                            resampling="cubic",
                            s2_bounds=s2_extent,
                        )
                        if poly_result.get("success", False):
                            poly_success = True
                            shutil.move(branch_poly_out, branch_local_out)

                    if not poly_success:
                        hb_local.update(
                            substage=f"Candidate {candidate_idx + 1} [{branch_name}]: polynomial failed, trying TPS"
                        )
                        tps_gcps_result = _build_tps_gcps_for_source_raster(
                            tie_points_df=merged_tp_result["merged_df"],
                            source_raster_path=branch_global_out,
                            min_gcps=max(6, int(min_tie_points)),
                        )
                        if tps_gcps_result.get("success", False):
                            tps_result = _apply_tps_warp_from_gcps(
                                input_raster=branch_global_out,
                                output_raster=branch_tps_out,
                                gcps=tps_gcps_result.get("gcps", []),
                                target_crs=s2_crs,
                                x_res=float(PRISMA_OUTPUT_RESOLUTION),
                                y_res=float(PRISMA_OUTPUT_RESOLUTION),
                                resampling="cubic",
                                nodata=PROCESSING_NODATA,
                                target_aligned_pixels=False,
                                target_extent=s2_extent,
                            )
                            if tps_result.get("success", False):
                                tps_success = True
                                shutil.move(branch_tps_out, branch_local_out)

            if not poly_success and not tps_success and POLYNOMIAL_FALLBACK_TO_AROSICS:
                hb_local.update(
                    substage=f"Candidate {candidate_idx + 1} [{branch_name}]: fallback local AROSICS"
                )
                fallback_kwargs = {
                    "grid_res": int(matcher_profile.get("local_grid_res", LOCAL_GRID_RES_M)),
                    "window_size": _coerce_window_size(matcher_profile.get("local_window_size"), (256, 256)),
                    "path_out": branch_local_out,
                    "fmt_out": "GTiff",
                    "out_crea_options": ["COMPRESS=LZW", "BIGTIFF=YES", "TILED=YES"],
                    "r_b4match": branch_ref_bidx,
                    "s_b4match": branch_match_bidx,
                    "max_shift": float(matcher_profile.get("local_max_shift", 50.0)),
                    "resamp_alg_deshift": "cubic",
                    "tieP_filter_level": int(matcher_profile.get("local_tieP_filter_level", 1)),
                    "outFillVal": branch_nodata,
                    "nodata": (0, branch_nodata),
                    "CPUs": CPUS_FOR_AROSICS,
                    "progress": True,
                    "v": False,
                    "q": False,
                    "ignore_errors": True,
                }
                fallback_max_iter = matcher_profile.get("local_max_iter")
                if fallback_max_iter is not None and _supports_constructor_kwarg(COREG_LOCAL, "max_iter"):
                    fallback_kwargs["max_iter"] = int(fallback_max_iter)
                CRL = COREG_LOCAL(s2_raster_path, branch_global_out, **fallback_kwargs)
                CRL.correct_shifts()

        tie_points_df = None
        tie_points_viz_df = None
        if CRL is not None:
            tie_points_df = getattr(CRL, "CoRegPoints_table", None)
            tie_points_viz_df = tie_points_df
        elif merged_tp_result.get("merged_df") is not None:
            tie_points_df = merged_tp_result["merged_df"]
            tie_points_viz_df = merged_tp_result.get("visualization_df", tie_points_df)

        tp_residuals = _compute_tiepoint_residuals(tie_points_df, pixel_size_m=30.0)
        local_tp_count = int(tp_residuals.get("n_tiepoints_used", 0) or 0)

        branch_output_path = branch_local_out if os.path.exists(branch_local_out) else branch_global_out
        if branch_output_path and not os.path.exists(branch_output_path):
            branch_output_path = None
        branch_output_validation = _validate_coreg_raster_content(
            branch_output_path,
            nodata=PROCESSING_NODATA,
            max_windows=0,
            stop_on_first_valid=True,
        )
        branch_output_ok = bool(branch_output_validation.get("ok", False))

        branch_phasecorr = _run_postwarp_phasecorr_qa(
            enabled=bool(postwarp_phasecorr_check),
            reference_raster_path=s2_raster_path,
            candidate_raster_path=branch_output_path,
            reference_band=int(branch_ref_bidx),
            candidate_band=int(max(1, min(branch_match_bidx, int(branch_wl.size)))),
            warn_threshold_px=float(postwarp_phasecorr_warn_threshold_px),
            reject_threshold_px=float(postwarp_phasecorr_reject_threshold_px),
            reject_bad=bool(postwarp_phasecorr_reject_bad),
            max_dim=int(postwarp_phasecorr_max_dim),
        )
        if branch_phasecorr.get("enabled", False) and branch_phasecorr.get("ok", False):
            if branch_phasecorr.get("reject", False):
                branch_output_ok = False

        branch_final_ok = (
            float(branch_validation.get("confidence", 0.0)) >= (float(min_accuracy) / 100.0)
            and int(local_tp_count) >= int(min_tie_points)
            and bool(branch_output_ok)
        )

        mb_counts = dict(multiband_result.get("tiepoint_counts", {}) or {}) if isinstance(multiband_result, dict) else {}
        for band_label, count in mb_counts.items():
            try:
                branch_multiband_counts[str(band_label)] = int(branch_multiband_counts.get(str(band_label), 0)) + int(count)
            except Exception:
                continue

        branch_stage_counts[branch_name] = dict(merged_tp_result.get("stage_counts", {}))
        branch_order_decisions[branch_name] = dict(polynomial_order_decision)
        branch_paths_for_recombine[branch_name] = str(branch_output_path) if branch_output_ok and branch_output_path else None
        branch_quality_flags.append(bool(branch_final_ok))
        branch_results[branch_name] = {
            "validation": dict(branch_validation),
            "global_success": bool(branch_global_success),
            "local_valid_tp_count": int(local_tp_count),
            "tp_residuals": dict(tp_residuals),
            "polynomial_warp_used": bool(poly_success),
            "tps_warp_used": bool(tps_success),
            "polynomial_order_decision": dict(polynomial_order_decision),
            "merged_tiepoint_stages": dict(merged_tp_result.get("stage_counts", {})),
            "multiband_tiepoint_counts": dict(mb_counts),
            "local_tiepoints_df": tie_points_df,
            "local_tiepoints_visualization_df": tie_points_viz_df,
            "output_path": branch_output_path,
            "output_content_valid": bool(branch_output_ok),
            "output_validation": dict(branch_output_validation),
            "postwarp_phasecorr_qa": dict(branch_phasecorr),
        }

    vnir_path = branch_paths_for_recombine.get("VNIR")
    swir_path = branch_paths_for_recombine.get("SWIR")
    recombined_path = os.path.join(folder_struct["temp"], f"{scene_name}_BRANCH_COMBINED_c{candidate_idx}.tif")
    recombine_result = {"ok": False, "error": "missing VNIR/SWIR branch output paths"}
    if vnir_path and swir_path:
        recombine_result = _recombine_detector_branches_windowed(
            vnir_raster_path=vnir_path,
            swir_raster_path=swir_path,
            output_raster_path=recombined_path,
            nodata=PROCESSING_NODATA,
            out_dtype=PROCESSING_DTYPE,
        )
    out["recombine_result"] = dict(recombine_result)

    candidate_output_path = recombined_path if recombine_result.get("ok", False) and os.path.exists(recombined_path) else None
    candidate_output_validation = _validate_coreg_raster_content(
        candidate_output_path,
        nodata=PROCESSING_NODATA,
        max_windows=0,
        stop_on_first_valid=True,
    )
    candidate_output_ok = bool(candidate_output_validation.get("ok", False))

    validation_by_branch = {name: dict(data.get("validation", {})) for name, data in branch_results.items()}
    confs = [float(v.get("confidence", 0.0)) for v in validation_by_branch.values()]
    shifts = [float(v.get("shift_m", 0.0)) for v in validation_by_branch.values()]
    msgs = [str(v.get("message", "")).strip() for v in validation_by_branch.values() if str(v.get("message", "")).strip()]

    merged_validation = {
        "is_valid": bool(all(v.get("is_valid", False) for v in validation_by_branch.values())),
        "confidence": float(min(confs)) if confs else 0.0,
        "shift_m": float(max(shifts)) if shifts else 0.0,
        "message": " | ".join(msgs) if msgs else "No branch validation messages.",
        "branch_validations": validation_by_branch,
        "ssim_before": None,
        "ssim_after": None,
        "ssim_delta": None,
    }
    try:
        before_vals = [float(v.get("ssim_before")) for v in validation_by_branch.values() if v.get("ssim_before") is not None]
        after_vals = [float(v.get("ssim_after")) for v in validation_by_branch.values() if v.get("ssim_after") is not None]
        delta_vals = [float(v.get("ssim_delta")) for v in validation_by_branch.values() if v.get("ssim_delta") is not None]
        if before_vals:
            merged_validation["ssim_before"] = float(np.mean(before_vals))
        if after_vals:
            merged_validation["ssim_after"] = float(np.mean(after_vals))
        if delta_vals:
            merged_validation["ssim_delta"] = float(np.mean(delta_vals))
    except Exception:
        pass

    merged_tp_frames = [
        branch_results[name].get("local_tiepoints_df")
        for name in ("VNIR", "SWIR")
        if branch_results.get(name, {}).get("local_tiepoints_df") is not None
        and len(branch_results.get(name, {}).get("local_tiepoints_df")) > 0
    ]
    tie_points_df = pd.concat(merged_tp_frames, ignore_index=True) if merged_tp_frames else None
    merged_tp_viz_frames = [
        branch_results[name].get("local_tiepoints_visualization_df")
        for name in ("VNIR", "SWIR")
        if branch_results.get(name, {}).get("local_tiepoints_visualization_df") is not None
        and len(branch_results.get(name, {}).get("local_tiepoints_visualization_df")) > 0
    ]
    tie_points_viz_df = (
        pd.concat(merged_tp_viz_frames, ignore_index=True)
        if merged_tp_viz_frames
        else tie_points_df
    )
    tp_residuals = _compute_tiepoint_residuals(tie_points_df, pixel_size_m=30.0)
    local_valid_tp_count = int(sum(int(branch_results.get(name, {}).get("local_valid_tp_count", 0)) for name in ("VNIR", "SWIR")))

    poly_used = bool(any(bool(branch_results.get(name, {}).get("polynomial_warp_used", False)) for name in ("VNIR", "SWIR")))
    tps_used = bool(any(bool(branch_results.get(name, {}).get("tps_warp_used", False)) for name in ("VNIR", "SWIR")))
    poly_order_used = int(
        min(
            int(branch_order_decisions.get(name, {}).get("order_used", preferred_polynomial_order))
            for name in ("VNIR", "SWIR")
        )
    )
    poly_decision = {
        "branches": {k: dict(v) for k, v in branch_order_decisions.items()},
        "order_used": int(poly_order_used),
        "n_gcps": int(sum(int(branch_order_decisions.get(k, {}).get("n_gcps", 0)) for k in branch_order_decisions)),
        "downgraded": bool(any(bool(branch_order_decisions.get(k, {}).get("downgraded", False)) for k in branch_order_decisions)),
        "reason": "; ".join([f"{k}: {branch_order_decisions.get(k, {}).get('reason', 'n/a')}" for k in ("VNIR", "SWIR")]),
        "tps_fallback_used": bool(tps_used),
    }

    out["validation"] = merged_validation
    out["final_quality_pass"] = bool(all(branch_quality_flags)) and bool(candidate_output_ok)
    out["local_valid_tp_count"] = int(local_valid_tp_count)
    out["tp_residuals"] = dict(tp_residuals)
    out["polynomial_warp_used"] = bool(poly_used)
    out["tps_warp_used"] = bool(tps_used)
    out["polynomial_order_decision"] = dict(poly_decision)
    out["polynomial_order_used"] = int(poly_order_used)
    out["local_tiepoints_df"] = tie_points_df
    out["local_tiepoints_visualization_df"] = tie_points_viz_df
    out["output_path"] = candidate_output_path
    out["output_content_valid"] = bool(candidate_output_ok)
    out["output_validation"] = dict(candidate_output_validation)
    out["postwarp_phasecorr_qa"] = {
        "branches": {name: dict(branch_results.get(name, {}).get("postwarp_phasecorr_qa", {})) for name in ("VNIR", "SWIR")}
    }
    out["merged_tiepoint_stages"] = dict(branch_stage_counts)
    out["multiband_tiepoint_counts"] = dict(branch_multiband_counts)
    out["branch_results"] = {k: dict(v) for k, v in branch_results.items()}
    return out


def _validate_coreg_raster_content(
    path: Optional[str],
    nodata: float = PROCESSING_NODATA,
    max_windows: int = 0,
    stop_on_first_valid: bool = False,
) -> Dict[str, Any]:
    """Validate that a coreg raster contains at least one finite non-nodata pixel."""
    out = {
        "ok": False,
        "error": None,
        "path": path,
        "valid_pixels": 0,
        "total_pixels": 0,
        "valid_fraction": 0.0,
        "bands_checked": 0,
        "windows_scanned": 0,
        "truncated": False,
    }
    try:
        if not path:
            out["error"] = "No raster path provided."
            return out
        if not os.path.exists(path):
            out["error"] = f"Raster missing: {path}"
            return out

        with rasterio.open(path) as src:
            if src.count < 1 or src.width < 1 or src.height < 1:
                out["error"] = "Invalid raster dimensions/count."
                return out

            nodata_val = src.nodata
            if nodata_val is None or not np.isfinite(float(nodata_val)):
                nodata_val = float(nodata)
            else:
                nodata_val = float(nodata_val)

            valid_pixels = 0
            total_pixels = 0
            windows_scanned = 0
            max_wins = int(max(0, int(max_windows)))
            found_valid = False
            for bidx in range(1, src.count + 1):
                for _, window in src.block_windows(bidx):
                    if max_wins > 0 and windows_scanned >= max_wins:
                        out["truncated"] = True
                        break
                    band = src.read(bidx, window=window).astype(np.float32, copy=False)
                    total_pixels += int(band.size)
                    valid = np.isfinite(band)
                    if np.isfinite(nodata_val):
                        valid &= (band != nodata_val)
                    vcount = int(np.count_nonzero(valid))
                    valid_pixels += vcount
                    windows_scanned += 1
                    if stop_on_first_valid and vcount > 0:
                        found_valid = True
                        out["truncated"] = True
                        break
                if found_valid:
                    break
                if max_wins > 0 and windows_scanned >= max_wins:
                    break

            out["bands_checked"] = int(src.count)
            out["windows_scanned"] = int(windows_scanned)
            out["valid_pixels"] = int(valid_pixels)
            out["total_pixels"] = int(total_pixels)
            if total_pixels > 0:
                out["valid_fraction"] = float(valid_pixels / total_pixels)
            out["ok"] = valid_pixels > 0
            if not out["ok"]:
                out["error"] = "Raster contains no finite non-nodata pixels."
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _probe_raster_valid_pixels(
    path: Optional[str],
    bands_1based: Sequence[int],
    nodata_fallback: float = PROCESSING_NODATA,
    sample_max_dim: int = 512,
) -> Dict[str, Any]:
    """Probe selected bands for finite, non-nodata pixels using optional decimation."""
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "path": path,
        "nodata": None,
        "band_count": 0,
        "sampled_bands": [],
        "band_valid_pixels": {},
        "sample_shape": None,
    }
    try:
        if not path:
            out["error"] = "No raster path provided."
            return out
        if not os.path.exists(path):
            out["error"] = f"Raster missing: {path}"
            return out

        with rasterio.open(path) as src:
            out["band_count"] = int(src.count)
            if src.count < 1:
                out["error"] = "Raster has no bands."
                return out

            nodata_val = src.nodata
            if nodata_val is None or not np.isfinite(float(nodata_val)):
                nodata_val = float(nodata_fallback)
            else:
                nodata_val = float(nodata_val)
            out["nodata"] = nodata_val

            sampled_bands: List[int] = []
            for b in bands_1based:
                try:
                    bidx = int(b)
                except Exception:
                    continue
                if bidx < 1 or bidx > int(src.count):
                    continue
                if bidx not in sampled_bands:
                    sampled_bands.append(bidx)
            if not sampled_bands:
                out["error"] = "No valid probe bands."
                return out
            out["sampled_bands"] = list(sampled_bands)

            src_w = int(src.width)
            src_h = int(src.height)
            max_dim = int(max(1, int(sample_max_dim)))
            longest = max(src_w, src_h)
            if longest > max_dim:
                scale = float(max_dim) / float(longest)
                out_w = max(1, int(round(src_w * scale)))
                out_h = max(1, int(round(src_h * scale)))
            else:
                out_w = src_w
                out_h = src_h
            out["sample_shape"] = (int(out_h), int(out_w))

            any_valid = False
            for bidx in sampled_bands:
                band = src.read(
                    bidx,
                    out_shape=(out_h, out_w),
                    resampling=Resampling.nearest,
                ).astype(np.float32, copy=False)
                valid = np.isfinite(band)
                if np.isfinite(nodata_val):
                    valid &= (band != nodata_val)
                valid_count = int(np.count_nonzero(valid))
                out["band_valid_pixels"][int(bidx)] = valid_count
                if valid_count > 0:
                    any_valid = True

            out["ok"] = bool(any_valid)
            if not out["ok"]:
                out["error"] = "No finite non-nodata pixels in sampled probe bands."
        return out
    except Exception as e:
        out["error"] = str(e)
        return out


def _validate_local_s2_stack_override(
    local_s2_stack_path: Optional[str],
    min_band_count: int = 6,
) -> Dict[str, Any]:
    """Validate a local Sentinel-2 stack override path for offline runs."""
    out: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "path": None,
        "crs": None,
        "band_count": 0,
    }
    path_raw = str(local_s2_stack_path or "").strip()
    if not path_raw:
        out["error"] = "No local S2 override path provided."
        return out

    path_abs = os.path.abspath(path_raw)
    out["path"] = path_abs
    if not os.path.isfile(path_abs):
        out["error"] = f"Local S2 stack override not found: {path_abs}"
        return out

    try:
        with rasterio.open(path_abs) as src:
            out["band_count"] = int(src.count)
            out["crs"] = src.crs
            if src.count < int(max(1, int(min_band_count))):
                out["error"] = (
                    f"Local S2 stack has too few bands ({int(src.count)} < {int(min_band_count)})."
                )
                return out
            if src.crs is None:
                out["error"] = "Local S2 stack CRS is missing."
                return out
    except Exception as e:
        out["error"] = f"Failed to open local S2 stack override: {e}"
        return out

    out["ok"] = True
    return out


def _summarize_crs(crs_obj: Any) -> str:
    """Return compact CRS summary suitable for footer text."""
    if crs_obj is None:
        return "n/a"
    try:
        epsg = crs_obj.to_epsg()
        if epsg is not None:
            return f"EPSG:{int(epsg)}"
    except Exception:
        pass
    txt = str(crs_obj).strip()
    if not txt:
        return "n/a"
    if len(txt) <= 40:
        return txt
    return f"{txt[:37]}..."


def _format_optional_float(value: Any, fmt: str = ".3f", na: str = "n/a") -> str:
    """Format float-like values safely for footer text."""
    try:
        v = float(value)
        if np.isfinite(v):
            return format(v, fmt)
    except Exception:
        pass
    return na


def _quicklook_gray_cmap() -> Any:
    """Build grayscale colormap where invalid pixels are dark and explicit."""
    try:
        base = matplotlib.colormaps.get_cmap("gray").resampled(256)
    except Exception:
        base = cm.get_cmap("gray", 256)
    cmap = ListedColormap(base(np.linspace(0.0, 1.0, 256)))
    cmap.set_bad(color="#1f1f1f")
    return cmap


def _compute_quicklook_display(
    band: np.ndarray,
    valid_mask: np.ndarray,
) -> Tuple[np.ndarray, Optional[float], Optional[float]]:
    """Compute p2/p98 display stretch over valid pixels and return display array."""
    disp = np.full(band.shape, np.nan, dtype=np.float32)
    if not np.any(valid_mask):
        return disp, None, None

    vals = band[valid_mask]
    try:
        lo = float(np.percentile(vals, 2))
        hi = float(np.percentile(vals, 98))
    except Exception:
        lo = float(np.min(vals))
        hi = float(np.max(vals))

    if not np.isfinite(lo):
        lo = float(np.min(vals))
    if not np.isfinite(hi):
        hi = float(np.max(vals))

    if hi > lo:
        disp[valid_mask] = np.clip((band[valid_mask] - lo) / (hi - lo), 0.0, 1.0)
    else:
        disp[valid_mask] = 0.5
    return disp, lo, hi


def _load_quicklook_band(
    raster_path: str,
    band_idx: int,
    nodata: float = PROCESSING_NODATA,
    max_quicklook_dim: int = DEFAULT_QUICKLOOK_MAX_DIM,
) -> Tuple[np.ndarray, np.ndarray, Dict[str, Any]]:
    """Load one band (optionally decimated) and return array, valid mask, and display metadata."""
    with rasterio.open(raster_path) as src:
        use_bidx = int(max(1, min(int(band_idx), int(src.count))))
        nodata_val = src.nodata
        if nodata_val is None or not np.isfinite(float(nodata_val)):
            nodata_val = float(nodata)
        else:
            nodata_val = float(nodata_val)

        src_w = int(src.width)
        src_h = int(src.height)
        max_dim = int(max(1, int(max_quicklook_dim)))
        longest = max(src_w, src_h)
        if longest > max_dim:
            scale = float(max_dim) / float(longest)
            out_w = max(1, int(round(src_w * scale)))
            out_h = max(1, int(round(src_h * scale)))
        else:
            out_w = src_w
            out_h = src_h

        band = src.read(
            use_bidx,
            out_shape=(out_h, out_w),
            resampling=Resampling.bilinear,
        ).astype(np.float32, copy=False)

        transform = src.transform
        try:
            px_x = abs(float(transform.a))
        except Exception:
            px_x = np.nan
        try:
            px_y = abs(float(transform.e))
        except Exception:
            px_y = np.nan

        info = {
            "raster_basename": os.path.basename(raster_path),
            "source_width": src_w,
            "source_height": src_h,
            "quicklook_width": out_w,
            "quicklook_height": out_h,
            "band_count": int(src.count),
            "dtype": str(src.dtypes[use_bidx - 1]) if src.dtypes else "unknown",
            "crs_summary": _summarize_crs(src.crs),
            "pixel_size_x": px_x,
            "pixel_size_y": px_y,
            "band_idx": use_bidx,
            "nodata": nodata_val,
            "x_scale": float(out_w) / float(src_w) if src_w > 0 else 1.0,
            "y_scale": float(out_h) / float(src_h) if src_h > 0 else 1.0,
            "decimated": bool(out_w != src_w or out_h != src_h),
        }

    valid = np.isfinite(band)
    if np.isfinite(nodata_val):
        valid &= (band != nodata_val)
    return band, valid, info


def _build_quicklook_footer_info(
    *,
    output_png: str,
    scene_name: Optional[str],
    raster_info: Dict[str, Any],
    band_wavelength_nm: Optional[float],
    stretch_lo: Optional[float],
    stretch_hi: Optional[float],
    valid_count: int,
    total_count: int,
    tiepoint_info: Optional[Dict[str, Any]] = None,
) -> List[str]:
    """Build compact footer lines for scene/tiepoint quicklooks."""
    scene_label = str(scene_name).strip() if scene_name is not None else ""
    if not scene_label:
        scene_label = os.path.splitext(os.path.basename(output_png))[0]

    src_w = int(raster_info.get("source_width", 0))
    src_h = int(raster_info.get("source_height", 0))
    ql_w = int(raster_info.get("quicklook_width", src_w))
    ql_h = int(raster_info.get("quicklook_height", src_h))
    band_idx = int(raster_info.get("band_idx", 1))
    band_count = int(raster_info.get("band_count", 0))
    dtype = str(raster_info.get("dtype", "unknown"))
    crs_summary = str(raster_info.get("crs_summary", "n/a"))
    px_x = _format_optional_float(raster_info.get("pixel_size_x"), ".3f")
    px_y = _format_optional_float(raster_info.get("pixel_size_y"), ".3f")
    nodata_txt = _format_optional_float(raster_info.get("nodata"), ".4g")
    lo_txt = _format_optional_float(stretch_lo, ".4g")
    hi_txt = _format_optional_float(stretch_hi, ".4g")
    wl_txt = _format_optional_float(band_wavelength_nm, ".1f")
    valid_pct = 0.0
    if total_count > 0:
        valid_pct = 100.0 * float(valid_count) / float(total_count)

    lines = [
        f"Scene: {scene_label} | Raster: {raster_info.get('raster_basename', 'n/a')} | PNG: {os.path.basename(output_png)}",
        (
            f"Size: {src_w}x{src_h}px (QL {ql_w}x{ql_h}) | Bands: {band_count} | dtype: {dtype} | "
            f"CRS: {crs_summary} | Pixel: {px_x} x {px_y}"
        ),
        (
            f"Display: B{band_idx}/{band_count} ({wl_txt} nm), p2/p98={lo_txt}/{hi_txt}, "
            f"nodata={nodata_txt}, valid={valid_count}/{total_count} ({valid_pct:.1f}%)"
        ),
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%MZ')}",
    ]

    if tiepoint_info is not None:
        outlier_counts = tiepoint_info.get("outlier_counts", {})
        outlier_parts: List[str] = []
        for key in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
            if key in outlier_counts:
                outlier_parts.append(f"{key.split('_')[0]}:{int(outlier_counts.get(key, 0))}")
        outlier_txt = ", ".join(outlier_parts) if outlier_parts else "n/a"
        lines.append(
            (
                f"Tie points: total={int(tiepoint_info.get('total', 0))}, "
                f"inliers plotted={int(tiepoint_info.get('inlier_plotted', 0))}, "
                f"outliers plotted={int(tiepoint_info.get('outlier_plotted', 0))}, outlier flags [{outlier_txt}]"
            )
        )
        lines.append(
            (
                "Residuals "
                f"({tiepoint_info.get('residual_source', 'unavailable')}): "
                f"mean={_format_optional_float(tiepoint_info.get('residual_mean_m'), '.2f')} m, "
                f"median={_format_optional_float(tiepoint_info.get('residual_median_m'), '.2f')} m, "
                f"RMSE={_format_optional_float(tiepoint_info.get('residual_rmse_m'), '.2f')} m, "
                f"P90={_format_optional_float(tiepoint_info.get('residual_p90_m'), '.2f')} m | "
                f"MIN_RELIABILITY_THRESHOLD={MIN_RELIABILITY_THRESHOLD:.1f}%"
            )
        )

    return lines


def _extract_tiepoint_plot_data(
    tie_points_df,
    x_scale: float,
    y_scale: float,
) -> Dict[str, Any]:
    """Extract inlier/outlier point coordinates and optional coloring payload."""
    empty = {
        "total": 0,
        "inlier_total": 0,
        "inlier_plotted": 0,
        "outlier_plotted": 0,
        "outlier_counts": {},
        "x_inlier": np.array([], dtype=np.float32),
        "y_inlier": np.array([], dtype=np.float32),
        "x_outlier": np.array([], dtype=np.float32),
        "y_outlier": np.array([], dtype=np.float32),
        "x_color": np.array([], dtype=np.float32),
        "y_color": np.array([], dtype=np.float32),
        "color_values": np.array([], dtype=np.float32),
        "color_label": None,
        "color_cmap": None,
    }
    if tie_points_df is None:
        return empty
    try:
        df = tie_points_df.copy()
    except Exception:
        return empty
    if not hasattr(df, "__len__") or len(df) < 1:
        return empty

    out = dict(empty)
    out["total"] = int(len(df))

    try:
        outlier_any = np.zeros(len(df), dtype=bool)
        outlier_counts: Dict[str, int] = {}
        for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
            if col in df.columns:
                try:
                    mask = df[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
                except Exception:
                    mask = np.asarray(df[col], dtype=bool)
                outlier_counts[col] = int(np.count_nonzero(mask))
                outlier_any |= mask
        out["outlier_counts"] = outlier_counts

        df_in = df[~outlier_any] if len(df) else df
        df_out = df[outlier_any] if len(df) else df.iloc[0:0]
        out["inlier_total"] = int(len(df_in))

        def _scaled_xy(frame):
            if "X_IM" not in frame.columns or "Y_IM" not in frame.columns or len(frame) < 1:
                return (
                    np.array([], dtype=float),
                    np.array([], dtype=float),
                    np.array([], dtype=bool),
                    np.array([], dtype=np.float32),
                    np.array([], dtype=np.float32),
                )
            x_raw = np.asarray(frame["X_IM"], dtype=float)
            y_raw = np.asarray(frame["Y_IM"], dtype=float)
            finite = np.isfinite(x_raw) & np.isfinite(y_raw)
            x_sc = (x_raw[finite] * float(x_scale)).astype(np.float32, copy=False)
            y_sc = (y_raw[finite] * float(y_scale)).astype(np.float32, copy=False)
            return x_raw, y_raw, finite, x_sc, y_sc

        x_in_raw, y_in_raw, in_xy_ok, x_in, y_in = _scaled_xy(df_in)
        _, _, _, x_out, y_out = _scaled_xy(df_out)
        out["x_inlier"] = x_in
        out["y_inlier"] = y_in
        out["x_outlier"] = x_out
        out["y_outlier"] = y_out
        out["inlier_plotted"] = int(x_in.size)
        out["outlier_plotted"] = int(x_out.size)

        if len(df_in) > 0 and x_in_raw.size == len(df_in):
            if "RELIABILITY" in df_in.columns:
                rel = np.asarray(df_in["RELIABILITY"], dtype=float)
                color_mask = in_xy_ok & np.isfinite(rel)
                color_vals = rel[color_mask]
                if color_vals.size > 1 and float(np.nanmax(color_vals)) > float(np.nanmin(color_vals)):
                    out["x_color"] = (x_in_raw[color_mask] * float(x_scale)).astype(np.float32, copy=False)
                    out["y_color"] = (y_in_raw[color_mask] * float(y_scale)).astype(np.float32, copy=False)
                    out["color_values"] = color_vals.astype(np.float32, copy=False)
                    out["color_label"] = "Reliability (%)"
                    out["color_cmap"] = "viridis"
            elif "ABS_SHIFT_M" in df_in.columns:
                resid = np.asarray(df_in["ABS_SHIFT_M"], dtype=float)
                color_mask = in_xy_ok & np.isfinite(resid)
                color_vals = resid[color_mask]
                if color_vals.size > 1 and float(np.nanmax(color_vals)) > float(np.nanmin(color_vals)):
                    out["x_color"] = (x_in_raw[color_mask] * float(x_scale)).astype(np.float32, copy=False)
                    out["y_color"] = (y_in_raw[color_mask] * float(y_scale)).astype(np.float32, copy=False)
                    out["color_values"] = color_vals.astype(np.float32, copy=False)
                    out["color_label"] = "Residual shift (m)"
                    out["color_cmap"] = "magma"
            elif "X_SHIFT_M" in df_in.columns and "Y_SHIFT_M" in df_in.columns:
                x_shift = np.asarray(df_in["X_SHIFT_M"], dtype=float)
                y_shift = np.asarray(df_in["Y_SHIFT_M"], dtype=float)
                resid = np.sqrt(x_shift**2 + y_shift**2)
                color_mask = in_xy_ok & np.isfinite(resid)
                color_vals = resid[color_mask]
                if color_vals.size > 1 and float(np.nanmax(color_vals)) > float(np.nanmin(color_vals)):
                    out["x_color"] = (x_in_raw[color_mask] * float(x_scale)).astype(np.float32, copy=False)
                    out["y_color"] = (y_in_raw[color_mask] * float(y_scale)).astype(np.float32, copy=False)
                    out["color_values"] = color_vals.astype(np.float32, copy=False)
                    out["color_label"] = "Residual shift (m)"
                    out["color_cmap"] = "magma"
    except Exception:
        return empty
    return out


def _coerce_float_triplet(
    value: Any,
    default: Tuple[float, float, float],
) -> Tuple[float, float, float]:
    """Parse three float values from config-like input."""
    if isinstance(value, str):
        parts = [p.strip() for p in value.split(",") if p.strip()]
    elif isinstance(value, (list, tuple, np.ndarray)):
        parts = list(value)
    else:
        parts = []
    if len(parts) >= 3:
        try:
            vals = tuple(float(parts[i]) for i in range(3))
            if all(np.isfinite(v) for v in vals):
                return vals  # type: ignore[return-value]
        except Exception:
            pass
    return tuple(float(v) for v in default)


def _coerce_percentiles(
    value: Any,
    default: Tuple[float, float] = (2.0, 98.0),
) -> Tuple[float, float]:
    """Parse percentile pair with bounds checks."""
    if isinstance(value, str):
        parts = [p.strip() for p in value.split(",") if p.strip()]
    elif isinstance(value, (list, tuple, np.ndarray)):
        parts = list(value)
    else:
        parts = []
    if len(parts) >= 2:
        try:
            low = float(parts[0])
            high = float(parts[1])
            if 0.0 <= low < high <= 100.0:
                return low, high
        except Exception:
            pass
    return float(default[0]), float(default[1])


def _pick_band_index_for_wavelength(target_nm: float, wavelengths_nm: np.ndarray) -> int:
    """Pick nearest 1-based band index for target wavelength."""
    arr = np.asarray(wavelengths_nm, dtype=float).reshape(-1)
    if arr.size < 1:
        return 1
    finite = np.isfinite(arr)
    if not np.any(finite):
        return 1
    idxs = np.where(finite)[0]
    nearest_local = int(np.argmin(np.abs(arr[idxs] - float(target_nm))))
    return int(idxs[nearest_local]) + 1


def _fallback_rgb_band_indices(
    band_count: int,
    targets_nm: Tuple[float, float, float],
) -> Tuple[int, int, int]:
    """Fallback RGB indices when wavelength metadata is unavailable."""
    if band_count <= 1:
        return 1, 1, 1
    if band_count == 2:
        return 2, 1, 1
    min_wl = 400.0
    max_wl = 2500.0
    raw = []
    for t in targets_nm:
        frac = (float(t) - min_wl) / (max_wl - min_wl)
        idx = 1 + int(round(frac * (band_count - 1)))
        raw.append(max(1, min(band_count, idx)))
    used = set()
    uniq: List[int] = []
    for idx in raw:
        if idx not in used:
            uniq.append(idx)
            used.add(idx)
            continue
        for radius in range(1, band_count):
            for candidate in (idx - radius, idx + radius):
                if 1 <= candidate <= band_count and candidate not in used:
                    uniq.append(candidate)
                    used.add(candidate)
                    break
            if len(uniq) == len(raw):
                break
    while len(uniq) < 3:
        uniq.append(min(band_count, len(uniq) + 1))
    return int(uniq[0]), int(uniq[1]), int(uniq[2])


def _pick_unique_band_index_for_wavelength(
    target_nm: float,
    wavelengths_nm: np.ndarray,
    used_zero_based: set,
) -> int:
    """Pick nearest 1-based band index not present in used_zero_based when possible."""
    arr = np.asarray(wavelengths_nm, dtype=float).reshape(-1)
    if arr.size < 1:
        return 1
    finite_idx = np.where(np.isfinite(arr))[0]
    if finite_idx.size < 1:
        return 1

    order = finite_idx[np.argsort(np.abs(arr[finite_idx] - float(target_nm)))]
    for idx0 in order:
        if int(idx0) not in used_zero_based:
            used_zero_based.add(int(idx0))
            return int(idx0) + 1

    # All candidates already used; allow reuse.
    idx0 = int(order[0])
    used_zero_based.add(idx0)
    return idx0 + 1


def _resolve_quicklook_rgb_bands(
    wavelengths_nm: Optional[np.ndarray],
    band_count: int,
    targets_nm: Tuple[float, float, float],
) -> Tuple[Tuple[int, int, int], List[Optional[float]], str]:
    """Resolve RGB band indices and representative wavelengths."""
    wl_arr = np.asarray(wavelengths_nm, dtype=float).reshape(-1) if wavelengths_nm is not None else np.array([])
    if wl_arr.size >= int(band_count) and np.any(np.isfinite(wl_arr[:band_count])):
        wl_match = np.asarray(wl_arr[:band_count], dtype=float)
        finite = wl_match[np.isfinite(wl_match)]
        wl_scale = 1.0
        rgb_source = "wavelength_metadata"
        # Some metadata sources expose wavelengths in micrometers instead of nanometers.
        if finite.size > 0:
            finite_abs_med = float(np.nanmedian(np.abs(finite)))
            if finite_abs_med < 10.0:
                wl_scale = 1000.0
                rgb_source = "wavelength_metadata_um_to_nm"
        wl_match_nm = wl_match * float(wl_scale)

        used: set = set()
        r_idx = _pick_unique_band_index_for_wavelength(targets_nm[0], wl_match_nm, used)
        g_idx = _pick_unique_band_index_for_wavelength(targets_nm[1], wl_match_nm, used)
        b_idx = _pick_unique_band_index_for_wavelength(targets_nm[2], wl_match_nm, used)
        indices = (int(r_idx), int(g_idx), int(b_idx))
        selected_wl = [
            float(wl_arr[r_idx - 1]) if np.isfinite(wl_arr[r_idx - 1]) else None,
            float(wl_arr[g_idx - 1]) if np.isfinite(wl_arr[g_idx - 1]) else None,
            float(wl_arr[b_idx - 1]) if np.isfinite(wl_arr[b_idx - 1]) else None,
        ]
        return indices, selected_wl, rgb_source

    indices = _fallback_rgb_band_indices(int(max(1, band_count)), targets_nm)
    logger.warning(
        _fmt_issue(
            "QUICKLOOK",
            "Wavelength metadata unavailable/incomplete; using deterministic fallback RGB band indices.",
        )
    )
    return indices, [None, None, None], "fallback_assumed_400_2500nm"


def _load_quicklook_rgb_composite(
    raster_path: str,
    rgb_band_indices: Tuple[int, int, int],
    rgb_band_wavelengths_nm: Optional[Sequence[Optional[float]]] = None,
    nodata: float = PROCESSING_NODATA,
    max_quicklook_dim: int = DEFAULT_QUICKLOOK_MAX_DIM,
    percentiles: Tuple[float, float] = (2.0, 98.0),
    gamma: float = 1.0,
    crop_to_valid: bool = False,
) -> Tuple[np.ndarray, np.ndarray, Dict[str, Any]]:
    """Load RGB composite with decimation and robust per-channel stretch."""
    p_low, p_high = _coerce_percentiles(percentiles, default=(2.0, 98.0))
    gamma_v = max(1e-6, float(gamma))

    with rasterio.open(raster_path) as src:
        src_w = int(src.width)
        src_h = int(src.height)
        band_count = int(src.count)
        use_indices = [int(max(1, min(band_count, int(b)))) for b in rgb_band_indices]

        nodata_val = src.nodata
        if nodata_val is None or not np.isfinite(float(nodata_val)):
            nodata_val = float(nodata)
        else:
            nodata_val = float(nodata_val)

        max_dim = int(max(1, int(max_quicklook_dim)))
        longest = max(src_w, src_h)
        if longest > max_dim:
            scale = float(max_dim) / float(longest)
            out_w = max(1, int(round(src_w * scale)))
            out_h = max(1, int(round(src_h * scale)))
        else:
            out_w = src_w
            out_h = src_h

        rgb_raw = src.read(
            indexes=use_indices,
            out_shape=(3, out_h, out_w),
            resampling=Resampling.bilinear,
        ).astype(np.float32, copy=False)

        transform = src.transform
        try:
            px_x = abs(float(transform.a))
        except Exception:
            px_x = np.nan
        try:
            px_y = abs(float(transform.e))
        except Exception:
            px_y = np.nan

        dtype_txt = str(src.dtypes[0]) if src.dtypes else "unknown"
        crs_summary = _summarize_crs(src.crs)

    valid = np.all(np.isfinite(rgb_raw), axis=0)
    if np.isfinite(nodata_val):
        valid &= np.all(rgb_raw != nodata_val, axis=0)

    stretched = np.zeros_like(rgb_raw, dtype=np.float32)
    channel_stats: List[Dict[str, Any]] = []
    for c in range(3):
        ch = rgb_raw[c]
        if np.any(valid):
            vals = ch[valid]
            try:
                lo = float(np.percentile(vals, p_low))
                hi = float(np.percentile(vals, p_high))
            except Exception:
                lo = float(np.min(vals))
                hi = float(np.max(vals))
            if not np.isfinite(lo):
                lo = float(np.min(vals))
            if not np.isfinite(hi):
                hi = float(np.max(vals))
            if hi > lo:
                stretched[c, valid] = np.clip((ch[valid] - lo) / (hi - lo), 0.0, 1.0)
            else:
                stretched[c, valid] = 0.5
        else:
            lo, hi = None, None
        wl_val: Optional[float] = None
        if rgb_band_wavelengths_nm is not None and len(rgb_band_wavelengths_nm) > c:
            try:
                candidate = rgb_band_wavelengths_nm[c]
                if candidate is not None and np.isfinite(float(candidate)):
                    wl_val = float(candidate)
            except Exception:
                wl_val = None
        channel_stats.append(
            {
                "band_idx": int(use_indices[c]),
                "wavelength_nm": wl_val,
                "p_low": lo,
                "p_high": hi,
            }
        )

    if np.any(valid) and gamma_v != 1.0:
        stretched[:, valid] = np.power(np.clip(stretched[:, valid], 0.0, 1.0), 1.0 / gamma_v)

    nodata_rgb = np.array([0.08, 0.08, 0.08], dtype=np.float32)
    for c in range(3):
        stretched[c, ~valid] = nodata_rgb[c]
    rgb = np.moveaxis(stretched, 0, -1)

    crop_x0 = 0
    crop_y0 = 0
    if bool(crop_to_valid) and np.any(valid):
        rows, cols = np.where(valid)
        if rows.size > 0 and cols.size > 0:
            margin = int(max(4, round(0.02 * max(valid.shape))))
            y0 = max(0, int(rows.min()) - margin)
            y1 = min(valid.shape[0], int(rows.max()) + margin + 1)
            x0 = max(0, int(cols.min()) - margin)
            x1 = min(valid.shape[1], int(cols.max()) + margin + 1)
            rgb = rgb[y0:y1, x0:x1, :]
            valid = valid[y0:y1, x0:x1]
            crop_x0 = int(x0)
            crop_y0 = int(y0)

    info = {
        "raster_basename": os.path.basename(raster_path),
        "source_width": src_w,
        "source_height": src_h,
        "quicklook_width": int(rgb.shape[1]),
        "quicklook_height": int(rgb.shape[0]),
        "quicklook_width_uncropped": int(out_w),
        "quicklook_height_uncropped": int(out_h),
        "band_count": int(band_count),
        "dtype": dtype_txt,
        "crs_summary": crs_summary,
        "pixel_size_x": px_x,
        "pixel_size_y": px_y,
        "nodata": nodata_val,
            "x_scale": float(out_w) / float(src_w) if src_w > 0 else 1.0,
            "y_scale": float(out_h) / float(src_h) if src_h > 0 else 1.0,
            "x_offset": int(crop_x0),
            "y_offset": int(crop_y0),
            "affine_a": float(transform.a),
            "affine_b": float(transform.b),
            "affine_d": float(transform.d),
            "affine_e": float(transform.e),
            "decimated": bool(out_w != src_w or out_h != src_h),
            "percentiles": (float(p_low), float(p_high)),
            "gamma": float(gamma_v),
            "channels": channel_stats,
        }
    return rgb, valid, info


def _footer_items_to_block(
    items: Sequence[Tuple[str, str]],
    key_width: int = 16,
    value_width: int = 58,
) -> str:
    """Format key/value footer rows into a compact aligned text block."""
    lines: List[str] = []
    for key, raw_val in items:
        val = str(raw_val)
        wrapped = textwrap.wrap(val, width=max(12, int(value_width))) or [""]
        lines.append(f"{key:<{key_width}} {wrapped[0]}")
        indent = " " * (int(key_width) + 1)
        for part in wrapped[1:]:
            lines.append(f"{indent}{part}")
    return "\n".join(lines)


def _draw_quicklook_footer(
    ax_info: Any,
    left_items: Sequence[Tuple[str, str]],
    right_items: Sequence[Tuple[str, str]],
) -> None:
    """Render two-column footer info bar with bold labels and aligned rows."""
    ax_info.set_xticks([])
    ax_info.set_yticks([])
    ax_info.set_facecolor("#eef1f4")
    for spine in ax_info.spines.values():
        spine.set_visible(True)
        spine.set_color("#cdd3da")
        spine.set_linewidth(0.8)

    def _expand_rows(items: Sequence[Tuple[str, str]], value_width: int) -> List[Tuple[str, str]]:
        rows: List[Tuple[str, str]] = []
        for key, raw_val in items:
            label = f"{str(key).strip()}:"
            wrapped = textwrap.wrap(str(raw_val), width=max(14, int(value_width))) or [""]
            rows.append((label, wrapped[0]))
            for part in wrapped[1:]:
                rows.append(("", part))
        return rows

    def _draw_column(rows: Sequence[Tuple[str, str]], label_x: float, value_x: float, y_positions: np.ndarray) -> None:
        for idx, (label, value) in enumerate(rows):
            y = float(y_positions[idx])
            if label:
                ax_info.text(
                    label_x,
                    y,
                    label,
                    transform=ax_info.transAxes,
                    va="center",
                    ha="left",
                    fontsize=8,
                    fontweight="bold",
                    family="DejaVu Sans",
                    color="#111827",
                )
            ax_info.text(
                value_x,
                y,
                str(value),
                transform=ax_info.transAxes,
                va="center",
                ha="left",
                fontsize=8,
                family="DejaVu Sans",
                color="#1f2937",
            )

    left_rows = _expand_rows(left_items, value_width=44)
    right_rows = _expand_rows(right_items, value_width=42)
    n_rows = max(1, len(left_rows), len(right_rows))
    y_positions = np.array([0.52], dtype=np.float32) if n_rows == 1 else np.linspace(0.90, 0.12, n_rows)
    _draw_column(left_rows, label_x=0.015, value_x=0.135, y_positions=y_positions)
    _draw_column(right_rows, label_x=0.515, value_x=0.645, y_positions=y_positions)


def _format_rgb_band_summary(channels: Sequence[Dict[str, Any]]) -> str:
    """Create concise RGB band summary for titles/footer."""
    labels = ["R", "G", "B"]
    parts: List[str] = []
    for i in range(min(3, len(channels))):
        ch = channels[i]
        bidx = int(ch.get("band_idx", i + 1))
        wl = ch.get("wavelength_nm")
        if wl is not None and np.isfinite(float(wl)):
            parts.append(f"{labels[i]}=B{bidx} ({float(wl):.1f}nm)")
        else:
            parts.append(f"{labels[i]}=B{bidx}")
    return ", ".join(parts) if parts else "RGB=n/a"


def _format_scene_date_label(scene_date_utc: Optional[Any]) -> str:
    """Format scene/acquisition date for quicklook titles and footer."""
    if scene_date_utc is None:
        return "n/a"
    if isinstance(scene_date_utc, datetime):
        dt = scene_date_utc
        if dt.tzinfo is not None:
            try:
                dt = dt.astimezone(timezone.utc)
            except Exception:
                pass
        return dt.strftime("%Y-%m-%d")

    txt = str(scene_date_utc).strip()
    if not txt:
        return "n/a"

    iso_txt = txt.replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(iso_txt)
        if dt.tzinfo is not None:
            dt = dt.astimezone(timezone.utc)
        return dt.strftime("%Y-%m-%d")
    except Exception:
        pass

    m = re.search(r"\b(\d{4}-\d{2}-\d{2})\b", txt)
    if m:
        return str(m.group(1))
    m = re.search(r"\b(\d{8})\b", txt)
    if m:
        raw = str(m.group(1))
        return f"{raw[0:4]}-{raw[4:6]}-{raw[6:8]}"
    return txt


def _resolve_tiepoint_colormap_name() -> str:
    """Prefer high-luminance ramps that contrast with natural-color backgrounds."""
    if HAS_MATPLOTLIB:
        try:
            if hasattr(matplotlib, "colormaps"):
                if "plasma" in matplotlib.colormaps:
                    return "plasma"
                if "magma" in matplotlib.colormaps:
                    return "magma"
        except Exception:
            pass
    return "viridis"


def _mute_tiepoint_background_rgb(rgb: np.ndarray) -> np.ndarray:
    """Reduce background visual noise so tiepoint markers stand out."""
    arr = np.asarray(rgb, dtype=np.float32)
    if arr.ndim != 3 or arr.shape[-1] < 3:
        return arr
    # Desaturate and lightly dim the background while preserving spatial context.
    lum = (
        0.2126 * arr[..., 0]
        + 0.7152 * arr[..., 1]
        + 0.0722 * arr[..., 2]
    ).astype(np.float32, copy=False)
    neutral = np.clip(0.82 * lum + 0.05, 0.0, 1.0)
    muted = 0.64 * arr + 0.36 * neutral[..., None]
    return np.clip(0.90 * muted, 0.0, 1.0).astype(np.float32, copy=False)


def _build_tiepoint_legend_handles(n_outlier: int, n_outside_valid: int) -> List[Any]:
    """Build tiepoint legend handles; inliers are explained by colorbar only."""
    if not HAS_MATPLOTLIB:
        return []
    handles: List[Any] = []
    if int(max(0, n_outlier)) > 0:
        handles.append(
            Line2D([0], [0], marker="x", color="#ff6b6b", linestyle="None", markersize=6, label="Outliers")
        )
    if int(max(0, n_outside_valid)) > 0:
        handles.append(
            Line2D(
                [0],
                [0],
                marker="^",
                color="none",
                markerfacecolor="#f4a261",
                markeredgecolor="#111827",
                markersize=6,
                label="Outside-valid",
            )
        )
    return handles


def _build_quicklook_footer_info_rgb(
    *,
    output_png: str,
    scene_name: Optional[str],
    scene_date_utc: Optional[Any],
    raster_info: Dict[str, Any],
    tiepoint_info: Optional[Dict[str, Any]] = None,
) -> Tuple[List[Tuple[str, str]], List[Tuple[str, str]]]:
    """Build two-column footer items for scene/tiepoint quicklooks."""
    scene_label = str(scene_name).strip() if scene_name is not None else ""
    if not scene_label:
        scene_label = os.path.splitext(os.path.basename(output_png))[0]
    scene_date_label = _format_scene_date_label(scene_date_utc)

    src_w = int(raster_info.get("source_width", 0))
    src_h = int(raster_info.get("source_height", 0))
    ql_w = int(raster_info.get("quicklook_width", src_w))
    ql_h = int(raster_info.get("quicklook_height", src_h))
    crs_summary = str(raster_info.get("crs_summary", "n/a"))

    left_items: List[Tuple[str, str]] = [
        ("Scene", scene_label),
        ("Date", scene_date_label),
        ("Size", f"{src_w}x{src_h}px  (QL {ql_w}x{ql_h})"),
        ("CRS", crs_summary),
    ]
    right_items: List[Tuple[str, str]] = []

    if tiepoint_info is not None:
        left_items.extend(
            [
                (
                    "Tie points",
                    (
                        f"total={int(tiepoint_info.get('total', 0))}, "
                        f"inliers={int(tiepoint_info.get('inlier_plotted', 0))}, "
                        f"outliers={int(tiepoint_info.get('outlier_plotted', 0))}"
                    ),
                ),
                ("Outside-valid", str(int(tiepoint_info.get("outside_valid_plotted", 0)))),
            ]
        )
        right_items.extend(
            [
                (
                    "Residuals",
                    (
                        f"{tiepoint_info.get('residual_source', 'unavailable')} | "
                        f"mean={_format_optional_float(tiepoint_info.get('residual_mean_m'), '.2f')}m, "
                        f"med={_format_optional_float(tiepoint_info.get('residual_median_m'), '.2f')}m, "
                        f"rmse={_format_optional_float(tiepoint_info.get('residual_rmse_m'), '.2f')}m, "
                        f"p90={_format_optional_float(tiepoint_info.get('residual_p90_m'), '.2f')}m"
                    ),
                ),
            ]
        )
    return left_items, right_items


def _choose_scalebar_length_m(max_length_m: float) -> float:
    """Choose a rounded scale-bar length not exceeding max_length_m."""
    if max_length_m <= 0:
        return 0.0
    exponent = int(np.floor(np.log10(max_length_m)))
    base = 10.0 ** exponent
    for mult in (5.0, 2.0, 1.0):
        candidate = mult * base
        if candidate <= max_length_m:
            return float(candidate)
    return float(base / 2.0)


def _add_quicklook_scalebar(
    ax: Any,
    image_shape: Tuple[int, int],
    pixel_size_x: Any,
    pixel_size_y: Any,
) -> None:
    """Draw a simple metric scale bar when pixel size is known."""
    try:
        px_x = abs(float(pixel_size_x))
        px_y = abs(float(pixel_size_y))
    except Exception:
        return
    if not (np.isfinite(px_x) and np.isfinite(px_y)):
        return
    pixel_size = float(0.5 * (px_x + px_y))
    if pixel_size <= 0:
        return
    h, w = int(image_shape[0]), int(image_shape[1])
    max_len_m = 0.22 * w * pixel_size
    bar_len_m = _choose_scalebar_length_m(max_len_m)
    if bar_len_m <= 0:
        return
    bar_len_px = bar_len_m / pixel_size
    x0 = 0.06 * w
    y0 = 0.93 * h
    x1 = x0 + bar_len_px
    ax.plot([x0, x1], [y0, y0], color="black", linewidth=4.5, solid_capstyle="butt", zorder=8)
    ax.plot([x0, x1], [y0, y0], color="white", linewidth=2.5, solid_capstyle="butt", zorder=9)
    label = f"{bar_len_m/1000:.1f} km" if bar_len_m >= 1000.0 else f"{int(round(bar_len_m))} m"
    ax.text(
        (x0 + x1) * 0.5,
        y0 - max(10.0, h * 0.02),
        label,
        ha="center",
        va="bottom",
        fontsize=8,
        color="white",
        bbox={"facecolor": "black", "alpha": 0.55, "pad": 2.5, "linewidth": 0},
        zorder=10,
    )


def _add_quicklook_north_arrow(
    ax: Any,
    image_shape: Tuple[int, int],
    raster_info: Dict[str, Any],
) -> None:
    """Draw north arrow near top-left using affine orientation when available."""
    h, w = int(image_shape[0]), int(image_shape[1])
    if h <= 1 or w <= 1:
        return

    # Default to a simple north-up arrow when affine orientation is unavailable.
    dir_x = 0.0
    dir_y = -1.0
    try:
        a = float(raster_info.get("affine_a"))
        b = float(raster_info.get("affine_b"))
        d = float(raster_info.get("affine_d"))
        e = float(raster_info.get("affine_e"))
        m = np.array([[a, b], [d, e]], dtype=float)
        det = float(np.linalg.det(m))
        if np.isfinite(det) and abs(det) > 1e-12:
            # Pixel-space direction (dx,dy) that corresponds to map-space north (0,+1).
            vec = np.linalg.solve(m, np.array([0.0, 1.0], dtype=float))
            vx = float(vec[0])
            vy = float(vec[1])
            norm = float(np.hypot(vx, vy))
            if np.isfinite(norm) and norm > 0:
                dir_x = vx / norm
                dir_y = vy / norm
    except Exception:
        pass

    arrow_len = float(max(22.0, min(56.0, 0.08 * min(h, w))))
    x0 = float(0.08 * w)
    y0 = float(0.14 * h)
    x1 = float(x0 + dir_x * arrow_len)
    y1 = float(y0 + dir_y * arrow_len)

    ax.annotate(
        "",
        xy=(x1, y1),
        xytext=(x0, y0),
        arrowprops={"arrowstyle": "-|>", "color": "black", "linewidth": 3.3, "alpha": 0.90},
        zorder=11,
    )
    ax.annotate(
        "",
        xy=(x1, y1),
        xytext=(x0, y0),
        arrowprops={"arrowstyle": "-|>", "color": "white", "linewidth": 1.8, "alpha": 0.95},
        zorder=12,
    )
    ax.text(
        x1 + dir_x * 6.0,
        y1 + dir_y * 6.0,
        "N",
        ha="center",
        va="center",
        fontsize=9,
        fontweight="bold",
        color="white",
        bbox={"facecolor": "black", "alpha": 0.55, "pad": 1.8, "linewidth": 0},
        zorder=13,
    )


def _extract_tiepoint_plot_data_rgb(
    tie_points_df,
    x_scale: float,
    y_scale: float,
    x_offset: int,
    y_offset: int,
    valid_mask: np.ndarray,
) -> Dict[str, Any]:
    """Extract tiepoint coordinates for inliers/outliers/outside-valid with optional color values."""
    empty = {
        "total": 0,
        "inlier_total": 0,
        "inlier_plotted": 0,
        "outlier_plotted": 0,
        "outside_valid_plotted": 0,
        "outlier_counts": {},
        "x_inlier": np.array([], dtype=np.float32),
        "y_inlier": np.array([], dtype=np.float32),
        "x_outlier": np.array([], dtype=np.float32),
        "y_outlier": np.array([], dtype=np.float32),
        "x_outside_valid": np.array([], dtype=np.float32),
        "y_outside_valid": np.array([], dtype=np.float32),
        "x_color": np.array([], dtype=np.float32),
        "y_color": np.array([], dtype=np.float32),
        "color_values": np.array([], dtype=np.float32),
        "color_label": None,
        "color_cmap": None,
    }
    if tie_points_df is None:
        return empty
    try:
        df = tie_points_df.copy()
    except Exception:
        return empty
    if not hasattr(df, "__len__") or len(df) < 1:
        return empty

    out = dict(empty)
    out["total"] = int(len(df))
    h = int(valid_mask.shape[0])
    w = int(valid_mask.shape[1])

    def _classify_points(frame):
        if "X_IM" not in frame.columns or "Y_IM" not in frame.columns or len(frame) < 1:
            return {
                "x_in": np.array([], dtype=np.float32),
                "y_in": np.array([], dtype=np.float32),
                "x_outside": np.array([], dtype=np.float32),
                "y_outside": np.array([], dtype=np.float32),
                "finite": np.array([], dtype=bool),
                "inside_valid": np.array([], dtype=bool),
            }
        x_raw = np.asarray(frame["X_IM"], dtype=float)
        y_raw = np.asarray(frame["Y_IM"], dtype=float)
        finite = np.isfinite(x_raw) & np.isfinite(y_raw)
        x_sc = (x_raw[finite] * float(x_scale) - float(x_offset)).astype(np.float32, copy=False)
        y_sc = (y_raw[finite] * float(y_scale) - float(y_offset)).astype(np.float32, copy=False)
        xi = np.rint(x_sc).astype(np.int64, copy=False)
        yi = np.rint(y_sc).astype(np.int64, copy=False)
        inside = (xi >= 0) & (xi < w) & (yi >= 0) & (yi < h)
        inside_valid = np.zeros(xi.shape, dtype=bool)
        if np.any(inside):
            inside_valid[inside] = valid_mask[yi[inside], xi[inside]]
        outside_valid = ~inside | ~inside_valid
        return {
            "x_in": x_sc[~outside_valid],
            "y_in": y_sc[~outside_valid],
            "x_outside": x_sc[outside_valid],
            "y_outside": y_sc[outside_valid],
            "finite": finite,
            "inside_valid": inside_valid,
        }

    try:
        tiepoint_cmap = _resolve_tiepoint_colormap_name()
        outlier_any = np.zeros(len(df), dtype=bool)
        outlier_counts: Dict[str, int] = {}
        for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
            if col in df.columns:
                try:
                    mask = df[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
                except Exception:
                    mask = np.asarray(df[col], dtype=bool)
                outlier_counts[col] = int(np.count_nonzero(mask))
                outlier_any |= mask
        out["outlier_counts"] = outlier_counts

        df_in = df[~outlier_any] if len(df) else df
        df_out = df[outlier_any] if len(df) else df.iloc[0:0]
        out["inlier_total"] = int(len(df_in))

        cls_in = _classify_points(df_in)
        cls_out = _classify_points(df_out)
        out["x_inlier"] = cls_in["x_in"]
        out["y_inlier"] = cls_in["y_in"]
        out["x_outlier"] = cls_out["x_in"]
        out["y_outlier"] = cls_out["y_in"]
        out["x_outside_valid"] = np.concatenate([cls_in["x_outside"], cls_out["x_outside"]]).astype(
            np.float32, copy=False
        )
        out["y_outside_valid"] = np.concatenate([cls_in["y_outside"], cls_out["y_outside"]]).astype(
            np.float32, copy=False
        )
        out["inlier_plotted"] = int(out["x_inlier"].size)
        out["outlier_plotted"] = int(out["x_outlier"].size)
        out["outside_valid_plotted"] = int(out["x_outside_valid"].size)

        if len(df_in) > 0 and "X_IM" in df_in.columns and "Y_IM" in df_in.columns:
            def _build_color(values: np.ndarray, label: str, cmap_name: str):
                finite = cls_in.get("finite", np.array([], dtype=bool))
                inside_valid = cls_in.get("inside_valid", np.array([], dtype=bool))
                if finite.size < 1 or values.size != len(df_in):
                    return
                finite_idx = np.where(finite)[0]
                if finite_idx.size != inside_valid.size:
                    return
                inside_full = np.zeros(len(df_in), dtype=bool)
                inside_full[finite_idx] = inside_valid
                use = inside_full & np.isfinite(values)
                if np.count_nonzero(use) < 2:
                    return
                vals = values[use].astype(np.float32, copy=False)
                if not (float(np.nanmax(vals)) > float(np.nanmin(vals))):
                    return
                x_raw = np.asarray(df_in["X_IM"], dtype=float)[use]
                y_raw = np.asarray(df_in["Y_IM"], dtype=float)[use]
                out["x_color"] = (x_raw * float(x_scale) - float(x_offset)).astype(np.float32, copy=False)
                out["y_color"] = (y_raw * float(y_scale) - float(y_offset)).astype(np.float32, copy=False)
                out["color_values"] = vals
                out["color_label"] = label
                out["color_cmap"] = cmap_name

            if "RELIABILITY" in df_in.columns:
                _build_color(np.asarray(df_in["RELIABILITY"], dtype=float), "Reliability (%)", tiepoint_cmap)
            elif "ABS_SHIFT_M" in df_in.columns:
                _build_color(np.asarray(df_in["ABS_SHIFT_M"], dtype=float), "Residual shift (m)", tiepoint_cmap)
            elif "X_SHIFT_M" in df_in.columns and "Y_SHIFT_M" in df_in.columns:
                x_shift = np.asarray(df_in["X_SHIFT_M"], dtype=float)
                y_shift = np.asarray(df_in["Y_SHIFT_M"], dtype=float)
                _build_color(np.sqrt(x_shift**2 + y_shift**2), "Residual shift (m)", tiepoint_cmap)
    except Exception:
        return empty
    return out


def _write_quicklook_png(
    raster_path: str,
    output_png: str,
    title: str,
    band_idx: int,
    nodata: float = PROCESSING_NODATA,
    scene_name: Optional[str] = None,
    scene_date_utc: Optional[Any] = None,
    band_wavelength_nm: Optional[float] = None,
    max_quicklook_dim: int = DEFAULT_QUICKLOOK_MAX_DIM,
    rgb_band_indices: Optional[Tuple[int, int, int]] = None,
    rgb_band_wavelengths_nm: Optional[Sequence[Optional[float]]] = None,
    quicklook_percentiles: Tuple[float, float] = (2.0, 98.0),
    quicklook_gamma: float = 1.0,
    quicklook_dpi: int = 220,
    quicklook_crop_to_valid: bool = False,
    quicklook_scalebar: bool = True,
) -> None:
    """Write RGB scene quicklook PNG."""
    if not HAS_MATPLOTLIB:
        raise RuntimeError("matplotlib is required for mandatory quicklook generation.")

    if rgb_band_indices is None:
        rgb_band_indices = (int(band_idx), int(band_idx), int(band_idx))
    rgb, valid, raster_info = _load_quicklook_rgb_composite(
        raster_path,
        rgb_band_indices=rgb_band_indices,
        rgb_band_wavelengths_nm=rgb_band_wavelengths_nm,
        nodata=nodata,
        max_quicklook_dim=max_quicklook_dim,
        percentiles=quicklook_percentiles,
        gamma=quicklook_gamma,
        crop_to_valid=quicklook_crop_to_valid,
    )
    if not np.any(valid):
        raise RuntimeError(f"No valid pixels available for quicklook: {raster_path}")

    raster_info["valid_count"] = int(np.count_nonzero(valid))
    raster_info["total_count"] = int(valid.size)
    left_items, right_items = _build_quicklook_footer_info_rgb(
        output_png=output_png,
        scene_name=scene_name,
        scene_date_utc=scene_date_utc,
        raster_info=raster_info,
    )
    scene_date_label = _format_scene_date_label(scene_date_utc)

    os.makedirs(os.path.dirname(output_png) or ".", exist_ok=True)
    fig = plt.figure(figsize=(10, 9.6), dpi=int(max(120, quicklook_dpi)), constrained_layout=True, facecolor="white")
    gs = fig.add_gridspec(2, 2, height_ratios=[20, 5], width_ratios=[20, 1.3], hspace=0.02, wspace=0.03)
    ax = fig.add_subplot(gs[0, 0])
    cax = fig.add_subplot(gs[0, 1])
    ax_info = fig.add_subplot(gs[1, :])
    try:
        ax.set_facecolor("#0b0b0b")
        ax.imshow(rgb, interpolation="nearest")
        ax.set_title(f"Quicklook | {scene_date_label}", fontsize=13, fontweight="semibold", color="#111827")
        ax.set_axis_off()
        cax.set_axis_off()
        if bool(quicklook_scalebar):
            _add_quicklook_scalebar(
                ax,
                image_shape=(int(rgb.shape[0]), int(rgb.shape[1])),
                pixel_size_x=raster_info.get("pixel_size_x"),
                pixel_size_y=raster_info.get("pixel_size_y"),
            )
        _add_quicklook_north_arrow(
            ax,
            image_shape=(int(rgb.shape[0]), int(rgb.shape[1])),
            raster_info=raster_info,
        )

        _draw_quicklook_footer(ax_info, left_items, right_items)
        fig.savefig(output_png, dpi=int(max(120, quicklook_dpi)), facecolor="white", transparent=False)
    finally:
        plt.close(fig)


def _write_tiepoint_quicklook_png(
    background_raster: str,
    output_png: str,
    title: str,
    band_idx: int,
    tie_points_df,
    nodata: float = PROCESSING_NODATA,
    scene_name: Optional[str] = None,
    scene_date_utc: Optional[Any] = None,
    band_wavelength_nm: Optional[float] = None,
    max_quicklook_dim: int = DEFAULT_QUICKLOOK_MAX_DIM,
    rgb_band_indices: Optional[Tuple[int, int, int]] = None,
    rgb_band_wavelengths_nm: Optional[Sequence[Optional[float]]] = None,
    quicklook_percentiles: Tuple[float, float] = (2.0, 98.0),
    quicklook_gamma: float = 1.0,
    quicklook_dpi: int = 220,
    quicklook_crop_to_valid: bool = False,
    quicklook_scalebar: bool = True,
) -> int:
    """Write tiepoint quicklook PNG on top of RGB composite."""
    if not HAS_MATPLOTLIB:
        raise RuntimeError("matplotlib is required for mandatory quicklook generation.")

    if rgb_band_indices is None:
        rgb_band_indices = (int(band_idx), int(band_idx), int(band_idx))
    rgb, valid, raster_info = _load_quicklook_rgb_composite(
        background_raster,
        rgb_band_indices=rgb_band_indices,
        rgb_band_wavelengths_nm=rgb_band_wavelengths_nm,
        nodata=nodata,
        max_quicklook_dim=max_quicklook_dim,
        percentiles=quicklook_percentiles,
        gamma=quicklook_gamma,
        crop_to_valid=quicklook_crop_to_valid,
    )
    if not np.any(valid):
        rgb = np.zeros_like(rgb, dtype=np.float32)
        rgb[..., :] = np.array([0.08, 0.08, 0.08], dtype=np.float32)

    tp_plot = _extract_tiepoint_plot_data_rgb(
        tie_points_df=tie_points_df,
        x_scale=float(raster_info.get("x_scale", 1.0)),
        y_scale=float(raster_info.get("y_scale", 1.0)),
        x_offset=int(raster_info.get("x_offset", 0)),
        y_offset=int(raster_info.get("y_offset", 0)),
        valid_mask=valid,
    )
    px_x = raster_info.get("pixel_size_x")
    px_y = raster_info.get("pixel_size_y")
    pixel_size_m = 30.0
    try:
        if np.isfinite(float(px_x)) and np.isfinite(float(px_y)):
            pixel_size_m = float(0.5 * (abs(float(px_x)) + abs(float(px_y))))
    except Exception:
        pass
    residuals = _compute_tiepoint_residuals(tie_points_df, pixel_size_m=pixel_size_m)
    tiepoint_info = {
        "total": int(tp_plot.get("total", 0)),
        "inlier_plotted": int(tp_plot.get("inlier_plotted", 0)),
        "outlier_plotted": int(tp_plot.get("outlier_plotted", 0)),
        "outside_valid_plotted": int(tp_plot.get("outside_valid_plotted", 0)),
        "outlier_counts": dict(tp_plot.get("outlier_counts", {})),
        "residual_source": residuals.get("residual_source"),
        "residual_mean_m": residuals.get("residual_mean_m"),
        "residual_median_m": residuals.get("residual_median_m"),
        "residual_rmse_m": residuals.get("residual_rmse_m"),
        "residual_p90_m": residuals.get("residual_p90_m"),
    }
    raster_info["valid_count"] = int(np.count_nonzero(valid))
    raster_info["total_count"] = int(valid.size)
    left_items, right_items = _build_quicklook_footer_info_rgb(
        output_png=output_png,
        scene_name=scene_name,
        scene_date_utc=scene_date_utc,
        raster_info=raster_info,
        tiepoint_info=tiepoint_info,
    )
    scene_date_label = _format_scene_date_label(scene_date_utc)

    os.makedirs(os.path.dirname(output_png) or ".", exist_ok=True)
    fig = plt.figure(figsize=(10, 9.6), dpi=int(max(120, quicklook_dpi)), constrained_layout=True, facecolor="white")
    gs = fig.add_gridspec(2, 2, height_ratios=[20, 5], width_ratios=[20, 1.3], hspace=0.02, wspace=0.03)
    ax = fig.add_subplot(gs[0, 0])
    cax = fig.add_subplot(gs[0, 1])
    ax_info = fig.add_subplot(gs[1, :])
    try:
        ax.set_facecolor("#0b0b0b")
        rgb_display = _mute_tiepoint_background_rgb(rgb)
        ax.imshow(rgb_display, interpolation="nearest")

        x_in = tp_plot.get("x_inlier", np.array([], dtype=np.float32))
        y_in = tp_plot.get("y_inlier", np.array([], dtype=np.float32))
        x_out = tp_plot.get("x_outlier", np.array([], dtype=np.float32))
        y_out = tp_plot.get("y_outlier", np.array([], dtype=np.float32))
        x_ov = tp_plot.get("x_outside_valid", np.array([], dtype=np.float32))
        y_ov = tp_plot.get("y_outside_valid", np.array([], dtype=np.float32))
        n_in = int(x_in.size)
        n_out = int(x_out.size)
        n_ov = int(x_ov.size)
        pt_size_base = float(max(6.0, min(18.0, 2000.0 / max(1, n_in))))
        pt_size = float(max(7.2, min(24.0, pt_size_base * 1.2)))
        fill_size = float(max(4.8, pt_size - 1.8))
        ring_white = float(fill_size + 4.8)
        ring_black = float(fill_size + 8.8)

        if n_in > 0:
            ax.scatter(
                x_in,
                y_in,
                s=ring_black,
                c="black",
                alpha=0.82,
                linewidths=0,
                marker="o",
                zorder=4,
            )
            ax.scatter(
                x_in,
                y_in,
                s=ring_white,
                c="white",
                alpha=0.98,
                linewidths=0,
                marker="o",
                zorder=5,
            )
            ax.scatter(
                x_in,
                y_in,
                s=fill_size,
                c="#5ec5f4",
                alpha=0.74,
                linewidths=0,
                marker="o",
                zorder=6,
            )

            color_values = tp_plot.get("color_values", np.array([], dtype=np.float32))
            x_color = tp_plot.get("x_color", np.array([], dtype=np.float32))
            y_color = tp_plot.get("y_color", np.array([], dtype=np.float32))
            color_scatter = None
            if color_values.size > 0 and x_color.size == color_values.size and y_color.size == color_values.size:
                color_scatter = ax.scatter(
                    x_color,
                    y_color,
                    s=fill_size,
                    c=color_values,
                    cmap=tp_plot.get("color_cmap", "viridis"),
                    alpha=0.96,
                    linewidths=0,
                    marker="o",
                    zorder=7,
                )
            if color_scatter is not None:
                cb = fig.colorbar(color_scatter, cax=cax)
                cb.set_label(str(tp_plot.get("color_label", "Tie-point value")), fontsize=8)
                cb.ax.tick_params(labelsize=7)
            else:
                cax.set_axis_off()
        else:
            cax.set_axis_off()

        if n_out > 0:
            ax.scatter(
                x_out,
                y_out,
                s=max(5.0, pt_size - 1.0),
                c="#ff6b6b",
                alpha=0.55,
                linewidths=0.9,
                marker="x",
                zorder=7,
            )
        if n_ov > 0:
            ax.scatter(
                x_ov,
                y_ov,
                s=pt_size + 6.0,
                c="#f4a261",
                alpha=0.85,
                linewidths=0.6,
                edgecolors="#111827",
                marker="^",
                zorder=7,
            )
        legend_handles = _build_tiepoint_legend_handles(n_outlier=n_out, n_outside_valid=n_ov)
        if legend_handles:
            ax.legend(
                handles=legend_handles,
                loc="lower right",
                frameon=True,
                framealpha=0.78,
                fontsize=8,
            )

        if n_in > 0:
            ax.set_title(
                f"Tie points quicklook | {scene_date_label}",
                fontsize=13,
                fontweight="semibold",
            )
        else:
            ax.text(
                0.5,
                0.5,
                "No usable tie points\n(placeholder quicklook)",
                transform=ax.transAxes,
                ha="center",
                va="center",
                fontsize=12,
                color="#f8f9fa",
                bbox={"facecolor": "black", "alpha": 0.65, "pad": 8},
            )
            ax.set_title(f"Tie points quicklook | {scene_date_label}", fontsize=13, fontweight="semibold")
        ax.set_axis_off()
        if bool(quicklook_scalebar):
            _add_quicklook_scalebar(
                ax,
                image_shape=(int(rgb.shape[0]), int(rgb.shape[1])),
                pixel_size_x=raster_info.get("pixel_size_x"),
                pixel_size_y=raster_info.get("pixel_size_y"),
            )
        _add_quicklook_north_arrow(
            ax,
            image_shape=(int(rgb.shape[0]), int(rgb.shape[1])),
            raster_info=raster_info,
        )

        _draw_quicklook_footer(ax_info, left_items, right_items)
        fig.savefig(output_png, dpi=int(max(120, quicklook_dpi)), facecolor="white", transparent=False)
    finally:
        plt.close(fig)

    return int(tp_plot.get("inlier_plotted", 0))


def _generate_mandatory_quicklooks(
    scene_name: str,
    quicklooks_dir: str,
    scene_raster_path: str,
    tie_points_df,
    generate_tiepoints_png: bool,
    wl: np.ndarray,
    target_wl_nm: float,
    max_quicklook_dim: int = DEFAULT_QUICKLOOK_MAX_DIM,
    quicklook_rgb_targets_nm: Tuple[float, float, float] = (660.0, 550.0, 480.0),
    quicklook_percentiles: Tuple[float, float] = (2.0, 98.0),
    quicklook_gamma: float = 1.0,
    quicklook_dpi: int = 220,
    quicklook_crop_to_valid: bool = False,
    quicklook_scalebar: bool = True,
    rgb_source_path: Optional[str] = None,
    scene_date_utc: Optional[Any] = None,
) -> Dict[str, Any]:
    """Generate mandatory scene/tiepoint quicklooks using RGB composite."""
    wl_arr = np.asarray(wl, dtype=float).reshape(-1)
    ql_percentiles = _coerce_percentiles(quicklook_percentiles, default=(2.0, 98.0))
    ql_targets = _coerce_float_triplet(quicklook_rgb_targets_nm, default=(660.0, 550.0, 480.0))

    rgb_base_path = scene_raster_path
    if rgb_source_path is not None and str(rgb_source_path).strip():
        if os.path.exists(str(rgb_source_path)):
            rgb_base_path = str(rgb_source_path)
        else:
            logger.warning(
                _fmt_issue(
                    "QUICKLOOK",
                    f"Configured quicklook_rgb_source_path not found: {rgb_source_path}; using coreg output raster.",
                )
            )

    with rasterio.open(rgb_base_path) as src:
        band_count = int(src.count)
    rgb_indices, rgb_wavelengths, rgb_source = _resolve_quicklook_rgb_bands(
        wl_arr if wl_arr.size > 0 else None,
        band_count=band_count,
        targets_nm=ql_targets,
    )

    if wl_arr.size > 0 and np.any(np.isfinite(wl_arr)):
        band_idx = _pick_band_index_for_wavelength(float(target_wl_nm), wl_arr)
        band_wl_nm = float(wl_arr[band_idx - 1]) if np.isfinite(wl_arr[band_idx - 1]) else None
    else:
        band_idx = int(rgb_indices[1])
        band_wl_nm = None

    os.makedirs(quicklooks_dir, exist_ok=True)
    out_scene = os.path.join(quicklooks_dir, f"{scene_name}_quicklook.png")
    out_tp = os.path.join(quicklooks_dir, f"{scene_name}_tiepoints.png") if generate_tiepoints_png else None

    _write_quicklook_png(
        rgb_base_path,
        out_scene,
        "Scene quicklook",
        band_idx,
        scene_name=scene_name,
        scene_date_utc=scene_date_utc,
        band_wavelength_nm=band_wl_nm,
        max_quicklook_dim=max_quicklook_dim,
        rgb_band_indices=rgb_indices,
        rgb_band_wavelengths_nm=rgb_wavelengths,
        quicklook_percentiles=ql_percentiles,
        quicklook_gamma=float(max(1e-6, quicklook_gamma)),
        quicklook_dpi=int(max(120, quicklook_dpi)),
        quicklook_crop_to_valid=bool(quicklook_crop_to_valid),
        quicklook_scalebar=bool(quicklook_scalebar),
    )
    tiepoint_count = 0
    if generate_tiepoints_png:
        tiepoint_count = _write_tiepoint_quicklook_png(
            rgb_base_path,
            out_tp,
            "Tie points quicklook",
            band_idx,
            tie_points_df,
            scene_name=scene_name,
            scene_date_utc=scene_date_utc,
            band_wavelength_nm=band_wl_nm,
            max_quicklook_dim=max_quicklook_dim,
            rgb_band_indices=rgb_indices,
            rgb_band_wavelengths_nm=rgb_wavelengths,
            quicklook_percentiles=ql_percentiles,
            quicklook_gamma=float(max(1e-6, quicklook_gamma)),
            quicklook_dpi=int(max(120, quicklook_dpi)),
            quicklook_crop_to_valid=bool(quicklook_crop_to_valid),
            quicklook_scalebar=bool(quicklook_scalebar),
        )

    outputs = {"scene": out_scene, "tiepoints": out_tp}
    for label, out_path in outputs.items():
        if out_path is None:
            continue
        if not os.path.exists(out_path):
            raise RuntimeError(f"Quicklook generation failed ({label}): {out_path}")

    return {
        "outputs": outputs,
        "band_idx": int(band_idx),
        "band_wavelength_nm": None if band_wl_nm is None else float(band_wl_nm),
        "rgb_band_indices": [int(rgb_indices[0]), int(rgb_indices[1]), int(rgb_indices[2])],
        "rgb_band_wavelength_nm": [
            None if rgb_wavelengths[0] is None else float(rgb_wavelengths[0]),
            None if rgb_wavelengths[1] is None else float(rgb_wavelengths[1]),
            None if rgb_wavelengths[2] is None else float(rgb_wavelengths[2]),
        ],
        "rgb_source": str(rgb_source),
        "tiepoints_enabled": bool(generate_tiepoints_png),
        "tiepoints_plotted": tiepoint_count,
    }


def _safe_remove_shapefile_family(shapefile_path: str) -> None:
    """Remove .shp family sidecars to allow deterministic overwrite."""
    if not shapefile_path:
        return
    shp = str(shapefile_path)
    stem, ext = os.path.splitext(shp)
    base = stem if ext.lower() == ".shp" else shp
    for suffix in (".shp", ".shx", ".dbf", ".prj", ".cpg", ".qix", ".fix"):
        target = f"{base}{suffix}"
        if os.path.exists(target):
            os.remove(target)


def _ensure_abs_shift_column(df_in):
    """Ensure ABS_SHIFT exists; derive from XY shift columns when needed."""
    if df_in is None:
        raise RuntimeError("Cannot derive ABS_SHIFT from empty tiepoint table.")
    try:
        df = df_in.copy()
    except Exception as exc:
        raise RuntimeError(f"Unable to copy tiepoint table: {exc}") from exc
    if not hasattr(df, "__len__") or len(df) < 1:
        raise RuntimeError("Cannot derive ABS_SHIFT: tiepoint table is empty.")

    basis = "existing_abs_shift"
    warning_message = None
    abs_shift = None

    if "ABS_SHIFT" in df.columns:
        try:
            abs_shift = np.asarray(df["ABS_SHIFT"], dtype=float)
            if not np.any(np.isfinite(abs_shift)):
                abs_shift = None
        except Exception:
            abs_shift = None

    if abs_shift is None and "ABS_SHIFT_M" in df.columns:
        try:
            abs_shift = np.asarray(df["ABS_SHIFT_M"], dtype=float)
            if np.any(np.isfinite(abs_shift)):
                basis = "abs_shift_m"
            else:
                abs_shift = None
        except Exception:
            abs_shift = None

    computed_m = None
    if "X_SHIFT_M" in df.columns and "Y_SHIFT_M" in df.columns:
        try:
            x_shift_m = np.asarray(df["X_SHIFT_M"], dtype=float)
            y_shift_m = np.asarray(df["Y_SHIFT_M"], dtype=float)
            computed_m = np.sqrt(x_shift_m**2 + y_shift_m**2)
        except Exception:
            computed_m = None

    computed_px = None
    if "X_SHIFT_PX" in df.columns and "Y_SHIFT_PX" in df.columns:
        try:
            x_shift_px = np.asarray(df["X_SHIFT_PX"], dtype=float)
            y_shift_px = np.asarray(df["Y_SHIFT_PX"], dtype=float)
            computed_px = np.sqrt(x_shift_px**2 + y_shift_px**2)
        except Exception:
            computed_px = None

    if abs_shift is None and computed_m is not None:
        abs_shift = computed_m
        basis = "xy_shift_m"
    elif abs_shift is None and computed_px is not None:
        abs_shift = computed_px
        basis = "xy_shift_px"
        warning_message = "ABS_SHIFT derived from X_SHIFT_PX/Y_SHIFT_PX because meter shifts are unavailable."

    if abs_shift is None:
        raise RuntimeError("Cannot derive ABS_SHIFT: missing ABS_SHIFT and XY shift columns.")

    abs_shift = np.asarray(abs_shift, dtype=float)
    if computed_m is not None:
        missing_abs = ~np.isfinite(abs_shift)
        if np.any(missing_abs):
            abs_shift = abs_shift.copy()
            abs_shift[missing_abs] = computed_m[missing_abs]
            basis = "xy_shift_m_fill"
    elif computed_px is not None:
        missing_abs = ~np.isfinite(abs_shift)
        if np.any(missing_abs):
            abs_shift = abs_shift.copy()
            abs_shift[missing_abs] = computed_px[missing_abs]
            basis = "xy_shift_px_fill"
            warning_message = "ABS_SHIFT partially filled from X_SHIFT_PX/Y_SHIFT_PX."

    df["ABS_SHIFT"] = abs_shift
    if "ABS_SHIFT_M" not in df.columns and basis.startswith("xy_shift_m"):
        df["ABS_SHIFT_M"] = abs_shift
    return df, basis, warning_message


def _export_displacement_shapefile_from_df(tie_points_df, shapefile_path: str, fallback_crs=None) -> Dict[str, Any]:
    """Export displacement vectors from a tiepoint table to a point shapefile."""
    if not HAS_GEOPANDAS:
        raise RuntimeError("GeoPandas is required to export displacement shapefiles.")
    if tie_points_df is None or len(tie_points_df) < 1:
        raise RuntimeError("No tie points available for displacement shapefile export.")

    df_abs, basis, warning_message = _ensure_abs_shift_column(tie_points_df)
    if isinstance(df_abs, gpd.GeoDataFrame):
        gdf = df_abs.copy()
        if gdf.crs is None and fallback_crs is not None:
            gdf = gdf.set_crs(fallback_crs, allow_override=True)
    elif "geometry" in df_abs.columns:
        gdf = gpd.GeoDataFrame(df_abs.copy(), geometry="geometry", crs=fallback_crs)
    else:
        if "X_MAP" not in df_abs.columns or "Y_MAP" not in df_abs.columns:
            raise RuntimeError("Cannot export merged tiepoints: missing X_MAP/Y_MAP columns.")
        x_map = np.asarray(df_abs["X_MAP"], dtype=float)
        y_map = np.asarray(df_abs["Y_MAP"], dtype=float)
        valid = np.isfinite(x_map) & np.isfinite(y_map)
        if not np.any(valid):
            raise RuntimeError("No finite X_MAP/Y_MAP coordinates available for shapefile export.")
        frame_valid = df_abs.loc[valid].copy()
        geom = gpd.points_from_xy(x_map[valid], y_map[valid])
        gdf = gpd.GeoDataFrame(frame_valid, geometry=geom, crs=fallback_crs)

    if len(gdf) < 1:
        raise RuntimeError("No valid tie points available for shapefile export.")

    os.makedirs(os.path.dirname(shapefile_path) or ".", exist_ok=True)
    _safe_remove_shapefile_family(shapefile_path)
    gdf.to_file(shapefile_path, driver="ESRI Shapefile")
    return {
        "path": shapefile_path,
        "count": int(len(gdf)),
        "abs_shift_basis": basis,
        "warning": warning_message,
    }


def _rank_displacement_vectors_for_plot(df):
    """Attach deterministic ranking fields used by displacement-plot dedup/sampling."""
    import pandas as pd

    work = df.copy()
    n_rows = int(len(work))

    def _numeric_with_default(column: str, default: float) -> np.ndarray:
        if column not in work.columns:
            return np.full(n_rows, float(default), dtype=float)
        values = pd.to_numeric(work[column], errors="coerce").to_numpy(dtype=float)
        invalid = ~np.isfinite(values)
        if np.any(invalid):
            values = values.copy()
            values[invalid] = float(default)
        return values

    work["_PLOT_SCORE"] = _numeric_with_default("QUALITY_SCORE", -np.inf)
    work["_PLOT_REL"] = _numeric_with_default("RELIABILITY", -np.inf)
    work["_PLOT_ABS"] = _numeric_with_default("ABS_SHIFT", np.inf)
    work["_PLOT_ERR"] = _numeric_with_default("LAST_ERR", np.inf)
    work["_PLOT_ROW"] = np.arange(n_rows, dtype=np.int64)
    work = work.sort_values(
        ["_PLOT_SCORE", "_PLOT_REL", "_PLOT_ABS", "_PLOT_ERR", "_PLOT_ROW"],
        ascending=[False, False, True, True, True],
        kind="mergesort",
    )
    return work


def _deduplicate_displacement_vectors_for_plot(
    df,
    image_rounding_px: float = 1.0,
    map_rounding_m: float = 1.0,
) -> Tuple[Any, str, int]:
    """Keep a single best vector per tie-point key for displacement plotting."""
    import pandas as pd

    work = _rank_displacement_vectors_for_plot(df)
    if len(work) == 0:
        return work, "none", 0

    def _clean_key_series(series) -> Tuple[pd.Series, pd.Series]:
        text = series.astype(str).str.strip()
        valid = series.notna() & (text != "") & ~text.str.lower().isin(["nan", "none", "null"])
        return text.astype(str), valid.astype(bool)

    key_source = "ROW_INDEX"
    key_series: Optional[pd.Series] = None

    # Prefer native tie-point identifiers when available; consensus IDs can merge nearby points.
    if "POINT_ID" in work.columns:
        key, valid = _clean_key_series(work["POINT_ID"])
        if bool(valid.any()):
            key_series = key.where(valid, "")
            key_source = "POINT_ID"

    if key_series is None and "CONSENSUS_GROUP_ID" in work.columns:
        key, valid = _clean_key_series(work["CONSENSUS_GROUP_ID"])
        if bool(valid.any()):
            key_series = key.where(valid, "")
            key_source = "CONSENSUS_GROUP_ID"

    if key_series is None and "X_IM" in work.columns and "Y_IM" in work.columns:
        round_px = float(max(0.1, image_rounding_px))
        x = pd.to_numeric(work["X_IM"], errors="coerce")
        y = pd.to_numeric(work["Y_IM"], errors="coerce")
        xq = np.round(x / round_px).astype("Int64")
        yq = np.round(y / round_px).astype("Int64")
        valid = xq.notna() & yq.notna()
        key_series = ("IM_" + xq.astype(str) + "_" + yq.astype(str)).where(valid, "")
        key_source = "ROUNDED_XY_IM"

    if key_series is None:
        x_map_col = "X_MAP" if "X_MAP" in work.columns else ("Map X" if "Map X" in work.columns else None)
        y_map_col = "Y_MAP" if "Y_MAP" in work.columns else ("Map Y" if "Map Y" in work.columns else None)
        if x_map_col is not None and y_map_col is not None:
            round_m = float(max(0.1, map_rounding_m))
            x = pd.to_numeric(work[str(x_map_col)], errors="coerce")
            y = pd.to_numeric(work[str(y_map_col)], errors="coerce")
            xq = np.round(x / round_m).astype("Int64")
            yq = np.round(y / round_m).astype("Int64")
            valid = xq.notna() & yq.notna()
            key_series = ("MAP_" + xq.astype(str) + "_" + yq.astype(str)).where(valid, "")
            key_source = "ROUNDED_XY_MAP"

    if key_series is None:
        key_series = pd.Series([""] * len(work), index=work.index, dtype=str)

    missing = key_series.isna() | (key_series.astype(str).str.strip() == "")
    if bool(missing.any()):
        fallback = pd.Series(np.arange(len(work)), index=work.index).astype(str)
        key_series = key_series.astype(str).where(~missing, "ROW_" + fallback)

    before = int(len(work))
    work["_PLOT_DEDUP_KEY"] = key_series.astype(str)
    work = work.drop_duplicates(subset="_PLOT_DEDUP_KEY", keep="first")
    dropped = int(before - len(work))

    drop_cols = [c for c in ["_PLOT_SCORE", "_PLOT_REL", "_PLOT_ABS", "_PLOT_ERR", "_PLOT_ROW", "_PLOT_DEDUP_KEY"] if c in work.columns]
    if drop_cols:
        work = work.drop(columns=drop_cols)
    return work, key_source, dropped


def _stratified_sample_displacement_vectors_for_plot(
    df,
    x_col: str,
    y_col: str,
    max_vectors: int,
    grid_rows: int = 6,
    grid_cols: int = 6,
) -> Tuple[Any, Dict[str, Any]]:
    """Apply spatially balanced subsampling for displacement-vector plotting."""
    import pandas as pd

    out: Dict[str, Any] = {
        "method": "none",
        "occupied_cells_total": 0,
        "occupied_cells_selected": 0,
        "warning": None,
    }
    if df is None or len(df) == 0:
        return df, out

    keep_n = int(max_vectors)
    if keep_n <= 0 or len(df) <= keep_n:
        return df, out

    if x_col not in df.columns or y_col not in df.columns:
        ranked = _rank_displacement_vectors_for_plot(df).head(keep_n).copy()
        drop_cols = [c for c in ["_PLOT_SCORE", "_PLOT_REL", "_PLOT_ABS", "_PLOT_ERR", "_PLOT_ROW"] if c in ranked.columns]
        if drop_cols:
            ranked = ranked.drop(columns=drop_cols)
        out["method"] = "top_quality_fallback"
        out["warning"] = f"Stratified sampling fallback: missing coordinate columns {x_col}/{y_col}."
        return ranked, out

    ranked = _rank_displacement_vectors_for_plot(df)
    x = pd.to_numeric(ranked[x_col], errors="coerce").to_numpy(dtype=float)
    y = pd.to_numeric(ranked[y_col], errors="coerce").to_numpy(dtype=float)
    finite = np.isfinite(x) & np.isfinite(y)
    if not np.any(finite):
        sampled = ranked.head(keep_n).copy()
        drop_cols = [c for c in ["_PLOT_SCORE", "_PLOT_REL", "_PLOT_ABS", "_PLOT_ERR", "_PLOT_ROW"] if c in sampled.columns]
        if drop_cols:
            sampled = sampled.drop(columns=drop_cols)
        out["method"] = "top_quality_fallback"
        out["warning"] = "Stratified sampling fallback: no finite plotting coordinates."
        return sampled, out

    valid = ranked.loc[finite].copy()
    xv = x[finite]
    yv = y[finite]

    gx = max(1, int(grid_cols))
    gy = max(1, int(grid_rows))
    x_min, x_max = float(np.min(xv)), float(np.max(xv))
    y_min, y_max = float(np.min(yv)), float(np.max(yv))
    x_span = max(1e-6, x_max - x_min)
    y_span = max(1e-6, y_max - y_min)
    cell_x = np.floor((xv - x_min) / x_span * gx).astype(int)
    cell_y = np.floor((yv - y_min) / y_span * gy).astype(int)
    cell_x = np.clip(cell_x, 0, gx - 1)
    cell_y = np.clip(cell_y, 0, gy - 1)
    valid["_PLOT_CELL"] = (cell_y * gx + cell_x).astype(int)
    out["occupied_cells_total"] = int(valid["_PLOT_CELL"].nunique())

    valid["_PLOT_CELL_RANK"] = valid.groupby("_PLOT_CELL", sort=False).cumcount()
    max_cell_rank = int(valid["_PLOT_CELL_RANK"].max()) if len(valid) > 0 else -1
    selected_indices: List[Any] = []
    for rank_i in range(max_cell_rank + 1):
        if len(selected_indices) >= keep_n:
            break
        level = valid[valid["_PLOT_CELL_RANK"] == rank_i]
        if len(level) == 0:
            continue
        remaining = keep_n - len(selected_indices)
        selected_indices.extend(level.head(remaining).index.tolist())

    sampled = valid.loc[selected_indices].copy()
    if len(sampled) < keep_n:
        remainder = valid.loc[~valid.index.isin(sampled.index)].head(keep_n - len(sampled))
        if len(remainder) > 0:
            sampled = pd.concat([sampled, remainder], axis=0)
    sampled = sampled.sort_values(
        ["_PLOT_SCORE", "_PLOT_REL", "_PLOT_ABS", "_PLOT_ERR", "_PLOT_ROW"],
        ascending=[False, False, True, True, True],
        kind="mergesort",
    )

    out["method"] = "stratified_grid"
    out["occupied_cells_selected"] = int(sampled["_PLOT_CELL"].nunique()) if len(sampled) > 0 else 0
    drop_cols = [c for c in ["_PLOT_SCORE", "_PLOT_REL", "_PLOT_ABS", "_PLOT_ERR", "_PLOT_ROW", "_PLOT_CELL", "_PLOT_CELL_RANK"] if c in sampled.columns]
    if drop_cols:
        sampled = sampled.drop(columns=drop_cols)
    return sampled, out


def _write_displacement_vector_cartography_png(
    visualization_df,
    output_png: str,
    scene_name: Optional[str] = None,
    scene_date_utc: Optional[Any] = None,
    max_vectors: int = 4000,
    basemap_raster_path: Optional[str] = None,
    basemap_rgb_band_indices: Optional[Tuple[int, int, int]] = None,
    quiver_cmap: str = "RdYlGn_r",
    arrow_len_min_frac: float = 0.030,
    arrow_len_max_frac: float = 0.070,
    sampling_grid_rows: int = 6,
    sampling_grid_cols: int = 6,
) -> Dict[str, Any]:
    """Render a raw analytical displacement-vector map as PNG."""
    if not HAS_MATPLOTLIB:
        raise RuntimeError("matplotlib is required for displacement cartography.")
    if visualization_df is None or len(visualization_df) < 1:
        raise RuntimeError("No tie points available for displacement cartography.")

    df_abs, abs_basis, warning_message = _ensure_abs_shift_column(visualization_df)
    count_raw = int(len(df_abs))
    dedup_key_source = "none"
    dedup_dropped = 0
    sampling_info: Dict[str, Any] = {
        "method": "none",
        "occupied_cells_total": 0,
        "occupied_cells_selected": 0,
        "warning": None,
    }
    arrow_len_max_rendered = 0.0
    arrow_len_max_frac_rendered = 0.0
    min_frac_used = float(max(1e-6, arrow_len_min_frac))
    max_frac_used = float(max(min_frac_used, arrow_len_max_frac))

    for col in ("L1_OUTLIER", "L2_OUTLIER", "L3_OUTLIER"):
        if col in df_abs.columns:
            try:
                out_mask = df_abs[col].astype("boolean").fillna(False).to_numpy(dtype=bool)
                df_abs = df_abs.loc[~out_mask].copy()
            except Exception:
                pass

    df_abs, dedup_key_source, dedup_dropped = _deduplicate_displacement_vectors_for_plot(df_abs)

    if len(df_abs) < 1:
        raise RuntimeError("No inlier tie points available for displacement cartography.")

    x_map_col = "X_MAP" if "X_MAP" in df_abs.columns else ("Map X" if "Map X" in df_abs.columns else None)
    y_map_col = "Y_MAP" if "Y_MAP" in df_abs.columns else ("Map Y" if "Map Y" in df_abs.columns else None)
    map_mode = (
        x_map_col is not None
        and y_map_col is not None
        and "X_SHIFT_M" in df_abs.columns
        and "Y_SHIFT_M" in df_abs.columns
    )
    image_mode = (
        "X_IM" in df_abs.columns
        and "Y_IM" in df_abs.columns
        and "X_SHIFT_PX" in df_abs.columns
        and "Y_SHIFT_PX" in df_abs.columns
    )

    if not map_mode and not image_mode:
        raise RuntimeError(
            "Cannot render displacement cartography: need map shifts (X/Y_MAP + X/Y_SHIFT_M) "
            "or image shifts (X/Y_IM + X/Y_SHIFT_PX)."
        )

    count_dedup = int(len(df_abs))
    max_vectors_int = int(max_vectors) if max_vectors is not None else 0
    if max_vectors_int > 0 and len(df_abs) > max_vectors_int:
        sample_x_col = str(x_map_col) if map_mode else "X_IM"
        sample_y_col = str(y_map_col) if map_mode else "Y_IM"
        df_abs, sampling_info = _stratified_sample_displacement_vectors_for_plot(
            df_abs,
            x_col=sample_x_col,
            y_col=sample_y_col,
            max_vectors=max_vectors_int,
            grid_rows=int(max(1, sampling_grid_rows)),
            grid_cols=int(max(1, sampling_grid_cols)),
        )

    if len(df_abs) < 1:
        raise RuntimeError("No tie points remained after deduplication/sampling for cartography.")

    count_sampled = int(len(df_abs))
    if map_mode:
        x = np.asarray(df_abs[str(x_map_col)], dtype=float)
        y = np.asarray(df_abs[str(y_map_col)], dtype=float)
        u = np.asarray(df_abs["X_SHIFT_M"], dtype=float)
        v = np.asarray(df_abs["Y_SHIFT_M"], dtype=float)
    else:
        x = np.asarray(df_abs["X_IM"], dtype=float)
        y = np.asarray(df_abs["Y_IM"], dtype=float)
        u = np.asarray(df_abs["X_SHIFT_PX"], dtype=float)
        v = np.asarray(df_abs["Y_SHIFT_PX"], dtype=float)

    abs_shift = np.asarray(df_abs["ABS_SHIFT"], dtype=float)
    finite = np.isfinite(x) & np.isfinite(y) & np.isfinite(u) & np.isfinite(v) & np.isfinite(abs_shift)
    if not np.any(finite):
        raise RuntimeError("No finite displacement vectors available for cartography.")
    x = x[finite]
    y = y[finite]
    u = u[finite]
    v = v[finite]
    abs_shift = abs_shift[finite]

    os.makedirs(os.path.dirname(output_png) or ".", exist_ok=True)
    fig, ax = plt.subplots(figsize=(10.5, 8.2), dpi=220, facecolor="white")
    basemap_used = False
    basemap_warning = None
    basemap_extent = None
    quiver_scale_used = 50.0
    try:
        ax.set_facecolor("#f4f7fb")

        if map_mode and basemap_raster_path:
            try:
                basemap_path = str(basemap_raster_path)
                if not os.path.exists(basemap_path):
                    basemap_warning = f"Basemap raster not found: {basemap_path}"
                else:
                    band_idx = 1
                    if basemap_rgb_band_indices is not None and len(basemap_rgb_band_indices) > 0:
                        try:
                            band_idx = int(basemap_rgb_band_indices[0])
                        except Exception:
                            band_idx = 1
                    band_arr, valid_arr, _info = _load_quicklook_band(
                        basemap_path,
                        band_idx=band_idx,
                        nodata=0.0,
                        max_quicklook_dim=DEFAULT_QUICKLOOK_MAX_DIM,
                    )
                    with rasterio.open(basemap_path) as src:
                        b = src.bounds
                        extent = (float(b.left), float(b.right), float(b.bottom), float(b.top))
                    basemap_extent = extent

                    if np.any(valid_arr):
                        valid_vals = np.asarray(band_arr[valid_arr], dtype=float)
                        try:
                            lo = float(np.nanpercentile(valid_vals, 2))
                            hi = float(np.nanpercentile(valid_vals, 98))
                        except Exception:
                            lo = float(np.nanmin(valid_vals))
                            hi = float(np.nanmax(valid_vals))
                        if np.isfinite(lo) and np.isfinite(hi) and hi > lo:
                            band_display = np.clip((band_arr - lo) / (hi - lo), 0.0, 1.0)
                        else:
                            band_display = band_arr
                    else:
                        band_display = band_arr

                    band_display = np.asarray(band_display, dtype=np.float32)
                    band_display[~valid_arr] = np.nan
                    ax.imshow(
                        band_display,
                        extent=extent,
                        origin="upper",
                        interpolation="bilinear",
                        alpha=0.8,
                        cmap="gray",
                        zorder=1,
                    )
                    basemap_used = True
            except Exception as exc:
                basemap_warning = f"Failed to load basemap raster for displacement cartography: {exc}"

        X = np.asarray(x, dtype=float)
        Y = np.asarray(y, dtype=float)
        U = np.asarray(u, dtype=float)
        V = np.asarray(v, dtype=float)
        C = np.asarray(abs_shift, dtype=float)
        U_plot = np.asarray(U, dtype=float)
        V_plot = np.asarray(V, dtype=float)

        # Visibility-normalized rendering:
        # project each vector to a bounded on-map length while preserving direction.
        span_x = float(np.nanmax(X) - np.nanmin(X)) if X.size > 0 else 0.0
        span_y = float(np.nanmax(Y) - np.nanmin(Y)) if Y.size > 0 else 0.0
        span_ref = float(max(span_x, span_y, 1.0))
        magnitudes = np.hypot(U, V) if U.size > 0 else np.array([], dtype=float)
        quiver_scale_used = 1.0
        min_visible = float(max(1e-12, span_ref * min_frac_used))
        max_visible = float(max(min_visible, span_ref * max_frac_used))
        if magnitudes.size > 0:
            finite_mag = np.asarray(magnitudes[np.isfinite(magnitudes)], dtype=float)
            if finite_mag.size > 0:
                pos_mag = np.asarray(finite_mag[finite_mag > 0.0], dtype=float)
                if pos_mag.size > 0:
                    mag_lo = float(np.nanpercentile(pos_mag, 20))
                    mag_hi = float(np.nanpercentile(pos_mag, 90))
                    if not np.isfinite(mag_lo):
                        mag_lo = float(np.nanmin(pos_mag))
                    if not np.isfinite(mag_hi):
                        mag_hi = float(np.nanmax(pos_mag))
                    if not np.isfinite(mag_lo):
                        mag_lo = 0.0
                    if not np.isfinite(mag_hi):
                        mag_hi = mag_lo

                    render_lengths = np.zeros_like(magnitudes, dtype=float)
                    pos_mask = np.isfinite(magnitudes) & (magnitudes > 0.0)
                    if bool(np.any(pos_mask)):
                        mag_vals = np.asarray(magnitudes[pos_mask], dtype=float)
                        if mag_hi > mag_lo:
                            mag_clip = np.clip(mag_vals, mag_lo, mag_hi)
                            rel = (mag_clip - mag_lo) / max(1e-12, (mag_hi - mag_lo))
                        else:
                            rel = np.ones_like(mag_vals, dtype=float)
                        render_lengths[pos_mask] = min_visible + rel * (max_visible - min_visible)

                        dir_u = np.zeros_like(U, dtype=float)
                        dir_v = np.zeros_like(V, dtype=float)
                        np.divide(U, magnitudes, out=dir_u, where=magnitudes > 0.0)
                        np.divide(V, magnitudes, out=dir_v, where=magnitudes > 0.0)
                        U_plot = dir_u * render_lengths
                        V_plot = dir_v * render_lengths
        rendered_lengths = np.hypot(U_plot, V_plot)
        if rendered_lengths.size > 0:
            finite_rendered = rendered_lengths[np.isfinite(rendered_lengths)]
            if finite_rendered.size > 0:
                arrow_len_max_rendered = float(np.max(finite_rendered))
                arrow_len_max_frac_rendered = float(arrow_len_max_rendered / max(1e-12, span_ref))

        quiver = ax.quiver(
            X,
            Y,
            U_plot,
            V_plot,
            C,
            cmap=str(quiver_cmap or "RdYlGn_r"),
            angles="xy",
            scale_units="xy",
            scale=float(quiver_scale_used),
            width=0.004,
            linewidths=0.35,
            edgecolor="black",
            alpha=0.95,
            zorder=5,
        )
        cbar = fig.colorbar(quiver, ax=ax, fraction=0.046, pad=0.04)
        cbar.set_label("")
        cbar.ax.tick_params(labelsize=8)

        # Ensure full arrow extensions are included in the rendered frame.
        x_head = X + U_plot
        y_head = Y + V_plot
        x_vals = np.concatenate([X, x_head])
        y_vals = np.concatenate([Y, y_head])
        if basemap_extent is not None:
            x_vals = np.concatenate([x_vals, np.array([basemap_extent[0], basemap_extent[1]], dtype=float)])
            y_vals = np.concatenate([y_vals, np.array([basemap_extent[2], basemap_extent[3]], dtype=float)])
        x_vals = x_vals[np.isfinite(x_vals)]
        y_vals = y_vals[np.isfinite(y_vals)]
        if x_vals.size > 0 and y_vals.size > 0:
            xmin = float(np.min(x_vals))
            xmax = float(np.max(x_vals))
            ymin = float(np.min(y_vals))
            ymax = float(np.max(y_vals))
            xpad = float(max(1e-9, 0.02 * max(1e-12, xmax - xmin)))
            ypad = float(max(1e-9, 0.02 * max(1e-12, ymax - ymin)))
            ax.set_xlim(xmin - xpad, xmax + xpad)
            ax.set_ylim(ymin - ypad, ymax + ypad)

        if map_mode:
            ax.set_aspect("equal", adjustable="box")
        else:
            ax.invert_yaxis()

        ax.set_title("displacement vectors (color = absolute shift [meters])", fontsize=14)
        ax.set_xlabel("")
        ax.set_ylabel("")
        ax.grid(True, linestyle=":", color="gray", linewidth=0.5)
        fig.tight_layout()
        fig.savefig(output_png, dpi=220, facecolor="white", transparent=False)
    finally:
        plt.close(fig)

    dedup_warning = None
    if dedup_dropped > 0:
        dedup_warning = (
            f"Deduplicated displacement vectors by {dedup_key_source}: "
            f"removed {dedup_dropped} duplicate row(s)."
        )
    warning_parts = [
        msg
        for msg in (
            warning_message,
            dedup_warning,
            sampling_info.get("warning"),
            basemap_warning,
        )
        if msg
    ]
    return {
        "path": output_png,
        "count": int(len(abs_shift)),
        "count_raw": int(count_raw),
        "count_dedup": int(count_dedup),
        "count_sampled": int(count_sampled),
        "mode": "map" if map_mode else "image",
        "abs_shift_basis": abs_basis,
        "warning": " | ".join(warning_parts) if warning_parts else None,
        "quiver_scale_used": float(quiver_scale_used),
        "basemap_used": bool(basemap_used),
        "dedup_key_source": dedup_key_source,
        "sampling_method": sampling_info.get("method", "none"),
        "occupied_cells_total": int(sampling_info.get("occupied_cells_total", 0) or 0),
        "occupied_cells_sampled": int(sampling_info.get("occupied_cells_selected", 0) or 0),
        "arrow_len_max_rendered": float(arrow_len_max_rendered),
        "arrow_len_max_frac_rendered": float(arrow_len_max_frac_rendered),
        "arrow_len_min_frac_target": float(min_frac_used),
        "arrow_len_max_frac_target": float(max_frac_used),
    }


def _coregister_prisma_ancillary_outputs(
    scene_name: str,
    folder_struct: Dict[str, str],
    s2_crs,
    save_pan: bool,
    save_quality_mask: bool,
    pan_data: Optional[np.ndarray],
    pan_geo_info: Optional[Dict[str, Any]],
    vnir_quality_data: Optional[np.ndarray],
    swir_quality_data: Optional[np.ndarray],
    lat_qm: Optional[np.ndarray],
    lon_qm: Optional[np.ndarray],
    best_candidate: Optional[Dict[str, Any]],
    hs_reference_raster_path: Optional[str] = None,
    s2_reference_raster_path: Optional[str] = None,
    matcher_profile: Optional[Dict[str, Any]] = None,
    pan_gcp_mode: str = "map_inverse",
    pan_map_dxdy_source: str = "auto",
    pan_target_aligned_pixels: bool = False,
    pan_use_synthetic_reference: bool = True,
    pan_min_points_for_poly2: int = 20,
    pan_local_window_size: Tuple[int, int] = (512, 512),
    pan_local_grid_res: int = LOCAL_GRID_RES_M,
    pan_local_max_shift: float = 220.0,
    pan_local_tieP_filter_level: int = 1,
    pan_local_max_iter: Optional[int] = None,
    pan_residual_check: bool = False,
    pan_residual_threshold_px: float = 0.5,
    pan_residual_max_dim: int = 1024,
    use_geolocation_mesh_affine: bool = False,
    geolocation_mesh_stride: int = 32,
) -> Dict[str, Any]:
    """Generate coregistered ancillary outputs (PAN and quality masks)."""
    result: Dict[str, Any] = {
        "status": "not_requested",
        "warnings": [],
        "outputs": {
            "pan": None,
            "quality_vnir": None,
            "quality_swir": None,
        },
        "warp_method": None,
        "tiepoints_used": 0,
        "pan_gcp_mode": _normalize_pan_gcp_mode(pan_gcp_mode),
        "pan_target_aligned_pixels": bool(pan_target_aligned_pixels),
        "pan_residual_check": {"enabled": bool(pan_residual_check), "ok": None},
    }
    if not (save_pan or save_quality_mask):
        return result

    result["status"] = "ok"
    base_tie_points_df = None if best_candidate is None else best_candidate.get("local_tiepoints_df")
    norm_pan_gcp_mode = _normalize_pan_gcp_mode(pan_gcp_mode)
    norm_pan_dxdy = _normalize_pan_dxdy_source(pan_map_dxdy_source)
    min_pan_points_for_poly2 = int(max(6, int(pan_min_points_for_poly2)))

    try:
        pan_grid_res = int(max(30, int(pan_local_grid_res)))
    except Exception:
        pan_grid_res = int(max(30, LOCAL_GRID_RES_M))
    try:
        pan_shift = float(max(5.0, float(pan_local_max_shift)))
    except Exception:
        pan_shift = 220.0
    try:
        pan_tiep_filter = int(max(0, int(pan_local_tieP_filter_level)))
    except Exception:
        pan_tiep_filter = 1
    pan_max_iter = pan_local_max_iter
    if pan_max_iter is None and isinstance(matcher_profile, dict):
        pan_max_iter = matcher_profile.get("local_max_iter")

    pan_dir = os.path.join(folder_struct["coreg"], "pan")
    quality_dir = os.path.join(folder_struct["coreg"], "quality")
    os.makedirs(pan_dir, exist_ok=True)
    os.makedirs(quality_dir, exist_ok=True)

    if save_pan:
        pan_out = os.path.join(pan_dir, f"{scene_name}_pan_coreg.tif")
        pan_src = os.path.join(folder_struct["temp"], f"{scene_name}_PAN_src.tif")
        synthetic_s2_pan_path = os.path.join(folder_struct["temp"], f"{scene_name}_S2_SYN_PAN.tif")
        pan_global_path = os.path.join(folder_struct["temp"], f"{scene_name}_PAN_GLOBAL_SYN.tif")
        pan_local_path = os.path.join(folder_struct["temp"], f"{scene_name}_PAN_LOCAL_SYN.tif")
        pan_temp_paths = [synthetic_s2_pan_path, pan_global_path, pan_local_path, pan_src]
        try:
            if pan_data is None or pan_geo_info is None:
                raise RuntimeError("PAN source data unavailable.")
            pan_arr = np.asarray(np.squeeze(pan_data), dtype=np.float32)
            if pan_arr.ndim != 2:
                raise RuntimeError(f"PAN array must be 2D; got shape {tuple(pan_arr.shape)}.")
            rows, cols = pan_arr.shape
            transform_pan = _estimate_transform_from_corner_coords(
                {
                    **dict(pan_geo_info),
                    "rows": rows,
                    "cols": cols,
                },
                s2_crs,
            )
            pan_clean = np.where(np.isfinite(pan_arr), pan_arr, PROCESSING_NODATA).astype(np.float32)
            _write_georeferenced_raster(
                pan_src,
                pan_clean[np.newaxis, :, :],
                s2_crs,
                transform_pan,
                dtype="float32",
                nodata=PROCESSING_NODATA,
            )
            pan_xres, pan_yres = _infer_raster_native_resolution(
                pan_src,
                fallback=float(pan_geo_info.get("pixel_size_m", 5.0) or 5.0),
            )
            pan_tie_points_df = None
            pan_warp_source = pan_src

            if bool(pan_use_synthetic_reference):
                if not s2_reference_raster_path or not os.path.exists(s2_reference_raster_path):
                    raise RuntimeError("Sentinel-2 reference stack unavailable for synthetic PAN branch.")

                synth_result = _create_synthetic_s2_pan(
                    s2_stack_path=s2_reference_raster_path,
                    output_path=synthetic_s2_pan_path,
                    s2_band_indices=(
                        int(MULTIBAND_S2_WAVELENGTHS["B02"]["stack_idx"]),
                        int(MULTIBAND_S2_WAVELENGTHS["B03"]["stack_idx"]),
                        int(MULTIBAND_S2_WAVELENGTHS["B04"]["stack_idx"]),
                        int(MULTIBAND_S2_WAVELENGTHS["B08"]["stack_idx"]),
                    ),
                )
                if not synth_result.get("ok", False):
                    raise RuntimeError(synth_result.get("error", "failed to create synthetic S2 PAN"))

                with rasterio.open(synthetic_s2_pan_path) as s2_pan_ref:
                    ref_nodata = s2_pan_ref.nodata
                    if ref_nodata is None or not np.isfinite(float(ref_nodata)):
                        ref_nodata = 0.0
                    else:
                        ref_nodata = float(ref_nodata)

                pan_ws = _resolve_pan_window_size_for_raster(
                    pan_src,
                    requested_window_size=pan_local_window_size,
                )
                pan_tp_result = _collect_pan_tiepoints_with_synthetic_reference(
                    synthetic_s2_pan_path=synthetic_s2_pan_path,
                    pan_source_path=pan_src,
                    pan_global_path=pan_global_path,
                    pan_local_path=pan_local_path,
                    ws=pan_ws,
                    grid_res=pan_grid_res,
                    max_shift=pan_shift,
                    tiep_filter_level=pan_tiep_filter,
                    local_max_iter=pan_max_iter,
                    reference_nodata=float(ref_nodata),
                    source_nodata=float(PROCESSING_NODATA),
                )
                for note in pan_tp_result.get("warnings", []):
                    result["warnings"].append(_fmt_issue("ANCILLARY", str(note)))
                pan_warp_source = str(pan_tp_result.get("global_path") or pan_src)
                pan_tie_points_df = pan_tp_result.get("tiepoints_df")

            if pan_tie_points_df is None or len(pan_tie_points_df) == 0:
                pan_tie_points_df = base_tie_points_df
                if pan_tie_points_df is not None and len(pan_tie_points_df) > 0:
                    result["warnings"].append(
                        _fmt_issue(
                            "ANCILLARY",
                            "PAN synthetic-reference tie points unavailable; falling back to base-scene tie points.",
                        )
                    )

            if pan_tie_points_df is None or len(pan_tie_points_df) == 0:
                raise RuntimeError("No tie-point table available for PAN ancillary warp.")

            merged_pan = _merge_tiepoints(
                [pan_tie_points_df],
                min_reliability=MIN_RELIABILITY_THRESHOLD,
                trim_by_residual=True,
                residual_mad_factor=3.0,
                min_band_support=1,
                allow_single_band_fallback=True,
                consensus_group_rounding_px=1.0,
                grid_rows=4,
                grid_cols=4,
                max_points_per_cell=4,
                min_points_required=max(_minimum_gcps_for_polynomial_order(1), min_pan_points_for_poly2),
                spatial_fallback_min_distance=60.0,
            )
            pan_selected_df = merged_pan.get("merged_df")
            if pan_selected_df is None or len(pan_selected_df) == 0:
                pan_selected_df = pan_tie_points_df

            result["tiepoints_used"] = int(len(pan_selected_df))
            if "fallback_notes" in merged_pan:
                for note in merged_pan.get("fallback_notes", []):
                    logger.info("PAN tie-point merge fallback: %s", note)

            pan_res = float(0.5 * (pan_xres + pan_yres))
            if norm_pan_gcp_mode == "scaled_image":
                if hs_reference_raster_path and os.path.exists(hs_reference_raster_path):
                    hs_xres, hs_yres = _infer_raster_native_resolution(hs_reference_raster_path, fallback=30.0)
                    hs_res = float(0.5 * (hs_xres + hs_yres))
                else:
                    hs_res = 30.0
                    result["warnings"].append(
                        _fmt_issue(
                            "ANCILLARY",
                            "HS reference path missing for PAN scaled-image GCP mode; using 30m fallback.",
                        )
                    )
                pan_gcps = build_pan_gcps_from_tiepoints(
                    tiepoints_df=pan_selected_df,
                    hs_pixel_size_m=hs_res,
                    pan_pixel_size_m=pan_res,
                    map_dx_dy_source=norm_pan_dxdy,
                    nodata=PROCESSING_NODATA,
                    crs_wkt_or_epsg=s2_crs.to_wkt() if hasattr(s2_crs, "to_wkt") else str(s2_crs),
                    pan_shape=(rows, cols),
                )
            else:
                pan_gcp_result = _build_tps_gcps_for_source_raster(
                    pan_selected_df,
                    pan_warp_source,
                    min_gcps=_minimum_gcps_for_polynomial_order(1),
                )
                if not pan_gcp_result.get("success", False):
                    raise RuntimeError(
                        pan_gcp_result.get("error_message", "failed to derive PAN polynomial GCPs")
                    )
                pan_gcps = pan_gcp_result.get("gcps", [])

            if len(pan_gcps) < _minimum_gcps_for_polynomial_order(1):
                raise RuntimeError(
                    f"Insufficient PAN GCPs for affine warp: {len(pan_gcps)} < {_minimum_gcps_for_polynomial_order(1)}"
                )

            poly_decision = _decide_polynomial_order(
                merged_df=pan_selected_df,
                preferred_order=2,
                auto_downgrade=True,
                min_gcps_order2=max(_minimum_gcps_for_polynomial_order(2), min_pan_points_for_poly2),
                min_cells_order2=4,
                grid_rows=4,
                grid_cols=4,
            )
            pan_order = int(max(1, int(poly_decision.get("order_used", 2))))
            if len(pan_gcps) < int(min_pan_points_for_poly2):
                pan_order = 1
                logger.warning(
                    "PAN GCP density low (%d < %d). Forcing order-1 affine warp (disabling order-2/TPS).",
                    int(len(pan_gcps)),
                    int(min_pan_points_for_poly2),
                )

            pan_warp = _apply_polynomial_warp(
                input_raster=pan_warp_source,
                output_raster=pan_out,
                gcps=pan_gcps,
                target_crs=s2_crs,
                polynomial_order=pan_order,
                output_resolution=pan_res,
                nodata=PROCESSING_NODATA,
                resampling="cubic",
            )
            if not pan_warp.get("success", False) and pan_order > 1:
                logger.warning(
                    "PAN order-%d warp failed (%s). Retrying with order-1 affine.",
                    int(pan_order),
                    pan_warp.get("error_message", "unknown error"),
                )
                pan_warp = _apply_polynomial_warp(
                    input_raster=pan_warp_source,
                    output_raster=pan_out,
                    gcps=pan_gcps,
                    target_crs=s2_crs,
                    polynomial_order=1,
                    output_resolution=pan_res,
                    nodata=PROCESSING_NODATA,
                    resampling="cubic",
                )
                pan_order = 1
            if not pan_warp.get("success", False):
                raise RuntimeError(pan_warp.get("error_message", "unknown PAN warp error"))

            pan_check = _validate_ancillary_raster(pan_out, s2_crs)
            if not pan_check.get("ok", False):
                raise RuntimeError(pan_check.get("error", "PAN output validation failed"))
            result["outputs"]["pan"] = pan_out
            result["warp_method"] = (
                f"poly_order_{int(pan_order)}_synthetic_s2_pan"
                if bool(pan_use_synthetic_reference)
                else f"poly_order_{int(pan_order)}_base_tiepoints"
            )
            result["pan_gcp_mode"] = norm_pan_gcp_mode

            if bool(pan_residual_check) and hs_reference_raster_path and os.path.exists(hs_reference_raster_path):
                residual = _estimate_translation_phasecorr(
                    reference_raster_path=hs_reference_raster_path,
                    candidate_raster_path=pan_out,
                    max_dim=int(max(128, int(pan_residual_max_dim))),
                    reference_band=1,
                    candidate_band=1,
                    reference_nodata=PROCESSING_NODATA,
                    candidate_nodata=PROCESSING_NODATA,
                )
                result["pan_residual_check"] = {
                    "enabled": True,
                    **residual,
                }
                if residual.get("ok", False):
                    shift_mag = float(residual.get("shift_magnitude_px", 0.0) or 0.0)
                    if shift_mag > float(max(0.0, pan_residual_threshold_px)):
                        result["warnings"].append(
                            _fmt_issue(
                                "ANCILLARY",
                                (
                                    f"PAN residual shift estimate is high ({shift_mag:.2f}px > "
                                    f"{float(pan_residual_threshold_px):.2f}px). "
                                    "Try pan_gcp_mode='scaled_image' and pan_target_aligned_pixels=True."
                                ),
                            )
                        )
                else:
                    result["warnings"].append(
                        _fmt_issue(
                            "ANCILLARY",
                            f"PAN residual check unavailable: {residual.get('error', 'unknown error')}",
                        )
                    )
        except Exception as e:
            result["status"] = "degraded"
            result["warnings"].append(_fmt_issue("ANCILLARY", f"PAN ancillary output failed: {e}"))
        finally:
            for tmp_path in pan_temp_paths:
                if not tmp_path:
                    continue
                if os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except Exception:
                        pass

    if save_quality_mask:
        if base_tie_points_df is None or len(base_tie_points_df) == 0:
            result["status"] = "degraded"
            result["warnings"].append(
                _fmt_issue(
                    "ANCILLARY",
                    "No accepted base-scene tie-point table available; skipped quality-mask ancillary coregistration.",
                )
            )
            return result
        quality_items = [
            ("VNIR", vnir_quality_data, "quality_vnir", f"{scene_name}_quality_vnir_coreg.tif"),
            ("SWIR", swir_quality_data, "quality_swir", f"{scene_name}_quality_swir_coreg.tif"),
        ]
        for label, mask_data, key, out_name in quality_items:
            src_path = os.path.join(folder_struct["temp"], f"{scene_name}_{label}_quality_src.tif")
            out_path = os.path.join(quality_dir, out_name)
            try:
                if mask_data is None:
                    raise RuntimeError(f"{label} quality matrix unavailable.")
                if lat_qm is None or lon_qm is None:
                    raise RuntimeError("Quality geolocation arrays unavailable.")

                lat_arr = np.asarray(lat_qm)
                lon_arr = np.asarray(lon_qm)
                if lat_arr.shape != lon_arr.shape or lat_arr.ndim != 2:
                    raise RuntimeError(
                        f"Invalid quality geolocation shape lat={lat_arr.shape}, lon={lon_arr.shape}."
                    )
                rows, cols = lat_arr.shape
                quality_bands = _prepare_bands_first(mask_data, rows, cols, f"{label} quality")
                quality_bands = np.where(
                    np.isfinite(quality_bands),
                    quality_bands,
                    255,
                )
                quality_bands = np.clip(quality_bands, 0, 255).astype(np.uint8)
                transform_qm = estimate_prisma_geotransform(
                    lon_arr,
                    lat_arr,
                    rows,
                    cols,
                    s2_crs,
                    use_geolocation_mesh=use_geolocation_mesh_affine,
                    geolocation_mesh_stride=geolocation_mesh_stride,
                )
                _write_georeferenced_raster(
                    src_path,
                    quality_bands,
                    s2_crs,
                    transform_qm,
                    dtype="uint8",
                    nodata=255,
                )
                qm_gcp_result = _build_tps_gcps_for_source_raster(base_tie_points_df, src_path, min_gcps=10)
                if not qm_gcp_result.get("success", False):
                    raise RuntimeError(
                        qm_gcp_result.get("error_message", f"failed to derive {label} TPS GCPs")
                    )
                qm_xres, qm_yres = _infer_raster_native_resolution(src_path, fallback=30.0)
                qm_warp = _apply_tps_warp_from_gcps(
                    src_path,
                    out_path,
                    qm_gcp_result.get("gcps", []),
                    s2_crs,
                    x_res=qm_xres,
                    y_res=qm_yres,
                    resampling="near",
                    nodata=255,
                )
                if not qm_warp.get("success", False):
                    raise RuntimeError(qm_warp.get("error_message", f"unknown {label} quality warp error"))
                qm_check = _validate_ancillary_raster(out_path, s2_crs)
                if not qm_check.get("ok", False):
                    raise RuntimeError(qm_check.get("error", f"{label} quality output validation failed"))
                result["outputs"][key] = out_path
            except Exception as e:
                result["status"] = "degraded"
                result["warnings"].append(
                    _fmt_issue("ANCILLARY", f"{label} quality ancillary output failed: {e}")
                )

    return result


def _build_band_display_labels(wl, band_detectors, label_precision=2):
    """Build display labels in the form 'DETECTOR | wavelength nm'."""
    labels = []
    if wl is None:
        return labels
    precision = _safe_parse_int(label_precision, 2, "metadata_label_precision")
    precision = min(6, max(0, precision))
    for i, w in enumerate(wl):
        detector = None
        if band_detectors is not None and i < len(band_detectors):
            detector_candidate = str(band_detectors[i]).upper().strip()
            if detector_candidate in {"VNIR", "SWIR"}:
                detector = detector_candidate
        if detector is None:
            detector = _infer_detector_from_wavelength(float(w))
        labels.append(f"{detector} | {float(w):.{precision}f} nm")
    return labels


def _infer_detector_from_wavelength(wavelength: float) -> str:
    """Infer detector class when explicit detector metadata is unavailable."""
    return "VNIR" if float(wavelength) < 1000.0 else "SWIR"


def _infer_detector_from_band_name(band_name: Optional[str], sensor_type: str) -> Optional[str]:
    """Infer detector class from band name/slot where possible."""
    if band_name is None:
        return None

    name = str(band_name).strip()
    if not name:
        return None

    name_lower = name.lower()
    if "swir" in name_lower:
        return "SWIR"
    if "vnir" in name_lower:
        return "VNIR"

    # EnMAP slot IDs: SWIR starts at ~92; lower slots are VNIR.
    if str(sensor_type).upper() == "ENMAP":
        match = re.search(r"\d+", name)
        if match:
            slot = int(match.group(0))
            if slot >= 92:
                return "SWIR"
            if slot > 0:
                return "VNIR"

    return None


def _infer_detector_fallback(wavelength: float, band_name: Optional[str], sensor_type: str) -> str:
    """Infer detector with tiered fallback: name hint -> wavelength."""
    by_name = _infer_detector_from_band_name(band_name, sensor_type)
    if by_name is not None:
        return by_name
    return _infer_detector_from_wavelength(wavelength)


def _safe_parse_int(value, default_value: int, field_name: str) -> int:
    """Safely parse integer metadata fields with warning fallback."""
    try:
        return int(value)
    except (TypeError, ValueError):
        logger.warning(
            f"Invalid {field_name} value '{value}'; using default {default_value}."
        )
        return int(default_value)


def _safe_parse_float(value, default_value: float, field_name: str) -> float:
    """Safely parse float fields with warning fallback."""
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        logger.warning(
            f"Invalid {field_name} value '{value}'; using default {default_value}."
        )
        return float(default_value)
    if not np.isfinite(parsed):
        logger.warning(
            f"Non-finite {field_name} value '{value}'; using default {default_value}."
        )
        return float(default_value)
    return float(parsed)


def _normalize_metadata_extension_level(level: Any) -> str:
    """Normalize metadata extension level to one of: none, stats, full."""
    token = str(level or "").strip().lower()
    if token in {"none", "stats", "full"}:
        return token
    if token:
        logger.warning(
            f"Invalid metadata_extension_level '{level}'; using default 'stats'."
        )
    return "stats"


def _normalize_metadata_stats_mode(mode: Any) -> str:
    """Normalize metadata stats mode to one of: exact, approx, none."""
    token = str(mode or "").strip().lower()
    if token in {"exact", "approx", "none"}:
        return token
    if token:
        logger.warning(
            f"Invalid metadata_stats_mode '{mode}'; using default 'exact'."
        )
    return "exact"


def _build_normalization_params_from_config(config: Dict[str, Any]) -> NormalizationParams:
    """Build normalized normalization parameters from runtime config."""
    mode = normalize_mode(
        config.get("normalization_mode", DEFAULT_CONFIG.get("normalization_mode", "none")),
        default="none",
    )
    params = NormalizationParams(
        mode=mode,
        p_low=_safe_parse_float(config.get("norm_p_low", 2.0), 2.0, "norm_p_low"),
        p_high=_safe_parse_float(config.get("norm_p_high", 98.0), 98.0, "norm_p_high"),
        clip=bool(config.get("norm_clip", True)),
        eps=_safe_parse_float(config.get("norm_eps", 1e-6), 1e-6, "norm_eps"),
        min_valid_pixels=_safe_parse_int(config.get("norm_min_valid_pixels", 1024), 1024, "norm_min_valid_pixels"),
        reservoir_size=_safe_parse_int(config.get("norm_reservoir_size", 8192), 8192, "norm_reservoir_size"),
        seed=_safe_parse_int(config.get("norm_seed", 1337), 1337, "norm_seed"),
        tile_size=_safe_parse_int(config.get("norm_tile_size", 256), 256, "norm_tile_size"),
    )
    params, sanitize_warnings = sanitize_normalization_params(params)
    for msg in sanitize_warnings:
        logger.warning(_fmt_issue("NORMALIZATION", msg))
    return params


def _prepare_output_metadata(n_bands, sensor_type, wl, fwhm, band_names, band_detectors):
    """Validate and sanitize metadata arrays so they align with output bands."""
    warnings = []
    errors = []

    if wl is None:
        errors.append("Missing wavelength metadata.")
        return {"ok": False, "warnings": warnings, "errors": errors}

    wl_arr = np.asarray(wl, dtype=float).reshape(-1)
    if len(wl_arr) != n_bands:
        errors.append(f"Wavelength count mismatch ({len(wl_arr)} vs {n_bands}).")
        return {"ok": False, "warnings": warnings, "errors": errors}

    if not np.all(np.isfinite(wl_arr)):
        errors.append("Wavelength array contains non-finite values.")
        return {"ok": False, "warnings": warnings, "errors": errors}

    fwhm_arr = None
    if fwhm is not None:
        fwhm_candidate = np.asarray(fwhm, dtype=float).reshape(-1)
        if len(fwhm_candidate) == n_bands:
            fwhm_arr = fwhm_candidate
        else:
            warnings.append(f"FWHM count mismatch ({len(fwhm_candidate)} vs {n_bands}); skipping FWHM metadata.")

    names_list = None
    if band_names is not None:
        names_candidate = [str(x) for x in band_names]
        if len(names_candidate) == n_bands:
            names_list = names_candidate
        elif len(names_candidate) > n_bands:
            warnings.append(
                f"Band-name count mismatch ({len(names_candidate)} vs {n_bands}); truncating extras."
            )
            names_list = names_candidate[:n_bands]
        else:
            warnings.append(
                f"Band-name count mismatch ({len(names_candidate)} vs {n_bands}); padding with canonical names."
            )
            sensor_tag = str(sensor_type).upper()
            names_list = names_candidate + [
                f"{sensor_tag}_{idx + 1:03d}" for idx in range(len(names_candidate), n_bands)
            ]
    if names_list is None:
        sensor_tag = str(sensor_type).upper()
        names_list = [f"{sensor_tag}_{idx + 1:03d}" for idx in range(n_bands)]

    detectors_list = None
    if band_detectors is not None:
        detectors_candidate = [str(x).upper() if x is not None else "" for x in band_detectors]
        if len(detectors_candidate) == n_bands:
            detectors_list = []
            for idx, detector in enumerate(detectors_candidate):
                if detector in ("VNIR", "SWIR"):
                    detectors_list.append(detector)
                else:
                    detectors_list.append(
                        _infer_detector_fallback(wl_arr[idx], names_list[idx], sensor_type)
                    )
        elif len(detectors_candidate) > n_bands:
            warnings.append(
                f"Detector count mismatch ({len(detectors_candidate)} vs {n_bands}); truncating extras."
            )
            detectors_list = []
            for idx, detector in enumerate(detectors_candidate[:n_bands]):
                if detector in ("VNIR", "SWIR"):
                    detectors_list.append(detector)
                else:
                    detectors_list.append(
                        _infer_detector_fallback(wl_arr[idx], names_list[idx], sensor_type)
                    )
        else:
            warnings.append(
                f"Detector count mismatch ({len(detectors_candidate)} vs {n_bands}); padding with inferred detectors."
            )
            detectors_list = []
            for idx in range(n_bands):
                if idx < len(detectors_candidate) and detectors_candidate[idx] in ("VNIR", "SWIR"):
                    detectors_list.append(detectors_candidate[idx])
                else:
                    detectors_list.append(
                        _infer_detector_fallback(wl_arr[idx], names_list[idx], sensor_type)
                    )
    if detectors_list is None:
        detectors_list = [
            _infer_detector_fallback(wl_arr[idx], names_list[idx], sensor_type)
            for idx in range(n_bands)
        ]

    return {
        "ok": True,
        "warnings": warnings,
        "errors": errors,
        "wl": wl_arr,
        "fwhm": fwhm_arr,
        "band_names": names_list,
        "band_detectors": detectors_list,
    }


def _is_critical_metadata_error(messages: List[str]) -> bool:
    """Return True if metadata errors indicate unsafe spectral metadata state."""
    critical_tokens = (
        "wavelength count mismatch",
        "missing wavelength metadata",
        "non-finite",
        "output band count mismatch",
    )
    for msg in messages:
        msg_l = str(msg).lower()
        if any(token in msg_l for token in critical_tokens):
            return True
    return False


def _envi_data_type_from_rasterio_dtype(raster_dtype):
    """Map rasterio dtype to ENVI numeric data type code."""
    dtype_key = np.dtype(raster_dtype).name.lower()
    return ENVI_DTYPE_MAP.get(dtype_key, ENVI_DTYPE_MAP["float32"])


def _envi_byte_order_from_rasterio_dtype(raster_dtype):
    """Compute ENVI byte-order field from dtype endianness."""
    byteorder = np.dtype(raster_dtype).byteorder
    if byteorder == ">":
        return 1
    if byteorder in ("<", "|"):
        return 0
    return 0 if sys.byteorder == "little" else 1


def _atomic_write_text(path, content, encoding="utf-8"):
    """Atomically write text content to disk."""
    target_dir = os.path.dirname(path) or "."
    fd, temp_path = tempfile.mkstemp(prefix=".tmp_", suffix=".txt", dir=target_dir, text=True)
    try:
        with os.fdopen(fd, "w", encoding=encoding, newline="\n") as f:
            f.write(content)
        os.replace(temp_path, path)
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass


def _write_geotiff_band_metadata(
    tif_path,
    sensor_type,
    wl,
    fwhm,
    band_names,
    band_detectors,
    metadata_schema_version=METADATA_SCHEMA_VERSION,
    normalization_mode="none",
    normalization_params: Optional[Dict[str, Any]] = None,
    remove_detector_overlap=False,
    label_precision=2,
    metadata_extension_level="stats",
    metadata_histogram_buckets=64,
):
    """Write per-band descriptions and tags into GeoTIFF for QGIS/ENVI visibility."""
    result = {"ok": False, "warnings": [], "errors": []}
    try:
        norm_params = dict(normalization_params or {})
        with rasterio.open(tif_path, "r+") as dst:
            n_bands = int(dst.count)
            prepared = _prepare_output_metadata(n_bands, sensor_type, wl, fwhm, band_names, band_detectors)
            if not prepared["ok"]:
                result["errors"].extend(prepared["errors"])
                return result

            result["warnings"].extend(prepared["warnings"])
            wl = prepared["wl"]
            fwhm = prepared["fwhm"]
            band_names = prepared["band_names"]
            band_detectors = prepared["band_detectors"]

            labels = _build_band_display_labels(wl, band_detectors, label_precision=label_precision)

            dst.update_tags(
                SENSOR=str(sensor_type),
                METADATA_SCHEMA_VERSION=str(metadata_schema_version),
                NORMALIZATION_MODE=str(normalization_mode),
                NORM_P_LOW=str(norm_params.get("p_low")),
                NORM_P_HIGH=str(norm_params.get("p_high")),
                NORM_CLIP=str(bool(norm_params.get("clip", True))).lower(),
                NORM_EPS=str(norm_params.get("eps")),
                NORM_MIN_VALID_PIXELS=str(norm_params.get("min_valid_pixels")),
                NORM_RESERVOIR_SIZE=str(norm_params.get("reservoir_size")),
                NORM_SEED=str(norm_params.get("seed")),
                NORM_TILE_SIZE=str(norm_params.get("tile_size")),
                NORM_ESTIMATOR="reservoir" if str(normalization_mode) == "percentile" else "none",
                REMOVE_OVERLAPPING_BANDS=str(bool(remove_detector_overlap)).lower(),
                METADATA_EXTENSION_LEVEL=_normalize_metadata_extension_level(metadata_extension_level),
                METADATA_HISTOGRAM_BUCKETS=str(
                    _safe_parse_int(metadata_histogram_buckets, 64, "metadata_histogram_buckets")
                ),
                METADATA_LABEL_PRECISION=str(
                    _safe_parse_int(label_precision, 2, "metadata_label_precision")
                ),
            )

            for i in range(n_bands):
                bidx = i + 1
                detector = str(band_detectors[i]).upper()
                label = labels[i] if i < len(labels) else f"{detector} | {float(wl[i]):.2f} nm"

                dst.set_band_description(bidx, label)

                tag_kwargs = {
                    "SENSOR": str(sensor_type),
                    "DETECTOR": detector,
                    "WAVELENGTH_NM": f"{float(wl[i]):.2f}",
                }
                if fwhm is not None and np.isfinite(float(fwhm[i])):
                    tag_kwargs["FWHM_NM"] = f"{float(fwhm[i]):.2f}"
                if band_names is not None:
                    tag_kwargs["SOURCE_BAND_NAME"] = str(band_names[i])
                dst.update_tags(bidx, **tag_kwargs)

            result["ok"] = True
            return result
    except Exception as e:
        result["errors"].append(str(e))
        return result


def _write_envi_header(
    tif_path,
    sensor_type,
    wl,
    fwhm,
    band_names,
    band_detectors,
    metadata_schema_version=METADATA_SCHEMA_VERSION,
    remove_detector_overlap=False,
    label_precision=2,
    metadata_extension_level="stats",
    metadata_histogram_buckets=64,
):
    """Write ENVI header file for the GeoTIFF.

    Keep this header strictly ENVI-standard for maximum compatibility.
    Extended/custom metadata must be written to GeoTIFF tags and PAM sidecars.
    """
    result = {"ok": False, "warnings": [], "errors": []}
    hdr_path = str(Path(tif_path).with_suffix(".hdr"))
    try:
        with rasterio.open(tif_path) as src:
            n_bands = int(src.count)
            prepared = _prepare_output_metadata(n_bands, sensor_type, wl, fwhm, band_names, band_detectors)
            if not prepared["ok"]:
                result["errors"].extend(prepared["errors"])
                return result

            result["warnings"].extend(prepared["warnings"])
            wl = prepared["wl"]
            fwhm = prepared["fwhm"]
            band_names = prepared["band_names"]
            band_detectors = prepared["band_detectors"]
            envi_data_type = _envi_data_type_from_rasterio_dtype(src.dtypes[0])
            envi_byte_order = _envi_byte_order_from_rasterio_dtype(src.dtypes[0])

            lines = ["ENVI", f"description = {{HyperCoreg output: {sensor_type}}}"]
            lines.append(f"samples = {src.width}")
            lines.append(f"lines = {src.height}")
            lines.append(f"bands = {src.count}")
            lines.append("header offset = 0")
            if Path(tif_path).suffix.lower() in {".tif", ".tiff"}:
                # ENVI opens GeoTIFF sidecars more reliably with explicit TIFF file type.
                lines.append("file type = TIFF")
            else:
                lines.append("file type = ENVI Standard")
            lines.append(f"data type = {envi_data_type}")
            lines.append("interleave = bsq")
            lines.append(f"byte order = {envi_byte_order}")
            if src.nodata is not None and np.isfinite(float(src.nodata)):
                lines.append(f"data ignore value = {float(src.nodata):.10g}")
            if wl is not None:
                wl_str = ", ".join([f"{w:.2f}" for w in wl])
                lines.append(f"wavelength = {{{wl_str}}}")
                lines.append("wavelength units = Nanometers")
            if fwhm is not None:
                fwhm_str = ", ".join([f"{f:.2f}" for f in fwhm])
                lines.append(f"fwhm = {{{fwhm_str}}}")
            if wl is not None:
                # Keep band names intentionally short to avoid oversized ENVI array fields.
                envi_band_names = [
                    f"Band {i + 1}"
                    for i, w in enumerate(wl)
                ]
                bn_str = ", ".join(envi_band_names)
                lines.append(f"band names = {{{bn_str}}}")
            elif band_names is not None:
                sanitized = [
                    str(name).replace("|", " ").replace(",", " ").strip()
                    for name in band_names
                ]
                bn_str = ", ".join(sanitized)
                lines.append(f"band names = {{{bn_str}}}")

        _atomic_write_text(hdr_path, "\n".join(lines), encoding="utf-8")
        logger.debug(f"ENVI header written: {hdr_path}")
        result["ok"] = True
        return result
    except Exception as e:
        result["errors"].append(str(e))
        return result


def _iter_valid_band_values(src: rasterio.io.DatasetReader, bidx: int, nodata_value: Optional[float]):
    """Yield finite, non-nodata pixel chunks for a given band."""
    for _, window in src.block_windows(bidx):
        band = src.read(bidx, window=window).astype(np.float64, copy=False)
        valid = np.isfinite(band)
        if nodata_value is not None and np.isfinite(nodata_value):
            valid &= (band != float(nodata_value))
        if np.any(valid):
            yield band[valid]


def _compute_band_statistics(
    src: rasterio.io.DatasetReader,
    bidx: int,
    nodata_value: Optional[float],
) -> Dict[str, Any]:
    """Compute exact band statistics using streaming block reads."""
    total_px = int(src.height) * int(src.width)
    valid_count = 0
    data_min = np.inf
    data_max = -np.inf
    sum_val = 0.0
    sum_sq = 0.0

    for vals in _iter_valid_band_values(src, bidx, nodata_value):
        valid_count += int(vals.size)
        local_min = float(np.min(vals))
        local_max = float(np.max(vals))
        if local_min < data_min:
            data_min = local_min
        if local_max > data_max:
            data_max = local_max
        sum_val += float(np.sum(vals, dtype=np.float64))
        sum_sq += float(np.sum(vals * vals, dtype=np.float64))

    if valid_count <= 0:
        return {
            "valid_count": 0,
            "invalid_count": total_px,
            "minimum": None,
            "maximum": None,
            "mean": None,
            "stddev": None,
            "valid_percent": 0.0,
        }

    mean = sum_val / float(valid_count)
    variance = max(0.0, (sum_sq / float(valid_count)) - (mean * mean))
    return {
        "valid_count": valid_count,
        "invalid_count": max(0, total_px - valid_count),
        "minimum": float(data_min),
        "maximum": float(data_max),
        "mean": float(mean),
        "stddev": float(np.sqrt(variance)),
        "valid_percent": float(100.0 * valid_count / max(1, total_px)),
    }


def _compute_band_statistics_approx(
    src: rasterio.io.DatasetReader,
    bidx: int,
    nodata_value: Optional[float],
    sample_windows: int = 128,
    seed: int = 1337,
) -> Tuple[Dict[str, Any], Optional[np.ndarray]]:
    """
    Approximate band statistics from deterministic sampled windows.

    Returns:
        (stats_dict, sampled_values_for_optional_histogram)
    """
    total_px = int(src.height) * int(src.width)
    target_windows = max(1, int(sample_windows))

    if src.block_shapes and len(src.block_shapes) >= bidx:
        block_h, block_w = src.block_shapes[bidx - 1]
        block_h = max(1, int(block_h))
        block_w = max(1, int(block_w))
    else:
        block_h = min(512, int(src.height))
        block_w = min(512, int(src.width))

    n_blocks_h = int(np.ceil(float(src.height) / float(block_h)))
    n_blocks_w = int(np.ceil(float(src.width) / float(block_w)))
    total_windows = max(1, n_blocks_h * n_blocks_w)
    step = max(1, int(np.ceil(float(total_windows) / float(target_windows))))
    offset = int((int(seed) + int(bidx)) % step)

    sampled_chunks: List[np.ndarray] = []
    valid_count = 0
    sampled_total = 0
    data_min = np.inf
    data_max = -np.inf
    sum_val = 0.0
    sum_sq = 0.0

    for widx, (_, window) in enumerate(src.block_windows(bidx)):
        if (widx % step) != offset:
            continue
        band = src.read(bidx, window=window).astype(np.float64, copy=False)
        valid = np.isfinite(band)
        if nodata_value is not None and np.isfinite(nodata_value):
            valid &= (band != float(nodata_value))
        vals = band[valid]
        if vals.size <= 0:
            continue
        vals = vals.astype(np.float64, copy=False)
        sampled_chunks.append(vals)
        sampled_total += int(vals.size)
        valid_count += int(vals.size)
        local_min = float(np.min(vals))
        local_max = float(np.max(vals))
        if local_min < data_min:
            data_min = local_min
        if local_max > data_max:
            data_max = local_max
        sum_val += float(np.sum(vals, dtype=np.float64))
        sum_sq += float(np.sum(vals * vals, dtype=np.float64))

    if valid_count <= 0:
        return ({
            "valid_count": 0,
            "invalid_count": total_px,
            "minimum": None,
            "maximum": None,
            "mean": None,
            "stddev": None,
            "valid_percent": 0.0,
            "approximate": True,
            "sampled_pixels": 0,
        }, None)

    mean = sum_val / float(valid_count)
    variance = max(0.0, (sum_sq / float(valid_count)) - (mean * mean))
    stats = {
        "valid_count": int(valid_count),
        "invalid_count": max(0, int(total_px - valid_count)),
        "minimum": float(data_min),
        "maximum": float(data_max),
        "mean": float(mean),
        "stddev": float(np.sqrt(variance)),
        "valid_percent": float(100.0 * valid_count / max(1, total_px)),
        "approximate": True,
        "sampled_pixels": int(sampled_total),
    }
    sampled_values = np.concatenate(sampled_chunks).astype(np.float64, copy=False) if sampled_chunks else None
    return stats, sampled_values


def _compute_band_histogram(
    src: rasterio.io.DatasetReader,
    bidx: int,
    nodata_value: Optional[float],
    hist_min: float,
    hist_max: float,
    buckets: int,
    valid_count: Optional[int] = None,
) -> Optional[Dict[str, Any]]:
    """Compute exact histogram counts from block reads."""
    if buckets <= 0:
        return None
    if not np.isfinite(hist_min) or not np.isfinite(hist_max):
        return None

    counts = np.zeros(int(buckets), dtype=np.int64)
    if float(hist_max) <= float(hist_min):
        if valid_count is not None:
            counts[0] = int(valid_count)
        else:
            counts[0] = int(_compute_band_statistics(src, bidx, nodata_value)["valid_count"])
    else:
        edges = np.linspace(float(hist_min), float(hist_max), int(buckets) + 1, dtype=np.float64)
        for vals in _iter_valid_band_values(src, bidx, nodata_value):
            hist, _ = np.histogram(vals, bins=edges)
            counts += hist.astype(np.int64, copy=False)

    return {
        "hist_min": float(hist_min),
        "hist_max": float(hist_max),
        "bucket_count": int(buckets),
        "counts": counts.tolist(),
    }


def _compute_histogram_from_values(
    values: Optional[np.ndarray],
    hist_min: float,
    hist_max: float,
    buckets: int,
) -> Optional[Dict[str, Any]]:
    """Compute histogram counts from sampled values."""
    if values is None or int(values.size) <= 0:
        return None
    if buckets <= 0:
        return None
    if not np.isfinite(hist_min) or not np.isfinite(hist_max):
        return None
    counts = np.zeros(int(buckets), dtype=np.int64)
    if float(hist_max) <= float(hist_min):
        counts[0] = int(values.size)
    else:
        edges = np.linspace(float(hist_min), float(hist_max), int(buckets) + 1, dtype=np.float64)
        hist, _ = np.histogram(values, bins=edges)
        counts += hist.astype(np.int64, copy=False)
    return {
        "hist_min": float(hist_min),
        "hist_max": float(hist_max),
        "bucket_count": int(buckets),
        "counts": counts.tolist(),
    }


def _append_pam_mdi(parent: ET.Element, key: str, value: Any) -> None:
    """Append one PAM MDI entry."""
    mdi = ET.SubElement(parent, "MDI", {"key": str(key)})
    mdi.text = str(value)


def _write_pam_aux_xml(
    tif_path: str,
    sensor_type: str,
    wl,
    fwhm,
    band_names,
    band_detectors,
    metadata_schema_version: int = METADATA_SCHEMA_VERSION,
    normalization_mode: str = "none",
    normalization_params: Optional[Dict[str, Any]] = None,
    remove_detector_overlap: bool = False,
    metadata_extension_level: str = "stats",
    metadata_stats_mode: str = "exact",
    metadata_stats_sample_windows: int = 128,
    metadata_stats_seed: int = 1337,
    metadata_histogram_buckets: int = 64,
    label_precision: int = 2,
    nodata_value: float = PROCESSING_NODATA,
    band_stats: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """Write a GDAL PAM .aux.xml sidecar with per-band stats and optional histograms."""
    result = {"ok": False, "warnings": [], "errors": [], "raster_passes": 0}
    aux_path = str(Path(tif_path).with_suffix(".aux.xml"))
    try:
        norm_params = dict(normalization_params or {})
        extension_level = _normalize_metadata_extension_level(metadata_extension_level)
        stats_mode = _normalize_metadata_stats_mode(metadata_stats_mode)
        sample_windows = max(1, _safe_parse_int(metadata_stats_sample_windows, 128, "metadata_stats_sample_windows"))
        stats_seed = _safe_parse_int(metadata_stats_seed, 1337, "metadata_stats_seed")
        include_stats = extension_level in {"stats", "full"} and stats_mode != "none"
        include_hist = extension_level == "full" and stats_mode != "none"
        if extension_level in {"stats", "full"} and stats_mode == "none":
            result["warnings"].append("metadata_stats_mode='none': skipping PAM STATISTICS_* fields.")
        hist_buckets = _safe_parse_int(
            metadata_histogram_buckets,
            64,
            "metadata_histogram_buckets",
        )
        hist_buckets = min(4096, max(2, hist_buckets))

        with rasterio.open(tif_path) as src:
            n_bands = int(src.count)
            prepared = _prepare_output_metadata(
                n_bands,
                sensor_type,
                wl,
                fwhm,
                band_names,
                band_detectors,
            )
            if not prepared["ok"]:
                result["errors"].extend(prepared["errors"])
                return result

            result["warnings"].extend(prepared["warnings"])
            wl_arr = prepared["wl"]
            fwhm_arr = prepared["fwhm"]
            names_list = prepared["band_names"]
            detectors_list = prepared["band_detectors"]
            labels = _build_band_display_labels(
                wl_arr,
                detectors_list,
                label_precision=label_precision,
            )

            pam_dataset = ET.Element("PAMDataset")
            dataset_metadata = ET.SubElement(pam_dataset, "Metadata")
            _append_pam_mdi(dataset_metadata, "SENSOR", str(sensor_type))
            _append_pam_mdi(dataset_metadata, "METADATA_SCHEMA_VERSION", int(metadata_schema_version))
            _append_pam_mdi(dataset_metadata, "METADATA_EXTENSION_LEVEL", extension_level)
            _append_pam_mdi(dataset_metadata, "METADATA_STATS_MODE", stats_mode)
            _append_pam_mdi(dataset_metadata, "METADATA_STATS_SAMPLE_WINDOWS", sample_windows)
            _append_pam_mdi(dataset_metadata, "METADATA_STATS_SEED", stats_seed)
            _append_pam_mdi(dataset_metadata, "NORMALIZATION_MODE", str(normalization_mode))
            _append_pam_mdi(dataset_metadata, "NORM_P_LOW", norm_params.get("p_low"))
            _append_pam_mdi(dataset_metadata, "NORM_P_HIGH", norm_params.get("p_high"))
            _append_pam_mdi(
                dataset_metadata,
                "NORM_CLIP",
                str(bool(norm_params.get("clip", True))).lower(),
            )
            _append_pam_mdi(dataset_metadata, "NORM_EPS", norm_params.get("eps"))
            _append_pam_mdi(dataset_metadata, "NORM_MIN_VALID_PIXELS", norm_params.get("min_valid_pixels"))
            _append_pam_mdi(dataset_metadata, "NORM_RESERVOIR_SIZE", norm_params.get("reservoir_size"))
            _append_pam_mdi(dataset_metadata, "NORM_SEED", norm_params.get("seed"))
            _append_pam_mdi(dataset_metadata, "NORM_TILE_SIZE", norm_params.get("tile_size"))
            _append_pam_mdi(
                dataset_metadata,
                "NORM_ESTIMATOR",
                "reservoir" if str(normalization_mode) == "percentile" else "none",
            )
            _append_pam_mdi(
                dataset_metadata,
                "REMOVE_OVERLAPPING_BANDS",
                str(bool(remove_detector_overlap)).lower(),
            )
            _append_pam_mdi(dataset_metadata, "METADATA_HISTOGRAM_BUCKETS", int(hist_buckets))
            _append_pam_mdi(
                dataset_metadata,
                "METADATA_LABEL_PRECISION",
                int(_safe_parse_int(label_precision, 2, "metadata_label_precision")),
            )

            nodata = src.nodata
            if nodata is None and nodata_value is not None:
                nodata = float(nodata_value)

            for i in range(n_bands):
                bidx = i + 1
                detector = str(detectors_list[i]).upper()
                band_label = labels[i] if i < len(labels) else f"{detector} | {float(wl_arr[i]):.2f} nm"
                pam_band = ET.SubElement(pam_dataset, "PAMRasterBand", {"band": str(bidx)})
                desc = ET.SubElement(pam_band, "Description")
                desc.text = band_label
                band_meta = ET.SubElement(pam_band, "Metadata")
                _append_pam_mdi(band_meta, "SENSOR", str(sensor_type))
                _append_pam_mdi(band_meta, "DETECTOR", detector)
                _append_pam_mdi(band_meta, "WAVELENGTH_NM", f"{float(wl_arr[i]):.2f}")
                if fwhm_arr is not None and np.isfinite(float(fwhm_arr[i])):
                    _append_pam_mdi(band_meta, "FWHM_NM", f"{float(fwhm_arr[i]):.2f}")
                _append_pam_mdi(band_meta, "SOURCE_BAND_NAME", str(names_list[i]))

                if include_stats:
                    sampled_values: Optional[np.ndarray] = None
                    approximate_stats = False
                    if band_stats is not None and i < len(band_stats):
                        stats = dict(band_stats[i])
                        approximate_stats = bool(stats.get("approximate", False))
                    else:
                        if stats_mode == "approx":
                            stats, sampled_values = _compute_band_statistics_approx(
                                src,
                                bidx,
                                nodata,
                                sample_windows=sample_windows,
                                seed=stats_seed,
                            )
                            approximate_stats = True
                            result["raster_passes"] += 1
                        else:
                            stats = _compute_band_statistics(src, bidx, nodata)
                            approximate_stats = False
                            result["raster_passes"] += 1
                    if stats["valid_count"] <= 0:
                        result["warnings"].append(
                            f"Band {bidx}: no valid pixels found for PAM statistics."
                        )
                    else:
                        _append_pam_mdi(band_meta, "STATISTICS_MINIMUM", f"{stats['minimum']:.10g}")
                        _append_pam_mdi(band_meta, "STATISTICS_MAXIMUM", f"{stats['maximum']:.10g}")
                        _append_pam_mdi(band_meta, "STATISTICS_MEAN", f"{stats['mean']:.10g}")
                        _append_pam_mdi(band_meta, "STATISTICS_STDDEV", f"{stats['stddev']:.10g}")
                        _append_pam_mdi(band_meta, "STATISTICS_VALID_PERCENT", f"{stats['valid_percent']:.8g}")
                        _append_pam_mdi(
                            band_meta,
                            "STATISTICS_APPROXIMATE",
                            "YES" if approximate_stats else "NO",
                        )

                        if include_hist:
                            if approximate_stats:
                                hist = _compute_histogram_from_values(
                                    sampled_values,
                                    float(stats["minimum"]),
                                    float(stats["maximum"]),
                                    hist_buckets,
                                )
                            else:
                                hist = _compute_band_histogram(
                                    src,
                                    bidx,
                                    nodata,
                                    float(stats["minimum"]),
                                    float(stats["maximum"]),
                                    hist_buckets,
                                    valid_count=int(stats["valid_count"]),
                                )
                                result["raster_passes"] += 1
                            if hist is not None:
                                histograms = ET.SubElement(pam_band, "Histograms")
                                hist_item = ET.SubElement(histograms, "HistItem")
                                ET.SubElement(hist_item, "HistMin").text = f"{hist['hist_min']:.10g}"
                                ET.SubElement(hist_item, "HistMax").text = f"{hist['hist_max']:.10g}"
                                ET.SubElement(hist_item, "BucketCount").text = str(hist["bucket_count"])
                                ET.SubElement(hist_item, "IncludeOutOfRange").text = "0"
                                ET.SubElement(hist_item, "Approximate").text = "0"
                                ET.SubElement(hist_item, "HistCounts").text = "|".join(
                                    str(int(v)) for v in hist["counts"]
                                )

            xml_bytes = ET.tostring(pam_dataset, encoding="utf-8")
            pretty = minidom.parseString(xml_bytes).toprettyxml(indent="  ", encoding="utf-8")
            _atomic_write_text(aux_path, pretty.decode("utf-8"), encoding="utf-8")
            result["ok"] = True
            return result
    except Exception as e:
        result["errors"].append(str(e))
        return result


def _write_shift_report(report_path, title, coreg_info, extra_lines=None):
    """Write shift report text file."""
    try:
        lines = [f"{'=' * 60}", f"{title}", f"{'=' * 60}"]
        lines.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("")
        if coreg_info:
            for key, val in coreg_info.items():
                lines.append(f"{key}: {val}")
        if extra_lines:
            for line in extra_lines:
                lines.append(f"{line}")
        _atomic_write_text(report_path, "\n".join(lines), encoding="utf-8")
        logger.debug(f"Shift report written: {report_path}")
    except Exception as e:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write shift report: {e}"))


def _write_per_scene_metrics_json(metrics_dict, json_path):
    """Write per-scene metrics to JSON file."""
    try:
        payload = json.dumps(metrics_dict, indent=2, default=str)
        _atomic_write_text(json_path, payload, encoding="utf-8")
        logger.debug(f"Metrics JSON written: {json_path}")
    except Exception as e:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write metrics JSON: {e}"))


def _write_scene_run_manifest(manifest_dict: Dict[str, Any], json_path: str) -> None:
    """Write per-scene run manifest JSON file."""
    try:
        payload = json.dumps(manifest_dict, indent=2, default=str)
        _atomic_write_text(json_path, payload, encoding="utf-8")
        logger.debug(f"Run manifest JSON written: {json_path}")
    except Exception as e:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write run manifest JSON: {e}"))


def _safe_dataset_file_stem(raw_value: Any, fallback: str = "scene") -> str:
    """Build a filesystem-safe ASCII stem for dataset workbook filenames."""
    text = str(raw_value).strip() if raw_value is not None else ""
    if not text:
        text = fallback
    text = text.encode("ascii", errors="ignore").decode("ascii")
    text = re.sub(r'[<>:"/\\|?*]+', "_", text)
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
    text = text.strip(" ._")
    return text or fallback


def _resolve_dataset_scene_name(metrics_dict: Dict[str, Any]) -> str:
    """Resolve a stable scene label used for folder_name and fallback workbook naming."""
    scene_name = metrics_dict.get("scene_name") or metrics_dict.get("folder_name")
    if scene_name:
        return str(scene_name)
    filename = metrics_dict.get("filename")
    if filename:
        return Path(str(filename)).stem
    return "scene"


def _build_dataset_xlsx_filename(scene_name: Optional[str], filename: Optional[str] = None) -> str:
    """Return the canonical per-scene dataset workbook filename."""
    label = scene_name
    if not label and filename:
        label = Path(str(filename)).stem
    safe_label = _safe_dataset_file_stem(label, fallback="scene")
    return f"{safe_label}_DATASET.xlsx"


def _infer_sensor_from_identifiers(
    hyp_type: Optional[str],
    scene_name: Optional[str],
    filename: Optional[str],
) -> Optional[str]:
    """Infer sensor label from explicit type, scene name, and filename patterns."""
    if hyp_type is not None:
        sensor = str(hyp_type).upper().strip()
        if sensor in {"PRISMA", "ENMAP"}:
            return sensor

    for candidate in (scene_name, filename):
        if candidate is None:
            continue
        text = str(candidate).upper()
        if "PRISMA" in text:
            return "PRISMA"
        if "ENMAP" in text or "SPECTRAL_IMAGE" in text:
            return "ENMAP"

    if filename:
        suffix = Path(str(filename)).suffix.lower()
        if suffix == ".he5":
            return "PRISMA"

    return None


def _build_dataset_row(metrics_dict: Dict[str, Any]) -> Dict[str, Any]:
    """Map scene metrics to one 43-column dataset row."""
    row = {column: None for column in DATASET_XLSX_COLUMNS}
    row["folder_name"] = _resolve_dataset_scene_name(metrics_dict)

    for column in DATASET_XLSX_COLUMNS:
        if column == "folder_name" or column.startswith("multiband_tiepoint_counts."):
            continue
        if column in metrics_dict:
            row[column] = metrics_dict.get(column)

    if row.get("status") is None and metrics_dict.get("status") is not None:
        row["status"] = metrics_dict.get("status")

    if row.get("notes") is None:
        for key in ("notes", "reason", "error", "message"):
            val = metrics_dict.get(key)
            if val is not None:
                row["notes"] = str(val)
                break

    if row.get("polynomial_n_gcps") is None:
        poly_decision = metrics_dict.get("polynomial_order_decision")
        if isinstance(poly_decision, dict) and poly_decision.get("n_gcps") is not None:
            row["polynomial_n_gcps"] = poly_decision.get("n_gcps")

    tp_residuals = metrics_dict.get("tp_residuals")
    if isinstance(tp_residuals, dict):
        for key in ("residual_mean_m", "residual_median_m", "residual_rmse_m", "residual_p90_m"):
            if row.get(key) is None and tp_residuals.get(key) is not None:
                row[key] = tp_residuals.get(key)

    multiband_counts = metrics_dict.get("multiband_tiepoint_counts")
    total_count = None
    if isinstance(multiband_counts, dict):
        running_total = 0
        running_total_valid = False
        for band in DATASET_MULTIBAND_COLUMNS:
            val = multiband_counts.get(band)
            if val is None:
                val = multiband_counts.get(band.lower())
            if val is not None:
                row[f"multiband_tiepoint_counts.{band}"] = val
                try:
                    running_total += int(val)
                    running_total_valid = True
                except Exception:
                    pass
        if multiband_counts.get("total") is not None:
            total_count = multiband_counts.get("total")
        elif multiband_counts.get("TOTAL") is not None:
            total_count = multiband_counts.get("TOTAL")
        elif running_total_valid:
            total_count = running_total
    elif multiband_counts is not None:
        total_count = multiband_counts

    if total_count is None and metrics_dict.get("multiband_tiepoint_counts_total") is not None:
        total_count = metrics_dict.get("multiband_tiepoint_counts_total")

    if total_count is not None:
        row["multiband_tiepoint_counts"] = total_count

    return row


def _write_single_scene_dataset_xlsx(metrics_dict: Dict[str, Any], xlsx_path: str) -> bool:
    """Write one-scene dataset workbook using the fixed 43-column schema."""
    try:
        from openpyxl import Workbook
    except ImportError:
        logger.debug("openpyxl not available, skipping per-scene dataset workbook.")
        return False

    try:
        parent = os.path.dirname(xlsx_path)
        if parent:
            os.makedirs(parent, exist_ok=True)

        row = _build_dataset_row(metrics_dict)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        for col_idx, column in enumerate(DATASET_XLSX_COLUMNS, 1):
            ws.cell(row=1, column=col_idx, value=column)
            ws.cell(row=2, column=col_idx, value=row.get(column))

        wb.save(xlsx_path)
        logger.info("Per-scene dataset workbook written: %s", xlsx_path)
        return True
    except Exception as e:
        logger.warning(_fmt_issue("REPORTS", f"Failed to write per-scene dataset workbook: {e}"))
        return False


def _build_failed_scene_metrics(hs_file: str, hyp_type: Optional[str], error_message: str) -> Dict[str, Any]:
    """Create a minimal metrics payload for fallback FAIL dataset workbook rows."""
    filename = os.path.basename(hs_file)
    inferred_sensor = _infer_sensor_from_identifiers(hyp_type, None, filename)
    scene_base = Path(filename).stem if filename else "scene"
    scene_name = f"{inferred_sensor}_{scene_base}" if inferred_sensor else scene_base
    return {
        "scene_name": scene_name,
        "filename": filename,
        "hyp_type": inferred_sensor,
        "status": "FAIL",
        "error": error_message,
        "notes": error_message,
    }


def _build_skip_result(
    hs_file: str,
    hyp_type: str,
    output_dir: str,
    scene_idx: int,
    scene_total: int,
    hs_time: Any,
    bbox: Any,
    config: Dict[str, Any],
    status: str,
    reason: str,
    normalization_mode: str,
    normalization_params: Dict[str, Any],
    build_overviews: bool,
    remove_detector_overlap_bands: bool,
    strict_metadata: bool,
    metadata_extension_level: str,
    metadata_stats_mode: str,
    metadata_stats_sample_windows: int,
    metadata_stats_seed: int,
    metadata_histogram_buckets: int,
    metadata_label_precision: int,
    validation_max_windows: int,
    extra_summary: Optional[Dict[str, Any]] = None,
    extra_metrics: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Create metrics + manifest for early scene skip conditions."""
    scene_name = f"{hyp_type.upper()}_{os.path.basename(hs_file)}"
    skip_manifest_path = os.path.join(output_dir, f"{scene_name}_run_manifest.json")
    skip_dataset_xlsx_path = os.path.join(
        output_dir,
        _build_dataset_xlsx_filename(scene_name=scene_name, filename=os.path.basename(hs_file)),
    )
    skip_metrics = {
        'status': status,
        'scene_name': scene_name,
        'filename': os.path.basename(hs_file),
        'hyp_type': hyp_type,
        'reason': reason,
        'metadata_schema_version': METADATA_SCHEMA_VERSION,
        'metadata_status': 'skipped',
        'metadata_warnings': [],
        'normalization_mode': str(normalization_mode),
        'normalization_params': dict(normalization_params),
        'build_overviews': bool(build_overviews),
        'remove_overlapping_bands': bool(remove_detector_overlap_bands),
        'strict_metadata': bool(strict_metadata),
        'metadata_extension_level': metadata_extension_level,
        'metadata_stats_mode': str(metadata_stats_mode),
        'metadata_stats_sample_windows': int(metadata_stats_sample_windows),
        'metadata_stats_seed': int(metadata_stats_seed),
        'metadata_histogram_buckets': metadata_histogram_buckets,
        'metadata_label_precision': metadata_label_precision,
        'validation_max_windows': int(validation_max_windows),
        'ancillary_status': 'not_processed',
        'ancillary_warp_method': None,
        'ancillary_tiepoints_used': 0,
        'ancillary_outputs': {"pan": None, "quality_vnir": None, "quality_swir": None},
        'ancillary_warnings': [],
        'run_manifest_path': skip_manifest_path,
        'displacement_vectors_path': None,
        'dataset_xlsx_path': skip_dataset_xlsx_path,
    }
    if extra_metrics:
        skip_metrics.update(dict(extra_metrics))

    skip_manifest = {
        "manifest_schema_version": RUN_MANIFEST_SCHEMA_VERSION,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "scene": {
            "scene_name": scene_name,
            "filename": os.path.basename(hs_file),
            "sensor_type": hyp_type,
            "scene_idx": int(scene_idx),
            "scene_total": int(scene_total),
        },
        "paths": {
            "output_dir": output_dir,
            "run_manifest": skip_manifest_path,
            "displacement_vectors_path": None,
            "dataset_xlsx": skip_dataset_xlsx_path,
        },
        "input": {
            "hs_file": hs_file,
            "acquisition_time": hs_time.isoformat() if isinstance(hs_time, datetime) else None,
            "bbox": list(bbox) if bbox is not None else None,
        },
        "config": _sanitize_config_for_manifest(config),
        "summary": {"reason": reason},
    }
    if extra_summary:
        skip_manifest["summary"].update(dict(extra_summary))

    _write_scene_run_manifest(skip_manifest, skip_manifest_path)
    _write_single_scene_dataset_xlsx(skip_metrics, skip_dataset_xlsx_path)
    return skip_metrics


def _normalize_raster_bands_01_inplace(tif_path, nodata=PROCESSING_NODATA):
    """Normalize each band to [0, 1] on valid pixels, preserving nodata."""
    t0 = perf_counter()
    with rasterio.open(tif_path, "r+") as dst:
        total_bands = int(dst.count)
        logger.info(f"Normalizing {total_bands} bands to [0,1]...")
        for bidx in range(1, dst.count + 1):
            bmin = np.inf
            bmax = -np.inf
            n_valid = 0
            for _, window in dst.block_windows(bidx):
                band = dst.read(bidx, window=window).astype(np.float32, copy=False)
                valid = np.isfinite(band) & (band != nodata)
                if not np.any(valid):
                    continue
                vals = band[valid]
                n_valid += int(vals.size)
                local_min = float(np.min(vals))
                local_max = float(np.max(vals))
                if local_min < bmin:
                    bmin = local_min
                if local_max > bmax:
                    bmax = local_max

            if n_valid == 0:
                logger.warning(f"Band {bidx}: no valid pixels found; skipping normalization.")
                continue

            if bmax <= bmin:
                logger.info(f"Band {bidx}: constant values ({bmin:.6g}); normalized to 0.0.")
                for _, window in dst.block_windows(bidx):
                    band = dst.read(bidx, window=window).astype(np.float32, copy=False)
                    valid = np.isfinite(band) & (band != nodata)
                    if np.any(valid):
                        band[valid] = 0.0
                        dst.write(band, bidx, window=window)
            else:
                logger.debug(f"Band {bidx}: normalized with min={bmin:.6g}, max={bmax:.6g}.")
                scale = (bmax - bmin)
                for _, window in dst.block_windows(bidx):
                    band = dst.read(bidx, window=window).astype(np.float32, copy=False)
                    valid = np.isfinite(band) & (band != nodata)
                    if np.any(valid):
                        band[valid] = (band[valid] - bmin) / scale
                        dst.write(band, bidx, window=window)
            if bidx % 25 == 0 or bidx == total_bands:
                logger.info(f"Normalization progress: {bidx}/{total_bands} bands")
    logger.info(f"Normalization finished in {perf_counter() - t0:.2f}s")


def _build_internal_overviews(
    tif_path: str,
    factors: Tuple[int, ...] = (2, 4, 8, 16, 32),
) -> Dict[str, Any]:
    """Build internal overviews for a GeoTIFF output."""
    result = {"ok": False, "levels": [], "warnings": [], "errors": []}
    try:
        with rasterio.open(tif_path, "r+") as dst:
            if dst.count < 1:
                result["errors"].append("Raster has no bands for overview generation.")
                return result

            valid_levels = [
                int(level)
                for level in factors
                if int(level) > 1 and (dst.width // int(level) >= 1) and (dst.height // int(level) >= 1)
            ]
            valid_levels = sorted(set(valid_levels))
            if not valid_levels:
                result["ok"] = True
                result["warnings"].append(
                    "No valid overview levels for raster dimensions; skipping overview generation."
                )
                return result

            dst.build_overviews(valid_levels, Resampling.average)
            dst.update_tags(ns="rio_overview", resampling="average")

        with rasterio.open(tif_path) as src:
            built_levels = list(src.overviews(1))
        result["ok"] = True
        result["levels"] = built_levels
        if not built_levels:
            result["warnings"].append("Overview generation completed but no overviews were reported.")
        return result
    except Exception as e:
        result["errors"].append(str(e))
        return result


def _strip_aux_matching_band_if_present(source_path, wl):
    """
    Ensure raster band count matches spectral metadata.

    If source has one extra trailing auxiliary band, write a temporary stripped
    raster with the last band removed and return its path.
    """
    if wl is None:
        raise RuntimeError(_fmt_issue("METADATA", "Missing wavelength metadata while preparing final output bands."))

    expected_bands = int(len(np.asarray(wl).reshape(-1)))
    if expected_bands <= 0:
        raise RuntimeError(
            _fmt_issue("METADATA", "Invalid wavelength metadata length while preparing final output bands.")
        )

    with rasterio.open(source_path) as src:
        src_count = int(src.count)
        if src_count == expected_bands:
            return source_path, None
        if src_count != expected_bands + 1:
            raise RuntimeError(
                _fmt_issue(
                    "METADATA",
                    f"Output band count mismatch before metadata write ({src_count} vs expected {expected_bands}).",
                )
            )

        target_dir = os.path.dirname(source_path) or "."
        fd, stripped_path = tempfile.mkstemp(prefix="coreg_stripped_", suffix=".tif", dir=target_dir)
        os.close(fd)
        if os.path.exists(stripped_path):
            os.remove(stripped_path)

    copy_result = stream_copy_raster_to_path(
        source_path=source_path,
        output_path=stripped_path,
        nodata_fallback=PROCESSING_NODATA,
        expected_band_count=expected_bands,
        tile_size=256,
        collect_band_stats=False,
    )
    if not copy_result.get("ok", False):
        err = "; ".join(list(copy_result.get("errors", []) or ["unknown strip copy error"]))
        raise RuntimeError(_fmt_issue("METADATA", f"Failed to stream-strip aux matching band: {err}"))

    logger.info(f"Removed auxiliary matching band from final source ({src_count} -> {expected_bands}).")
    return stripped_path, stripped_path


def _save_precoreg_output(
    source_path: str,
    output_path: str,
    wl: np.ndarray,
    source_bands_1based: Optional[Sequence[int]] = None,
) -> str:
    """Save pre-coreg raster output, applying optional EnMAP band selection/reorder."""
    if not source_path or not os.path.exists(source_path):
        raise RuntimeError(_fmt_issue("PRE_COREG", f"Missing source pre-coreg raster: {source_path}"))

    prepared_source = source_path
    stripped_temp_path = None
    try:
        expected_bands = int(len(np.asarray(wl).reshape(-1)))
        if source_bands_1based is not None:
            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            copy_result = stream_copy_raster_to_path(
                source_path=source_path,
                output_path=output_path,
                nodata_fallback=PROCESSING_NODATA,
                expected_band_count=expected_bands,
                source_bands_1based=source_bands_1based,
                tile_size=256,
                collect_band_stats=False,
            )
            if not copy_result.get("ok", False):
                err = "; ".join(list(copy_result.get("errors", []) or ["unknown pre-coreg copy error"]))
                raise RuntimeError(_fmt_issue("PRE_COREG", f"Failed to save pre-coreg output: {err}"))
            if not os.path.exists(output_path):
                raise RuntimeError(_fmt_issue("PRE_COREG", f"Failed to write pre-coreg output: {output_path}"))
            return output_path

        prepared_source, stripped_temp_path = _strip_aux_matching_band_if_present(source_path, wl)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        copy_result = stream_copy_raster_to_path(
            source_path=prepared_source,
            output_path=output_path,
            nodata_fallback=PROCESSING_NODATA,
            expected_band_count=expected_bands,
            source_bands_1based=None,
            tile_size=256,
            collect_band_stats=False,
        )
        if not copy_result.get("ok", False):
            err = "; ".join(list(copy_result.get("errors", []) or ["unknown pre-coreg copy error"]))
            raise RuntimeError(_fmt_issue("PRE_COREG", f"Failed to save pre-coreg output: {err}"))
        if not os.path.exists(output_path):
            raise RuntimeError(_fmt_issue("PRE_COREG", f"Failed to write pre-coreg output: {output_path}"))
        return output_path
    finally:
        if stripped_temp_path and os.path.exists(stripped_temp_path):
            try:
                os.remove(stripped_temp_path)
            except Exception:
                pass


def _finalize_coreg_output(
    source_path,
    output_path,
    hyp_type,
    wl,
    fwhm,
    band_names,
    band_detectors,
    source_bands_1based: Optional[Sequence[int]] = None,
    remove_source=True,
    normalization_params: Optional[NormalizationParams] = None,
    remove_detector_overlap=False,
    strict_metadata=True,
    metadata_extension_level="stats",
    metadata_stats_mode="exact",
    metadata_stats_sample_windows=128,
    metadata_stats_seed=1337,
    metadata_histogram_buckets=64,
    metadata_label_precision=2,
    build_overviews=False,
    timing_logs=True,
):
    """Finalize coregistration output with metadata."""
    t0_total = perf_counter()
    t_prepare = 0.0
    t_transfer = 0.0
    t_normalize = 0.0
    t_overviews = 0.0
    t_metadata = 0.0
    t_validation = 0.0
    try:
        params = normalization_params or NormalizationParams(mode="none")
        params, sanitize_warnings = sanitize_normalization_params(params)
        normalization_mode = normalize_mode(params.mode, default="none")
        stats_mode = _normalize_metadata_stats_mode(metadata_stats_mode)
        stats_sample_windows = max(
            1,
            _safe_parse_int(metadata_stats_sample_windows, 128, "metadata_stats_sample_windows"),
        )
        stats_seed = _safe_parse_int(metadata_stats_seed, 1337, "metadata_stats_seed")
        normalization_cfg = {
            "mode": normalization_mode,
            "p_low": float(params.p_low),
            "p_high": float(params.p_high),
            "clip": bool(params.clip),
            "eps": float(params.eps),
            "min_valid_pixels": int(params.min_valid_pixels),
            "reservoir_size": int(params.reservoir_size),
            "seed": int(params.seed),
            "tile_size": int(params.tile_size),
        }

        t0 = perf_counter()
        if wl is None:
            raise RuntimeError(_fmt_issue("METADATA", "Missing wavelength metadata for finalize output bands."))
        expected_bands = int(len(np.asarray(wl).reshape(-1)))
        if expected_bands <= 0:
            raise RuntimeError(_fmt_issue("METADATA", "Invalid wavelength metadata length for finalize output bands."))
        with rasterio.open(source_path) as src_check:
            src_count = int(src_check.count)
        selected_source_bands: Optional[List[int]] = None
        if source_bands_1based is not None:
            try:
                selected_source_bands = [int(b) for b in source_bands_1based]
            except Exception:
                raise RuntimeError(
                    _fmt_issue("METADATA", f"Invalid source_bands_1based: {source_bands_1based}")
                )
            if not selected_source_bands:
                raise RuntimeError(_fmt_issue("METADATA", "source_bands_1based cannot be empty."))
            invalid = [b for b in selected_source_bands if b < 1 or b > src_count]
            if invalid:
                raise RuntimeError(
                    _fmt_issue(
                        "METADATA",
                        f"source_bands_1based contains out-of-range indexes for source count {src_count}: "
                        f"{invalid[:5]}",
                    )
                )
            if len(selected_source_bands) != expected_bands:
                raise RuntimeError(
                    _fmt_issue(
                        "METADATA",
                        f"source_bands_1based length mismatch ({len(selected_source_bands)} vs expected "
                        f"{expected_bands}).",
                    )
                )
            needs_strip = False
            needs_reorder = selected_source_bands != list(range(1, expected_bands + 1))
        elif src_count == expected_bands:
            needs_strip = False
            needs_reorder = False
        elif src_count == expected_bands + 1:
            needs_strip = True
            needs_reorder = False
        else:
            raise RuntimeError(
                _fmt_issue(
                    "METADATA",
                    f"Output band count mismatch before finalize ({src_count} vs expected {expected_bands}).",
                )
            )
        t_prepare = perf_counter() - t0

        t0 = perf_counter()
        rewrite_result: Dict[str, Any] = {
            "ok": True,
            "mode": normalization_mode,
            "band_statistics": [],
            "output_validation": {},
            "warnings": [],
            "raster_passes": 0,
            "source_bands": src_count,
            "output_bands": expected_bands,
        }
        needs_normalize = normalization_mode != "none"
        needs_rewrite = bool(needs_normalize or needs_strip or needs_reorder)
        metadata_extension = _normalize_metadata_extension_level(metadata_extension_level)
        metadata_stats_requested = metadata_extension in {"stats", "full"} and stats_mode != "none"
        collect_rewrite_stats = metadata_stats_requested and stats_mode == "exact"

        if needs_rewrite:
            if needs_normalize:
                logger.info(
                    "Applying normalization mode '%s' (p_low=%.3g, p_high=%.3g, clip=%s).",
                    normalization_mode,
                    float(params.p_low),
                    float(params.p_high),
                    bool(params.clip),
                )
                rewrite_result = normalize_raster_to_path(
                    source_path,
                    output_path,
                    params=params,
                    nodata_fallback=PROCESSING_NODATA,
                    expected_band_count=expected_bands,
                    source_bands_1based=selected_source_bands,
                )
                if not rewrite_result.get("ok", False):
                    err_msg = "; ".join(list(rewrite_result.get("errors", []) or ["unknown normalization error"]))
                    raise RuntimeError(_fmt_issue("NORMALIZATION", f"Failed to normalize output: {err_msg}"))
                t_normalize = float(rewrite_result.get("timings", {}).get("total_s", 0.0))
            else:
                rewrite_result = stream_copy_raster_to_path(
                    source_path,
                    output_path,
                    nodata_fallback=PROCESSING_NODATA,
                    expected_band_count=expected_bands,
                    source_bands_1based=selected_source_bands,
                    tile_size=int(params.tile_size),
                    collect_band_stats=bool(collect_rewrite_stats),
                )
                if not rewrite_result.get("ok", False):
                    err_msg = "; ".join(list(rewrite_result.get("errors", []) or ["unknown copy error"]))
                    raise RuntimeError(_fmt_issue("FINALIZE", f"Failed streaming copy/strip: {err_msg}"))
                if needs_strip:
                    logger.info(f"Dropped trailing auxiliary matching band ({src_count} -> {expected_bands}) in copy pass.")
                if needs_reorder:
                    logger.info("Applied source band reordering in finalize copy pass.")
                t_transfer = float(rewrite_result.get("timings", {}).get("copy_s", 0.0))
        else:
            if source_path != output_path:
                if remove_source:
                    try:
                        shutil.move(source_path, output_path)
                    except Exception:
                        shutil.copy2(source_path, output_path)
                        if os.path.exists(source_path):
                            os.remove(source_path)
                else:
                    shutil.copy2(source_path, output_path)
            t_transfer = perf_counter() - t0
            sanitize_result = _sanitize_raster_nonfinite_inplace(output_path, nodata=PROCESSING_NODATA)
            if not sanitize_result.get("ok", False):
                raise RuntimeError(
                    _fmt_issue(
                        "FINALIZE",
                        f"Failed post-copy non-finite sanitation: {sanitize_result.get('error', 'unknown error')}",
                    )
                )
            replaced_nf = int(sanitize_result.get("nonfinite_replaced", 0))
            if replaced_nf > 0:
                msg = (
                    f"Normalization mode 'none': replaced {replaced_nf} non-finite output pixels with nodata."
                )
                logger.warning(_fmt_issue("FINALIZE", msg))
                rewrite_result["warnings"].append(msg)
            rewrite_result["output_validation"] = {}
            rewrite_result["band_statistics"] = []
            rewrite_result["raster_passes"] = 0

        if remove_source and source_path != output_path and os.path.exists(source_path):
            os.remove(source_path)

        t0 = perf_counter()
        all_warnings: List[str] = []
        for msg in sanitize_warnings:
            all_warnings.append(_fmt_issue("NORMALIZATION", msg))
        for msg in list(rewrite_result.get("warnings", [])):
            all_warnings.append(_fmt_issue("NORMALIZATION", str(msg)))
        metadata_status = "ok"
        band_stats = list(rewrite_result.get("band_statistics", []))
        output_validation = dict(rewrite_result.get("output_validation", {}))
        raster_passes = int(rewrite_result.get("raster_passes", 0))

        if build_overviews:
            logger.info("Building internal GeoTIFF overviews (levels: 2,4,8,16,32).")
            t0_ovr = perf_counter()
            overview_result = _build_internal_overviews(output_path, factors=(2, 4, 8, 16, 32))
            t_overviews = perf_counter() - t0_ovr
            for warning_msg in overview_result.get("warnings", []):
                wrapped = _fmt_issue("OVERVIEWS", warning_msg)
                logger.warning(wrapped)
                all_warnings.append(wrapped)
            if overview_result.get("ok", False):
                levels_txt = overview_result.get("levels", [])
                if levels_txt:
                    logger.info(f"Internal overviews created: {levels_txt}")
            else:
                errors = list(overview_result.get("errors", []) or ["Unknown overview generation error"])
                error_msg = "; ".join(errors)
                if strict_metadata:
                    raise RuntimeError(_fmt_issue("OVERVIEWS", f"Failed to build overviews: {error_msg}"))
                degraded_msg = _fmt_issue("OVERVIEWS", f"Overview generation degraded: {error_msg}")
                logger.warning(degraded_msg)
                all_warnings.append(degraded_msg)
                metadata_status = "degraded"

        geotiff_result = _write_geotiff_band_metadata(
            output_path,
            hyp_type,
            wl,
            fwhm,
            band_names,
            band_detectors,
            metadata_schema_version=METADATA_SCHEMA_VERSION,
            normalization_mode=normalization_mode,
            normalization_params=normalization_cfg,
            remove_detector_overlap=remove_detector_overlap,
            label_precision=metadata_label_precision,
            metadata_extension_level=metadata_extension_level,
            metadata_histogram_buckets=metadata_histogram_buckets,
        )
        all_warnings.extend(list(geotiff_result.get("warnings", [])))
        if not geotiff_result.get("ok", False):
            geo_errors = list(geotiff_result.get("errors", []) or ["Unknown GeoTIFF metadata error"])
            error_msg = "; ".join(geo_errors)
            if strict_metadata or _is_critical_metadata_error(geo_errors):
                raise RuntimeError(_fmt_issue("METADATA", f"Failed to write GeoTIFF metadata: {error_msg}"))
            logger.warning(_fmt_issue("METADATA", f"GeoTIFF metadata degraded: {error_msg}"))
            all_warnings.append(_fmt_issue("METADATA", f"GeoTIFF metadata degraded: {error_msg}"))
            metadata_status = "degraded"

        envi_result = _write_envi_header(
            output_path,
            hyp_type,
            wl,
            fwhm,
            band_names,
            band_detectors,
            metadata_schema_version=METADATA_SCHEMA_VERSION,
            remove_detector_overlap=remove_detector_overlap,
            label_precision=metadata_label_precision,
            metadata_extension_level=metadata_extension_level,
            metadata_histogram_buckets=metadata_histogram_buckets,
        )
        all_warnings.extend(list(envi_result.get("warnings", [])))
        if not envi_result.get("ok", False):
            envi_errors = list(envi_result.get("errors", []) or ["Unknown ENVI metadata error"])
            hdr_path = str(Path(output_path).with_suffix(".hdr"))
            if os.path.exists(hdr_path):
                try:
                    os.remove(hdr_path)
                except Exception as cleanup_err:
                    logger.warning(
                        _fmt_issue("METADATA", f"Failed to clean stale ENVI header after write failure: {cleanup_err}")
                    )
            error_msg = "; ".join(envi_errors)
            if strict_metadata:
                raise RuntimeError(_fmt_issue("METADATA", f"Failed to write ENVI metadata: {error_msg}"))
            logger.warning(_fmt_issue("METADATA", f"ENVI metadata degraded: {error_msg}"))
            all_warnings.append(_fmt_issue("METADATA", f"ENVI metadata degraded: {error_msg}"))
            metadata_status = "degraded"

        aux_path = str(Path(output_path).with_suffix(".aux.xml"))
        if metadata_extension == "none":
            if os.path.exists(aux_path):
                try:
                    os.remove(aux_path)
                except Exception as cleanup_err:
                    warning_msg = _fmt_issue(
                        "METADATA",
                        f"Failed to remove stale PAM metadata sidecar: {cleanup_err}",
                    )
                    if strict_metadata:
                        raise RuntimeError(warning_msg)
                    logger.warning(warning_msg)
                    all_warnings.append(warning_msg)
                    metadata_status = "degraded"
        else:
            pam_result = _write_pam_aux_xml(
                output_path,
                hyp_type,
                wl,
                fwhm,
                band_names,
                band_detectors,
                metadata_schema_version=METADATA_SCHEMA_VERSION,
                normalization_mode=normalization_mode,
                normalization_params=normalization_cfg,
                remove_detector_overlap=remove_detector_overlap,
                metadata_extension_level=metadata_extension,
                metadata_stats_mode=stats_mode,
                metadata_stats_sample_windows=stats_sample_windows,
                metadata_stats_seed=stats_seed,
                metadata_histogram_buckets=metadata_histogram_buckets,
                label_precision=metadata_label_precision,
                nodata_value=PROCESSING_NODATA,
                band_stats=band_stats if band_stats else None,
            )
            all_warnings.extend(list(pam_result.get("warnings", [])))
            if not pam_result.get("ok", False):
                pam_errors = list(pam_result.get("errors", []) or ["Unknown PAM metadata error"])
                error_msg = "; ".join(pam_errors)
                if strict_metadata or _is_critical_metadata_error(pam_errors):
                    raise RuntimeError(_fmt_issue("METADATA", f"Failed to write PAM metadata: {error_msg}"))
                logger.warning(_fmt_issue("METADATA", f"PAM metadata degraded: {error_msg}"))
                all_warnings.append(_fmt_issue("METADATA", f"PAM metadata degraded: {error_msg}"))
                metadata_status = "degraded"
            raster_passes += int(pam_result.get("raster_passes", 0))
        t_metadata = perf_counter() - t0

        if all_warnings:
            for warning_msg in all_warnings:
                logger.warning(_fmt_issue("METADATA", f"Metadata warning: {warning_msg}"))
            if metadata_status != "degraded":
                metadata_status = "degraded"

        logger.info(f"Finalized output: {output_path}")
        if timing_logs:
            logger.info(
                "Finalize timings (s): prepare_source=%.2f, transfer=%.2f, normalize=%.2f, "
                "validation=%.2f, overviews=%.2f, metadata=%.2f, total=%.2f",
                t_prepare, t_transfer, t_normalize, t_validation, t_overviews, t_metadata, perf_counter() - t0_total
            )
            logger.info("Finalize raster passes: %d", int(raster_passes))
        return {
            "metadata_status": metadata_status,
            "metadata_warnings": all_warnings,
            "metadata_schema_version": METADATA_SCHEMA_VERSION,
            "normalization_mode": normalization_mode,
            "normalization_params": normalization_cfg,
            "band_statistics": band_stats,
            "output_validation": output_validation,
            "raster_passes": int(raster_passes),
            "timings": {
                "prepare_source_s": t_prepare,
                "transfer_s": t_transfer,
                "normalize_s": t_normalize,
                "validation_s": t_validation,
                "overviews_s": t_overviews,
                "metadata_s": t_metadata,
                "total_s": perf_counter() - t0_total,
            },
        }
    finally:
        pass


def _cleanup_temp_folder(temp_folder_path, keep_temp_files=False):
    """Clean up temporary folder."""
    if keep_temp_files:
        return False
    if not os.path.isdir(temp_folder_path):
        return False
    try:
        shutil.rmtree(temp_folder_path, ignore_errors=True)
        return True
    except Exception:
        return False


def _promote_s2_stack(source_path, target_path, temp_root=None):
    """
    Promote selected S2 stack to reference location.

    Use move when source is inside temp workspace to avoid an extra full copy.
    """
    if not source_path or not target_path:
        return
    if os.path.abspath(source_path) == os.path.abspath(target_path):
        return

    target_dir = os.path.dirname(target_path) or "."
    os.makedirs(target_dir, exist_ok=True)

    source_abs = os.path.abspath(source_path)
    target_abs = os.path.abspath(target_path)
    temp_abs = os.path.abspath(temp_root) if temp_root else None

    should_move = False
    if temp_abs:
        try:
            should_move = os.path.commonpath([source_abs, temp_abs]) == temp_abs
        except ValueError:
            should_move = False

    if should_move:
        try:
            shutil.move(source_abs, target_abs)
            return
        except Exception as e:
            logger.warning(f"Could not move S2 stack to reference location; falling back to copy ({e})")

    shutil.copy2(source_abs, target_abs)


def _emit_progress(
    progress_callback: Optional[Callable[[Dict[str, Any]], None]],
    stage: str,
    scene_idx: int = 1,
    scene_total: int = 1,
    status: str = "running",
    **extra: Any,
) -> None:
    """Emit progress payload to optional callback."""
    if progress_callback is None:
        return
    payload = {
        "stage": str(stage),
        "scene_idx": int(scene_idx),
        "scene_total": max(1, int(scene_total)),
        "status": str(status),
    }
    if extra:
        payload.update(extra)
    try:
        progress_callback(payload)
    except Exception as e:
        logger.debug(f"Progress callback failed: {e}")


class _ProgressHeartbeat:
    """Emit periodic progress updates while long-running operations are in flight."""

    def __init__(
        self,
        progress_callback: Optional[Callable[[Dict[str, Any]], None]],
        stage: str,
        *,
        scene_idx: int = 1,
        scene_total: int = 1,
        substage: Optional[str] = None,
        interval_s: float = PROGRESS_HEARTBEAT_INTERVAL_S,
        log_heartbeats: bool = False,
        **extra: Any,
    ) -> None:
        self.progress_callback = progress_callback
        self.stage = str(stage)
        self.scene_idx = int(scene_idx)
        self.scene_total = int(scene_total)
        self.substage = str(substage) if substage else None
        self.interval_s = max(0.1, float(interval_s))
        self.log_heartbeats = bool(log_heartbeats)
        self.extra = dict(extra) if extra else {}

        self._start_t: Optional[float] = None
        self._count = 0
        self._stop_event = threading.Event()
        self._thread: Optional[threading.Thread] = None
        self._stopped = False

    def _elapsed_s(self) -> float:
        if self._start_t is None:
            return 0.0
        return max(0.0, perf_counter() - self._start_t)

    def _emit(self, *, heartbeat: bool, status: str = "running", **extra: Any) -> None:
        payload: Dict[str, Any] = {
            "elapsed_s": round(self._elapsed_s(), 1),
            "heartbeat": bool(heartbeat),
        }
        if self.substage:
            payload["substage"] = self.substage
        if self.extra:
            payload.update(self.extra)
        if extra:
            payload.update(extra)
        _emit_progress(
            self.progress_callback,
            self.stage,
            scene_idx=self.scene_idx,
            scene_total=self.scene_total,
            status=status,
            **payload,
        )

    def _run(self) -> None:
        while not self._stop_event.wait(self.interval_s):
            self._count += 1
            self._emit(heartbeat=True, heartbeat_count=self._count)

    def start(self) -> "_ProgressHeartbeat":
        self._start_t = perf_counter()
        self._emit(heartbeat=False)
        if self.progress_callback is None:
            return self
        self._thread = threading.Thread(
            target=self._run,
            name="hypercoreg_progress_heartbeat",
            daemon=True,
        )
        self._thread.start()
        return self

    def update(self, *, substage: Optional[str] = None, **extra: Any) -> None:
        if substage is not None:
            self.substage = str(substage)
        self._emit(heartbeat=False, **extra)

    def stop(self, *, status: str = "running", **extra: Any) -> None:
        if self._stopped:
            return
        self._stopped = True
        self._stop_event.set()
        if self._thread is not None and self._thread.is_alive():
            self._thread.join(timeout=min(1.0, self.interval_s))
        self._emit(heartbeat=False, status=status, **extra)

    def __enter__(self) -> "_ProgressHeartbeat":
        return self.start()

    def __exit__(self, exc_type, exc, tb) -> bool:
        self.stop()
        return False


def run_coregistration(
    hs_file: str,
    hyp_type: str,
    output_dir: str,
    config: Dict[str, Any],
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
    scene_idx: int = 1,
    scene_total: int = 1,
) -> Dict[str, Any]:
    """
    Run coregistration on a single hyperspectral file.

    This is the main entry point for processing a single PRISMA or EnMAP file.

    Args:
        hs_file: Path to hyperspectral file (.he5 for PRISMA, SPECTRAL_IMAGE.TIF/.BSQ for EnMAP)
        hyp_type: Sensor type ("PRISMA" or "ENMAP")
        output_dir: Output directory for results
        config: Configuration dictionary with processing parameters
        progress_callback: Optional callback receiving progress events
        scene_idx: Current scene index (1-based)
        scene_total: Total number of scenes

    Returns:
        dict: Metrics dictionary with processing results

    Raises:
        SentinelNotFoundError: If no suitable Sentinel-2 reference is found
        ValueError: If input file is invalid
    """
    _emit_progress(progress_callback, "Initializing", scene_idx=scene_idx, scene_total=scene_total)
    log_section_header(f"COREGISTRATION: {os.path.basename(hs_file)}")

    logger.info(f"Sensor type: {hyp_type}")
    logger.info(f"Output directory: {output_dir}")

    os.makedirs(output_dir, exist_ok=True)

    # Extract config parameters with defaults
    days_window = config.get('days_window', 30)
    min_overlap = config.get('min_overlap', 0.5)
    max_cloud = config.get('max_cloud', 20)
    local_s2_stack_path = str(config.get('local_s2_stack_path', '') or '').strip() or None
    max_input_cloud = config.get('max_input_cloud_cover', 70)
    min_accuracy = config.get('min_accuracy', 70.0)
    max_displacement = config.get('max_displacement', 350.0)
    residual_threshold = config.get('residual_threshold', 25.0)
    min_tie_points = config.get('min_tie_points', 10)
    max_s2_candidates = config.get('max_s2_candidates', 3)
    residual_mad_factor = config.get('residual_mad_factor', 3.0)
    s2_ref_band = config.get('s2_ref_band', 8)
    prefer_fixed_band_pairs = bool(
        config.get('prefer_fixed_band_pairs', DEFAULT_CONFIG.get('prefer_fixed_band_pairs', True))
    )
    fixed_band_pairs_by_sensor = config.get(
        'fixed_band_pairs_by_sensor',
        DEFAULT_CONFIG.get('fixed_band_pairs_by_sensor', {}),
    )
    bandpair_wavelength_window_nm = _safe_parse_float(
        config.get('bandpair_wavelength_window_nm', DEFAULT_CONFIG.get('bandpair_wavelength_window_nm', 20.0)),
        float(DEFAULT_CONFIG.get('bandpair_wavelength_window_nm', 20.0)),
        'bandpair_wavelength_window_nm',
    )
    bandpair_wavelength_window_nm = max(0.1, bandpair_wavelength_window_nm)
    min_band_support = _safe_parse_int(
        config.get('min_band_support', DEFAULT_CONFIG.get('min_band_support', 2)),
        int(DEFAULT_CONFIG.get('min_band_support', 2)),
        'min_band_support',
    )
    min_band_support = max(1, min_band_support)
    allow_single_band_fallback = bool(
        config.get('allow_single_band_fallback', DEFAULT_CONFIG.get('allow_single_band_fallback', True))
    )
    consensus_group_rounding_px = _safe_parse_float(
        config.get('consensus_group_rounding_px', DEFAULT_CONFIG.get('consensus_group_rounding_px', 1.0)),
        float(DEFAULT_CONFIG.get('consensus_group_rounding_px', 1.0)),
        'consensus_group_rounding_px',
    )
    consensus_group_rounding_px = max(0.1, consensus_group_rounding_px)
    spatial_grid_rows = _safe_parse_int(
        config.get('spatial_stratification_grid_rows', DEFAULT_CONFIG.get('spatial_stratification_grid_rows', 4)),
        int(DEFAULT_CONFIG.get('spatial_stratification_grid_rows', 4)),
        'spatial_stratification_grid_rows',
    )
    spatial_grid_cols = _safe_parse_int(
        config.get('spatial_stratification_grid_cols', DEFAULT_CONFIG.get('spatial_stratification_grid_cols', 4)),
        int(DEFAULT_CONFIG.get('spatial_stratification_grid_cols', 4)),
        'spatial_stratification_grid_cols',
    )
    spatial_grid_rows = max(1, spatial_grid_rows)
    spatial_grid_cols = max(1, spatial_grid_cols)
    max_points_per_cell = _safe_parse_int(
        config.get('max_points_per_cell', DEFAULT_CONFIG.get('max_points_per_cell', 3)),
        int(DEFAULT_CONFIG.get('max_points_per_cell', 3)),
        'max_points_per_cell',
    )
    max_points_per_cell = max(1, max_points_per_cell)
    preferred_polynomial_order = _safe_parse_int(
        config.get('preferred_polynomial_order', DEFAULT_CONFIG.get('preferred_polynomial_order', 2)),
        int(DEFAULT_CONFIG.get('preferred_polynomial_order', 2)),
        'preferred_polynomial_order',
    )
    preferred_polynomial_order = 1 if preferred_polynomial_order <= 1 else 2
    auto_downgrade_polynomial_order = bool(
        config.get(
            'auto_downgrade_polynomial_order',
            DEFAULT_CONFIG.get('auto_downgrade_polynomial_order', True),
        )
    )
    min_gcps_order2 = _safe_parse_int(
        config.get('min_gcps_order2', DEFAULT_CONFIG.get('min_gcps_order2', 12)),
        int(DEFAULT_CONFIG.get('min_gcps_order2', 12)),
        'min_gcps_order2',
    )
    min_cells_order2 = _safe_parse_int(
        config.get('min_cells_order2', DEFAULT_CONFIG.get('min_cells_order2', 6)),
        int(DEFAULT_CONFIG.get('min_cells_order2', 6)),
        'min_cells_order2',
    )
    min_gcps_order2 = max(6, min_gcps_order2)
    min_cells_order2 = max(1, min_cells_order2)
    postwarp_phasecorr_check = bool(
        config.get('postwarp_phasecorr_check', DEFAULT_CONFIG.get('postwarp_phasecorr_check', False))
    )
    postwarp_phasecorr_warn_threshold_px = _safe_parse_float(
        config.get(
            'postwarp_phasecorr_warn_threshold_px',
            DEFAULT_CONFIG.get('postwarp_phasecorr_warn_threshold_px', 1.0),
        ),
        float(DEFAULT_CONFIG.get('postwarp_phasecorr_warn_threshold_px', 1.0)),
        'postwarp_phasecorr_warn_threshold_px',
    )
    postwarp_phasecorr_warn_threshold_px = max(0.0, postwarp_phasecorr_warn_threshold_px)
    postwarp_phasecorr_reject_threshold_px = _safe_parse_float(
        config.get(
            'postwarp_phasecorr_reject_threshold_px',
            DEFAULT_CONFIG.get('postwarp_phasecorr_reject_threshold_px', 3.0),
        ),
        float(DEFAULT_CONFIG.get('postwarp_phasecorr_reject_threshold_px', 3.0)),
        'postwarp_phasecorr_reject_threshold_px',
    )
    postwarp_phasecorr_reject_threshold_px = max(
        postwarp_phasecorr_warn_threshold_px,
        postwarp_phasecorr_reject_threshold_px,
    )
    postwarp_phasecorr_reject_bad = bool(
        config.get(
            'postwarp_phasecorr_reject_bad',
            DEFAULT_CONFIG.get('postwarp_phasecorr_reject_bad', False),
        )
    )
    postwarp_phasecorr_max_dim = _safe_parse_int(
        config.get(
            'postwarp_phasecorr_max_dim',
            DEFAULT_CONFIG.get('postwarp_phasecorr_max_dim', 1024),
        ),
        int(DEFAULT_CONFIG.get('postwarp_phasecorr_max_dim', 1024)),
        'postwarp_phasecorr_max_dim',
    )
    postwarp_phasecorr_max_dim = max(128, postwarp_phasecorr_max_dim)
    use_geolocation_mesh_affine = bool(
        config.get('use_geolocation_mesh_affine', DEFAULT_CONFIG.get('use_geolocation_mesh_affine', False))
    )
    geolocation_mesh_stride = _safe_parse_int(
        config.get('geolocation_mesh_stride', DEFAULT_CONFIG.get('geolocation_mesh_stride', 32)),
        int(DEFAULT_CONFIG.get('geolocation_mesh_stride', 32)),
        'geolocation_mesh_stride',
    )
    geolocation_mesh_stride = max(1, geolocation_mesh_stride)
    save_pre = bool(config.get('save_pre', False))
    save_pan = config.get('save_pan', False)
    save_quality_mask = config.get('save_quality_mask', False)
    gen_tiepoint_pngs = config.get('gen_tiepoint_pngs', False)
    save_displacement_vectors = bool(
        config.get('save_displacement_vectors', DEFAULT_CONFIG.get('save_displacement_vectors', True))
    )
    keep_temp_files = config.get('keep_temp_files', False)
    allow_gui_prompt = config.get('allow_gui_prompt', False)
    prompt_userpass_fn = config.get('prompt_userpass_fn')
    remove_detector_overlap_bands = config.get('remove_detector_overlap_bands', False)
    normalization_params = _build_normalization_params_from_config(config)
    normalization_mode = str(normalization_params.mode)
    build_overviews = config.get('build_overviews', False)
    strict_metadata = config.get('strict_metadata', True)
    metadata_extension_level = _normalize_metadata_extension_level(
        config.get('metadata_extension_level', 'stats')
    )
    metadata_stats_mode = _normalize_metadata_stats_mode(
        config.get('metadata_stats_mode', 'exact')
    )
    enmap_metadata_stats_mode = _normalize_metadata_stats_mode(
        config.get('enmap_metadata_stats_mode', 'none')
    )
    if str(hyp_type).upper() == "ENMAP":
        metadata_stats_mode = enmap_metadata_stats_mode
    metadata_stats_sample_windows = _safe_parse_int(
        config.get('metadata_stats_sample_windows', 128),
        128,
        "metadata_stats_sample_windows",
    )
    metadata_stats_sample_windows = max(1, metadata_stats_sample_windows)
    metadata_stats_seed = _safe_parse_int(
        config.get('metadata_stats_seed', 1337),
        1337,
        "metadata_stats_seed",
    )
    metadata_histogram_buckets = _safe_parse_int(
        config.get('metadata_histogram_buckets', 64),
        64,
        "metadata_histogram_buckets",
    )
    metadata_histogram_buckets = min(4096, max(2, metadata_histogram_buckets))
    metadata_label_precision = _safe_parse_int(
        config.get('metadata_label_precision', 2),
        2,
        "metadata_label_precision",
    )
    metadata_label_precision = min(6, max(0, metadata_label_precision))
    validation_max_windows = _safe_parse_int(
        config.get('validation_max_windows', 0),
        0,
        "validation_max_windows",
    )
    validation_max_windows = max(0, validation_max_windows)
    quicklook_max_dim = _safe_parse_int(
        config.get('quicklook_max_dim', DEFAULT_QUICKLOOK_MAX_DIM),
        DEFAULT_QUICKLOOK_MAX_DIM,
        "quicklook_max_dim",
    )
    quicklook_max_dim = max(256, quicklook_max_dim)
    quicklook_dpi = _safe_parse_int(
        config.get('quicklook_dpi', 220),
        220,
        "quicklook_dpi",
    )
    quicklook_dpi = max(120, min(400, quicklook_dpi))
    quicklook_rgb_targets_nm = _coerce_float_triplet(
        config.get('quicklook_rgb_targets_nm', (660.0, 550.0, 480.0)),
        default=(660.0, 550.0, 480.0),
    )
    quicklook_percentiles = _coerce_percentiles(
        config.get('quicklook_percentiles', (2.0, 98.0)),
        default=(2.0, 98.0),
    )
    try:
        quicklook_gamma = float(config.get('quicklook_gamma', 1.0))
    except Exception:
        quicklook_gamma = 1.0
    quicklook_gamma = max(1e-6, quicklook_gamma)
    quicklook_crop_to_valid = bool(config.get('quicklook_crop_to_valid', False))
    quicklook_scalebar = bool(config.get('quicklook_scalebar', True))
    quicklook_rgb_source_path = config.get('quicklook_rgb_source_path')
    if quicklook_rgb_source_path is not None:
        quicklook_rgb_source_path = str(quicklook_rgb_source_path).strip() or None
    pan_gcp_mode = _normalize_pan_gcp_mode(config.get('pan_gcp_mode', 'map_inverse'))
    pan_map_dxdy_source = _normalize_pan_dxdy_source(config.get('pan_map_dxdy_source', 'auto'))
    pan_target_aligned_pixels = bool(config.get('pan_target_aligned_pixels', False))
    pan_use_synthetic_reference = bool(
        config.get('pan_use_synthetic_reference', DEFAULT_CONFIG.get('pan_use_synthetic_reference', True))
    )
    pan_min_points_for_poly2 = _safe_parse_int(
        config.get('pan_min_points_for_poly2', DEFAULT_CONFIG.get('pan_min_points_for_poly2', 20)),
        int(DEFAULT_CONFIG.get('pan_min_points_for_poly2', 20)),
        'pan_min_points_for_poly2',
    )
    pan_min_points_for_poly2 = max(6, pan_min_points_for_poly2)
    pan_local_window_size = _coerce_window_size(
        config.get('pan_local_window_size', DEFAULT_CONFIG.get('pan_local_window_size', (512, 512))),
        default=(512, 512),
    )
    pan_residual_check = bool(config.get('pan_residual_check', False))
    pan_residual_threshold_px = _safe_parse_float(
        config.get('pan_residual_threshold_px', 0.5),
        0.5,
        'pan_residual_threshold_px',
    )
    pan_residual_threshold_px = max(0.0, pan_residual_threshold_px)
    pan_residual_max_dim = _safe_parse_int(
        config.get('pan_residual_max_dim', 1024),
        1024,
        'pan_residual_max_dim',
    )
    pan_residual_max_dim = max(128, pan_residual_max_dim)
    defer_temp_cleanup_gui = config.get('defer_temp_cleanup_gui', False)
    timing_logs = config.get('timing_logs', True)
    progress_heartbeat_interval_s = _safe_parse_float(
        config.get('progress_heartbeat_interval_s', PROGRESS_HEARTBEAT_INTERVAL_S),
        PROGRESS_HEARTBEAT_INTERVAL_S,
        'progress_heartbeat_interval_s',
    )
    progress_heartbeat_interval_s = max(1.0, progress_heartbeat_interval_s)
    matcher_profile = _resolve_sensor_matcher_profile(config, hyp_type)
    pan_local_grid_res = _safe_parse_int(
        config.get(
            'pan_local_grid_res',
            DEFAULT_CONFIG.get('pan_local_grid_res', matcher_profile.get("local_grid_res", LOCAL_GRID_RES_M)),
        ),
        int(DEFAULT_CONFIG.get('pan_local_grid_res', matcher_profile.get("local_grid_res", LOCAL_GRID_RES_M))),
        'pan_local_grid_res',
    )
    pan_local_grid_res = max(30, pan_local_grid_res)
    pan_local_max_shift = _safe_parse_float(
        config.get(
            'pan_local_max_shift',
            DEFAULT_CONFIG.get('pan_local_max_shift', matcher_profile.get("local_max_shift", 220.0)),
        ),
        float(DEFAULT_CONFIG.get('pan_local_max_shift', matcher_profile.get("local_max_shift", 220.0))),
        'pan_local_max_shift',
    )
    pan_local_max_shift = max(5.0, pan_local_max_shift)
    pan_local_tieP_filter_level = _safe_parse_int(
        config.get(
            'pan_local_tieP_filter_level',
            DEFAULT_CONFIG.get(
                'pan_local_tieP_filter_level',
                matcher_profile.get("local_tieP_filter_level", 1),
            ),
        ),
        int(
            DEFAULT_CONFIG.get(
                'pan_local_tieP_filter_level',
                matcher_profile.get("local_tieP_filter_level", 1),
            )
        ),
        'pan_local_tieP_filter_level',
    )
    pan_local_tieP_filter_level = max(0, pan_local_tieP_filter_level)
    pan_local_max_iter_raw = config.get(
        'pan_local_max_iter',
        DEFAULT_CONFIG.get('pan_local_max_iter', matcher_profile.get("local_max_iter")),
    )
    try:
        pan_local_max_iter = None if pan_local_max_iter_raw in (None, "", False) else max(1, int(pan_local_max_iter_raw))
    except Exception:
        pan_local_max_iter = None

    s2_ref_wl = S2_BANDS.get(s2_ref_band, {}).get('wavelength', S2_BAND08_CENTER_WL_NM)
    s2_stack_band_map = {2: 1, 3: 2, 4: 3, 8: 4, 11: 5, 12: 6}
    s2_stack_idx = s2_stack_band_map.get(s2_ref_band, 4)

    logger.info(f"Parameters: days_window={days_window}, max_cloud={max_cloud}, "
                f"residual_threshold={residual_threshold}, "
                f"normalization_mode={normalization_mode}, "
                f"build_overviews={build_overviews}, "
                f"metadata_extension_level={metadata_extension_level}, "
                f"metadata_stats_mode={metadata_stats_mode}")
    logger.info(
        "Tiepoint controls: prefer_fixed_band_pairs=%s, min_band_support=%d, "
        "allow_single_band_fallback=%s, spatial_grid=%dx%d, max_points_per_cell=%d",
        bool(prefer_fixed_band_pairs),
        int(min_band_support),
        bool(allow_single_band_fallback),
        int(spatial_grid_rows),
        int(spatial_grid_cols),
        int(max_points_per_cell),
    )
    logger.info(
        "Polynomial controls: preferred_order=%d, auto_downgrade=%s, min_gcps_order2=%d, min_cells_order2=%d",
        int(preferred_polynomial_order),
        bool(auto_downgrade_polynomial_order),
        int(min_gcps_order2),
        int(min_cells_order2),
    )
    logger.info(
        "Sensor matcher profile: sensor=%s source=%s local_shift=%.1f global_attempts=%s",
        matcher_profile.get("sensor"),
        matcher_profile.get("global_profile_source"),
        float(matcher_profile.get("local_max_shift", 0.0)),
        [
            {"ws": list(item.get("ws", ())), "max_shift": float(item.get("max_shift", 0.0))}
            for item in matcher_profile.get("global_attempt_ladder", [])
        ],
    )
    logger.info(
        "Post-warp QA: enabled=%s warn_threshold_px=%.2f reject_enabled=%s reject_threshold_px=%.2f",
        bool(postwarp_phasecorr_check),
        float(postwarp_phasecorr_warn_threshold_px),
        bool(postwarp_phasecorr_reject_bad),
        float(postwarp_phasecorr_reject_threshold_px),
    )
    logger.info(
        "PRISMA affine mode: use_geolocation_mesh_affine=%s stride=%d",
        bool(use_geolocation_mesh_affine),
        int(geolocation_mesh_stride),
    )
    logger.info(
        "PAN ancillary config: pan_gcp_mode=%s, pan_map_dxdy_source=%s, pan_target_aligned_pixels=%s, "
        "pan_use_synthetic_reference=%s, pan_min_points_for_poly2=%d, pan_local_ws=%s, pan_local_grid=%d, "
        "pan_local_max_shift=%.1f, pan_residual_check=%s",
        pan_gcp_mode,
        pan_map_dxdy_source,
        bool(pan_target_aligned_pixels),
        bool(pan_use_synthetic_reference),
        int(pan_min_points_for_poly2),
        tuple(pan_local_window_size),
        int(pan_local_grid_res),
        float(pan_local_max_shift),
        bool(pan_residual_check),
    )
    logger.debug("Progress heartbeat interval (s): %.1f", progress_heartbeat_interval_s)
    logger.info(f"Tiepoint quicklook PNGs enabled: {bool(gen_tiepoint_pngs)}")
    logger.info(f"Displacement vector outputs enabled: {bool(save_displacement_vectors)}")

    # READ HYPERSPECTRAL DATA
    _emit_progress(progress_callback, "Reading hyperspectral data", scene_idx=scene_idx, scene_total=scene_total)
    log_section_header("READING HYPERSPECTRAL DATA")

    pan_data, pan_geo_info = None, None
    vnir_quality_data, swir_quality_data, lat_qm, lon_qm = None, None, None, None
    enmap_sort_idx = None
    enmap_meta_merged: Dict[str, Any] = {}
    enmap_processing_source_path = hs_file

    if hyp_type == "PRISMA":
        cube, wl, hs_time, bbox, lat, lon, fwhm, band_names, band_detectors = read_prisma_cube_and_meta(
            hs_file,
            remove_detector_overlap=remove_detector_overlap_bands
        )
        extended_meta = extract_prisma_extended_metadata(hs_file)

        if save_pan:
            pan_data, pan_geo_info = read_prisma_pan_and_geo(hs_file)
        if save_quality_mask:
            vnir_quality_data, swir_quality_data, lat_qm, lon_qm = read_prisma_quality_mask(hs_file)
    else:
        # EnMAP workflow
        if Path(hs_file).suffix.lower() == ".bsq":
            hdr_candidates = [Path(hs_file).with_suffix(".hdr"), Path(hs_file).with_suffix(".HDR")]
            if not any(candidate.exists() for candidate in hdr_candidates):
                raise ValueError(
                    f"EnMAP BSQ input requires an ENVI header sidecar (.hdr): {hs_file}"
                )

        enmap_meta_path = find_enmap_metadata_for_spectral_image(hs_file)
        if not enmap_meta_path:
            raise ValueError(f"EnMAP metadata XML not found for {hs_file}")

        enmap_meta = read_enmap_metadata(enmap_meta_path, spectral_image_path=hs_file)
        injected_meta: Dict[str, Any] = {}
        hs_suffix = Path(hs_file).suffix.lower()
        if hs_suffix in {".tif", ".tiff"}:
            inject_result = inject_metadata_into_raster(hs_file, enmap_meta)
            if not inject_result.get("ok", False):
                logger.warning(
                    _fmt_issue(
                        "METADATA",
                        f"Failed to inject EnMAP metadata into raster header: {inject_result.get('error', 'unknown error')}",
                    )
                )
            else:
                logger.info(
                    _fmt_issue(
                        "METADATA",
                        (
                            "Injected EnMAP metadata into source raster "
                            f"(band_tags={int(inject_result.get('bands_updated', 0))}, "
                            f"dataset_tags={int(inject_result.get('dataset_tags_updated', 0))})."
                        ),
                    )
                )
            injected_meta = read_enmap_metadata_from_raster(hs_file)
        elif hs_suffix == ".bsq":
            logger.info(
                _fmt_issue(
                    "METADATA",
                    (
                        "Deferring EnMAP metadata injection until BSQ source is converted "
                        f"to temporary GeoTIFF: {hs_file}"
                    ),
                )
            )
        else:
            logger.info(
                _fmt_issue(
                    "METADATA",
                    f"Skipping EnMAP metadata injection for unsupported source extension: {hs_file}",
                )
            )

        enmap_meta_merged = dict(enmap_meta)
        if isinstance(injected_meta, dict):
            for key, val in injected_meta.items():
                if val is not None:
                    enmap_meta_merged[key] = val

        hs_time = enmap_meta_merged.get('acquisition_time')
        bbox = enmap_meta_merged.get('bbox')
        wl_raw = enmap_meta_merged.get('wavelengths')
        fwhm_raw = enmap_meta_merged.get('fwhm')
        band_names_raw = enmap_meta_merged.get('band_names')

        if hs_time is None:
            raise ValueError(f"EnMAP metadata missing acquisition_time: {enmap_meta_path}")
        if bbox is None or len(bbox) != 4:
            logger.warning(
                "EnMAP metadata missing/invalid bbox in XML (%s). "
                "Attempting raster-derived geographic bbox fallback.",
                enmap_meta_path,
            )
            bbox = derive_enmap_bbox_from_raster(hs_file)
            if bbox is not None and len(bbox) == 4:
                logger.warning(
                    "Using EnMAP raster-derived geographic bbox fallback: "
                    "(%.6f, %.6f, %.6f, %.6f)",
                    float(bbox[0]),
                    float(bbox[1]),
                    float(bbox[2]),
                    float(bbox[3]),
                )
            else:
                raise ValueError(
                    "EnMAP metadata missing/invalid bbox and raster fallback failed: "
                    f"{enmap_meta_path}"
                )
        if wl_raw is None or len(wl_raw) == 0:
            raise ValueError(f"EnMAP metadata missing wavelengths: {enmap_meta_path}")

        band_table, enmap_sort_idx = build_enmap_band_table(
            wl_raw,
            fwhm_raw,
            band_names_raw,
            remove_detector_overlap=remove_detector_overlap_bands,
            enmap_processing_version=enmap_meta_merged.get("enmap_processing_version"),
        )
        wl = band_table.wavelengths
        fwhm = band_table.fwhm
        band_names = band_table.band_names
        band_detectors = band_table.detectors
        cube = None
        diagnose_raster(hs_file, "INPUT ENMAP")
        extended_meta = {
            'prisma_id': enmap_meta_merged.get('enmap_id'),
            'prisma_date': enmap_meta_merged.get('enmap_date'),
            'prisma_cloud_pct': enmap_meta_merged.get('prisma_cloud_pct'),
            'enmap_cloud_pct': enmap_meta_merged.get('enmap_cloud_pct'),
            'enmap_haze_pct': enmap_meta_merged.get('enmap_haze_pct'),
            'enmap_cirrus_pct': enmap_meta_merged.get('enmap_cirrus_pct'),
            'enmap_snow_pct': enmap_meta_merged.get('enmap_snow_pct'),
            'enmap_water_pct': enmap_meta_merged.get('enmap_water_pct'),
            'sun_azimuth_angle': enmap_meta_merged.get('sun_azimuth_angle'),
            'sun_elevation_angle': enmap_meta_merged.get('sun_elevation_angle'),
            'sun_zenith_angle': enmap_meta_merged.get('sun_zenith_angle'),
            'across_offnadir_angle': enmap_meta_merged.get('across_offnadir_angle'),
            'along_offnadir_angle': enmap_meta_merged.get('along_offnadir_angle'),
            'scene_azimuth_angle': enmap_meta_merged.get('scene_azimuth_angle'),
            'observation_angle': enmap_meta_merged.get('observation_angle'),
        }
        if extended_meta.get('prisma_cloud_pct') is None and enmap_meta_merged.get('enmap_cloud_pct') is not None:
            extended_meta['prisma_cloud_pct'] = enmap_meta_merged.get('enmap_cloud_pct')
        matching_probe_bidx = int(np.argmin(np.abs(wl - s2_ref_wl))) + 1
        probe_bands = sorted({1, int(max(1, matching_probe_bidx)), int(max(1, len(wl)))})
        integrity = _probe_raster_valid_pixels(
            hs_file,
            bands_1based=probe_bands,
            nodata_fallback=PROCESSING_NODATA,
            sample_max_dim=512,
        )
        if not integrity.get("ok", False):
            hs_path = str(hs_file)
            hs_path_obj = Path(hs_path)
            hdr_upper = hs_path_obj.with_suffix(".HDR")
            hdr_lower = hs_path_obj.with_suffix(".hdr")
            aux_sidecar = Path(f"{hs_path}.aux.xml")
            enp_sidecar = Path(f"{hs_path}.enp")
            band_stats_txt = ", ".join(
                [f"b{int(k)}={int(v)}" for k, v in sorted((integrity.get("band_valid_pixels") or {}).items())]
            ) or "none"
            raise RuntimeError(
                _fmt_issue(
                    "HS_PREP",
                    (
                        "EnMAP source integrity check failed: no valid source pixels found "
                        f"(path={hs_path}, nodata={integrity.get('nodata')}, "
                        f"bands={integrity.get('band_count')}, probes=[{band_stats_txt}], "
                        f"sidecars: HDR={bool(hdr_upper.exists() or hdr_lower.exists())}, "
                        f"AUX_XML={bool(aux_sidecar.exists())}, ENP={bool(enp_sidecar.exists())})."
                    ),
                )
            )

    # Check cloud threshold
    cloud_passed, cloud_pct, cloud_reason = check_cloud_threshold(extended_meta, max_input_cloud, hyp_type)
    if not cloud_passed:
        logger.warning(f"Skipping image due to cloud threshold: {cloud_reason}")
        _emit_progress(
            progress_callback,
            "Skipped (cloud threshold)",
            scene_idx=scene_idx,
            scene_total=scene_total,
            status="done",
        )
        return _build_skip_result(
            hs_file=hs_file,
            hyp_type=hyp_type,
            output_dir=output_dir,
            scene_idx=scene_idx,
            scene_total=scene_total,
            hs_time=hs_time,
            bbox=bbox,
            config=config,
            status="SKIPPED_CLOUD",
            reason=cloud_reason or "Cloud threshold exceeded",
            normalization_mode=normalization_mode,
            normalization_params={
                "mode": normalization_mode,
                "p_low": float(normalization_params.p_low),
                "p_high": float(normalization_params.p_high),
                "clip": bool(normalization_params.clip),
                "eps": float(normalization_params.eps),
                "min_valid_pixels": int(normalization_params.min_valid_pixels),
                "reservoir_size": int(normalization_params.reservoir_size),
                "seed": int(normalization_params.seed),
                "tile_size": int(normalization_params.tile_size),
            },
            build_overviews=bool(build_overviews),
            remove_detector_overlap_bands=remove_detector_overlap_bands,
            strict_metadata=bool(strict_metadata),
            metadata_extension_level=metadata_extension_level,
            metadata_stats_mode=metadata_stats_mode,
            metadata_stats_sample_windows=metadata_stats_sample_windows,
            metadata_stats_seed=metadata_stats_seed,
            metadata_histogram_buckets=metadata_histogram_buckets,
            metadata_label_precision=metadata_label_precision,
            validation_max_windows=validation_max_windows,
            extra_summary={
                "cloud_pct": cloud_pct,
                "cloud_threshold_pct": max_input_cloud,
            },
            extra_metrics={
                "cloud_pct": cloud_pct,
                "threshold": max_input_cloud,
            },
        )

    # Create scene identifiers
    sensor_tag = hyp_type.upper()
    date_tag = hs_time.strftime("%y%m%d")
    unique_hex = uuid.uuid4().hex[:3].upper()
    scene_name = f"{sensor_tag}_{date_tag}_{unique_hex}"

    folder_struct = _create_scene_folder_structure(output_dir, sensor_tag, date_tag, unique_hex)
    scene_folder = folder_struct['scene_root']

    if hyp_type == "ENMAP":
        enmap_processing_source_path = hs_file
        if Path(hs_file).suffix.lower() == ".bsq":
            prepared_source_path = os.path.join(
                folder_struct["temp"],
                f"{scene_name}_ENMAP_SOURCE_PREP.tif",
            )
            prep_result = _prepare_enmap_processing_source(
                source_path=hs_file,
                output_path=prepared_source_path,
                metadata=enmap_meta_merged,
                inject_metadata=True,
            )
            if not prep_result.get("ok", False):
                raise RuntimeError(
                    _fmt_issue(
                        "HS_PREP",
                        f"Failed EnMAP BSQ preparation: {prep_result.get('error', 'unknown error')}",
                    )
                )
            enmap_processing_source_path = str(prep_result.get("path", prepared_source_path))
            if prep_result.get("metadata_injection_attempted", False):
                if prep_result.get("metadata_injected", False):
                    logger.info(
                        _fmt_issue(
                            "METADATA",
                            f"Injected EnMAP metadata into prepared source raster: {enmap_processing_source_path}",
                        )
                    )
                else:
                    logger.warning(
                        _fmt_issue(
                            "METADATA",
                            (
                                "Failed metadata injection for prepared EnMAP source raster: "
                                f"{prep_result.get('metadata_injection_error', 'unknown error')}"
                            ),
                        )
                    )
            injected_prepared_meta = prep_result.get("metadata_readback")
            if isinstance(injected_prepared_meta, dict):
                for key, val in injected_prepared_meta.items():
                    if val is not None:
                        enmap_meta_merged[key] = val

    detector_plan = _build_detector_branch_plan(
        wl=wl,
        fwhm=fwhm,
        band_names=band_names,
        band_detectors=band_detectors,
        sensor_type=hyp_type,
    )
    if not detector_plan.get("ok", False):
        raise RuntimeError(
            _fmt_issue(
                "HS_PREP",
                f"Failed detector branch planning: {detector_plan.get('error', 'unknown error')}",
            )
        )
    for warn_msg in detector_plan.get("warnings", []):
        logger.warning(_fmt_issue("HS_PREP", str(warn_msg)))

    bidx = int(np.argmin(np.abs(wl - s2_ref_wl))) + 1
    logger.info(f"Matching HS band: {bidx} ({wl[bidx - 1]:.1f}nm) -> S2 B{s2_ref_band:02d} ({s2_ref_wl:.0f}nm)")

    # Output paths
    s2_path = os.path.join(folder_struct['reference'], f"{scene_name}_Sentinel2_stack.tif")
    coreg_out = os.path.join(folder_struct['coreg'], f"{scene_name}_coreg.tif")
    shift_report_path = os.path.join(folder_struct['reports'], f"{scene_name}_shift_report.txt")
    metrics_json_path = os.path.join(folder_struct['reports'], f"{scene_name}_metrics.json")
    manifest_json_path = os.path.join(folder_struct['reports'], f"{scene_name}_run_manifest.json")
    dataset_xlsx_path = os.path.join(folder_struct['reports'], f"{scene_name}_DATASET.xlsx")
    displacement_vectors_path = os.path.join(
        folder_struct['reports'], f"{scene_name}_displacement_vectors.shp"
    )
    displacement_cartography_path = os.path.join(
        folder_struct['quicklooks'], f"{scene_name}_displacement_vectors.png"
    )

    # GET SENTINEL-2 DATA
    _emit_progress(progress_callback, "Fetching Sentinel-2 reference", scene_idx=scene_idx, scene_total=scene_total)
    log_section_header("FETCHING SENTINEL-2 REFERENCE")
    s2_candidates_list = []
    s2_crs = None
    selected_s2_metadata = {}
    session = config.get("_cdse_session")
    if session is not None:
        logger.info("Reusing CDSE session for this run.")

    if local_s2_stack_path:
        local_s2_validation = _validate_local_s2_stack_override(
            local_s2_stack_path,
            min_band_count=6,
        )
        if not local_s2_validation.get("ok", False):
            raise RuntimeError(
                _fmt_issue(
                    "S2_REF",
                    str(local_s2_validation.get("error", "invalid local S2 stack override")),
                )
            )
        local_s2_path = str(local_s2_validation["path"])
        logger.info("Using local S2 override stack: %s", local_s2_path)
        s2_candidates_list.append({'type': 'local_override', 'path': local_s2_path, 'Name': 'Local Override'})
        s2_crs = local_s2_validation.get("crs")
    elif os.path.isfile(s2_path):
        logger.info(f"S2 stack exists: {s2_path}")
        s2_candidates_list.append({'type': 'local', 'path': s2_path, 'Name': 'Local File'})
        with rasterio.open(s2_path) as src:
            s2_crs = src.crs
    else:
        if session is None:
            session = _create_cdse_session_with_retry(
                allow_gui_prompt=allow_gui_prompt,
                prompt_userpass_fn=prompt_userpass_fn,
            )
        config["_cdse_session"] = session
        items, session = _query_s2_with_retry(
            session,
            hs_time,
            bbox,
            days_window,
            max_cloud,
            allow_gui_prompt=allow_gui_prompt,
            prompt_userpass_fn=prompt_userpass_fn,
        )
        config["_cdse_session"] = session
        if items:
            s2_candidates_list = _rank_s2_candidates(items, hs_time, bbox, min_overlap)

    if not s2_candidates_list:
        _emit_progress(progress_callback, "Failed: no Sentinel-2 candidates", scene_idx=scene_idx, scene_total=scene_total, status="error")
        raise SentinelNotFoundError("No suitable Sentinel-2 candidates found.")

    # Process candidates
    final_success = False
    candidate_results = []
    final_validation = None
    best_candidate = None
    accepted_source_path = None
    pre_coreg_output_path = None
    quicklook_outputs = {"scene": None, "tiepoints": None}
    quicklook_status = "pending"
    quicklook_metadata: Dict[str, Any] = {}
    displacement_vectors_written_path: Optional[str] = None
    enmap_source_cache_by_crs: Dict[str, Dict[str, Any]] = {}
    final_output_validation: Dict[str, Any] = {
        "ok": False,
        "valid_pixels": 0,
        "total_pixels": 0,
        "valid_fraction": 0.0,
    }

    for cand_idx, s2_candidate in enumerate(s2_candidates_list[:max_s2_candidates]):
        logger.info(f"\n{'#' * 60}")
        logger.info(f"CANDIDATE {cand_idx + 1}/{min(max_s2_candidates, len(s2_candidates_list))}")

        try:
            # Prepare S2 stack
            if s2_candidate.get('type') in {'local', 'local_override'}:
                current_s2_path = s2_candidate['path']
                with rasterio.open(current_s2_path) as src:
                    s2_crs = src.crs
                selected_s2_metadata = {'product_id': os.path.basename(current_s2_path)}
            else:
                current_s2_path = os.path.join(folder_struct['temp'], f"{scene_name}_S2_cand{cand_idx}.tif")
                s2_prod_name = s2_candidate.get('Name', 'Unknown')
                selected_s2_metadata = {'product_id': s2_prod_name}

                if not os.path.exists(current_s2_path):
                    if session is None:
                        session = _create_cdse_session_with_retry(
                            allow_gui_prompt=allow_gui_prompt,
                            prompt_userpass_fn=prompt_userpass_fn,
                        )
                        config["_cdse_session"] = session
                    zip_path, session = _download_s2_product(
                        session, s2_candidate, folder_struct['temp'],
                        allow_gui_prompt=allow_gui_prompt,
                        progress_callback=progress_callback,
                        scene_idx=scene_idx,
                        scene_total=scene_total,
                        prompt_userpass_fn=prompt_userpass_fn,
                    )
                    config["_cdse_session"] = session
                    with _ProgressHeartbeat(
                        progress_callback,
                        "Building Sentinel-2 stack",
                        scene_idx=scene_idx,
                        scene_total=scene_total,
                        substage=f"Candidate {cand_idx + 1}: preparing stack",
                        interval_s=progress_heartbeat_interval_s,
                    ):
                        current_s2_path, s2_crs = _build_s2_stack(
                            zip_path,
                            bbox,
                            current_s2_path,
                            progress_callback=progress_callback,
                            scene_idx=scene_idx,
                            scene_total=scene_total,
                        )
                    try:
                        os.remove(zip_path)
                    except Exception:
                        pass
                else:
                    with rasterio.open(current_s2_path) as src:
                        s2_crs = src.crs

            # Prepare hyperspectral for coregistration
            matching_band_bidx = int(bidx)
            candidate_source_bands_1based: Optional[List[int]] = None
            candidate_band_limit = int(max(1, len(wl)))
            skip_source_sanitize = False
            if hyp_type == "PRISMA":
                hs_transform = estimate_prisma_geotransform(
                    lon,
                    lat,
                    cube.shape[0],
                    cube.shape[1],
                    s2_crs,
                    use_geolocation_mesh=use_geolocation_mesh_affine,
                    geolocation_mesh_stride=geolocation_mesh_stride,
                )
                temp_hs_path = os.path.join(folder_struct['temp'], f"{scene_name}_temp_georef_c{cand_idx}.tif")

                n_bands_total = cube.shape[2]
                profile = {
                    'driver': 'GTiff', 'width': cube.shape[1], 'height': cube.shape[0],
                    'count': n_bands_total, 'dtype': PROCESSING_DTYPE, 'crs': s2_crs,
                    'transform': hs_transform, 'tiled': True, 'BIGTIFF': 'YES', 'nodata': PROCESSING_NODATA
                }

                with rasterio.open(temp_hs_path, 'w', **profile) as dst:
                    cube_bip = np.transpose(cube, (2, 0, 1)).astype(PROCESSING_DTYPE)
                    dst.write(cube_bip, indexes=list(range(1, cube.shape[2] + 1)))
                candidate_band_limit = int(n_bands_total)
            else:
                # EnMAP
                with _ProgressHeartbeat(
                    progress_callback,
                    "Preparing EnMAP raster",
                    scene_idx=scene_idx,
                    scene_total=scene_total,
                    substage=f"Candidate {cand_idx + 1}: preparing source raster",
                    interval_s=progress_heartbeat_interval_s,
                ) as hb_enmap:
                    s2_crs_key = str(s2_crs)
                    cache_entry = enmap_source_cache_by_crs.get(s2_crs_key)
                    resolved_enmap = None
                    enmap_band_count = 0
                    if (
                        cache_entry is not None
                        and os.path.exists(str(cache_entry.get("path", "")))
                        and isinstance(cache_entry.get("resolved_enmap"), dict)
                    ):
                        temp_hs_path = str(cache_entry["path"])
                        resolved_enmap = cache_entry.get("resolved_enmap")
                        enmap_band_count = int(cache_entry.get("enmap_band_count", 0))
                        skip_source_sanitize = True
                        hb_enmap.update(substage=f"Candidate {cand_idx + 1}: reusing cached EnMAP source prep")
                        logger.info(
                            _fmt_issue(
                                "HS_PREP",
                                (
                                    f"Candidate {cand_idx + 1}: reused EnMAP source prep cache for CRS "
                                    f"'{s2_crs_key}' ({enmap_band_count} bands)."
                                ),
                            )
                        )
                    else:
                        temp_hs_path = enmap_processing_source_path
                        skip_source_sanitize = True

                        hb_enmap.update(substage=f"Candidate {cand_idx + 1}: resolving band mapping")
                        with rasterio.open(temp_hs_path) as src:
                            enmap_band_count = int(src.count)

                        resolved_enmap = _resolve_enmap_band_selection(
                            raster_band_count=enmap_band_count,
                            wl=wl,
                            fwhm=fwhm,
                            band_names=band_names,
                            band_detectors=band_detectors,
                            sort_idx=enmap_sort_idx,
                        )
                    for warn_msg in resolved_enmap.get("warnings", []):
                        logger.warning(_fmt_issue("HS_PREP", warn_msg))
                    if not resolved_enmap.get("ok", False):
                        raise RuntimeError(
                            _fmt_issue(
                                "HS_PREP",
                                f"Failed EnMAP band harmonization for candidate {cand_idx + 1}: "
                                f"{resolved_enmap.get('error', 'unknown error')}",
                            )
                        )

                    if len(resolved_enmap["wl"]) != len(wl):
                        wl = np.asarray(resolved_enmap["wl"], dtype=float).reshape(-1)
                        fwhm = resolved_enmap.get("fwhm")
                        band_names = resolved_enmap.get("band_names")
                        band_detectors = resolved_enmap.get("band_detectors")
                        enmap_sort_idx = resolved_enmap.get("sort_idx")
                        detector_plan = _build_detector_branch_plan(
                            wl=wl,
                            fwhm=fwhm,
                            band_names=band_names,
                            band_detectors=band_detectors,
                            sensor_type=hyp_type,
                        )
                        if not detector_plan.get("ok", False):
                            raise RuntimeError(
                                _fmt_issue(
                                    "HS_PREP",
                                    "Detector branch planning failed after EnMAP harmonization update: "
                                    f"{detector_plan.get('error', 'unknown error')}",
                                )
                            )
                        for warn_msg in detector_plan.get("warnings", []):
                            logger.warning(_fmt_issue("HS_PREP", str(warn_msg)))
                        bidx = int(np.argmin(np.abs(wl - s2_ref_wl))) + 1
                        logger.warning(
                            _fmt_issue(
                                "HS_PREP",
                                f"Adjusted EnMAP spectral metadata to {len(wl)} bands to match raster preparation.",
                            )
                        )
                    candidate_source_bands_1based = [int(b) for b in resolved_enmap["source_bands_1based"]]
                    if len(candidate_source_bands_1based) != len(wl):
                        raise RuntimeError(
                            _fmt_issue("HS_PREP", "EnMAP source band mapping length mismatch.")
                        )
                    candidate_band_limit = len(wl)

            if skip_source_sanitize:
                logger.info(
                    _fmt_issue(
                        "HS_PREP",
                        f"Candidate {cand_idx + 1}: skipped non-finite sanitation for source raster.",
                    )
                )
            else:
                sanitize_result = _sanitize_raster_nonfinite_inplace(temp_hs_path, nodata=PROCESSING_NODATA)
                if not sanitize_result.get("ok", False):
                    raise RuntimeError(
                        _fmt_issue(
                            "HS_PREP",
                            f"Failed non-finite sanitation for candidate {cand_idx + 1}: "
                            f"{sanitize_result.get('error', 'unknown error')}",
                        )
                    )
                replaced_nf = int(sanitize_result.get("nonfinite_replaced", 0))
                if replaced_nf > 0:
                    logger.warning(
                        _fmt_issue(
                            "HS_PREP",
                            f"Candidate {cand_idx + 1}: replaced {replaced_nf} non-finite source pixels with nodata.",
                        )
                    )

            if hyp_type == "ENMAP":
                enmap_source_cache_by_crs[s2_crs_key] = {
                    "path": temp_hs_path,
                    "resolved_enmap": resolved_enmap,
                    "enmap_band_count": int(enmap_band_count),
                }

                if candidate_source_bands_1based is None:
                    raise RuntimeError(_fmt_issue("HS_PREP", "Missing EnMAP source band mapping."))

                vrt_mapped_path = os.path.join(folder_struct['temp'], f"{scene_name}_bands_c{cand_idx}.vrt")
                vrt_res = _build_vrt_with_band_order(
                    source_path=temp_hs_path,
                    output_path=vrt_mapped_path,
                    source_bands_1based=candidate_source_bands_1based,
                )
                if not vrt_res.get("ok", False):
                    logger.warning(
                        _fmt_issue(
                            "HS_PREP",
                            (
                                f"Candidate {cand_idx + 1}: VRT band mapping failed "
                                f"({vrt_res.get('error', 'unknown error')}); falling back to GeoTIFF reorder copy."
                            ),
                        )
                    )
                    reordered_path = os.path.join(folder_struct['temp'], f"{scene_name}_reordered_c{cand_idx}.tif")
                    reorder_res = _stream_copy_raster_with_band_order(
                        temp_hs_path, reordered_path, candidate_source_bands_1based, compress="LZW"
                    )
                    if not reorder_res.get("ok", False):
                        raise RuntimeError(
                            _fmt_issue(
                                "HS_PREP",
                                (
                                    f"Early band mapping failed (VRT + fallback copy): "
                                    f"{reorder_res.get('error', 'unknown error')}"
                                ),
                            )
                        )
                    temp_hs_path = reordered_path
                    logger.info(
                        _fmt_issue(
                            "HS_PREP",
                            (
                                f"Candidate {cand_idx + 1}: applied fallback GeoTIFF band reordering "
                                f"(bands={len(candidate_source_bands_1based)})."
                            ),
                        )
                    )
                else:
                    temp_hs_path = vrt_mapped_path
                    logger.info(
                        _fmt_issue(
                            "HS_PREP",
                            (
                                f"Candidate {cand_idx + 1}: applied virtual EnMAP band mapping via VRT "
                                f"(bands={len(candidate_source_bands_1based)})."
                            ),
                        )
                    )

                bidx = int(np.argmin(np.abs(wl - s2_ref_wl))) + 1
                matching_band_bidx = bidx
                candidate_band_limit = len(wl)
                candidate_source_bands_1based = None

            # Ensure reference and target are in the same projection for AROSICS.
            coreg_s2_path = current_s2_path
            coreg_s2_crs = s2_crs
            if hyp_type == "ENMAP":
                with rasterio.open(temp_hs_path) as hs_src_for_crs:
                    hs_crs = hs_src_for_crs.crs
                if hs_crs is None:
                    raise RuntimeError(
                        _fmt_issue(
                            "S2_REF",
                            f"Candidate {cand_idx + 1}: EnMAP source has no CRS; cannot align reference stack.",
                        )
                    )
                if not _crs_equivalent(coreg_s2_crs, hs_crs):
                    aligned_s2_path = os.path.join(
                        folder_struct['temp'],
                        f"{scene_name}_S2_cand{cand_idx}_aligned_to_hs.tif",
                    )
                    align_result = _reproject_reference_stack_to_target_crs(
                        source_path=current_s2_path,
                        target_crs=hs_crs,
                        output_path=aligned_s2_path,
                    )
                    if not align_result.get("ok", False):
                        raise RuntimeError(
                            _fmt_issue(
                                "S2_REF",
                                (
                                    f"Candidate {cand_idx + 1}: failed to reproject Sentinel-2 reference stack "
                                    f"to EnMAP CRS: {align_result.get('error', 'unknown error')}"
                                ),
                            )
                        )
                    coreg_s2_path = str(align_result.get("path", aligned_s2_path))
                    coreg_s2_crs = hs_crs
                    logger.info(
                        _fmt_issue(
                            "S2_REF",
                            (
                                f"Candidate {cand_idx + 1}: reprojected Sentinel-2 stack to EnMAP CRS for AROSICS "
                                f"({align_result.get('source_crs')} -> {align_result.get('target_crs')})."
                            ),
                        )
                    )

            current_s2_path = coreg_s2_path
            s2_crs = coreg_s2_crs
            candidate_pre_source_path = temp_hs_path
            branch_eval = _process_detector_branch_candidate(
                candidate_idx=cand_idx,
                scene_name=scene_name,
                sensor_tag=sensor_tag,
                date_tag=date_tag,
                detector_plan=detector_plan,
                source_raster_path=candidate_pre_source_path,
                s2_raster_path=current_s2_path,
                s2_crs=s2_crs,
                hyp_type=hyp_type,
                folder_struct=folder_struct,
                matcher_profile=matcher_profile,
                prefer_fixed_band_pairs=prefer_fixed_band_pairs,
                fixed_band_pairs_by_sensor=fixed_band_pairs_by_sensor,
                bandpair_wavelength_window_nm=bandpair_wavelength_window_nm,
                min_tie_points=min_tie_points,
                min_accuracy=min_accuracy,
                max_displacement=max_displacement,
                residual_mad_factor=residual_mad_factor,
                min_band_support=min_band_support,
                allow_single_band_fallback=allow_single_band_fallback,
                consensus_group_rounding_px=consensus_group_rounding_px,
                spatial_grid_rows=spatial_grid_rows,
                spatial_grid_cols=spatial_grid_cols,
                max_points_per_cell=max_points_per_cell,
                preferred_polynomial_order=preferred_polynomial_order,
                auto_downgrade_polynomial_order=auto_downgrade_polynomial_order,
                min_gcps_order2=min_gcps_order2,
                min_cells_order2=min_cells_order2,
                progress_callback=progress_callback,
                scene_idx=scene_idx,
                scene_total=scene_total,
                progress_heartbeat_interval_s=progress_heartbeat_interval_s,
                postwarp_phasecorr_check=postwarp_phasecorr_check,
                postwarp_phasecorr_warn_threshold_px=postwarp_phasecorr_warn_threshold_px,
                postwarp_phasecorr_reject_threshold_px=postwarp_phasecorr_reject_threshold_px,
                postwarp_phasecorr_reject_bad=postwarp_phasecorr_reject_bad,
                postwarp_phasecorr_max_dim=postwarp_phasecorr_max_dim,
                s2_ref_stack_idx=s2_stack_idx,
                s2_ref_wavelength_nm=s2_ref_wl,
                config=config,
            )

            validation = dict(branch_eval.get("validation", {}))
            final_quality_pass = bool(branch_eval.get("final_quality_pass", False))
            candidate_results.append({
                "idx": cand_idx,
                "validation": validation,
                "global_success": bool(validation.get("is_valid", False)),
                "final_quality_pass": bool(final_quality_pass),
                "temp_global_path": branch_eval.get("temp_global_path"),
                "temp_local_path": branch_eval.get("temp_local_path"),
                "s2_path": current_s2_path,
                "s2_crs": s2_crs,
                "local_valid_tp_count": int(branch_eval.get("local_valid_tp_count", 0)),
                "tp_residuals": dict(branch_eval.get("tp_residuals", {})),
                "polynomial_warp_used": bool(branch_eval.get("polynomial_warp_used", False)),
                "tps_warp_used": bool(branch_eval.get("tps_warp_used", False)),
                "polynomial_order_decision": dict(branch_eval.get("polynomial_order_decision", {})),
                "polynomial_order_used": int(branch_eval.get("polynomial_order_used", preferred_polynomial_order)),
                "local_tiepoints_df": branch_eval.get("local_tiepoints_df"),
                "local_tiepoints_visualization_df": branch_eval.get("local_tiepoints_visualization_df"),
                "pre_coreg_source_path": candidate_pre_source_path,
                "pre_coreg_source_bands_1based": list(branch_eval.get("pre_coreg_source_bands_1based", [])),
                "source_bands_1based": None,
                "output_path": branch_eval.get("output_path"),
                "output_content_valid": bool(branch_eval.get("output_content_valid", False)),
                "output_validation": dict(branch_eval.get("output_validation", {})),
                "postwarp_phasecorr_qa": dict(branch_eval.get("postwarp_phasecorr_qa", {})),
                "merged_tiepoint_stages": dict(branch_eval.get("merged_tiepoint_stages", {})),
                "multiband_tiepoint_counts": dict(branch_eval.get("multiband_tiepoint_counts", {})),
                "branch_results": dict(branch_eval.get("branch_results", {})),
                "recombine_result": dict(branch_eval.get("recombine_result", {})),
                "s2_metadata": dict(selected_s2_metadata),
                "output_wl": branch_eval.get("output_wl"),
                "output_fwhm": branch_eval.get("output_fwhm"),
                "output_band_names": list(branch_eval.get("output_band_names", []) or []),
                "output_band_detectors": list(branch_eval.get("output_band_detectors", []) or []),
            })

            if final_quality_pass:
                logger.info(f"\nCANDIDATE {cand_idx + 1} ACCEPTED!")
                accepted_source_path = branch_eval.get("output_path")
                if not accepted_source_path:
                    raise RuntimeError("Accepted candidate produced no output raster.")
                if current_s2_path != s2_path:
                    _promote_s2_stack(current_s2_path, s2_path, temp_root=folder_struct['temp'])
                final_success = True
                best_candidate = candidate_results[-1]
                final_validation = best_candidate.get('validation')
                selected_s2_metadata = best_candidate.get('s2_metadata', selected_s2_metadata)
                break
            continue

        except Exception as e:
            logger.error(f"Critical error with candidate {cand_idx + 1}: {e}")
            import traceback
            traceback.print_exc()

    # Fallback: use best available candidate
    if not final_success and candidate_results:
        valid_candidates = []
        for cand in candidate_results:
            output_path = cand.get("output_path")
            if not output_path or not os.path.exists(output_path):
                continue
            if not cand.get("output_content_valid", False):
                continue
            valid_candidates.append(cand)

        valid_candidates.sort(key=lambda x: x.get('validation', {}).get('confidence', 0), reverse=True)
        if valid_candidates:
            best = valid_candidates[0]
            accepted_source_path = best.get('output_path')
            best_s2_path = best.get('s2_path')
            if best_s2_path and best_s2_path != s2_path and os.path.exists(best_s2_path):
                _promote_s2_stack(best_s2_path, s2_path, temp_root=folder_struct['temp'])
            final_success = True
            best_candidate = best
            final_validation = best_candidate.get('validation')
            selected_s2_metadata = best_candidate.get('s2_metadata', selected_s2_metadata)
            logger.warning(f"Using best available candidate #{best['idx'] + 1}")
        else:
            logger.error(_fmt_issue("QUALITY", "No candidates produced valid non-empty outputs."))

    if best_candidate is not None:
        output_wl = best_candidate.get("output_wl")
        if output_wl is not None:
            wl = np.asarray(output_wl, dtype=float).reshape(-1)
        output_fwhm = best_candidate.get("output_fwhm")
        if output_fwhm is not None:
            fwhm = np.asarray(output_fwhm, dtype=float).reshape(-1)
        output_band_names = best_candidate.get("output_band_names")
        if output_band_names:
            band_names = list(output_band_names)
        output_band_detectors = best_candidate.get("output_band_detectors")
        if output_band_detectors:
            band_detectors = list(output_band_detectors)

    # Finalize output
    _emit_progress(progress_callback, "Finalizing output", scene_idx=scene_idx, scene_total=scene_total)
    status_code = "SUCCESS" if final_success else "FAIL"
    final_source_path = accepted_source_path if accepted_source_path and os.path.exists(accepted_source_path) else None
    metadata_status = "failed"
    metadata_warnings = []
    metadata_schema_version = METADATA_SCHEMA_VERSION
    finalize_timings = {}
    finalize_raster_passes = 0
    ancillary_result: Dict[str, Any] = {
        "status": "not_requested",
        "warnings": [],
        "outputs": {"pan": None, "quality_vnir": None, "quality_swir": None},
        "warp_method": None,
        "tiepoints_used": 0,
    }

    if final_source_path and os.path.exists(final_source_path):
        t0_finalize = perf_counter()
        with _ProgressHeartbeat(
            progress_callback,
            "Finalizing output",
            scene_idx=scene_idx,
            scene_total=scene_total,
            substage="Writing final raster and metadata",
            interval_s=progress_heartbeat_interval_s,
        ):
            finalize_result = _finalize_coreg_output(
                final_source_path,
                coreg_out,
                hyp_type,
                wl,
                fwhm,
                band_names,
                band_detectors,
                source_bands_1based=(
                    best_candidate.get("source_bands_1based") if isinstance(best_candidate, dict) else None
                ),
                remove_source=True,
                normalization_params=normalization_params,
                build_overviews=bool(build_overviews),
                remove_detector_overlap=remove_detector_overlap_bands,
                strict_metadata=bool(strict_metadata),
                metadata_extension_level=metadata_extension_level,
                metadata_stats_mode=metadata_stats_mode,
                metadata_stats_sample_windows=metadata_stats_sample_windows,
                metadata_stats_seed=metadata_stats_seed,
                metadata_histogram_buckets=metadata_histogram_buckets,
                metadata_label_precision=metadata_label_precision,
                timing_logs=timing_logs,
            )
        finalize_elapsed_s = perf_counter() - t0_finalize
        if isinstance(finalize_result, dict):
            metadata_status = finalize_result.get("metadata_status", metadata_status)
            metadata_warnings = list(finalize_result.get("metadata_warnings", []))
            metadata_schema_version = _safe_parse_int(
                finalize_result.get("metadata_schema_version", metadata_schema_version),
                metadata_schema_version,
                "metadata_schema_version",
            )
            finalize_timings = dict(finalize_result.get("timings", {}))
            finalize_raster_passes = _safe_parse_int(
                finalize_result.get("raster_passes", finalize_raster_passes),
                finalize_raster_passes,
                "finalize_raster_passes",
            )
            normalization_mode = str(finalize_result.get("normalization_mode", normalization_mode))
            if finalize_result.get("normalization_params"):
                # Persist sanitized params in downstream metrics/manifest.
                norm_cfg = dict(finalize_result.get("normalization_params", {}))
                normalization_params = NormalizationParams(
                    mode=normalize_mode(norm_cfg.get("mode", normalization_mode), default=normalization_mode),
                    p_low=float(norm_cfg.get("p_low", normalization_params.p_low)),
                    p_high=float(norm_cfg.get("p_high", normalization_params.p_high)),
                    clip=bool(norm_cfg.get("clip", normalization_params.clip)),
                    eps=float(norm_cfg.get("eps", normalization_params.eps)),
                    min_valid_pixels=_safe_parse_int(
                        norm_cfg.get("min_valid_pixels", normalization_params.min_valid_pixels),
                        normalization_params.min_valid_pixels,
                        "norm_min_valid_pixels",
                    ),
                    reservoir_size=_safe_parse_int(
                        norm_cfg.get("reservoir_size", normalization_params.reservoir_size),
                        normalization_params.reservoir_size,
                        "norm_reservoir_size",
                    ),
                    seed=_safe_parse_int(
                        norm_cfg.get("seed", normalization_params.seed),
                        normalization_params.seed,
                        "norm_seed",
                    ),
                    tile_size=_safe_parse_int(
                        norm_cfg.get("tile_size", normalization_params.tile_size),
                        normalization_params.tile_size,
                        "norm_tile_size",
                    ),
                )
        if timing_logs:
            logger.info("Post-accept finalize elapsed: %.2fs", finalize_elapsed_s)

        if isinstance(finalize_result, dict) and finalize_result.get("output_validation"):
            final_output_validation = dict(finalize_result.get("output_validation", {}))
        elif best_candidate is not None and best_candidate.get("output_validation"):
            # Reuse candidate-stage content validation when finalize did not rewrite pixels.
            final_output_validation = dict(best_candidate.get("output_validation", {}))
            final_output_validation["path"] = coreg_out
        else:
            final_output_validation = _validate_coreg_raster_content(
                coreg_out,
                nodata=PROCESSING_NODATA,
                max_windows=validation_max_windows,
                stop_on_first_valid=False,
            )
        if not final_output_validation.get("ok", False):
            raise RuntimeError(
                _fmt_issue(
                    "QUALITY",
                    f"Final output is empty/invalid: {final_output_validation.get('error', 'unknown error')}",
                )
            )

        if save_pre:
            if best_candidate is None or not best_candidate.get("pre_coreg_source_path"):
                raise RuntimeError(_fmt_issue("PRE_COREG", "No candidate pre-coreg source available."))
            pre_coreg_output_path = os.path.join(folder_struct['inputs'], f"{scene_name}_pre_coreg.tif")
            _save_precoreg_output(
                best_candidate.get("pre_coreg_source_path"),
                pre_coreg_output_path,
                wl,
                source_bands_1based=(
                    best_candidate.get("pre_coreg_source_bands_1based")
                    if isinstance(best_candidate, dict)
                    else None
                ),
            )

        if (
            save_displacement_vectors
            and best_candidate is not None
            and best_candidate.get("local_tiepoints_df") is not None
        ):
            accepted_tp_df = best_candidate.get("local_tiepoints_df")
            visualization_tp_df = best_candidate.get("local_tiepoints_visualization_df", accepted_tp_df)
            if HAS_GEOPANDAS:
                try:
                    export_result = _export_displacement_shapefile_from_df(
                        tie_points_df=accepted_tp_df,
                        shapefile_path=displacement_vectors_path,
                        fallback_crs=s2_crs,
                    )
                    displacement_vectors_written_path = str(export_result.get("path"))
                    if export_result.get("warning"):
                        logger.warning(_fmt_issue("REPORTS", str(export_result.get("warning"))))
                    logger.info(
                        "Exported accepted displacement vectors (%d points, basis=%s): %s",
                        int(export_result.get("count", 0)),
                        export_result.get("abs_shift_basis"),
                        displacement_vectors_written_path,
                    )
                except Exception as shp_exc:
                    logger.warning(
                        _fmt_issue("REPORTS", f"Failed to export accepted displacement vectors: {shp_exc}")
                    )
            else:
                try:
                    _, abs_basis, abs_warning = _ensure_abs_shift_column(accepted_tp_df)
                    if abs_warning:
                        logger.warning(_fmt_issue("REPORTS", abs_warning))
                    logger.warning(
                        _fmt_issue(
                            "REPORTS",
                            "GeoPandas is unavailable; displacement vectors shapefile skipped "
                            f"(ABS_SHIFT basis={abs_basis}).",
                        )
                    )
                except Exception as abs_exc:
                    logger.warning(
                        _fmt_issue(
                            "REPORTS",
                            "GeoPandas is unavailable and ABS_SHIFT derivation failed; "
                            f"displacement vectors shapefile skipped ({abs_exc}).",
                        )
                    )

            try:
                carto_result = _write_displacement_vector_cartography_png(
                    visualization_df=visualization_tp_df,
                    output_png=displacement_cartography_path,
                    scene_name=scene_name,
                    scene_date_utc=hs_time,
                    basemap_raster_path=s2_path,
                    basemap_rgb_band_indices=(3, 2, 1),
                    quiver_cmap="RdYlGn_r",
                )
                if carto_result.get("warning"):
                    logger.warning(_fmt_issue("QUICKLOOK", str(carto_result.get("warning"))))
                logger.info(
                    "Exported displacement cartography (%d vectors, mode=%s): %s",
                    int(carto_result.get("count", 0)),
                    carto_result.get("mode"),
                    displacement_cartography_path,
                )
            except Exception as carto_exc:
                logger.warning(
                    _fmt_issue(
                        "QUICKLOOK",
                        f"Failed to export displacement cartography PNG: {carto_exc}",
                    )
                )

        quicklook_status = "running"
        quicklook_metadata = _generate_mandatory_quicklooks(
            scene_name=scene_name,
            scene_date_utc=hs_time,
            quicklooks_dir=folder_struct["quicklooks"],
            scene_raster_path=coreg_out,
            tie_points_df=None if best_candidate is None else best_candidate.get("local_tiepoints_df"),
            generate_tiepoints_png=bool(gen_tiepoint_pngs),
            wl=wl,
            target_wl_nm=s2_ref_wl,
            max_quicklook_dim=quicklook_max_dim,
            quicklook_rgb_targets_nm=quicklook_rgb_targets_nm,
            quicklook_percentiles=quicklook_percentiles,
            quicklook_gamma=quicklook_gamma,
            quicklook_dpi=quicklook_dpi,
            quicklook_crop_to_valid=quicklook_crop_to_valid,
            quicklook_scalebar=quicklook_scalebar,
            rgb_source_path=quicklook_rgb_source_path,
        )
        quicklook_outputs = dict(quicklook_metadata.get("outputs", {}))
        quicklook_status = "ok"

        if hyp_type == "PRISMA" and (save_pan or save_quality_mask):
            ancillary_result = _coregister_prisma_ancillary_outputs(
                scene_name=scene_name,
                folder_struct=folder_struct,
                s2_crs=s2_crs,
                save_pan=bool(save_pan),
                save_quality_mask=bool(save_quality_mask),
                pan_data=pan_data,
                pan_geo_info=pan_geo_info,
                vnir_quality_data=vnir_quality_data,
                swir_quality_data=swir_quality_data,
                lat_qm=lat_qm,
                lon_qm=lon_qm,
                best_candidate=best_candidate,
                hs_reference_raster_path=coreg_out,
                s2_reference_raster_path=s2_path,
                matcher_profile=matcher_profile,
                pan_gcp_mode=pan_gcp_mode,
                pan_map_dxdy_source=pan_map_dxdy_source,
                pan_target_aligned_pixels=pan_target_aligned_pixels,
                pan_use_synthetic_reference=pan_use_synthetic_reference,
                pan_min_points_for_poly2=pan_min_points_for_poly2,
                pan_local_window_size=pan_local_window_size,
                pan_local_grid_res=pan_local_grid_res,
                pan_local_max_shift=pan_local_max_shift,
                pan_local_tieP_filter_level=pan_local_tieP_filter_level,
                pan_local_max_iter=pan_local_max_iter,
                pan_residual_check=pan_residual_check,
                pan_residual_threshold_px=pan_residual_threshold_px,
                pan_residual_max_dim=pan_residual_max_dim,
                use_geolocation_mesh_affine=use_geolocation_mesh_affine,
                geolocation_mesh_stride=geolocation_mesh_stride,
            )
            for anc_warning in ancillary_result.get("warnings", []):
                logger.warning(anc_warning)
    else:
        status_code = "FAIL"
        _emit_progress(progress_callback, "Failed: no valid output produced", scene_idx=scene_idx, scene_total=scene_total, status="error")
        raise RuntimeError("Coregistration failed: no valid output produced.")

    # Build metrics
    bbox_top_left_x = None
    bbox_top_left_y = None
    bbox_bottom_right_x = None
    bbox_bottom_right_y = None
    if bbox is not None and len(bbox) == 4:
        try:
            xmin, ymin, xmax, ymax = [float(v) for v in bbox]
            bbox_top_left_x = xmin
            bbox_top_left_y = ymax
            bbox_bottom_right_x = xmax
            bbox_bottom_right_y = ymin
        except Exception:
            pass

    tp_residuals = best_candidate.get("tp_residuals", {}) if best_candidate else {}
    poly_decision = best_candidate.get("polynomial_order_decision", {}) if best_candidate else {}
    notes_text = final_validation.get("message") if isinstance(final_validation, dict) else None

    metrics_dict = {
        'scene_name': scene_name,
        'filename': os.path.basename(hs_file),
        'hyp_type': hyp_type,
        'prisma_id': extended_meta.get('prisma_id'),
        'prisma_date': extended_meta.get('prisma_date'),
        'prisma_cloud_pct': extended_meta.get('prisma_cloud_pct'),
        'prisma_sea_pct': extended_meta.get('prisma_sea_pct'),
        'observation_angle': extended_meta.get('observation_angle'),
        'rel_azimuth_angle': extended_meta.get('rel_azimuth_angle'),
        'sun_azimuth_angle': extended_meta.get('sun_azimuth_angle'),
        'solar_zenith_angle': extended_meta.get('solar_zenith_angle'),
        'view_zenith_angle': extended_meta.get('view_zenith_angle'),
        'bbox_top_left_x': bbox_top_left_x,
        'bbox_top_left_y': bbox_top_left_y,
        'bbox_bottom_right_x': bbox_bottom_right_x,
        'bbox_bottom_right_y': bbox_bottom_right_y,
        's2_product_id': selected_s2_metadata.get('product_id'),
        's2_date': selected_s2_metadata.get('date'),
        's2_cloud_cover_pct': selected_s2_metadata.get('cloud_cover_pct'),
        'tie_points_count': best_candidate.get('local_valid_tp_count', 0) if best_candidate else 0,
        'accuracy_pct': (final_validation.get('confidence', 0.0) * 100) if final_validation else 0.0,
        'residual_mean_m': tp_residuals.get('residual_mean_m'),
        'residual_median_m': tp_residuals.get('residual_median_m'),
        'residual_rmse_m': tp_residuals.get('residual_rmse_m'),
        'residual_p90_m': tp_residuals.get('residual_p90_m'),
        'polynomial_warp_used': best_candidate.get('polynomial_warp_used', False) if best_candidate else False,
        'polynomial_order_used': best_candidate.get('polynomial_order_used') if best_candidate else None,
        'polynomial_n_gcps': poly_decision.get('n_gcps') if isinstance(poly_decision, dict) else None,
        'polynomial_order_decision': poly_decision if isinstance(poly_decision, dict) else {},
        'merged_tiepoint_stages': best_candidate.get('merged_tiepoint_stages', {}) if best_candidate else {},
        'multiband_tiepoint_counts': (
            best_candidate.get('multiband_tiepoint_counts', {}) if best_candidate else {}
        ),
        'postwarp_phasecorr_qa': best_candidate.get('postwarp_phasecorr_qa', {}) if best_candidate else {},
        'status': status_code,
        'notes': notes_text,
        'rmse_global': final_validation.get('rmse_global') if isinstance(final_validation, dict) else None,
        'rmse_local': final_validation.get('rmse_local') if isinstance(final_validation, dict) else None,
        'rmse_improvement_pct': (
            final_validation.get('rmse_improvement_pct') if isinstance(final_validation, dict) else None
        ),
        'spatial_spread_score': (
            best_candidate.get('spatial_spread_score') if isinstance(best_candidate, dict) else None
        ),
        'hull_bbox_ratio': best_candidate.get('hull_bbox_ratio') if isinstance(best_candidate, dict) else None,
        'is_clustered': best_candidate.get('is_clustered') if isinstance(best_candidate, dict) else None,
        'quality_tier': best_candidate.get('quality_tier') if isinstance(best_candidate, dict) else None,
        'quality_score': best_candidate.get('quality_score') if isinstance(best_candidate, dict) else None,
        'ssim_before': final_validation.get('ssim_before') if isinstance(final_validation, dict) else None,
        'ssim_after': final_validation.get('ssim_after') if isinstance(final_validation, dict) else None,
        'ssim_delta': final_validation.get('ssim_delta') if isinstance(final_validation, dict) else None,
        'output_path': coreg_out,
        'output_content_valid': bool(final_output_validation.get("ok", False)),
        'output_valid_pixels': final_output_validation.get("valid_pixels"),
        'output_total_pixels': final_output_validation.get("total_pixels"),
        'output_valid_fraction': final_output_validation.get("valid_fraction"),
        'output_validation_windows_scanned': final_output_validation.get("windows_scanned"),
        'output_validation_truncated': bool(final_output_validation.get("truncated", False)),
        'pre_coreg_output': pre_coreg_output_path,
        'quicklook_status': quicklook_status,
        'quicklook_outputs': dict(quicklook_outputs),
        'quicklook_band_idx': quicklook_metadata.get("band_idx"),
        'quicklook_band_wavelength_nm': quicklook_metadata.get("band_wavelength_nm"),
        'quicklook_rgb_band_indices': quicklook_metadata.get("rgb_band_indices"),
        'quicklook_rgb_band_wavelength_nm': quicklook_metadata.get("rgb_band_wavelength_nm"),
        'quicklook_rgb_source': quicklook_metadata.get("rgb_source"),
        'quicklook_tiepoints_enabled': bool(quicklook_metadata.get("tiepoints_enabled", False)),
        'quicklook_tiepoints_plotted': quicklook_metadata.get("tiepoints_plotted", 0),
        'metadata_schema_version': metadata_schema_version,
        'metadata_status': metadata_status,
        'metadata_warnings': metadata_warnings,
        'normalization_mode': normalization_mode,
        'normalization_params': {
            "mode": str(normalization_mode),
            "p_low": float(normalization_params.p_low),
            "p_high": float(normalization_params.p_high),
            "clip": bool(normalization_params.clip),
            "eps": float(normalization_params.eps),
            "min_valid_pixels": int(normalization_params.min_valid_pixels),
            "reservoir_size": int(normalization_params.reservoir_size),
            "seed": int(normalization_params.seed),
            "tile_size": int(normalization_params.tile_size),
        },
        'build_overviews': bool(build_overviews),
        'remove_overlapping_bands': bool(remove_detector_overlap_bands),
        'strict_metadata': bool(strict_metadata),
        'metadata_extension_level': metadata_extension_level,
        'metadata_stats_mode': metadata_stats_mode,
        'metadata_stats_sample_windows': metadata_stats_sample_windows,
        'metadata_stats_seed': metadata_stats_seed,
        'metadata_histogram_buckets': metadata_histogram_buckets,
        'metadata_label_precision': metadata_label_precision,
        'validation_max_windows': validation_max_windows,
        'pan_gcp_mode': pan_gcp_mode,
        'pan_map_dxdy_source': pan_map_dxdy_source,
        'pan_target_aligned_pixels': bool(pan_target_aligned_pixels),
        'pan_use_synthetic_reference': bool(pan_use_synthetic_reference),
        'pan_min_points_for_poly2': int(pan_min_points_for_poly2),
        'pan_local_window_size': list(pan_local_window_size),
        'pan_local_grid_res': int(pan_local_grid_res),
        'pan_local_max_shift': float(pan_local_max_shift),
        'pan_local_tieP_filter_level': int(pan_local_tieP_filter_level),
        'pan_local_max_iter': pan_local_max_iter,
        'pan_residual_check_enabled': bool(pan_residual_check),
        'pan_residual_threshold_px': float(pan_residual_threshold_px),
        'pan_residual_estimate': dict(ancillary_result.get("pan_residual_check", {})),
        'ancillary_status': ancillary_result.get("status", "not_requested"),
        'ancillary_warp_method': ancillary_result.get("warp_method"),
        'ancillary_tiepoints_used': ancillary_result.get("tiepoints_used", 0),
        'ancillary_outputs': dict(ancillary_result.get("outputs", {})),
        'ancillary_warnings': list(ancillary_result.get("warnings", [])),
        'finalize_timings': finalize_timings,
        'finalize_raster_passes': finalize_raster_passes,
        'run_manifest_path': manifest_json_path,
        'dataset_xlsx_path': dataset_xlsx_path,
        'displacement_vectors_path': displacement_vectors_written_path,
    }

    # Cleanup
    t0_cleanup = perf_counter()
    cleanup_mode = "immediate"
    cleanup_elapsed_s = 0.0
    if defer_temp_cleanup_gui and not keep_temp_files:
        logger.info("Deferring temp cleanup in background (GUI mode).")
        cleanup_thread = threading.Thread(
            target=_cleanup_temp_folder,
            args=(folder_struct['temp'], False),
            daemon=True,
            name=f"cleanup_{scene_name}",
        )
        cleanup_thread.start()
        cleanup_mode = "deferred"
    else:
        _cleanup_temp_folder(folder_struct['temp'], keep_temp_files=keep_temp_files)
        cleanup_elapsed_s = perf_counter() - t0_cleanup
    if timing_logs:
        if cleanup_mode == "immediate":
            logger.info("Cleanup timing (s): %.2f", cleanup_elapsed_s)
        else:
            logger.info("Cleanup timing: deferred")

    metrics_dict["cleanup_mode"] = cleanup_mode
    metrics_dict["cleanup_elapsed_s"] = cleanup_elapsed_s

    # Write reports
    _write_shift_report(shift_report_path, f"COREG RESULT ({sensor_tag} {date_tag})",
                        final_validation, [f"Status: {status_code}"])
    _write_per_scene_metrics_json(metrics_dict, metrics_json_path)
    run_manifest = {
        "manifest_schema_version": RUN_MANIFEST_SCHEMA_VERSION,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": status_code,
        "scene": {
            "scene_name": scene_name,
            "filename": os.path.basename(hs_file),
            "sensor_type": hyp_type,
            "scene_idx": int(scene_idx),
            "scene_total": int(scene_total),
        },
        "input": {
            "hs_file": hs_file,
            "acquisition_time": hs_time.isoformat() if isinstance(hs_time, datetime) else None,
            "bbox": list(bbox) if bbox is not None else None,
        },
        "paths": {
            "scene_root": scene_folder,
            "reference_stack": s2_path,
            "output_raster": coreg_out,
            "pre_coreg_raster": pre_coreg_output_path,
            "quicklook_outputs": dict(quicklook_outputs),
            "ancillary_outputs": dict(ancillary_result.get("outputs", {})),
            "shift_report": shift_report_path,
            "metrics_json": metrics_json_path,
            "dataset_xlsx": dataset_xlsx_path,
            "run_manifest": manifest_json_path,
            "displacement_vectors_path": displacement_vectors_written_path,
        },
        "selection": {
            "selected_s2": selected_s2_metadata,
            "candidate_count": len(s2_candidates_list),
            "best_candidate_index": (best_candidate.get("idx", -1) + 1) if best_candidate else None,
        },
        "processing": {
            "metadata_status": metadata_status,
            "metadata_warnings": metadata_warnings,
            "normalization_mode": normalization_mode,
            "normalization_params": {
                "mode": str(normalization_mode),
                "p_low": float(normalization_params.p_low),
                "p_high": float(normalization_params.p_high),
                "clip": bool(normalization_params.clip),
                "eps": float(normalization_params.eps),
                "min_valid_pixels": int(normalization_params.min_valid_pixels),
                "reservoir_size": int(normalization_params.reservoir_size),
                "seed": int(normalization_params.seed),
                "tile_size": int(normalization_params.tile_size),
            },
            "build_overviews": bool(build_overviews),
            "remove_overlapping_bands": bool(remove_detector_overlap_bands),
            "strict_metadata": bool(strict_metadata),
            "metadata_extension_level": metadata_extension_level,
            "metadata_stats_mode": metadata_stats_mode,
            "metadata_stats_sample_windows": metadata_stats_sample_windows,
            "metadata_stats_seed": metadata_stats_seed,
            "metadata_histogram_buckets": metadata_histogram_buckets,
            "metadata_label_precision": metadata_label_precision,
            "validation_max_windows": validation_max_windows,
            "pan_gcp_mode": pan_gcp_mode,
            "pan_map_dxdy_source": pan_map_dxdy_source,
            "pan_target_aligned_pixels": bool(pan_target_aligned_pixels),
            "pan_use_synthetic_reference": bool(pan_use_synthetic_reference),
            "pan_min_points_for_poly2": int(pan_min_points_for_poly2),
            "pan_local_window_size": list(pan_local_window_size),
            "pan_local_grid_res": int(pan_local_grid_res),
            "pan_local_max_shift": float(pan_local_max_shift),
            "pan_local_tieP_filter_level": int(pan_local_tieP_filter_level),
            "pan_local_max_iter": pan_local_max_iter,
            "pan_residual_check_enabled": bool(pan_residual_check),
            "pan_residual_threshold_px": float(pan_residual_threshold_px),
            "pan_residual_estimate": dict(ancillary_result.get("pan_residual_check", {})),
            "prefer_fixed_band_pairs": bool(prefer_fixed_band_pairs),
            "fixed_band_pairs_by_sensor": fixed_band_pairs_by_sensor,
            "bandpair_wavelength_window_nm": float(bandpair_wavelength_window_nm),
            "min_band_support": int(min_band_support),
            "allow_single_band_fallback": bool(allow_single_band_fallback),
            "consensus_group_rounding_px": float(consensus_group_rounding_px),
            "spatial_stratification_grid_rows": int(spatial_grid_rows),
            "spatial_stratification_grid_cols": int(spatial_grid_cols),
            "max_points_per_cell": int(max_points_per_cell),
            "preferred_polynomial_order": int(preferred_polynomial_order),
            "auto_downgrade_polynomial_order": bool(auto_downgrade_polynomial_order),
            "min_gcps_order2": int(min_gcps_order2),
            "min_cells_order2": int(min_cells_order2),
            "postwarp_phasecorr_check": bool(postwarp_phasecorr_check),
            "postwarp_phasecorr_warn_threshold_px": float(postwarp_phasecorr_warn_threshold_px),
            "postwarp_phasecorr_reject_threshold_px": float(postwarp_phasecorr_reject_threshold_px),
            "postwarp_phasecorr_reject_bad": bool(postwarp_phasecorr_reject_bad),
            "postwarp_phasecorr_max_dim": int(postwarp_phasecorr_max_dim),
            "use_geolocation_mesh_affine": bool(use_geolocation_mesh_affine),
            "geolocation_mesh_stride": int(geolocation_mesh_stride),
            "matcher_profile_selected": {
                "sensor": matcher_profile.get("sensor"),
                "global_profile_source": matcher_profile.get("global_profile_source"),
                "global_attempt_ladder": [
                    {"ws": list(item.get("ws", ())), "max_shift": float(item.get("max_shift", 0.0))}
                    for item in matcher_profile.get("global_attempt_ladder", [])
                ],
                "local_max_shift": float(matcher_profile.get("local_max_shift", 0.0)),
                "local_grid_res": int(matcher_profile.get("local_grid_res", LOCAL_GRID_RES_M)),
                "local_window_size": list(_coerce_window_size(matcher_profile.get("local_window_size"), (256, 256))),
                "local_tieP_filter_level": int(matcher_profile.get("local_tieP_filter_level", 1)),
                "local_max_iter": matcher_profile.get("local_max_iter"),
            },
            "quicklook_status": quicklook_status,
            "quicklook_band_idx": quicklook_metadata.get("band_idx"),
            "quicklook_band_wavelength_nm": quicklook_metadata.get("band_wavelength_nm"),
            "quicklook_rgb_band_indices": quicklook_metadata.get("rgb_band_indices"),
            "quicklook_rgb_band_wavelength_nm": quicklook_metadata.get("rgb_band_wavelength_nm"),
            "quicklook_rgb_source": quicklook_metadata.get("rgb_source"),
            "quicklook_tiepoints_enabled": bool(quicklook_metadata.get("tiepoints_enabled", False)),
            "quicklook_tiepoints_plotted": quicklook_metadata.get("tiepoints_plotted", 0),
            "output_content_valid": bool(final_output_validation.get("ok", False)),
            "output_valid_fraction": final_output_validation.get("valid_fraction"),
            "output_valid_pixels": final_output_validation.get("valid_pixels"),
            "output_total_pixels": final_output_validation.get("total_pixels"),
            "output_validation_windows_scanned": final_output_validation.get("windows_scanned"),
            "output_validation_truncated": bool(final_output_validation.get("truncated", False)),
            "ancillary_status": ancillary_result.get("status", "not_requested"),
            "ancillary_warp_method": ancillary_result.get("warp_method"),
            "ancillary_tiepoints_used": ancillary_result.get("tiepoints_used", 0),
            "ancillary_warnings": list(ancillary_result.get("warnings", [])),
            "cleanup_mode": cleanup_mode,
            "cleanup_elapsed_s": cleanup_elapsed_s,
            "finalize_timings": finalize_timings,
            "finalize_raster_passes": finalize_raster_passes,
        },
        "metrics": metrics_dict,
        "config": _sanitize_config_for_manifest(config),
    }
    _write_scene_run_manifest(run_manifest, manifest_json_path)
    _write_single_scene_dataset_xlsx(metrics_dict, dataset_xlsx_path)

    log_section_header("WORKFLOW COMPLETE")
    logger.info(f"Output: {coreg_out}")
    _emit_progress(progress_callback, "Completed", scene_idx=scene_idx, scene_total=scene_total, status="done")

    return metrics_dict


def run_batch_coregistration(
    input_dir: str,
    output_dir: str,
    config: Dict[str, Any],
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, Any]:
    """
    Run batch coregistration on a directory of hyperspectral files.

    Args:
        input_dir: Directory containing hyperspectral files
        output_dir: Output directory for results
        config: Configuration dictionary with processing parameters

    Returns:
        dict: Results summary with succeeded, failed, and skipped files
    """
    import glob

    log_section_header("BATCH COREGISTRATION")
    _emit_progress(progress_callback, "Scanning input scenes", scene_idx=1, scene_total=1)

    logger.info(f"Input directory: {input_dir}")
    logger.info(f"Output directory: {output_dir}")

    os.makedirs(output_dir, exist_ok=True)

    # Find all hyperspectral files (recursive)
    all_files = []
    for pattern in [
        "*.he5",
        "*.HE5",
        "*-SPECTRAL_IMAGE.tif",
        "*-SPECTRAL_IMAGE.TIF",
        "*-SPECTRAL_IMAGE.tiff",
        "*-SPECTRAL_IMAGE.TIFF",
        "*-SPECTRAL_IMAGE.bsq",
        "*-SPECTRAL_IMAGE.BSQ",
    ]:
        all_files.extend(glob.glob(os.path.join(input_dir, "**", pattern), recursive=True))

    # Normalize and deduplicate
    all_files = list(dict.fromkeys(os.path.normcase(os.path.normpath(f)) for f in all_files))

    if not all_files:
        logger.error("No valid hyperspectral files found.")
        _emit_progress(progress_callback, "No valid scenes found", scene_idx=1, scene_total=1, status="error")
        return {
            'total': 0, 'succeeded': [], 'failed': [], 'skipped': [],
            'metrics': [], 'errors': {}, 'failed_details': []
        }

    logger.info(f"Found {len(all_files)} files")
    _emit_progress(progress_callback, "Starting batch processing", scene_idx=1, scene_total=len(all_files))

    results = {
        'total': len(all_files),
        'succeeded': [],
        'failed': [],
        'skipped': [],
        'metrics': [],
        'errors': {},
        'skip_reasons': {},
        'failed_details': [],
    }
    shared_session = config.get("_cdse_session")

    for i, hs_file in enumerate(all_files, 1):
        logger.info(f"\n[{i}/{len(all_files)}] Processing: {os.path.basename(hs_file)}")
        _emit_progress(progress_callback, "Initializing", scene_idx=i, scene_total=len(all_files))

        scene_config = dict(config)
        scene_hyp_type: Optional[str] = None
        if shared_session is not None:
            scene_config["_cdse_session"] = shared_session

        try:
            hyp_type = detect_hyp_type(hs_file)
            scene_hyp_type = hyp_type
            metrics_dict = run_coregistration(
                hs_file,
                hyp_type,
                output_dir,
                scene_config,
                progress_callback=progress_callback,
                scene_idx=i,
                scene_total=len(all_files),
            )

            if isinstance(metrics_dict, dict) and str(metrics_dict.get('status', '')).startswith('SKIPPED_'):
                fname = os.path.basename(hs_file)
                results['skipped'].append(metrics_dict)
                results['skip_reasons'][fname] = metrics_dict.get('reason', 'Scene skipped')
                _emit_progress(
                    progress_callback,
                    "Scene skipped",
                    scene_idx=i,
                    scene_total=len(all_files),
                    substage=fname,
                    scene_status="skipped",
                )
            else:
                results['succeeded'].append(os.path.basename(hs_file))
                results['metrics'].append(metrics_dict)
                _emit_progress(
                    progress_callback,
                    "Scene completed",
                    scene_idx=i,
                    scene_total=len(all_files),
                    substage=os.path.basename(hs_file),
                    scene_status="success",
                )

        except SentinelNotFoundError as e:
            msg = f"No Sentinel-2 match: {e}"
            fname = os.path.basename(hs_file)
            results['failed'].append(fname)
            results['errors'][fname] = msg
            failed_metrics = _build_failed_scene_metrics(hs_file, scene_hyp_type, msg)
            fallback_dataset_path = os.path.join(
                output_dir,
                _build_dataset_xlsx_filename(
                    scene_name=failed_metrics.get("scene_name"),
                    filename=failed_metrics.get("filename"),
                ),
            )
            failed_metrics["dataset_xlsx_path"] = fallback_dataset_path
            _write_single_scene_dataset_xlsx(failed_metrics, fallback_dataset_path)
            results["failed_details"].append(failed_metrics)
            logger.warning(msg)
            _emit_progress(
                progress_callback,
                "Scene failed (continuing batch)",
                scene_idx=i,
                scene_total=len(all_files),
                substage=fname,
                scene_status="failed",
                error=msg,
            )

        except Exception as e:
            msg = str(e)
            fname = os.path.basename(hs_file)
            results['failed'].append(fname)
            results['errors'][fname] = msg
            failed_metrics = _build_failed_scene_metrics(hs_file, scene_hyp_type, msg)
            fallback_dataset_path = os.path.join(
                output_dir,
                _build_dataset_xlsx_filename(
                    scene_name=failed_metrics.get("scene_name"),
                    filename=failed_metrics.get("filename"),
                ),
            )
            failed_metrics["dataset_xlsx_path"] = fallback_dataset_path
            _write_single_scene_dataset_xlsx(failed_metrics, fallback_dataset_path)
            results["failed_details"].append(failed_metrics)
            logger.error(f"Failed: {e}")
            _emit_progress(
                progress_callback,
                "Scene failed (continuing batch)",
                scene_idx=i,
                scene_total=len(all_files),
                substage=fname,
                scene_status="failed",
                error=msg,
            )
        finally:
            shared_session = scene_config.get("_cdse_session", shared_session)
            if shared_session is not None:
                config["_cdse_session"] = shared_session

    # Summary
    log_section_header("BATCH SUMMARY")
    logger.info(
        f"Total: {results['total']}, Success: {len(results['succeeded'])}, "
        f"Skipped: {len(results['skipped'])}, Failed: {len(results['failed'])}"
    )

    # Write summary file
    summary_path = os.path.join(output_dir, "batch_summary.txt")
    with open(summary_path, 'w') as f:
        f.write("Batch Processing Summary\n")
        f.write("=" * 60 + "\n")
        f.write(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write(f"Total: {results['total']}\n")
        f.write(f"  Succeeded: {len(results['succeeded'])}\n")
        f.write(f"  Skipped: {len(results['skipped'])}\n")
        f.write(f"  Failed: {len(results['failed'])}\n\n")

        if results['failed']:
            f.write("Failed files:\n")
            f.write("-" * 40 + "\n")
            for fname in results['failed']:
                f.write(f"  {fname}: {results['errors'].get(fname, 'Unknown')}\n")

    logger.info(f"Summary written: {summary_path}")

    # Write Excel summary if openpyxl is available
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        xlsx_path = os.path.join(output_dir, "batch_summary.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Coregistration Results"

        headers = ['Scene', 'Status', 'Sensor', 'Tie Points', 'Accuracy %', 'RMSE (m)', 'Poly Warp']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", fill_type="solid")

        row = 2
        for m in results['metrics']:
            ws.cell(row=row, column=1, value=m.get('scene_name', ''))
            ws.cell(row=row, column=2, value=m.get('status', ''))
            ws.cell(row=row, column=3, value=m.get('hyp_type', ''))
            ws.cell(row=row, column=4, value=m.get('tie_points_count', 0))
            ws.cell(row=row, column=5, value=m.get('accuracy_pct', 0))
            ws.cell(row=row, column=6, value=m.get('residual_rmse_m'))
            ws.cell(row=row, column=7, value='Yes' if m.get('polynomial_warp_used') else 'No')
            row += 1

        for skipped in results['skipped']:
            ws.cell(row=row, column=1, value=skipped.get('scene_name', ''))
            ws.cell(row=row, column=2, value='SKIPPED')
            ws.cell(row=row, column=3, value=skipped.get('hyp_type', ''))
            row += 1

        wb.save(xlsx_path)
        logger.info(f"Excel summary written: {xlsx_path}")
    except ImportError:
        logger.debug("openpyxl not available, skipping Excel summary")

    _emit_progress(
        progress_callback,
        "Batch completed",
        scene_idx=len(all_files),
        scene_total=len(all_files),
        status="done",
    )
    return results


